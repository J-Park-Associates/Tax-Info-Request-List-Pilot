"""The agreement check every test in the suite pays for once.

`tracker/ledger.py` is the engagement's own record, and since decision 88 it
is what ``read_index()`` and ``load_manifest()`` answer from wherever it has
anything to say. What keeps that honest is this: after **every** test, every
engagement folder the test left under its ``tmp_path`` that carries a record
is folded back and compared with what the *workbooks* say. One fixture, and
the whole suite - every drop sorted, every parked file filed, every scan,
every locked-Excel sidecar, every one of the storage readings in the decision
log - becomes a test of the record too.

**Compared against the workbooks' own readings, never the live ones.** The
live readers now believe the record, so ``read_index()`` and
``load_manifest()`` here would be the record compared with itself - a check
that can only pass. This fixture therefore reads
``filer.read_index_from_workbook()`` and ``manifest.statuses_from_workbook()``
(the reading every reader made before decision 88, kept public for exactly
this) and compares the record with those.

What is compared, and what is not:

- The fold of the record's index-shaped events against the workbook reading
  with its snapshot overlay: the same identities, **in the same order**, and
  for each the same row, field for field. Order matters now that the fold is
  what ``read_index()`` returns and every write rebuilds the workbook from
  it. Only when the record carries at least one index-shaped event - an
  engagement whose record holds nothing but a scan has had no index write go
  through a writer that records.
- The statuses the record's ``scanned`` events add up to against the
  workbook's scanner columns with the pending sidecar overlaid, for every
  identifier the record has ever seen a status for. An identifier the
  manifest no longer carries is passed over: a row deleted or renamed in
  Excel since the scan is dropped by ``_apply_updates`` and by
  ``with_pending`` alike, and the record is right that it was written.

An engagement with no record at all is skipped - a test that hand-builds a
legacy workbook and never writes through the writers has nothing to agree
with - and so is one whose index the tracker refuses to read, which is itself
what that test is about.

**And the view agrees with the readers** (decision 89). Wherever a test left
a view behind, its Index sheet is compared with ``read_index()`` and its
Requests sheet with ``load_manifest()``, cell for cell, through the same two
row builders the writer uses (``view.index_row``, ``view.request_row``) so
the comparison cannot drift into a second copy of those rules. Here the
*live* readers are the right side to compare against: the view is a
derivation of them, not another record of the same facts, and the whole
claim is that what a person opens says what the readers say.

A view ``view_state()`` does not call current is passed over, because that
is exactly what a pass reports as ``view_stale``: somebody had the workbook
open, the replace did not land, and the view on disk is honestly one pass
behind. Passing it over is not a hole - the state itself is asserted by
``tests/test_view.py``, including the Windows test that really holds the
file the way Excel does.

Looking changes nothing: the workbooks are read with ``quarantine=False``, so
a sidecar this fixture cannot parse is reported and left exactly where the
test put it.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from tracker import ledger, view
from tracker.filer import (
    INDEX_FILENAME,
    FilingError,
    ledger_key,
    read_index,
    read_index_from_workbook,
)
from tracker.manifest import HEADERS, SHEET_NAME, ManifestError, load_manifest, statuses_from_workbook
from tracker.scaffold import MANIFEST_FILENAME

#: Engagements the comparison passes over, by folder name, with the decision
#: that will make them agree. Empty: the fixture found nothing the writers and
#: the index disagreed about.
KNOWN_DISAGREEMENTS: dict[str, str] = {}


def _index_rows(engagement_dir):
    """The index as the *workbook* reads it, by the identity the record keys
    rows on. Never ``read_index()``, which answers from the record."""
    return {
        ledger_key(entry): {f: getattr(entry, f) for f in entry.__dataclass_fields__}
        for entry in read_index_from_workbook(engagement_dir / INDEX_FILENAME, quarantine=False)
    }


def _manifest_statuses(engagement_dir):
    """The *workbook's* scanner columns with the pending sidecar overlaid.
    Never ``load_manifest()``, which answers from the record."""
    path = engagement_dir / MANIFEST_FILENAME
    if not path.is_file():
        return None
    try:
        recorded = statuses_from_workbook(path)
    except ManifestError:
        return None        # a manifest the loader refuses; that is what the test is about
    return {identifier.lower(): status for identifier, status in recorded.items()}


def _compare_rows(name, folded, actual):
    assert set(folded) == set(actual), (
        f"{name}: the record and the index do not hold the same originals; "
        f"only in the record: {sorted(set(folded) - set(actual))}; "
        f"only in the index: {sorted(set(actual) - set(folded))}"
    )
    assert list(folded) == list(actual), (
        f"{name}: the record and the index hold the same originals in different "
        f"orders; the record says {list(folded)}, the index says {list(actual)}"
    )
    for key, row in actual.items():
        recorded = folded[key]
        for field, value in row.items():
            assert recorded.get(field) == value, (
                f"{name}: row {key!r}, field {field!r}: the record says "
                f"{recorded.get(field)!r}, the index says {value!r}"
            )


def _compare_statuses(name, recorded, actual):
    for identifier, status in recorded.items():
        expected = actual.get(identifier.lower())
        if expected is None:
            continue           # the row is not in the manifest any more
        assert status == expected, (
            f"{name}: {identifier} - the record says {status!r}, the manifest says {expected!r}"
        )


def _check(engagement_dir) -> None:
    name = engagement_dir.name
    if name in KNOWN_DISAGREEMENTS:
        return
    events = ledger.read_events(engagement_dir)
    if any(event.get(ledger.EVENT_KEY) in ledger.ROW_EVENTS for event in events):
        try:
            actual = _index_rows(engagement_dir)
        except FilingError:
            return             # an index the tracker refuses to read; that is the test
        _compare_rows(name, ledger.fold(events), actual)
    recorded = ledger.statuses(events)
    if recorded:
        actual = _manifest_statuses(engagement_dir)
        if actual is not None:
            _compare_statuses(name, recorded, actual)


def _sheet_rows(path, title):
    """One sheet of the view as text, header row dropped. None if it is not there."""
    wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=True)
    try:
        if title not in wb.sheetnames:
            return None
        return [["" if cell is None else str(cell) for cell in row]
                for row in wb[title].iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()


def _check_view(engagement_dir) -> None:
    """The view a test left behind says what the readers say, cell for cell."""
    name = engagement_dir.name
    if view.view_state(engagement_dir) != view.CURRENT:
        return          # behind: the replace did not land, which is view_stale
    path = view.path_for(engagement_dir)

    entries = read_index(engagement_dir / INDEX_FILENAME, quarantine=False)
    rows = _sheet_rows(path, view.INDEX_SHEET)
    assert rows is not None, f"{name}: the view has no index sheet"
    assert rows == [view.index_row(entry) for entry in entries], (
        f"{name}: the view's index is not what the reader says"
    )

    items = load_manifest(engagement_dir / MANIFEST_FILENAME)
    rows = _sheet_rows(path, SHEET_NAME)
    assert rows is not None, f"{name}: the view has no request sheet"
    assert rows == [[view.request_row(item)[header] for header in HEADERS] for item in items], (
        f"{name}: the view's requests are not what the reader says"
    )


@pytest.fixture(autouse=True)
def the_record_agrees_with_the_workbooks(tmp_path):
    """After every test: every engagement it wrote a record for still agrees,
    and every view it left behind says what the readers say."""
    yield
    for path in sorted(tmp_path.rglob(ledger.LEDGER_FILENAME)):
        _check(path.parent)
    for path in sorted(tmp_path.rglob(view.VIEW_FILENAME)):
        _check_view(path.parent)
