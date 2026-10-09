"""Tests for tracker/content_check.py — tier-3 rules, all local files."""

import datetime as dt

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter

from tests.conftest import make_engagement
from tracker import reasons, store
from tracker.content_check import (
    CACHE_VERSION,
    ContentCache,
    check_content,
    evaluate_rules,
    extract_text,
    has_content_rules,
    rules_fingerprint,
)
from tracker.locking import engagement_lock
from tracker.manifest import RequestItem


def text_pdf(path, text: str):
    """Hand-build a minimal one-page PDF with a real text layer."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    bodies = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        4: b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        5: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for n in sorted(bodies):
        offsets[n] = len(out)
        out += b"%d 0 obj\n" % n + bodies[n] + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(bodies) + 1)
    for n in sorted(bodies):
        out += b"%010d 00000 n \n" % offsets[n]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(bodies) + 1,
        xref_at,
    )
    path.write_bytes(bytes(out))
    return path


def item(**kwargs) -> RequestItem:
    return RequestItem(identifier="A01", document="Doc", **kwargs)


# ------------------------------------------------------------------ rules ----


def test_no_rules_short_circuits(tmp_path, monkeypatch):
    # With no content rules the file must never even be opened.
    monkeypatch.setattr(
        "tracker.content_check.extract_text",
        lambda p: (_ for _ in ()).throw(AssertionError("file was read")),
    )
    result = check_content(tmp_path / "does-not-even-exist.pdf", item())
    assert result.ok
    assert not has_content_rules(item())
    assert has_content_rules(item(required_keywords=("x",)))


def test_required_keywords(tmp_path):
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    assert check_content(pdf, item(required_keywords=("Chase",))).ok
    result = check_content(pdf, item(required_keywords=("Chase", "Wells Fargo")))
    assert not result.ok
    assert "'Wells Fargo'" in result.reason and result.code == reasons.WRONG_DOCUMENT.code


def test_keywords_case_insensitive(tmp_path):
    pdf = text_pdf(tmp_path / "s.pdf", "CHASE BANK statement for the account")
    assert check_content(pdf, item(required_keywords=("chase", "Statement"))).ok


def test_any_keywords(tmp_path):
    pdf = text_pdf(tmp_path / "s.pdf", "Monthly account summary for December")
    assert check_content(pdf, item(any_keywords=("statement", "account summary"))).ok
    result = check_content(pdf, item(any_keywords=("invoice", "receipt")))
    assert not result.ok and "invoice, receipt" in result.reason


def test_date_pattern(tmp_path):
    pdf = text_pdf(tmp_path / "s.pdf", "Statement period: December 2025")
    assert check_content(pdf, item(date_pattern=r"(?i)december\s+2025")).ok
    result = check_content(pdf, item(date_pattern=r"(?i)january\s+2026"))
    assert not result.ok and result.code == reasons.WRONG_PERIOD.code


def test_evaluate_rules_lists_all_missing():
    result = evaluate_rules("nothing here", item(required_keywords=("Alpha", "Beta")))
    assert "'Alpha'" in result.reason and "'Beta'" in result.reason


# -------------------------------------------------------------- extractors ----


def test_xlsx_extraction_with_dates(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Payroll Register Q4 2025"
    ws["B1"] = dt.date(2025, 12, 31)
    xlsx = tmp_path / "payroll.xlsx"
    wb.save(xlsx)

    assert check_content(xlsx, item(required_keywords=("payroll",))).ok
    # Date cells match both ISO and US-style patterns.
    assert check_content(xlsx, item(date_pattern=r"2025-12-31")).ok
    assert check_content(xlsx, item(date_pattern=r"12/31/2025")).ok


def test_csv_extraction(tmp_path):
    f = tmp_path / "gl.csv"
    f.write_text("date,memo\n2025-12-01,Chase transfer\n", encoding="utf-8")
    assert check_content(f, item(required_keywords=("Chase",))).ok


def test_unsupported_extension(tmp_path):
    f = tmp_path / "notes.docx"
    f.write_bytes(b"binary")
    result = check_content(f, item(required_keywords=("x",)))
    assert not result.ok and not result.extractable
    assert ".docx" in result.reason and result.code == reasons.UNCHECKABLE_TYPE.code


def test_image_only_pdf_without_ocr(tmp_path, monkeypatch):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)  # valid PDF, zero text
    scan = tmp_path / "scan.pdf"
    with scan.open("wb") as fh:
        writer.write(fh)
    monkeypatch.setattr("tracker.content_check._ocr_pdf", lambda p: None)

    result = check_content(scan, item(required_keywords=("Chase",)))
    assert not result.ok and not result.extractable
    assert result.code == reasons.NO_TEXT_LAYER.code


def test_ocr_text_used_when_available(tmp_path, monkeypatch):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    scan = tmp_path / "scan.pdf"
    with scan.open("wb") as fh:
        writer.write(fh)
    monkeypatch.setattr(
        "tracker.content_check._ocr_pdf", lambda p: "Chase Bank December 2025"
    )
    assert check_content(scan, item(required_keywords=("Chase",))).ok


# --------------------------------------- photos, and the page upright ----
#
# Decision 127. Every image below is drawn here with Pillow and thrown
# away with tmp_path: no binary is committed and no client's photo is
# anywhere near the suite. The engine itself is not needed - what is
# claimed is which pixels the reading is handed, so the reader is a fake
# (the real one reads in tests/test_ocr.py).


def photo(path, words: str = "Form W-2 Wage and Tax Statement", size=(900, 200)):
    """A photo of a document, in a form's own words."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", size, "white")
    ImageDraw.Draw(image).text(
        (20, 60), words, font=ImageFont.load_default(size=40), fill="black",
    )
    image.save(path)
    return path


def catching_the_reading(monkeypatch, reading: str = ""):
    """Hold on to the image the reader (``tracker.ocr.read_page``) is handed."""
    from tracker import ocr

    seen = {}

    def reader(image, *, name=""):
        seen["image"] = image
        return reading

    monkeypatch.setattr(ocr, "read_page", reader)
    return seen


def test_a_photo_is_read_by_ocr_like_a_scanned_pdf_and_never_has_a_text_layer(tmp_path, monkeypatch):
    """An image has no text layer to be disappointed by, so it takes the
    path a scanned PDF takes: unrouted when OCR is withheld, read by OCR
    when it is not, and the reading is OCR's - which is what keeps
    decision 50's stricter rule over it."""
    from tracker.content_check import extract, extract_text

    picture = photo(tmp_path / "w-2.png")
    assert extract_text(picture) is None                   # never a text layer

    withheld = extract(picture, ocr=False)
    assert withheld.needs_ocr and withheld.text is None    # a scan, with no words yet

    monkeypatch.setattr("tracker.content_check._ocr_image", lambda p: "Chase Bank December 2025")
    reading = extract(picture)
    assert reading.from_ocr and not reading.needs_ocr
    assert "Chase" in reading.text
    assert check_content(picture, item(required_keywords=("Chase",))).ok

    # And the machine's two answers are the machine's, for a photo as for a scan.
    monkeypatch.setattr("tracker.content_check._ocr_image", lambda p: None)
    no_engine = check_content(picture, item(required_keywords=("Chase",)))
    assert no_engine.transient and no_engine.code == reasons.NO_TEXT_LAYER.code
    monkeypatch.setattr("tracker.content_check._ocr_image", lambda p: "   ")
    assert (check_content(picture, item(required_keywords=("Chase",))).code
            == reasons.NO_TEXT_AFTER_OCR.code)


def test_exif_rotation_is_undone_before_reading(tmp_path, monkeypatch):
    """A phone records which way it was held rather than turning the pixels,
    so the commonest sideways photo there is comes upright for free. It is
    undone first, before anything else looks at the image."""
    from PIL import Image, ImageDraw

    from tracker.content_check import _ocr_image

    marked = Image.new("L", (200, 100), 255)
    ImageDraw.Draw(marked).rectangle([0, 0, 19, 19], fill=0)   # a dark corner, top-left
    tag = Image.Exif()
    tag[0x0112] = 6                                            # "turn me a quarter clockwise"
    picture = tmp_path / "held sideways.png"
    marked.save(picture, exif=tag)

    seen = catching_the_reading(monkeypatch)
    _ocr_image(picture)

    handed = seen["image"]
    assert handed.size == (100, 200)                           # turned, not merely flagged
    assert handed.getpixel((handed.width - 1, 0)) == (0, 0, 0)     # the corner came round
    assert handed.getpixel((0, 0)) == (255, 255, 255)
    assert handed.mode == "RGB"                                # in colour, as the reader was measured
    assert Image.open(picture).size == (200, 100)              # the file on disk is untouched


def test_a_page_is_handed_to_the_reader_as_it_comes_and_never_turned(tmp_path, monkeypatch):
    """Decision 169, ruling 6 as re-ruled: nothing turns a page before it is
    read. RapidOCR reads the words of a sideways page itself (the claim is
    in tests/test_ocr.py); a scanned PDF's page and a photo are handed to
    it exactly as they stand, in memory, and the file is never touched."""
    from pypdf import PdfWriter

    from tracker.content_check import _ocr_image, _ocr_pdf

    writer = PdfWriter()
    writer.add_blank_page(width=792, height=612)               # landscape, as a sideways scan is
    scan = tmp_path / "scan.pdf"
    with scan.open("wb") as fh:
        writer.write(fh)
    seen = catching_the_reading(monkeypatch, "Form 1099-R")

    assert _ocr_pdf(scan).strip() == "Form 1099-R"
    handed = seen["image"]
    assert handed.width > handed.height                        # landscape still: not turned
    assert handed.mode == "RGB"
    assert scan.read_bytes()[:5] == b"%PDF-"                   # the file itself, untouched

    picture = photo(tmp_path / "landscape.png", size=(400, 200))
    assert _ocr_image(picture) == "Form 1099-R"
    assert seen["image"].size == (400, 200)                    # as it stands


def test_the_cache_version_moved_so_a_verdict_read_sideways_is_read_again(tmp_path):
    """A verdict reached at version 9 may have been reached on nonsense: the
    page it was read from was never turned upright. The cache is
    disposable, so the version moves and the first pass after this
    decision re-reads those scans once. (Decision 141 moved it again, for
    the notice's first-page phrases, decision 169 once more: a new reader
    means new verdicts on every scan and photo, and decision 190 again: a
    verdict carries its cause's code.)"""
    assert CACHE_VERSION == 13

    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    fingerprint = rules_fingerprint(rule)
    conn = store.connect()
    engagement_id = store._engagement_row(conn, engagement)["id"]
    for version, digest in ((9, "sideways"), (11, "read by tesseract")):
        conn.execute(
            f"INSERT INTO {store.VERDICTS_TABLE} (engagement_id, digest, fingerprint, version, verdict) "
            "VALUES (?, ?, ?, ?, ?)", (engagement_id, digest, fingerprint, version, '{"ok": true}'))

    assert ContentCache(engagement).get_by_digest("sideways", fingerprint) is None
    assert ContentCache(engagement).get_by_digest("read by tesseract", fingerprint) is None


def test_a_slow_reading_is_said_and_never_cut_short(tmp_path, monkeypatch):
    """The number the reader benchmark and the owner's speed ceiling are
    measured in. Nothing acts on it: a reading that takes its time is
    finished, its words are the words the row is decided on, and the run's
    own line says which document took them."""
    import time as clock

    from tracker.content_check import extract
    from tracker.runner import SLOW_READING_NOTE, SLOW_READING_SECONDS

    def a_slow_reader(path):
        clock.sleep(0.05)
        return "Chase Bank December 2025, every word of it"

    monkeypatch.setattr("tracker.content_check._ocr_image", a_slow_reader)
    monkeypatch.setattr("tracker.runner.SLOW_READING_SECONDS", 0.01)

    reading = extract(photo(tmp_path / "slow.png"))
    assert reading.seconds >= 0.05                             # timed
    assert reading.text.endswith("every word of it")           # and read to the end

    run = an_engagement_run(slowest=[("slow.png", reading.seconds)])
    assert SLOW_READING_NOTE.format(name="slow.png", seconds=reading.seconds) in run.summary()
    assert SLOW_READING_SECONDS == 20.0                        # the shipped threshold

    quiet = an_engagement_run(slowest=[("quick.pdf", 0.001)])
    assert "slow reading" not in quiet.summary()


def an_engagement_run(**kwargs):
    """One run row, enough of one to read its summary line."""
    from pathlib import Path

    from tracker.registry import Engagement
    from tracker.runner import EngagementRun

    return EngagementRun(engagement=Engagement(path=Path("Smith 2025")), **kwargs)


# ------------------------------------------------------------------- cache ----


def counting_extractor(monkeypatch):
    calls = {"n": 0}
    real = extract_text

    def wrapped(path):
        calls["n"] += 1
        return real(path)

    monkeypatch.setattr("tracker.content_check.extract_text", wrapped)
    return calls


def an_engagement(tmp_path, *rules):
    """An engagement the store holds, whose request list is ``rules``."""
    return make_engagement(tmp_path / "Clients", list(rules), household="Smith Family",
                           scaffold=False)


def cache_rows(engagement):
    """The engagement's cache as the store holds it at CACHE_VERSION: (memos, verdicts)."""
    return store.cached_verdicts(store.connect(), engagement, version=CACHE_VERSION)


def saved(cache, engagement):
    with engagement_lock(engagement):
        cache.save()


def test_a_verdict_is_served_from_the_store_and_the_document_read_once(tmp_path, monkeypatch):
    calls = counting_extractor(monkeypatch)
    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    cache = ContentCache(engagement)

    assert check_content(pdf, rule, cache).ok
    assert check_content(pdf, rule, cache).ok
    assert calls["n"] == 1                       # second call served from the cache

    saved(cache, engagement)
    reopened = ContentCache(engagement)          # loaded from the store, not from memory
    assert check_content(pdf, rule, reopened).ok
    assert calls["n"] == 1                       # survives a save and a fresh load

    text_pdf(pdf, "Chase Bank Statement December 2025 v2")      # file changed -> re-extract
    assert check_content(pdf, rule, reopened).ok
    assert calls["n"] == 2

    stricter = item(required_keywords=("Chase", "Wells Fargo"))
    assert not check_content(pdf, stricter, reopened).ok  # rules changed
    assert calls["n"] == 3
    assert rules_fingerprint(rule) != rules_fingerprint(stricter)


def test_the_cache_is_keyed_on_content_not_path(tmp_path, monkeypatch):
    # A routed drop and its working copy are the same bytes under two
    # names; the second name must be a hit, not a second read.
    calls = counting_extractor(monkeypatch)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    rule = item(required_keywords=("Chase",))
    cache = ContentCache()
    assert check_content(pdf, rule, cache).ok
    copy = tmp_path / "A01 - Bank Statement.pdf"
    copy.write_bytes(pdf.read_bytes())
    assert check_content(copy, rule, cache).ok
    assert calls["n"] == 1
    assert cache.digest_of(copy) == cache.digest_of(pdf)


def test_a_verdict_row_at_another_cache_version_is_ignored_and_dropped_on_save(tmp_path):
    # Version 6 stored verdicts with no evidence at all; reading one back
    # would say a verdict had no reason behind it, which is worse than
    # re-extracting once. Version 7 was read before decision 85 taught
    # says() that a form heading its line with its printed title and its
    # year names itself, so a version-7 verdict would say a one-copy W-2
    # is nobody's form. Version 8 was read before decision 90 let a
    # keyword name alternatives, so a version-8 verdict read the "|" and
    # the "+" in one as words to look for and found neither. Version 9 was
    # read before decision 127 turned a page upright, so a verdict on a
    # scan that came through it may have been reached on nonsense. The
    # version is carried per row now (decision 107), and it still means
    # all that. Version 10 was read before decision 141 kept a notice's
    # header phrases to its first page, version 11 by the reader before
    # decision 169's, and version 12 carried no cause's code (decision 190).
    assert CACHE_VERSION == 13
    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    fingerprint = rules_fingerprint(rule)
    conn = store.connect()
    engagement_id = store._engagement_row(conn, engagement)["id"]
    conn.execute(
        f"INSERT INTO {store.VERDICTS_TABLE} (engagement_id, digest, fingerprint, version, verdict) "
        "VALUES (?, ?, ?, ?, ?)", (engagement_id, "deadbeef", fingerprint, 6, '{"ok": true}'))

    cache = ContentCache(engagement)
    assert cache.get_by_digest("deadbeef", fingerprint) is None      # not this version's
    assert check_content(pdf, rule, cache).evidence                  # read, with its evidence
    saved(cache, engagement)

    versions = [row[0] for row in conn.execute(
        f"SELECT version FROM {store.VERDICTS_TABLE} WHERE engagement_id = ?", (engagement_id,))]
    assert versions == [CACHE_VERSION]                               # the old row is gone
    _memos, verdicts = cache_rows(engagement)
    assert "deadbeef" not in verdicts and len(verdicts) == 1


def test_a_save_that_fails_leaves_the_previous_verdicts(tmp_path, monkeypatch):
    import sqlite3

    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    cache = ContentCache(engagement)
    check_content(pdf, rule, cache)
    saved(cache, engagement)
    before = cache_rows(engagement)
    assert before[0] and before[1]

    real = store.connect()

    class FailsOnTheVerdict:
        """The process's connection, whose verdict write raises mid-transaction."""

        def execute(self, sql, *args):
            if f"INTO {store.VERDICTS_TABLE}" in sql:
                raise sqlite3.OperationalError("disk I/O error")
            return real.execute(sql, *args)

        def __getattr__(self, name):
            return getattr(real, name)

    monkeypatch.setattr(store, "connect", lambda path=None: FailsOnTheVerdict())
    text_pdf(pdf, "Chase Bank Statement December 2025 v2")        # a memo and a verdict to write
    check_content(pdf, rule, cache)
    with pytest.raises(sqlite3.OperationalError, match="disk I/O error"):
        saved(cache, engagement)
    monkeypatch.undo()

    assert cache_rows(engagement) == before                          # the memo write rolled back too
    assert real.in_transaction is False


def test_a_cache_with_no_engagement_never_touches_the_store(tmp_path, monkeypatch):
    def never(path=None):
        raise AssertionError("the store was asked for")

    monkeypatch.setattr(store, "connect", never)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    rule = item(required_keywords=("Chase",))
    cache = ContentCache()
    assert check_content(pdf, rule, cache).ok
    assert cache.get(pdf, rules_fingerprint(rule)) is not None
    cache.save()                                                     # nothing to touch, no lock needed


def test_a_cache_for_an_engagement_the_store_does_not_hold_is_refused_by_name(tmp_path):
    folder = tmp_path / "Clients" / "Nobody 2025"
    folder.mkdir(parents=True)
    with pytest.raises(store.StoreError, match="does not hold this engagement") as raised:
        ContentCache(folder)
    assert folder.name in str(raised.value)


def test_a_transient_verdict_is_still_never_cached(tmp_path, monkeypatch):
    from tracker.content_check import Extraction

    monkeypatch.setattr("tracker.content_check.extract", lambda path, ocr=True: Extraction(
        None, reason=reasons.NO_TEXT_LAYER.format(), extractable=False, transient=True))
    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    pdf = text_pdf(tmp_path / "s.pdf", "a scan")
    cache = ContentCache(engagement)

    verdict = check_content(pdf, rule, cache)
    assert verdict.transient and not verdict.ok
    assert cache.get(pdf, rules_fingerprint(rule)) is None
    saved(cache, engagement)
    memos, verdicts = cache_rows(engagement)
    assert memos and verdicts == {}                                  # the digest, never the verdict


def test_the_cache_save_refuses_outside_the_engagement_lock(tmp_path):
    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    cache = ContentCache(engagement)
    check_content(pdf, rule, cache)

    with pytest.raises(store.StoreError, match="engagement lock"):
        cache.save()
    assert cache_rows(engagement) == ({}, {})
    saved(cache, engagement)                                         # the same save, under the lock
    assert cache_rows(engagement) != ({}, {})


def test_own_forms_is_none_for_one_family_and_the_named_set_for_two():
    from tests.samples import scanned_1098_lines, scanned_w2_lines
    from tracker.content_check import MULTI_FORM_FAMILIES, own_forms, self_named_forms

    one = "\n".join(scanned_w2_lines(2025))
    two = "\n".join(scanned_w2_lines(2025) + scanned_1098_lines(2025))
    assert MULTI_FORM_FAMILIES == 2
    assert own_forms(self_named_forms(one)) is None                # the ordinary reading
    assert own_forms(self_named_forms("a letter about nothing in particular")) is None
    assert own_forms(self_named_forms(two)) == set(self_named_forms(two)) == {"w2", "1098"}


def test_on_every_corpus_form_the_scans_miss_verdict_equals_the_routers_kept_verdict():
    """For every placement in tests/irs/: route with a memory cache, then
    read the same file through check_content() with an empty cache on the
    row(s) it filed under. The kept verdict and the miss verdict are one
    verdict - ok, reason and evidence - so a rebuilt store reaches what
    the router filed on."""
    from tests.test_irs_forms import EXPECT, IRS
    from tracker.manifest import validated
    from tracker.router import route_file
    from tracker.templates import template_items

    catalogs: dict = {}

    def rows(form, year):
        if (form, year) not in catalogs:
            from dataclasses import replace
            catalogs[(form, year)] = [replace(i, min_size_kb=0, date_pattern="")
                                      for i in validated(template_items(form, year=year))]
        return catalogs[(form, year)]

    compared = 0
    for pdf, form, year, expected in EXPECT:
        if expected is None:
            continue
        items = rows(form, year)
        kept = ContentCache()
        routing = route_file(IRS / pdf, items, cache=kept)
        assert routing.identifier == expected, (pdf, form, routing.reason)
        for identifier in (routing.identifier, *routing.also):
            row = next(i for i in items if i.identifier == identifier)
            remembered = kept.get(IRS / pdf, rules_fingerprint(row))
            assert remembered is not None, (pdf, form, identifier)
            miss = check_content(IRS / pdf, row, ContentCache())
            assert (miss.ok, miss.reason, miss.evidence) == (
                remembered.ok, remembered.reason, remembered.evidence), (pdf, form, identifier)
            compared += 1
    assert compared >= len([e for e in EXPECT if e[3] is not None])


def test_cache_prune(tmp_path):
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement page")
    rule = item(required_keywords=("Chase",))
    cache = ContentCache()
    check_content(pdf, rule, cache)
    assert cache.get(pdf, rules_fingerprint(rule)) is not None

    cache.prune(existing=set())                   # file no longer in any folder
    assert cache.get(pdf, rules_fingerprint(rule)) is None


def _many_page_pdf(path, texts):
    """One page per entry in ``texts``, each with a real text layer."""
    pages = []
    for text in texts:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        pages.append(content)
    n = len(pages)
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>"}
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n)).encode()
    objects[2] = b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % n
    font_id = 3 + 2 * n
    for i, content in enumerate(pages):
        page_id, stream_id = 3 + 2 * i, 4 + 2 * i
        objects[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (stream_id, font_id)
        )
        objects[stream_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content)
    objects[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objects):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objects[num] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for num in sorted(objects):
        out += b"%010d 00000 n \n" % offsets[num]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))
    return path


def test_only_the_first_pages_of_a_pdf_are_read(tmp_path):
    from tracker.content_check import extract_text

    texts = [f"page {i + 1} filler" for i in range(15)]
    texts[2] = "Form W-2 early"
    texts[12] = "Form 1098 late"
    text = extract_text(_many_page_pdf(tmp_path / "long.pdf", texts))
    assert "W-2 early" in text
    assert "1098 late" not in text
    assert "page 10 filler" in text and "page 11 filler" not in text


# ------------------------------------------------ a form's own number ----


@pytest.mark.parametrize("text, own", [
    ("Form 1040 (2025)", "1040"), ("Form 1099-DIV (Rev. January 2024)", "1099div"), ("941 for 2026:", "941"),
    ("Form W-2 Wage and Tax Statement 2025", "w2"),        # plain: nothing quotes it, it is named in its own right
])
def test_a_form_names_itself_with_its_year_or_revision(text, own):
    from tracker.content_check import _PLAIN_WEIGHT, _mentions

    (key, _, weight), = list(_mentions(text.lower()))
    assert key == own and weight >= _PLAIN_WEIGHT


@pytest.mark.parametrize("text", [
    "(Form 1040)", "Form 1040 or 1040-SR", "Form 1040 instructions", "Form 1040, line 8", "Form 1040 for individuals",
    "attach Form 1098", "Forms W-2", "reported on Form 1099-INT", "a 1098 - for my return", "see Form 1099-NEC, to use",
    "such as Form 1040)", "Interest and dividends: attach Forms 1099-INT and 1099-DIV",
])
def test_a_form_quoted_in_a_sentence_is_a_reference(text):
    from tracker.content_check import _REFERENCE_WEIGHT, _mentions

    assert all(weight == _REFERENCE_WEIGHT for _, _, weight in _mentions(text.lower()))


def test_the_dominant_form_is_the_heaviest_then_the_first_mentioned():
    from tracker.content_check import dominant_forms

    # A W-2's instructions name Form 1040 many times; the W-2 names itself with its year.
    assert dominant_forms("Form W-2 (2025) " + "see Form 1040 line 1. " * 30) == {"w2"}
    # Two forms named in their own right, equally: the first mentioned.
    assert dominant_forms("Form 1099-INT (2025) Interest Income\nForm 1099-DIV (2025) Dividends") == {"1099int"}
    # Once in passing, beside another form: nobody's.
    assert dominant_forms("You must file Form 941, Form 940 and Form W-2.") == set()
    # A number that is not a known family counts only after "Form", and takes a variant only by a dash.
    assert dominant_forms("Form 4562 (2025) Depreciation and Amortization") == {"4562"}
    assert dominant_forms("4562 4562 4562 (2025)") == set()
    assert dominant_forms("Form 1125-E (Rev. October 2016) Compensation of Officers") == {"1125e"}
    assert dominant_forms("Form 4562 depreciation\nForm 4562 depreciation") == {"4562"}   # "4562 depreciation" is not 4562-DEPR


def test_a_form_told_to_the_reader_is_a_reference_whatever_follows_it():
    # The tenth reading: "attach Form 1098 (2025)" read as a form naming
    # itself, because the year was looked at before the word in front.
    from tracker.content_check import _REFERENCE_WEIGHT, _mentions, says

    for text in ("attach Form 1098 (2025)", "see Form 1098 for 2025 mailed separately", "Forms W-2 2025"):
        assert all(weight == _REFERENCE_WEIGHT for _, _, weight in _mentions(text.lower())), text
    assert not says("2025 Individual Income Tax Organizer\nMortgage interest paid: attach Form 1098 (2025)", "1098")
    assert not says("Annual Escrow Account Disclosure Statement\nFor your deduction, see Form 1098 for 2025.", "1098")


def test_a_title_that_lists_three_forms_names_none_of_them():
    from tracker.content_check import _title_forms, says

    checklist = ("2025 Individual Income Tax Organizer\nIncome Documents Checklist\n"
                 "Form W-2 - Wage and Tax Statement\nForm 1098 - Mortgage Interest Statement\n"
                 "Form 1099-INT - Interest Income\nForm 1099-DIV - Dividends and Distributions")
    assert _title_forms(checklist.lower()) == set()
    assert not says(checklist, "1098") and not says(checklist, "1099-int")
    # A transmittal lists its forms with commas, and every form after one
    # continues the list (the thirteenth reading), which leaves the first
    # naming itself - and the firm's letter is still none of the rows'
    # documents, because none of them asks for a W-2 by its number alone.
    transmittal = "Enclosed please find: Form W-2 2025, Form 1098 2025, Form 1099-INT 2025"
    assert _title_forms(transmittal.lower()) == {"w2"}
    assert not says(transmittal, "1098") and not says(transmittal, "1099-int")
    # Families, not numbers: a broker's consolidated 1099 names the family and
    # its parts and is one family's document, which the rows then contest.
    composite = ("vanguard brokerage services\n2025 consolidated form 1099 - account 8812-4455\n"
                 "form 1099-int   interest income\nform 1099-div   dividends and distributions\n"
                 "form 1099-b   proceeds from broker and barter exchange transactions")
    assert _title_forms(composite) == {"1099int", "1099div", "1099b"}   # "Form 1099 - Account" is set off by its dash
    assert says(composite, "1099-int") and says(composite, "1099-b")
    # A checklist sets a form off from its title with a dash; no form names itself that way.
    assert _title_forms("form 1099-int - interest income\nform 1099-div - dividends and distributions") == set()
    assert _title_forms("form 1099-int   interest income   2025\nform 1099-div   dividends   2025") == {"1099int", "1099div"}
    assert _title_forms("form w-2 - 2025 wage and tax statement") == {"w2"}     # a dash before the year is still dated
    # A sentence wraps where it will: the word before a form may end the line above.
    assert not says("2025 organizer\nmortgage interest paid: please attach\nform 1098 (2025) from each lender", "1098")
    assert not says("annual escrow account disclosure\nfor your deduction, see\nform 1098 for 2025 mailed separately", "1098")


def test_a_form_number_in_the_title_counts_only_when_named_in_its_own_right():
    from tracker.content_check import says

    attention = ("Attention: Which Revision To Use for Which Year. For all forms that we do not issue annually "
                 "(such as Form 1099-NEC), we issue the revision to use for the next calendar year. ")
    assert not says(attention + "Form 1098-T Tuition Statement 2025", "1099-nec")
    assert says(attention + "Form 1098-T Tuition Statement 2025", "1098-t")
    assert not says("2025 Tax Organizer\nMortgage interest: attach Form 1098\nRetirement: attach Forms 1099-R", "1098")
    assert says("Form 1098 Mortgage Interest Statement 2025", "1098")
    assert says("Attached: Form 1099-INT, Form 1098 and W-2 for 2025 tax prep\nForm W-2 Wage and Tax Statement", "w-2")
    assert not says("Subject: Your 2025 Form 1099-R is ready\nForm 5498 IRA Contribution Information 2025", "1099-r")


def test_a_single_dash_set_off_form_beside_its_year_names_itself():
    # Decision 69 made a dash after a form number a checklist's line, which
    # parked a payer's own substitute form: "2025 Form 1099-INT - Interest
    # Income" is how a bank prints its title. A menu *lists* - two such
    # lines or more in the title window - and a single one, dated on its
    # own line, is a form naming itself.
    from tracker.content_check import _title_forms, says

    assert says("2025 Form 1099-INT - Interest Income\nAlly Bank  1 Interest income 1,842.55", "1099-int")
    assert says("Form 1098 - Mortgage Interest Statement  2025\n1 Mortgage interest received 14,220.19", "1098")
    menu = ("2025 Individual Income Tax Organizer - Interest and Dividend Income\n"
            "Form 1099-INT - Interest Income\nForm 1099-DIV - Dividends and Distributions")
    assert _title_forms(menu.lower()) == set()
    # No year beside it: a checklist's line, as before.
    assert not says("Income Documents\nForm 1099-INT - Interest Income\nPayer  Amount", "1099-int")
    # An ask before it is an ask, year or no year.
    assert not says("2025 Organizer\nPlease attach Form 1098 - Mortgage Interest Statement 2025", "1098")


def test_a_form_after_a_comma_continues_a_list():
    # A cover's last form carries no punctuation after it, so it was read as
    # named in its own right and blinded the form printed behind it.
    from tracker.content_check import says

    cover = ("Charles Schwab  2025 Tax Reporting Package\n"
             "Enclosed: Form 1099-INT, Form 1098, Form 5498\n"
             "Form 1099-INT Interest Income 2025\n1 Interest income 1,842.55")
    assert not says(cover, "5498") and not says(cover, "1098")
    assert says(cover, "1099-int")


def test_a_menu_line_names_neither_the_form_nor_its_title():
    # The form and its title on a checklist line are one reference: the
    # firm's organizer sets "Form 1099-B" off from "Proceeds From Broker
    # and Barter Exchange Transactions" with a dash, and the row that keys
    # on the broker's words filed the organizer page as the client's 1099-B.
    from tracker.content_check import says

    page = ("2025 Individual Income Tax Organizer - Investment Income\n"
            "Form 1099-B - Proceeds From Broker and Barter Exchange Transactions\n"
            "Form 1099-DIV - Dividends and Distributions\n"
            "Please list every brokerage account below.  Broker  Account number")
    assert not says(page, "1099-b")
    assert not says(page, "proceeds from broker")
    # A comma sets a form off from its title the same way.
    assert not says("Form 1099-INT, Interest Income\nPayer  Amount", "interest income")
    # The form's own page is not a menu, so its title is its own.
    assert says("Form 1099-B Proceeds From Broker and Barter Exchange Transactions 2025\n"
                "1d Proceeds 24,318.55", "proceeds from broker")


def test_a_document_does_not_ask_for_itself():
    # An ask governs the rest of its sentence, because one ask names several
    # documents; across a line break it governs only the wrap.
    from tracker.content_check import says

    organizer = ("2025 Individual Income Tax Organizer - Investment Income\n"
                 "Did you sell any stocks, bonds or mutual funds in 2025?\n"
                 "If yes, attach your brokerage statement and any realized gain and loss report.")
    assert not says(organizer, "brokerage statement")
    assert not says(organizer, "realized gain and loss")
    fiduciary = ("2025 Fiduciary Organizer - Investment Income\n"
                 "Attach the year-end account statement for each brokerage account.\n"
                 "Attach the realized gain and loss report for the year.")
    assert not says(fiduciary, "year-end account statement")
    assert not says(fiduciary, "realized gain and loss")
    assert not says("Mortgage interest: please attach\nyour brokerage statement", "brokerage statement")
    # The broker's own statement says it plainly, and a line further up than
    # the wrap is another label, not the rest of a request: a return's
    # "Attach Forms W-2G and 1099-R if tax was withheld" two lines above its
    # jurat asks for nothing on the jurat's line.
    assert says("Charles Schwab & Co., Inc.\nYear-End Brokerage Statement January 1, 2025 - December 31, 2025\n"
                "Realized Gain and Loss Summary", "brokerage statement")
    assert says("Attach Forms W-2G and 1099-R if tax was withheld\n1e Taxable dependent care benefits\n"
                "Sign Here Under penalties of perjury, I declare", "under penalties of perjury")


def test_a_keyword_may_name_alternatives_and_the_words_one_of_them_wants_together():
    # Decision 90. Required Keywords is an AND over its cells and Any
    # Keywords an OR over theirs, and neither can say "these three words,
    # or this one" - which is what the prior-year return row has to say,
    # because a standalone state return shares none of the federal
    # return's lines, jurat included.
    from tracker.content_check import contains_keyword, says

    prior_return = ("individual income tax return + filing status + under penalties of perjury"
                    " | resident income tax return")
    federal = ("Form 1040 2024 U.S. Individual Income Tax Return\n"
               "Filing Status Single Married filing jointly\n"
               "Sign Here Under penalties of perjury, I declare")
    state = "TAXABLE YEAR FORM\n2024 California Resident Income Tax Return 540"
    assert says(federal, prior_return)
    assert says(state, prior_return)
    # Two of the three is not the federal alternative: the firm's own
    # organizer says both and signs nothing (decision 65).
    assert not says("2024 Individual Income Tax Organizer\nFiling Status Single", prior_return)
    # A keyword holding neither character is one phrase, exactly as before.
    assert says("Trial Balance - Year-End 2025", "trial balance")
    # A file name is read the same way, phrase by phrase.
    assert contains_keyword("2024 california resident income tax return", prior_return)
    assert not contains_keyword("2024 individual income tax return", prior_return)


def test_the_alternatives_of_a_keyword_are_read_where_a_keyword_cell_is_read():
    # In the module that owns how a keyword cell is read, so the catalog
    # and the matcher cannot disagree about what one keyword means.
    from tracker.manifest import keyword_alternatives

    assert keyword_alternatives("trial balance") == (("trial balance",),)
    assert keyword_alternatives("a + b | c") == (("a", "b"), ("c",))
    assert keyword_alternatives("  spaced  +  out  ") == (("spaced", "out"),)
    assert keyword_alternatives("|") == ()


def test_the_any_keywords_are_read_apart_from_every_other_rule():
    # The year a Period implies is a check on a document a keyword already
    # matched, so the router can ask which rows a document's words fit
    # without reading a verdict's sentence back.
    from tracker.content_check import any_keyword_matched

    childcare = item(any_keywords=("child care statement",), date_pattern=r"(?i)\b2025\b")
    assert any_keyword_matched("2024 Child Care Statement", childcare)
    assert not any_keyword_matched("2024 Tuition Statement", childcare)
    assert not any_keyword_matched("anything at all", item(required_keywords=("w-2",)))


# ------------------------------------------------ a keyword's words on one line ----


def test_a_keywords_words_are_on_one_line_or_wrap_as_a_heading():
    from tracker.content_check import contains_keyword

    # Two labels on two lines are not one phrase.
    assert not contains_keyword("19 Distributions\nSchedule K-1 (Form 1065) 2025", "distribution schedule")
    assert not contains_keyword("16 Items affecting shareholder basis\nSchedule K-1 (Form 1120-S)", "shareholder basis schedule")
    # A heading wraps from the start of its line, two words or more on the first.
    assert contains_keyword("Balance Sheet\nAs of December 31, 2025", "balance sheet as of")
    assert contains_keyword("  Fixed Asset\n  Schedule", "fixed asset schedule")
    assert contains_keyword("Statement of\nFinancial Position", "statement of financial position")
    # One word over another is two labels; the 1098's one-word-per-line title is known by its number.
    assert not contains_keyword("Mortgage\nInterest\nStatement", "mortgage interest statement")
    assert not contains_keyword("Distributions\nSchedule", "distribution schedule")
    # A phrase that begins mid-line and runs on is a sentence, not a heading.
    assert not contains_keyword("the items affecting shareholder basis\nschedule", "shareholder basis schedule")
    # One line, as before: spaces, tabs, a no-break space, a dash.
    assert contains_keyword("Fixed\xa0Asset\tSchedule", "fixed asset schedule")
    assert contains_keyword("fixed-asset-schedule.pdf", "fixed asset schedule")
    assert not contains_keyword("W\n2", "w-2")


def test_a_sheets_row_is_a_line_and_its_cells_are_set_apart(tmp_path):
    from tracker.content_check import contains_keyword

    wb = Workbook()
    ws = wb.active
    ws.append(["Fixed Asset", "Schedule"])
    ws.append(["Distributions"])
    ws.append(["Schedule"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    text = extract_text(path)
    assert "Fixed Asset\tSchedule" in text
    assert contains_keyword(text, "fixed asset schedule")
    assert not contains_keyword(text, "distribution schedule")


# --------------------------------------------------------------- evidence ----


#: A W-2 as a reader meets it: its number and title at the top, its own
#: number again at the foot of page 1 the way the IRS sets it, and a box
#: label deep on page 2 that no footer and no title carries.
_FILLER = "Wage and Tax Statement for the employee named below. " * 12
TWO_PAGE_W2 = (
    "Form W-2 (2025) Wage and Tax Statement\n" + _FILLER
    + "\nemployer identification number 00-0000000\nForm W-2 (2025)\n"
    + "\f" + "dependent care benefits paid in the year\n" + _FILLER
    + "\nsafe harbor election\nCopy B to be filed with the return\n"
)


def test_a_title_self_mention_is_evidence_from_the_title():
    from tracker.content_check import RULE_REQUIRED, WHERE_TITLE

    result = evaluate_rules(TWO_PAGE_W2, item(required_keywords=("W-2", "wage and tax statement")))
    assert result.ok
    assert [(e.rule, e.term, e.where, e.page) for e in result.evidence] == [
        (RULE_REQUIRED, "W-2", WHERE_TITLE, 1),
        (RULE_REQUIRED, "wage and tax statement", WHERE_TITLE, 1),
    ]


def test_a_footer_hit_names_the_footer_and_the_page_it_was_on():
    from tracker.content_check import WHERE_FOOTER

    result = evaluate_rules(TWO_PAGE_W2, item(required_keywords=("safe harbor election",)))
    assert result.ok
    assert [(e.where, e.page) for e in result.evidence] == [(WHERE_FOOTER, 2)]
    # Page 1's own footer is a footer too, not "the first page".
    first = evaluate_rules(TWO_PAGE_W2, item(required_keywords=("employer identification number",)))
    assert [(e.where, e.page) for e in first.evidence] == [(WHERE_FOOTER, 1)]


def test_a_deep_any_keyword_says_it_was_deep_and_which_page():
    from tracker.content_check import RULE_ANY, WHERE_DEEP

    result = evaluate_rules(TWO_PAGE_W2, item(any_keywords=("dependent care benefits", "tips")))
    assert result.ok
    # Only the keyword that matched leaves evidence; "tips" said nothing.
    assert [(e.rule, e.term, e.where, e.page) for e in result.evidence] == [
        (RULE_ANY, "dependent care benefits", WHERE_DEEP, 2),
    ]


def test_the_date_rule_records_the_period_it_checked_never_the_document():
    from tracker.content_check import RULE_DATE

    rule = item(period="TY2025", date_pattern=r"(?i)\b2025\b",
                required_keywords=("wage and tax statement",))
    result = evaluate_rules(TWO_PAGE_W2, rule)
    assert result.ok
    dates = [e for e in result.evidence if e.rule == RULE_DATE]
    # The row's own Period, which a person can read; the regex it derives is
    # in the reason when the check fails, and no word of the document is here.
    assert [e.term for e in dates] == ["TY2025"]


def test_a_failing_verdict_keeps_what_did_match():
    # The verdict and its sentence are untouched; what matched is the lead a
    # person works the parked file from.
    result = evaluate_rules(
        TWO_PAGE_W2, item(required_keywords=("wage and tax statement", "1099-INT")),
    )
    assert not result.ok and "'1099-INT'" in result.reason
    assert [e.term for e in result.evidence] == ["wage and tax statement"]


def test_the_compact_evidence_parses_back_to_what_was_formatted():
    from tracker.content_check import (
        RULE_FILENAME,
        WHERE_TITLE,
        Evidence,
        format_evidence,
        parse_evidence,
    )

    record = {
        "C01": evaluate_rules(TWO_PAGE_W2, item(
            period="TY2025", date_pattern=r"(?i)\b2025\b",
            required_keywords=("W-2", "safe harbor election"),
            any_keywords=("dependent care benefits",),
        )).evidence,
        "L01": (Evidence(RULE_FILENAME, "1098-t", WHERE_TITLE),),
    }
    written = format_evidence(record)
    assert "C01: W-2@title:1 required" in written
    assert "L01: 1098-t@title filename" in written   # a name has no page
    assert parse_evidence(written) == record
    assert parse_evidence("") == {}


def test_a_verdict_read_back_from_the_store_brings_its_evidence(tmp_path):
    from tracker.content_check import RULE_REQUIRED, WHERE_TITLE

    rule = item(required_keywords=("Chase",))
    engagement = an_engagement(tmp_path, rule)
    pdf = text_pdf(tmp_path / "s.pdf", "Chase Bank Statement December 2025")
    cache = ContentCache(engagement)
    assert check_content(pdf, rule, cache).ok
    saved(cache, engagement)

    reopened = ContentCache(engagement)
    hit = reopened.get(pdf, rules_fingerprint(rule))
    assert hit is not None and hit.ok
    assert [(e.rule, e.term, e.where) for e in hit.evidence] == [
        (RULE_REQUIRED, "Chase", WHERE_TITLE),
    ]


# ---------------------------------------------- decision 137: the security review ----


def test_a_long_text_file_is_read_to_the_cap_and_says_so(tmp_path, monkeypatch):
    """Decision 137 (M5): a ``.csv``/``.tsv``/``.txt`` is read up to its first
    ``TEXT_READ_CAP_MB`` and no further, and a verdict reached on the cut
    text says the rest was never read."""
    import tracker.content_check as content_check

    monkeypatch.setattr(content_check, "TEXT_READ_CAP_MB", 1)
    cap = 1024 * 1024
    early = tmp_path / "early.csv"
    early.write_bytes(b"Form 1099-B proceeds\n" + b"x," * cap)
    late = tmp_path / "late.csv"
    late.write_bytes(b"x," * cap + b"\nForm 1099-B proceeds\n")

    reading = content_check.extract(early)
    assert reading.cut and len(reading.text) == cap
    rules = item(required_keywords=("1099-B",))
    assert check_content(early, rules).ok                    # found in the part that was read

    verdict = check_content(late, rules)
    assert not verdict.ok
    assert verdict.code == reasons.WRONG_DOCUMENT.code
    assert verdict.reason.endswith(reasons.TEXT_CUT.format(limit=1))

    whole = tmp_path / "short.txt"
    whole.write_text("Form 1099-B proceeds", encoding="utf-8")
    assert not content_check.extract(whole).cut


def test_a_giant_pdf_page_is_rendered_within_the_pixel_budget(tmp_path, monkeypatch):
    """Decision 137 (B1, the ruling's first half): each page is rendered at
    ``min(2, sqrt(budget / page area))``. A maximum-size PDF page (14,400
    points square) at the old fixed scale of 2 is some 830 million pixels,
    handed to Pillow past its own guard. An ordinary page keeps scale 2."""
    import tracker.content_check as content_check

    pypdfium2 = pytest.importorskip("pypdfium2")

    def pdf(path, width, height):
        writer = PdfWriter()
        writer.add_blank_page(width=width, height=height)
        with path.open("wb") as handle:
            writer.write(handle)
        return path

    asked: list[tuple[float, float]] = []

    class Stop(Exception):
        pass

    def render(page, *, scale=1.0, **_kwargs):
        width, height = page.get_size()
        asked.append((width * scale, height * scale))
        raise Stop("rendering stops here: the scale is what is under test")

    monkeypatch.setattr(pypdfium2.PdfPage, "render", render)

    for path in (pdf(tmp_path / "giant.pdf", 14400, 14400), pdf(tmp_path / "letter.pdf", 612, 792)):
        with pytest.raises(Stop):
            content_check._ocr_pdf(path)
    (giant_w, giant_h), (letter_w, letter_h) = asked
    assert giant_w * giant_h <= content_check.PIXEL_BUDGET * 1.0001
    assert giant_w * giant_h > content_check.PIXEL_BUDGET * 0.99       # read smaller, not refused
    assert (letter_w, letter_h) == (612 * 2.0, 792 * 2.0)              # an ordinary page as before


def test_a_photo_past_the_pixel_budget_is_read_smaller(tmp_path, monkeypatch):
    """Decision 137 (B1): a photo is checked against the same budget before
    it is decoded, and read reduced to it rather than refused."""
    from PIL import Image

    import tracker.content_check as content_check
    from tracker import ocr

    monkeypatch.setattr(content_check, "PIXEL_BUDGET", 10_000)
    seen = []
    monkeypatch.setattr(ocr, "read_page", lambda image, **k: seen.append(image.size) or "words")
    for name in ("photo.png", "photo.jpg"):
        Image.new("RGB", (400, 300), "white").save(tmp_path / name)
        assert content_check._ocr_image(tmp_path / name) == "words"
    for width, height in seen:
        assert width * height <= 10_000


def test_a_date_pattern_runs_line_by_line_and_skips_lines_over_the_limit():
    """Decision 137 (L3): a staff Date Pattern is compiled and run one line
    at a time, and a line longer than ``DATE_LINE_MAX`` is skipped, so a
    typo that backtracks is bounded by one short line rather than the whole
    OCR text. Where it matched is still found in the whole text."""
    import tracker.content_check as content_check

    rules = item(period="TY2025", date_pattern=r"2025")
    text = "Statement\n" + "2025 " * 200 + "\nPeriod ending 12/31/2025\n"
    verdict = evaluate_rules(text, rules)
    assert verdict.ok
    assert content_check.date_pattern_at(r"2025", text) == text.index("12/31/2025") + 6

    only_long = "x" + "2025 " * 200
    assert content_check.date_pattern_at(r"2025", only_long) is None
    assert not evaluate_rules(only_long, rules).ok
    assert content_check.date_pattern_at(r"(unclosed", "anything") is None
    # A catastrophic pattern over a long run of what it backtracks on
    # finishes at once: the line holding the run is past the limit.
    evil = "a" * 5000 + "!"
    assert content_check.date_pattern_at(r"(a+)+b", evil) is None


def test_a_picture_past_pillows_guard_is_a_kept_too_large_verdict_not_a_retry(tmp_path, monkeypatch):
    """Decision 137 (B1, the owner's ruling on phase 1): a picture Pillow
    refuses as a decompression bomb is a size rule, like the file-size
    ceiling - "Too large to read ... A person looks at it." - kept as the
    verdict, parked for a person, and not read again until the file changes.
    Pillow's guard itself stays on."""
    from PIL import Image

    import tracker.content_check as content_check
    from tracker.content_check import ContentCache
    from tracker.router import route_file
    from tracker.validators import check_file

    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 1_000)          # 2,000 px is a bomb now
    photo = tmp_path / "photo.png"
    Image.new("L", (100, 100), "white").save(photo)
    rules = item(allowed_extensions=("pdf",), any_keywords=("w-2",), min_size_kb=0)

    assert check_file(photo, rules).code == reasons.TOO_LARGE.code   # tier 2: ours, not the client's
    reading = content_check.extract(photo)
    assert reading.text is None and not reading.transient
    assert reading.reason == "Too large to read (0 megapixels). A person looks at it."
    routed = route_file(photo, [rules])
    assert routed.identifier is None and routed.code == reasons.TOO_LARGE.code

    cache = ContentCache()
    first = check_content(photo, rules, cache)
    assert not first.ok and first.code == reasons.TOO_LARGE.code

    read: list = []

    def reading(path):
        read.append(path)
        return None

    monkeypatch.setattr(content_check, "_ocr_image", reading)
    assert check_content(photo, rules, cache) == first                 # not retried
    assert read == []
    Image.new("L", (100, 101), "white").save(photo)                     # the file changed
    check_content(photo, rules, cache)
    assert read == [photo]                                              # read again


def test_no_page_starts_past_the_safety_stop_and_the_stopped_reading_is_kept(tmp_path, monkeypatch):
    """Decision 137 (B1.2; the owner's Q-1 of 2026-09-24), as decision 169
    (ruling 5) words it: a reading gets a minute a page and ten minutes a
    document, rendering included. RapidOCR takes no timeout, so the stop
    is kept between pages: **no page starts past it** - the page before
    overran its minute, or the document its ten - and a page under way is
    bounded by the child's end at the document's stop (tests/test_ocr.py).
    A reading that reaches the stop is abandoned - "The reader stopped
    after N minutes on this file. A person reads it." - kept as the
    verdict, parked for a person, and not read again until the file
    changes. A slow reading under the stop is finished as before. The
    clock is moved, not waited for."""
    from pypdf import PdfWriter

    import tracker.content_check as content_check
    from tracker import ocr
    from tracker.content_check import ContentCache
    from tracker.router import route_file

    pypdfium2 = pytest.importorskip("pypdfium2")
    now = [1000.0]
    monkeypatch.setattr(content_check, "_clock", lambda: now[0])
    assert (content_check.READING_STOP_PAGE_SECONDS,
            content_check.READING_STOP_DOCUMENT_SECONDS) == (60.0, 600.0)
    writer = PdfWriter()
    for _ in range(10):
        writer.add_blank_page(width=612, height=792)
    scan = tmp_path / "scan.pdf"
    with scan.open("wb") as handle:
        writer.write(handle)
    real_render = pypdfium2.PdfPage.render
    monkeypatch.setattr(pypdfium2.PdfPage, "render",
                        lambda page, **kwargs: real_render(page, scale=0.1))

    # Slow but under the stop: every page finished and read.
    def pages_taking(*seconds):
        took = iter(seconds)
        read = []

        def reader(image, *, name=""):
            read.append(name)
            now[0] += next(took)
            return "Form W-2"
        return reader, read

    reader, read = pages_taking(*[59.0] * 10)
    monkeypatch.setattr(ocr, "read_page", reader)
    assert content_check.extract_by_ocr(scan).text.count("Form W-2") == 10
    assert len(read) == 10

    # The second page overruns its minute: the third never starts.
    reader, read = pages_taking(30.0, 61.0, 1.0)
    monkeypatch.setattr(ocr, "read_page", reader)
    stopped = content_check.extract_by_ocr(scan)
    said = "The reader stopped after 2 minutes on this file. A person reads it."
    assert stopped.text is None and stopped.reason == said and not stopped.transient
    assert len(read) == 2                                     # page 3 never reached the reader

    reader, read = pages_taking(30.0, 61.0, 1.0)
    monkeypatch.setattr(ocr, "read_page", reader)
    rules = item(allowed_extensions=("pdf",), any_keywords=("w-2",), min_size_kb=0)
    routed = route_file(scan, [rules])
    assert routed.identifier is None and routed.reason == said

    reader, read = pages_taking(30.0, 61.0, 1.0)
    monkeypatch.setattr(ocr, "read_page", reader)
    cache = ContentCache()
    verdict = check_content(scan, rules, cache)
    assert not verdict.ok and verdict.code == reasons.READING_STOPPED.code
    monkeypatch.setattr(ocr, "read_page", lambda *a, **k: pytest.fail("a kept verdict was read again"))
    assert check_content(scan, rules, cache) == verdict       # not retried

    # A document: the render counts against the budget. Two pages render
    # in half a minute each; the third's render alone runs past its minute,
    # and the stop comes before the reader is asked about it.
    renders = iter([30.0, 30.0, 61.0])

    def slow_render(page, **kwargs):
        now[0] += next(renders)
        return real_render(page, scale=0.1)

    asked = []
    monkeypatch.setattr(pypdfium2.PdfPage, "render", slow_render)
    monkeypatch.setattr(ocr, "read_page", lambda image, **k: asked.append(1) or "page")
    stopped = content_check.extract_by_ocr(scan)
    assert stopped.reason == "The reader stopped after 2 minutes on this file. A person reads it."
    assert asked == [1, 1]                                    # page 3 never reached the reader
    # A page's minute is never more than what is left of the document's ten.
    stop = content_check._SafetyStop()
    now[0] = stop.started + 590.0
    stop.page()
    assert stop.remaining() == 10.0
    now[0] = stop.started + 601.0
    with pytest.raises(content_check.ReadingStopped):
        stop.page()                                           # no page starts past the ten
    assert content_check._STOP is None                       # nothing outlives the reading


def test_only_a_staff_date_pattern_runs_line_by_line():
    """Decision 137's review (F7): the line-by-line bound is for a pattern a
    person typed. The one the tracker derives from the Period is its own and
    runs over the whole text as it always did, so a month and its year split
    by a line break still say the year."""
    import re

    from tracker.manifest import derived_date_pattern

    derived = derived_date_pattern("TY2025")
    split = "Statement period\nDecember\n2025 and more\n"
    assert derived and re.search(derived, split) is not None       # the premise
    firm = item(period="TY2025", date_pattern=derived, date_pattern_derived=True)
    assert evaluate_rules(split, firm).ok

    typed = item(period="TY2025", date_pattern=r"December\s2025", date_pattern_derived=False)
    assert not evaluate_rules(split, typed).ok                      # a person's: line by line
    assert evaluate_rules("December 2025\n", typed).ok


def test_a_pattern_the_record_calls_derived_runs_line_by_line_unless_it_is():
    """Decision 187 (the review's M2): "derived" is what the pattern is -
    the Period's own, recomputed - never what the record's flag claims. A
    forged pattern flagged derived runs line by line like any typed one."""
    split = "Statement period\nDecember\n2025 and more\n"
    forged = item(period="TY2025", date_pattern=r"December\s2025", date_pattern_derived=True)
    assert not evaluate_rules(split, forged).ok                     # line by line, whatever the flag
    assert evaluate_rules("December 2025\n", forged).ok


# ------------------------------------ a reading the pass can stop (150) ----
#
# Decision 150. The pass reads each document in a child process and waits
# for it at most the document's stop; the claims below run a real child.
# The suite reads in its own process everywhere else (tests/conftest.py),
# so each of these turns the child back on. The stops are patched down to
# seconds - never minutes - and the child is handed a stand-in reader from
# tests/child_readers.py, because a patch made here never reaches it.

#: The stop these claims patch in, in seconds. A child starts in about half
#: a second on the office machine; this leaves a slow build runner room to
#: reach the stage it is claimed to be stopped in.
STOP_IN_TESTS = 5.0
#: What a reading stopped at a stop under a minute says.
STOPPED = "The reader stopped after 1 minute on this file. A person reads it."
CRASHED = "The reader could not read this file (it stopped unexpectedly). A person reads it."
#: What pypdf (decision 153, 6.19.0) raises on 150's page-tree bomb, refusing it itself:
#: its class, which is all of it a reason says since decision 190.
BOMB_REFUSED = "(LimitReachedError)"


@pytest.fixture
def in_a_child(monkeypatch):
    """The reading in a child process, as the pass makes it, at the real stops."""
    import tracker.content_check as content_check

    monkeypatch.setattr(content_check, "READ_IN_A_CHILD", True)
    return monkeypatch


@pytest.fixture
def a_short_stop(in_a_child):
    """The child, stopped at :data:`STOP_IN_TESTS` for a file and a photo alike.
    Only the pass's own wait is shortened: the child's own stop (decision 137,
    B1.2) is still ten minutes, so what ends the reading is the child being
    ended."""
    import tracker.content_check as content_check

    in_a_child.setattr(content_check, "READING_STOP_DOCUMENT_SECONDS", STOP_IN_TESTS)
    in_a_child.setattr(content_check, "READING_STOP_PAGE_SECONDS", STOP_IN_TESTS)
    return in_a_child


def no_child_left():
    """No reading's child of this process is still running."""
    import multiprocessing

    return multiprocessing.active_children() == []


def running(pid: int) -> bool:
    """Whether the process ``pid`` is still running, asked of the system."""
    import os
    import sys
    from pathlib import Path

    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.OpenProcess.restype = wintypes.HANDLE
        handle = kernel32.OpenProcess(0x1000, False, pid)     # PROCESS_QUERY_LIMITED_INFORMATION
        if not handle:
            return False
        try:
            code = wintypes.DWORD()
            kernel32.GetExitCodeProcess(handle, ctypes.byref(code))
            return code.value == 259                          # STILL_ACTIVE
        finally:
            kernel32.CloseHandle(handle)
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    stat = Path(f"/proc/{pid}/stat")
    try:
        return stat.read_text().rsplit(")", 1)[1].split()[0] != "Z"   # a zombie has ended
    except (OSError, IndexError):
        return True


def eventually(check, within: float = 15.0) -> bool:
    import time

    deadline = time.monotonic() + within
    while time.monotonic() < deadline:
        if check():
            return True
        time.sleep(0.1)
    return check()


def test_a_text_layer_that_never_finishes_is_stopped_and_abandoned(tmp_path, a_short_stop):
    """The PDF text layer runs before any OCR, where decision 137's stop
    never reached: a few megabytes of drawing commands kept pdfplumber busy
    for hours. The reading is now ended at the document's stop, and it is
    the abandoned reading - the same sentence, kept by the cache and not
    read again until the file changes."""
    import time

    import tracker.content_check as content_check
    from tests import child_readers

    a_short_stop.setattr(content_check, "_CHILD_READER", child_readers.a_text_layer_that_never_finishes)
    pdf = text_pdf(tmp_path / "drawing.pdf", "Form W-2 Wage and Tax Statement 2025")
    rules = item(required_keywords=("W-2",))
    cache = ContentCache()

    started = time.monotonic()
    verdict = check_content(pdf, rules, cache)
    took = time.monotonic() - started

    assert child_readers.reached(pdf, "text-layer").is_file()      # stopped inside the text layer
    assert STOP_IN_TESTS <= took < STOP_IN_TESTS + 45
    assert verdict == content_check.ContentResult(ok=False, reason=STOPPED, extractable=False,
                                                     code=reasons.READING_STOPPED.code)
    assert no_child_left()
    a_short_stop.setattr(content_check, "_read_in_a_child",
                         lambda *a, **k: pytest.fail("a kept verdict was read again"))
    assert check_content(pdf, rules, cache) == verdict              # kept, not retried


def test_a_render_that_never_finishes_is_stopped_and_abandoned(tmp_path, a_short_stop):
    """One page's render ran inside the reading, and decision 137's stop
    only noticed an overrun once it returned. Ended at the stop now: the
    router parks the file on the abandoned reading's sentence, and the
    scanner keeps that verdict."""
    import tracker.content_check as content_check
    from tests import child_readers
    from tracker.router import route_file

    pytest.importorskip("pypdfium2")
    a_short_stop.setattr(content_check, "_CHILD_READER", child_readers.a_render_that_never_finishes)
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)                    # a scan: no text layer
    scan = tmp_path / "scan.pdf"
    with scan.open("wb") as handle:
        writer.write(handle)
    rules = item(allowed_extensions=("pdf",), required_keywords=("W-2",), min_size_kb=0)

    routed = route_file(scan, [rules])
    assert child_readers.reached(scan, "render").is_file()          # stopped inside the render
    assert routed.identifier is None and routed.reason == STOPPED
    assert no_child_left()

    cache = ContentCache()
    verdict = check_content(scan, rules, cache)
    assert verdict == content_check.ContentResult(ok=False, reason=STOPPED, extractable=False,
                                                     code=reasons.READING_STOPPED.code)
    a_short_stop.setattr(content_check, "_read_in_a_child",
                         lambda *a, **k: pytest.fail("a kept verdict was read again"))
    assert check_content(scan, rules, cache) == verdict
    assert no_child_left()


def test_a_reader_that_crashes_is_a_failed_reading_not_a_dead_pass(tmp_path, in_a_child):
    """A crash in pdfium or the reader used to end the pass's own process,
    and every return after the file went unsorted. The child ends instead:
    the file is a reading that failed - parked with its own sentence, not
    the abandoned one - and the pass goes on to the next document."""
    import tracker.content_check as content_check
    from tests import child_readers
    from tests.conftest import named_page, sort
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.layout import inbox_of

    in_a_child.setattr(content_check, "_CHILD_READER", child_readers.a_reader_that_dies_on_a_crash)
    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(tmp_path, [w2])
    page_pdf(inbox_of(engagement) / "a crash.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))
    page_pdf(inbox_of(engagement) / "the next one.pdf",
             named_page("Form W-2 Wage and Tax Statement 2025 Employer copy"))

    report = sort(engagement, today=dt.date(2026, 7, 1))

    [parked] = report.review
    assert parked.original_name == "a crash.pdf" and parked.reason == CRASHED
    assert parked.code == reasons.READING_CRASHED.code and reasons.READING_CRASHED.firm_side
    [filed] = report.filed
    assert filed.original_name == "the next one.pdf" and filed.identifier == "A01"
    assert no_child_left()


def test_a_reading_on_time_returns_what_it_returned_in_process(tmp_path, in_a_child):
    """On time, the child hands back exactly the reading the pass's own
    process makes - text, OCR or not, every reason, and tier 2's open test
    beside it (the review's ruling) - for the synthetic
    pile: the samples every client sends (text PDFs, a workbook, a CSV, a
    Word file, a Google shortcut, a bitmap), IRS forms, and a photo and a
    scan read by OCR where this machine has it. Only the seconds differ."""
    from dataclasses import replace
    from pathlib import Path

    from PIL import Image

    import tracker.content_check as content_check
    from tests.samples import build_samples

    pile = tmp_path / "pile"
    pile.mkdir()
    build_samples(pile)
    (pile / "statement.csv").write_text("Date,Amount\n2025-12-31,100\n", encoding="utf-8")
    picture = photo(pile / "w-2.png")
    Image.open(picture).convert("RGB").save(pile / "scanned w-2.pdf")
    irs = Path(__file__).resolve().parent / "irs"
    documents = sorted(pile.iterdir()) + [irs / name for name in ("fw2.pdf", "f1099div.pdf", "f1098.pdf")]

    assert len(documents) >= 12
    for document in documents:
        in_process = content_check.open_and_read(document)
        in_child = content_check.judge_bounded(document, content_check.Questions()).extraction
        assert in_child.opened is not None, document.name
        words_let_go = content_check.unjudged(in_process).extraction
        assert replace(in_child, seconds=0.0) == replace(words_let_go, seconds=0.0), document.name
    assert no_child_left()


def test_a_cached_verdict_starts_no_child(tmp_path, in_a_child):
    """A child is started on a cache miss only. A pass reads a drop once to
    route it, and the scan that follows finds the router's kept verdict for
    the working copy; a second pass starts nothing at all."""
    from pathlib import Path

    import tracker.content_check as content_check
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.layout import inbox_of
    from tracker.registry import discover_engagements
    from tracker.runner import REMINDERS_NEVER, run_registry

    started = []
    real = content_check._read_in_a_child

    def counted(path, questions):
        started.append(Path(path).name)
        return real(path, questions)

    in_a_child.setattr(content_check, "_read_in_a_child", counted)
    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(tmp_path, [w2])
    page_pdf(inbox_of(engagement) / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))

    first = run_registry(discover_engagements(tmp_path), today=dt.date(2026, 7, 1),
                         reminders=REMINDERS_NEVER)
    assert [run.filed for run in first.runs] == [1]
    assert started == ["w2.pdf"]                       # routed once; the scan found the verdict
    run_registry(discover_engagements(tmp_path), today=dt.date(2026, 7, 2), reminders=REMINDERS_NEVER)
    assert started == ["w2.pdf"]                       # a second pass starts nothing

    loose = page_pdf(tmp_path / "loose.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))
    rules = item(required_keywords=("W-2",))
    cache = ContentCache()
    assert check_content(loose, rules, cache).ok
    assert check_content(loose, rules, cache).ok
    assert started == ["w2.pdf", "loose.pdf"]          # one child for the miss, none for the hit
    assert no_child_left()


def test_no_child_is_left_running_after_a_stop(tmp_path, a_short_stop):
    """Ending the reading ends everything it started. A reader that had
    started a process of its own would leave it reading a client's page
    with nobody waiting for it if only the child were ended - so the whole
    tree goes, on Windows as elsewhere."""
    import tracker.content_check as content_check
    from tests import child_readers

    a_short_stop.setattr(content_check, "_CHILD_READER",
                         child_readers.a_reader_that_starts_a_reader_of_its_own)
    pdf = text_pdf(tmp_path / "page.pdf", "Form W-2 2025")

    reading = content_check.judge_bounded(pdf, content_check.Questions()).extraction

    assert reading.text is None and reading.reason == STOPPED and not reading.transient
    helper = child_readers.reached(pdf, "helper")
    assert helper.is_file()                                          # it had started its own
    pid = int(helper.read_text(encoding="utf-8"))
    assert eventually(lambda: not running(pid)), f"process {pid} outlived the stop"
    assert no_child_left()


#: A pass, as a process of its own, that reads one document in a child
#: with a stand-in reader that starts a helper (a process of the reader's
#: own) and never finishes. ``lifeline`` as the second argument takes the job object
#: away, so the lifeline alone is what is proved.
A_PASS_READING = """
import sys
from pathlib import Path
sys.path[:0] = [sys.argv[3]]
from tests import child_readers
from tracker import content_check, ocr
content_check._CHILD_READER = child_readers.a_reader_that_starts_a_reader_of_its_own
if sys.argv[2] == "lifeline":
    ocr._kill_on_close_job = lambda pid, **_limits: None
content_check.judge_bounded(Path(sys.argv[1]), content_check.Questions())
"""


@pytest.mark.parametrize("protection", ["job and lifeline", "lifeline"])
def test_a_reading_child_dies_with_its_pass(tmp_path, protection):
    """The designer's ruling on the build: the reading's child never
    outlives its pass. Task Scheduler's stop kills the pass outright - no
    code of the pass's runs - and the child must go with it, within a few
    seconds, not whenever its own reading ends. Here the pass is a real
    process, killed the way the scheduler kills it, mid-read.

    Two things do it. The lifeline: the child ends itself when the pass's
    end of a pipe closes, which the pass's death closes. And on Windows a
    job object that kills everything in it - the child and what it started
    - when the pass's handle on it closes. Each is proved: all of it, and
    the lifeline alone."""
    import subprocess
    import sys
    import time
    from pathlib import Path

    from tests import child_readers

    if protection == "job and lifeline" and sys.platform != "win32":
        pytest.skip("the job object is Windows'")
    repo = str(Path(__file__).resolve().parents[1])
    pdf = text_pdf(tmp_path / "page.pdf", "Form W-2 2025")
    the_pass = subprocess.Popen([sys.executable, "-c", A_PASS_READING, str(pdf),
                                 "lifeline" if protection == "lifeline" else "all", repo])
    waiting, helper = child_readers.reached(pdf, "waiting"), child_readers.reached(pdf, "helper")
    try:
        assert eventually(waiting.is_file, within=60), "the child never started reading"
        child = int(waiting.read_text(encoding="utf-8"))
        started = int(helper.read_text(encoding="utf-8"))
        assert running(child) and running(started)                 # mid-read

        the_pass.kill()                                            # the scheduler's stop: no code runs
        the_pass.wait(30)
        killed = time.monotonic()

        assert eventually(lambda: not running(child), within=10), "the child outlived its pass"
        assert time.monotonic() - killed < 10
        if protection == "job and lifeline":
            assert eventually(lambda: not running(started), within=10), "what it started outlived it"
    finally:
        if the_pass.poll() is None:
            the_pass.kill()
        for mark in (waiting, helper):
            if mark.is_file():
                _stop_if_running(int(mark.read_text(encoding="utf-8")))


def _stop_if_running(pid: int) -> None:
    """Leave nothing behind, whatever the test found (the lifeline alone
    ends the child, and on Windows not what it started)."""
    import os
    import signal
    import subprocess
    import sys

    if not running(pid):
        return
    if sys.platform == "win32":
        subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True, check=False)
    else:
        os.kill(pid, signal.SIGKILL)


def page_tree_bomb(path, depth: int = 40):
    """A PDF whose page tree is ``depth`` levels deep, each level naming its
    one child twice: a few kilobytes that claim 2**depth pages, and a page
    count that walks every one of them (the review of decision 150, its
    generator's logic; every byte is made here). Padded with a comment line
    so no row's minimum size refuses it before it is opened."""
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>"}
    for level in range(depth):
        number = 2 + level
        objects[number] = b"<< /Type /Pages /Kids [%d 0 R %d 0 R] /Count %d >>" % (
            number + 1, number + 1, 2 ** (depth - level))
    leaf = depth + 2
    objects[leaf] = b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] >>" % (depth + 1)
    out = bytearray(b"%PDF-1.4\n%" + b"x" * 200_000 + b"\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (leaf + 1)
    out += b"".join(b"%010d 00000 n \n" % offsets[number] for number in range(1, leaf + 1))
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (leaf + 1, xref)
    path.write_bytes(bytes(out))
    return path


def test_a_page_tree_bomb_cannot_stall_the_pass(tmp_path, a_short_stop):
    """The review of decision 150's blocker: tier 2's open test - pypdf
    counting a PDF's pages - ran in the pass's own process, and a page
    tree of a few kilobytes counted to a trillion. The open test now runs
    in the reading's child, beside the reading, and since decision 153
    pypdf (6.16.1 on) refuses the bomb itself, in seconds: the pass
    finishes well inside the stop, the bomb parks as an unreadable PDF
    with pypdf's limit named, the next document files, and the scanner's
    open test gives the same refusal, in seconds, each time it asks. The
    stop's own bound on the open test, and its kept verdict, are claimed
    without pypdf, by test_an_open_test_that_never_returns_is_stopped."""
    import time
    from pathlib import Path

    import tracker.content_check as content_check
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.filer import read_index
    from tracker.layout import inbox_of
    from tracker.registry import discover_engagements
    from tracker.runner import REMINDERS_NEVER, run_registry

    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(tmp_path, [w2])
    page_tree_bomb(inbox_of(engagement) / "bomb.pdf")
    page_pdf(inbox_of(engagement) / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))

    started = time.monotonic()
    report = run_registry(discover_engagements(tmp_path), today=dt.date(2026, 7, 1),
                          reminders=REMINDERS_NEVER)
    took = time.monotonic() - started

    assert took < 60                                     # the pass finished, in seconds
    rows = {row.original_name: row for row in read_index(engagement)}
    assert rows["bomb.pdf"].identifier == ""             # parked, refused by pypdf itself
    assert BOMB_REFUSED in rows["bomb.pdf"].reason
    assert rows["bomb.pdf"].reason != STOPPED            # not by the stop
    assert rows["w2.pdf"].identifier == "A01"            # and the next document filed
    assert [run.filed for run in report.runs] == [1]
    assert no_child_left()

    # The scanner's open test: the child's verdict, kept by the fingerprint.
    bomb = page_tree_bomb(tmp_path / "loose bomb.pdf")
    cache = ContentCache()
    started = time.monotonic()
    verdict = content_check.open_verdict(bomb, cache)
    assert time.monotonic() - started < STOP_IN_TESTS    # seconds, not the stop
    assert verdict.startswith("not a readable PDF") and BOMB_REFUSED in verdict
    # Not kept: pdfium cannot load the bomb either, so the reading beside
    # the refusal is OCR's transient failure, and a transient reading is
    # never kept. The next ask opens it again - in seconds, to the same
    # refusal.
    from tracker.validators import check_file
    started = time.monotonic()
    refused = check_file(bomb, w2, open_test=lambda path: content_check.open_verdict(Path(path), cache))
    assert time.monotonic() - started < STOP_IN_TESTS
    assert not refused.ok and refused.reason == verdict
    assert no_child_left()


def test_an_open_test_that_never_returns_is_stopped(tmp_path, a_short_stop):
    """Decision 153: pypdf now refuses 150's page-tree bomb before the stop
    comes, so the open test's own bound gets a claim no library upgrade
    can remove. A PDF whose open test never returns - pypdf's open blocked
    for ever inside the child, the reading untouched - is ended at the
    (patched) stop: the router parks it on the stop's sentence, the child
    and everything it started are gone, and the scanner's open test is the
    same kept verdict."""
    import time
    from pathlib import Path

    import tracker.content_check as content_check
    from tests import child_readers
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.router import route_file
    from tracker.validators import check_file

    a_short_stop.setattr(content_check, "_CHILD_READER", child_readers.a_pdf_open_that_never_finishes)
    pdf = page_pdf(tmp_path / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))
    rules = item(allowed_extensions=("pdf",), required_keywords=("W-2",), min_size_kb=0)

    started = time.monotonic()
    routed = route_file(pdf, [rules])
    took = time.monotonic() - started

    assert child_readers.reached(pdf, "pdf-open").is_file()       # stopped inside the open test
    assert not child_readers.reached(pdf, "text-layer").exists()
    assert STOP_IN_TESTS <= took < STOP_IN_TESTS + 45
    assert routed.identifier is None and routed.reason == STOPPED
    assert no_child_left()

    cache = ContentCache()
    assert content_check.open_verdict(pdf, cache) == STOPPED
    a_short_stop.setattr(content_check, "_read_in_a_child",
                         lambda *a, **k: pytest.fail("a kept open test was made again"))
    refused = check_file(pdf, rules, open_test=lambda path: content_check.open_verdict(Path(path), cache))
    assert not refused.ok and refused.reason == STOPPED
    assert no_child_left()


def _a_reader_only_the_pass_has(monkeypatch):
    """A reader the child can never import: its module is in the pass's
    process only, so the child ends while it is still being set up - the
    shape of a frozen build missing a module, or a machine that cannot
    start the reader - before it has said "started"."""
    import sys
    import types

    module = types.ModuleType("only_in_the_pass_150")

    def reader(path, *, ocr=True):          # never runs: the child cannot find it
        raise AssertionError("unreachable")

    reader.__module__ = module.__name__
    reader.__qualname__ = "reader"
    module.reader = reader
    monkeypatch.setitem(sys.modules, module.__name__, module)
    return reader


def test_a_reader_that_cannot_start_keeps_no_verdict(tmp_path, in_a_child):
    """The review's second fix. A reader that cannot start is the machine's
    fault, not the file's: a death before the child says "started" - or a
    child that could not be created at all - keeps nothing against the
    file. Nothing is kept, a drop is neither decided nor recorded - it
    rests in the year's folder with no row, for the next pass - and the
    pass warns once, however many files it met. Only a death after
    "started" is kept (the crash claim above)."""
    import multiprocessing.context

    import tracker.content_check as content_check
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.filer import read_index
    from tracker.layout import inbox_of, originals_of
    from tracker.registry import discover_engagements
    from tracker.runner import REMINDERS_NEVER, append_log, format_report, run_registry
    from tracker.validators import sha256_of

    unavailable = ("The reader could not start on this machine, so this file was not read. "
                   "It is read again on the next pass.")
    in_a_child.setattr(content_check, "_CHILD_READER", _a_reader_only_the_pass_has(in_a_child))
    rules = item(required_keywords=("W-2",))
    page = text_pdf(tmp_path / "page.pdf", "Form W-2 2025")

    # Dies before "started": transient, and nothing is kept.
    reading = content_check.judge_bounded(page, content_check.Questions()).extraction
    assert reading.text is None and reading.transient and reading.reason == unavailable
    cache = ContentCache()
    verdict = check_content(page, rules, cache)
    assert verdict.transient and verdict.reason == unavailable
    assert cache.get(page, rules_fingerprint(rules)) is None
    assert content_check.open_verdict(page, cache) == unavailable
    assert cache.get(page, content_check.OPEN_TEST_FINGERPRINT) is None

    # Could not even be created: the same.
    def no_process(_process):
        raise OSError("no more processes")

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(multiprocessing.context.SpawnProcess, "_Popen", staticmethod(no_process))
        assert content_check.judge_bounded(page, content_check.Questions()).extraction.reason == unavailable
    content_check.readers_that_could_not_start()          # the claims above were not a pass

    # A pass: two documents wait for the next pass - no row, nothing kept
    # about either - and the pass says it once.
    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(tmp_path / "root", [w2])
    for name in ("one.pdf", "two.pdf"):              # two documents, not one sent twice
        page_pdf(inbox_of(engagement) / name, named_page(f"Form W-2 Wage and Tax Statement 2025 {name}"))
    report = run_registry(discover_engagements(tmp_path / "root"), today=dt.date(2026, 7, 1),
                          reminders=REMINDERS_NEVER)

    assert read_index(engagement) == []                  # not decided, not recorded
    assert sorted(path.name for path in originals_of(engagement).iterdir()) == ["one.pdf", "two.pdf"]
    assert [run.review for run in report.runs] == [0]
    assert reasons.READER_UNAVAILABLE.firm_side
    [warning] = report.warnings
    assert "could not start on this machine for 2 file(s)" in warning
    assert format_report(report).count(warning) == 1
    log = append_log(tmp_path / "runs.log", report).read_text(encoding="utf-8")
    assert warning not in log and "reader-could-not-start=1" in log     # a code (decision 186)
    _memos, verdicts = store.cached_verdicts(store.connect(), engagement, version=CACHE_VERSION)
    for original in originals_of(engagement).iterdir():
        assert sha256_of(original) not in verdicts, original.name      # nothing kept
    assert no_child_left()


def test_a_drop_the_reader_could_not_start_on_is_routed_next_pass(tmp_path, in_a_child):
    """The re-review's blocker. A machine fault must leave nothing
    permanent, and a routing decision is as permanent as a kept verdict: a
    Needs Review row made the original a non-stray for ever, so a good W-2
    waited for a person after the machine was fixed. Pass 1's reader cannot
    start: the W-2 is moved out of the inbox but decided and recorded
    nowhere. Pass 2's reader works: the W-2 is a stray in the year's folder,
    routed where it lies, and filed."""
    import tracker.content_check as content_check
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.filer import read_index
    from tracker.layout import inbox_of, originals_of
    from tracker.registry import discover_engagements
    from tracker.runner import REMINDERS_NEVER, run_registry

    root = tmp_path / "root"
    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(root, [w2])
    page_pdf(inbox_of(engagement) / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))

    with pytest.MonkeyPatch.context() as broken:           # the machine is broken
        broken.setattr(content_check, "_CHILD_READER", _a_reader_only_the_pass_has(broken))
        first = run_registry(discover_engagements(root), today=dt.date(2026, 7, 1),
                             reminders=REMINDERS_NEVER)
    [warning] = first.warnings
    assert "could not start on this machine for 1 file(s)" in warning and "w2.pdf" in warning
    assert read_index(engagement) == []                    # no row: not decided, not recorded
    assert [path.name for path in originals_of(engagement).iterdir()] == ["w2.pdf"]

    second = run_registry(discover_engagements(root), today=dt.date(2026, 7, 2),
                          reminders=REMINDERS_NEVER)       # the machine is fixed
    assert second.warnings == []
    [row] = read_index(engagement)
    assert row.original_name == "w2.pdf" and row.identifier == "A01"
    assert [run.filed for run in second.runs] == [1]
    assert no_child_left()


def test_a_resent_file_the_reader_could_not_start_on_waits_for_the_next_pass(
        tmp_path, in_a_child, monkeypatch):
    """The final review's note. The re-send road reads again too: a W-2
    already on record whose working copy was deleted by hand, sent again,
    goes through ``_sort_one`` - the digest is known - and must be read
    before it is filed again. Pass 1's reader cannot start: no new row,
    nothing re-filed, no working copy made, one warning; the re-send rests
    in the year's folder. Pass 2's reader works: it is filed again. (Since
    decision 157 a deleted copy is made again from its original first; the
    re-file road is the one a copy takes while its original cannot be read,
    as here - the first W-2's original still syncing.)"""
    import tracker.content_check as content_check
    import tracker.filer as filer_module
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.filer import read_index
    from tracker.layout import inbox_of, locate, originals_of
    from tracker.registry import discover_engagements
    from tracker.runner import REMINDERS_NEVER, run_registry

    root = tmp_path / "root"
    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(root, [w2])
    page = named_page("Form W-2 Wage and Tax Statement 2025")
    page_pdf(inbox_of(engagement) / "w2.pdf", page)
    run_registry(discover_engagements(root), today=dt.date(2026, 7, 1), reminders=REMINDERS_NEVER)
    [filed] = read_index(engagement)
    prepared = engagement / filed.prepared_location
    prepared.unlink()                                      # the working copy, deleted by hand
    syncing = locate(engagement, filed.pbc_location)
    real = filer_module.is_cloud_placeholder
    monkeypatch.setattr(filer_module, "is_cloud_placeholder", lambda p: p == syncing or real(p))

    page_pdf(inbox_of(engagement) / "w2 again.pdf", page)  # the same document, sent again
    with pytest.MonkeyPatch.context() as broken:           # the machine is broken
        broken.setattr(content_check, "_CHILD_READER", _a_reader_only_the_pass_has(broken))
        first = run_registry(discover_engagements(root), today=dt.date(2026, 7, 2),
                             reminders=REMINDERS_NEVER)
    [warning] = first.warnings
    assert "could not start on this machine for 1 file(s)" in warning and "w2 again.pdf" in warning
    assert read_index(engagement) == [filed]               # no new row: not decided, not recorded
    assert not prepared.exists()                           # nothing re-filed
    assert sorted(path.name for path in originals_of(engagement).iterdir()) == ["w2 again.pdf", "w2.pdf"]

    second = run_registry(discover_engagements(root), today=dt.date(2026, 7, 3),
                          reminders=REMINDERS_NEVER)       # the machine is fixed
    assert second.warnings == []
    first_row, again = read_index(engagement)
    assert first_row == filed
    assert again.original_name == "w2 again.pdf" and again.identifier == "A01"
    assert "re-filed: the earlier copy" in again.reason     # the re-send road, not a new arrival
    assert [run.filed for run in second.runs] == [1]
    assert no_child_left()


def test_a_photo_the_open_test_cannot_finish_is_stopped(tmp_path, in_a_child):
    """The re-review's note 6: the open test's photo branch - Pillow and
    pillow-heif opening the picture, native code - runs in the reading's
    child like the PDF branch. A photo whose open never finishes is ended
    at a photo's stop (patched to seconds; a file's stays ten minutes, so
    the photo's is the one that ended it): the router parks it on the
    stop's sentence, and the scanner's open test is the same kept verdict."""
    import time
    from pathlib import Path

    from PIL import Image

    import tracker.content_check as content_check
    from tests import child_readers
    from tracker.router import route_file
    from tracker.validators import check_file

    in_a_child.setattr(content_check, "READING_STOP_PAGE_SECONDS", STOP_IN_TESTS)
    in_a_child.setattr(content_check, "_CHILD_READER", child_readers.a_photo_open_that_never_finishes)
    photo = tmp_path / "w2 photo.jpg"
    Image.new("RGB", (200, 100), "white").save(photo)
    rules = item(allowed_extensions=("pdf",), required_keywords=("W-2",), min_size_kb=0)

    started = time.monotonic()
    routed = route_file(photo, [rules])
    took = time.monotonic() - started

    assert child_readers.reached(photo, "photo-open").is_file()   # stopped inside the photo open
    assert STOP_IN_TESTS <= took < STOP_IN_TESTS + 45
    assert routed.identifier is None and routed.reason == STOPPED
    assert no_child_left()

    cache = ContentCache()
    assert content_check.open_verdict(photo, cache) == STOPPED
    in_a_child.setattr(content_check, "_read_in_a_child",
                       lambda *a, **k: pytest.fail("a kept open test was made again"))
    refused = check_file(photo, rules, open_test=lambda path: content_check.open_verdict(Path(path), cache))
    assert not refused.ok and refused.reason == STOPPED
    assert no_child_left()


# ------------------------------------------------ decision 146 -------------


@pytest.mark.parametrize("lines, carries", [
    (["Charles Schwab", "2025 Consolidated Form 1099", "Form 1099-DIV Dividends and Distributions",
      "Form 1099-INT Interest Income", "Form 1099-B Proceeds From Broker and Barter Exchange Transactions"],
     True),
    (["Form 1099-B Proceeds From Broker and Barter Exchange Transactions 2025", "1d Proceeds 24,318.55"], True),
    (["Vanguard Brokerage Services", "2025 Consolidated Form 1099 - Account 8812-4455",
      "Form 1099-INT   Interest Income", "Form 1099-DIV   Dividends and Distributions"], False),
    (["Harborline Savings Bank", "Consolidated Statement December 2025",
      "Form 1099-INT Interest Income 2025"], False),
    (["Your 2025 Consolidated Form 1099 is now available online.",
      "Log in to view your Form 1099-B and Form 1099-DIV."], False),
    (["2025 Individual Income Tax Organizer - Investment Income",
      "Form 1099-B - Proceeds From Broker and Barter Exchange Transactions",
      "Form 1099-DIV - Dividends and Distributions"], False),
], ids=["d146-consolidated", "d146-loose-1099b", "d146-int-div-only", "d146-bank",
        "d146-notice", "d146-organizer"])
def test_a_1099b_section_is_read_as_a_form_label_not_as_the_word_consolidated(lines, carries):
    """d146. The predicate reads the 1099-B the way a ``1099-b`` keyword is
    read - named in its own right in the title, or the page's own number -
    so a notice that mentions it and a checklist that lists it carry none,
    and "consolidated" decides nothing either way."""
    from tracker.content_check import carries_a_1099b_section

    assert carries_a_1099b_section("\n".join(lines)) is carries


def test_running_out_of_memory_is_never_a_retry_or_a_corrupt_file(tmp_path, monkeypatch):
    """SPEC-169 section 9: a MemoryError - the reading child past its
    memory limit - is not turned into a transient "OCR failed" (retried
    every pass) nor into "the file is corrupt": it goes up to the child,
    which says so and ends, and the pass handles it as a crash."""
    import pypdfium2
    from PIL import Image

    from tracker import content_check, ocr

    def out_of_memory(*_args, **_kwargs):
        raise MemoryError

    photo = tmp_path / "photo.png"
    Image.new("RGB", (40, 40), "white").save(photo)
    monkeypatch.setattr(ocr, "read_page", out_of_memory)
    with pytest.raises(MemoryError):
        content_check.extract_by_ocr(photo)

    scan = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with scan.open("wb") as fh:
        writer.write(fh)
    monkeypatch.setattr(pypdfium2.PdfPage, "render", out_of_memory)
    with pytest.raises(MemoryError):
        content_check.extract_by_ocr(scan)

    monkeypatch.setattr(content_check, "extract_text", out_of_memory)
    with pytest.raises(MemoryError):
        content_check._extract(scan, ocr=False)


# Decision 178: a reading is bounded in what it hands back, whatever it
# reads.


def _the_old_line_of(low: str, start: int, end: int) -> tuple[int, int]:
    """``_line_of`` as it was before decision 178, walking the whole text
    once for every kind of break: the answer the new one must give."""
    breaks = "\r\n\f\v"
    left = max(low.rfind(ch, 0, start) for ch in breaks) + 1
    right = min((i for i in (low.find(ch, end) for ch in breaks) if i >= 0), default=len(low))
    return left, right


def test_the_line_a_mention_is_on_is_found_as_before_without_walking_the_text():
    """Two thousand random texts - with every kind of break, with ``\r``
    line ends only, and with no break at all, short and past the first
    window the look back reaches - each asked about several mentions in a
    row, so the line remembered from one is tried on the next."""
    import random

    from tracker.content_check import _line_of

    rng = random.Random(163)
    alphabets = ("ab w-2 \r\n\f\v", "ab w-2 \r", "ab w-2 ")
    for n in range(2000):
        alphabet = alphabets[n % len(alphabets)]
        size = rng.randint(0, 40) if n % 4 else rng.randint(200, 1500)
        low = "".join(rng.choice(alphabet) for _ in range(size))
        for _ in range(5):
            start = rng.randint(0, len(low))
            end = rng.randint(start, min(len(low), start + 12))
            assert _line_of(low, start, end) == _the_old_line_of(low, start, end), (low, start, end)


def test_the_line_finder_is_quick_on_a_text_with_no_line_feed():
    """The review's S-5: the first quick look back was anchored on the
    nearest line feed, so a text with none before a mention - one line, or
    ``\r``-only line ends - was walked to its start once for every kind of
    break, for every mention. Three texts of about 2 MB, two thousand
    mentions each: every answer is the old one, and all of them together
    take under half a second (the old walk takes about two seconds on the
    office machine, and the look back anchored on the line feed as long on
    the first two texts; the new one about a hundredth of that)."""
    import time

    from tracker.content_check import _line_of

    word = "w-2 wages "
    texts = {
        "one line": word * 200_000,
        "carriage returns only": (word * 7 + "\r") * 28_000,
        "line feeds": (word * 7 + "\n") * 28_000,
    }
    for kind, low in texts.items():
        step = len(low) // 2000
        mentions = [(i * step, i * step + 3) for i in range(2000)]
        began = time.perf_counter()
        found = [_line_of(low, start, end) for start, end in mentions]
        took = time.perf_counter() - began
        assert took < 0.5, (kind, took)
        assert found == [_the_old_line_of(low, start, end) for start, end in mentions], kind


def test_a_workbook_reads_to_the_budget_and_says_it_was_cut(tmp_path, monkeypatch):
    """A 14.7 KB workbook whose one shared string sat in twenty cells read
    to two hundred million characters, all handed to the pass. A workbook
    now stops at the reading's budget, and the reading says it was cut, as a
    text file past its cap always did."""
    import tracker.content_check as content_check
    from tests.samples import sheet_xlsx

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", 500)
    long = sheet_xlsx(tmp_path / "ledger.xlsx", [["Fixed Asset Schedule", "x" * 40]] * 200)
    reading = content_check.extract(long, ocr=False)
    assert reading.cut and len(reading.text) == 500
    assert reading.text.startswith("Sheet") and "Fixed Asset Schedule" in reading.text
    short = sheet_xlsx(tmp_path / "short.xlsx", [["Fixed Asset Schedule", "2025"]])
    assert not content_check.extract(short, ocr=False).cut


def test_one_wide_workbook_row_stops_at_the_budget(tmp_path, monkeypatch):
    """A row can hold 16,384 cells of 32,767 characters each: checked only
    between rows, one wide row was read whole before the budget was asked.
    The reader asks after every cell, so it hands back no more than the
    budget and the cell that crossed it."""
    import tracker.content_check as content_check
    from tests.samples import sheet_xlsx

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", 500)
    wide = sheet_xlsx(tmp_path / "wide.xlsx", [["x" * 40] * 200])
    raw = content_check._extract_xlsx(wide)
    assert 500 < len(raw) <= 500 + 41
    reading = content_check.extract(wide, ocr=False)
    assert reading.cut and len(reading.text) == 500


def test_the_cut_mark_is_exact_at_the_budget(tmp_path, monkeypatch):
    """A reading exactly as long as the budget is whole and not cut; one
    character past it is cut. And a reading whose first lines land exactly
    on the budget, with more lines after, is cut, never shortened in
    silence."""
    import tracker.content_check as content_check
    from tests.samples import sheet_xlsx

    book = sheet_xlsx(tmp_path / "book.xlsx", [["a" * 30], ["b" * 30], ["c" * 30]])
    whole = content_check.extract(book, ocr=False).text
    head = whole[: whole.index("c")].rstrip("\n")          # the title and two rows

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", len(whole))
    exact = content_check.extract(book, ocr=False)
    assert not exact.cut and exact.text == whole

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", len(whole) - 1)
    past = content_check.extract(book, ocr=False)
    assert past.cut and past.text == whole[:-1]

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", len(head))
    early = content_check.extract(book, ocr=False)
    assert early.cut and early.text == head


def test_a_pdf_reads_to_the_budget_and_says_it_was_cut(tmp_path, monkeypatch):
    import tracker.content_check as content_check
    from tests.samples import text_pdf

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", 300)
    pdf = text_pdf(tmp_path / "long.pdf", [f"Line {n} of a long statement 2025" for n in range(40)])
    reading = content_check.extract(pdf, ocr=False)
    assert reading.cut and len(reading.text) == 300 and reading.text.startswith("Line 0")


def test_a_workbook_packed_with_bzip2_is_never_unpacked(tmp_path):
    """A workbook is a zip, and openpyxl unpacks whatever packing a part
    names: a 19 KB .xlsx whose sheet was bzip2 exhausted memory the way the
    zip did. Only stored and deflated parts are read; anything else fails
    the reading with the sentence that says why, before a byte is unpacked
    (``UnknownPacking``, whose message is the firm's own sentence and so is
    said whole, decisions 189 and 190)."""
    import zipfile

    import tracker.content_check as content_check
    from tests.samples import sheet_xlsx

    plain = sheet_xlsx(tmp_path / "plain.xlsx", [["Fixed Asset Schedule", "2025"]])
    packed = tmp_path / "packed.xlsx"
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(packed, "w") as target:
        for info in source.infolist():
            method = zipfile.ZIP_BZIP2 if info.filename.endswith("sheet1.xml") else info.compress_type
            target.writestr(info.filename, source.read(info), compress_type=method)
    assert "Fixed Asset Schedule" in content_check.extract(plain, ocr=False).text
    reading = content_check.extract(packed, ocr=False)
    assert reading.text is None and not reading.extractable
    assert content_check.UNKNOWN_PACKING in reading.reason


# ------------------------------------------ the judgment in the reader child ----
#
# Decision 189: the rules, the form scan and the name run in the reader child
# with the reading (content_check.judge), and what comes back is a Judgment
# that carries no word of the document. Every document here is made by the
# test that reads it, and every name on it is made up.

#: A word no rule, catalog or spelling holds: if it is anywhere in what
#: crosses the pipe, the document's words crossed.
MARKER = "Quendrilvax"


def _questions_about(*items, names=()):
    from tracker.content_check import Questions, row_question

    return Questions(rows=tuple(row_question(one) for one in items), names=tuple(names))


def test_no_document_word_crosses_the_pipe(tmp_path, in_a_child):
    """The reader child hands the pass a Judgment, and the Judgment is the
    firm's words only: the rows' keywords and reason sentences, form
    numbers, the firm's spellings. A document holding a word nobody else
    holds - on its title line, beside the name, on the dated line - is
    judged against rows that match and rows that fail, and a name question
    that confirms and one that does not; the word is nowhere in the answer
    as it crossed, pickled."""
    import pickle

    import tracker.content_check as content_check
    from tests.conftest import TEST_CLIENT
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.content_check import NameQuestion

    page = page_pdf(tmp_path / "w2.pdf", "\n".join([
        f"Form W-2 Wage and Tax Statement 2025 {MARKER}",
        f"Employee {TEST_CLIENT} {MARKER}",
        f"Employer {MARKER} Widgets 2025-12-31",
    ]))
    rows = (item(required_keywords=("W-2",), date_pattern=r"2025"),
            item(required_keywords=("1098",)),
            item(any_keywords=("wage and tax statement", "interest income")))
    names = (NameQuestion(own=(TEST_CLIENT,)), NameQuestion(own=("Pat Nobody",),
                                                            others=(("the other return", (TEST_CLIENT,)),)))
    questions = _questions_about(*rows, names=names)

    judged = content_check.judge_bounded(page, questions)

    assert isinstance(judged, content_check.Judgment)
    assert judged.extraction.text == "" and judged.has_words
    assert judged.row(rows[0]).verdict.ok and not judged.row(rows[1]).verdict.ok
    assert judged.name(names[0]).matched == TEST_CLIENT
    blob = pickle.dumps(judged)
    for shape in (MARKER, MARKER.lower(), MARKER.upper()):
        assert shape.encode("utf-8") not in blob
    assert MARKER.lower() not in repr(judged).lower()


def test_a_parser_error_crosses_the_pipe_as_its_class_only(tmp_path, in_a_child):
    """A reader that raises quotes what it choked on: openpyxl names the
    cell's value, the operating system names the path. Neither crosses the
    pipe (decision 189, M1): a workbook whose number cell holds a word
    nobody else holds, in a folder named with it, is judged in the real
    child, and the word and the path are in no pickled Judgment and in no
    reason the router parks the file with - the class alone is said."""
    import pickle
    import re
    import zipfile

    import tracker.content_check as content_check
    from tests.samples import sheet_xlsx
    from tracker.router import route_file

    folder = tmp_path / f"{MARKER} household"
    folder.mkdir()
    book = sheet_xlsx(folder / "statement.xlsx", [[1]])
    with zipfile.ZipFile(book) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    parts["xl/worksheets/sheet1.xml"] = re.sub(
        r'<c r="A1"[^>]*>.*?</c>', f'<c r="A1" t="n"><v>{MARKER}</v></c>', sheet).encode()
    with zipfile.ZipFile(book, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in parts.items():
            target.writestr(name, data)
    row = item(any_keywords=("W-2",), allowed_extensions=("xlsx",), min_size_kb=0)

    judged = content_check.judge_bounded(book, _questions_about(row))

    assert judged.extraction.error == "ValueError"
    blob = pickle.dumps(judged)
    for word in (MARKER, str(folder)):
        assert word.encode("utf-8") not in blob
    routing = route_file(book, [row], judgment=judged)
    assert "ValueError" in routing.reason
    assert MARKER not in routing.reason and str(folder) not in routing.reason


def test_no_client_text_is_matched_in_the_pass_own_process(tmp_path, in_a_child):
    """With the reader in its child, the pass's own process never matches a
    word: every function that reads a document's words - the rules, the
    form scan, the self-named forms, the 1099-B test, the judgment itself
    and the name check - is made to fail in this process, and a pass over
    a household still files the W-2 and the 1099-INT and scans both as
    received."""
    import tracker.content_check as content_check
    import tracker.names as names_module
    from tests.conftest import named_page
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.layout import inbox_of
    from tracker.manifest import Status, load_manifest
    from tracker.registry import discover_engagements
    from tracker.runner import REMINDERS_NEVER, run_registry

    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    interest = RequestItem(identifier="A02", document="1099-INT", period="TY2025",
                           allowed_extensions=("pdf",), min_size_kb=0,
                           any_keywords=("1099-INT", "interest income"))
    engagement = make_engagement(tmp_path, [w2, interest])
    page_pdf(inbox_of(engagement) / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))
    page_pdf(inbox_of(engagement) / "interest.pdf", named_page("Form 1099-INT Interest Income 2025"))

    def matched_in_the_pass(*_args, **_kwargs):
        raise AssertionError("a document's words were matched in the pass's own process")

    for name in ("evaluate_rules", "says", "own_forms", "dominant_forms", "self_named_forms",
                 "required_matched", "any_keyword_matched", "carries_a_1099b_section",
                 "judgment_of"):
        in_a_child.setattr(content_check, name, matched_in_the_pass)
    in_a_child.setattr(names_module, "check_name", matched_in_the_pass)
    in_a_child.setattr(names_module, "read_names", matched_in_the_pass)

    report = run_registry(discover_engagements(tmp_path), today=dt.date(2026, 7, 1),
                          reminders=REMINDERS_NEVER)

    [run] = report.runs
    assert run.filed == 2 and not run.error, run
    status = {row.identifier: row.status for row in load_manifest(engagement)}
    assert status == {"A01": Status.RECEIVED, "A02": Status.RECEIVED}
    assert no_child_left()


def test_a_catastrophic_date_pattern_costs_one_stop_and_the_next_pass_takes_the_kept_verdict(
        tmp_path, a_short_stop):
    """A typed Date Pattern that backtracks - ``(\\d+)+x`` over forty digits
    runs about a day - used to run in the pass's own process, where nothing
    could stop it. It runs in the child now, under the document's stop:
    the first pass parks the file with the stop's sentence and keeps that
    verdict against the file's bytes and the row's rules, and the next pass
    starts no judgment for it at all."""
    import time
    from pathlib import Path

    import tracker.content_check as content_check
    from tests.conftest import named_page, sort
    from tracker.filer import read_index
    from tracker.layout import inbox_of

    started = []
    real = content_check._read_in_a_child

    def counted(path, questions):
        started.append(Path(path).name)
        return real(path, questions)

    a_short_stop.setattr(content_check, "_read_in_a_child", counted)
    # Decision 187 now refuses this pattern where it is typed and where the
    # record is read. The stop is the bound that does not depend on reading
    # a pattern (187's residual), so this test lets one through to prove it.
    import tracker.manifest as manifest
    import tracker.records as records
    a_short_stop.setattr(manifest, "date_pattern_problem", lambda pattern: "")
    a_short_stop.setattr(records, "date_pattern_problem", lambda pattern: "")
    row = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("txt",),
                      min_size_kb=0, required_keywords=("W-2",), date_pattern=r"(\d+)+x")
    engagement = make_engagement(tmp_path, [row])
    body = named_page("Form W-2 Wage and Tax Statement 2025\n" + "1" * 40) + "\n"
    (inbox_of(engagement) / "w2.txt").write_text(body, encoding="utf-8")

    began = time.monotonic()
    first = sort(engagement, today=dt.date(2026, 7, 1))
    took = time.monotonic() - began

    [parked] = first.review
    assert parked.original_name == "w2.txt" and parked.reason == STOPPED
    assert STOP_IN_TESTS <= took < STOP_IN_TESTS + 45
    assert started == ["w2.txt"]
    assert no_child_left()

    second = sort(engagement, today=dt.date(2026, 7, 2))
    assert started == ["w2.txt"]                        # the next pass judged nothing
    assert second.review == [] and [e.reason for e in read_index(engagement)] == [STOPPED]

    # The kept verdict is the file's, under the row's rules: the same bytes
    # anywhere - a working copy a person makes of it - are not judged again.
    loose = tmp_path / "copy of w2.txt"
    loose.write_text(body, encoding="utf-8")
    a_short_stop.setattr(content_check, "_read_in_a_child",
                         lambda *a, **k: pytest.fail("a kept verdict was judged again"))
    kept = check_content(loose, row, ContentCache(engagement))
    assert kept == content_check.ContentResult(ok=False, reason=STOPPED, extractable=False,
                                               code=content_check.reasons.READING_STOPPED.code)
    # The record now carries a pattern decision 187 refuses on read. The
    # read-back check after every test is about honest records, so this
    # test takes its forged one away.
    from tracker import ledger
    ledger.path_for(engagement).unlink()


def test_a_one_line_workbook_parks_within_the_stop_and_never_stalls_the_pass(tmp_path, a_short_stop):
    """The council's proof line. A workbook of one row - about twenty
    million characters, each cell a form's own name, which the form scan
    must weigh one by one - is judged in the child, and the stop ends it:
    it parks on the stop's sentence, the W-2 dropped beside it is filed in
    the same pass, and the pass takes the stop and not the scan's time."""
    import time

    from tests.conftest import named_page, sort
    from tests.samples import sheet_xlsx
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.content_check import READING_CHAR_BUDGET
    from tracker.layout import inbox_of

    w2 = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                     min_size_kb=0, required_keywords=("W-2",))
    ledger = RequestItem(identifier="A02", document="Interest ledger", period="TY2025",
                         allowed_extensions=("xlsx",), min_size_kb=0,
                         required_keywords=("1099-INT",), any_keywords=("interest income",))
    engagement = make_engagement(tmp_path, [w2, ledger])
    cell = "Form 1099-INT 2025 " * 1500
    cells = READING_CHAR_BUDGET // len(cell) - 1              # one row, just inside the budget
    sheet_xlsx(inbox_of(engagement) / "a ledger.xlsx", [[cell] * cells])
    page_pdf(inbox_of(engagement) / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))

    began = time.monotonic()
    report = sort(engagement, today=dt.date(2026, 7, 1))
    took = time.monotonic() - began

    [parked] = report.review
    assert parked.original_name == "a ledger.xlsx" and parked.reason == STOPPED
    [filed] = report.filed
    assert filed.original_name == "w2.pdf" and filed.identifier == "A01"
    assert took < STOP_IN_TESTS + 45
    assert no_child_left()


def test_a_transient_judgment_is_neither_kept_nor_recorded(tmp_path, in_a_child):
    """SPEC-156's seam: the one place a judgment says the machine decided
    it is ``Judgment.extraction.transient``, set where the reading ran. A
    judgment whose reader could not start carries it and no answers; the
    pass records nothing for the drop and keeps no verdict for its bytes,
    and the scanner keeps neither its open test nor its row verdict."""
    import tracker.content_check as content_check
    from tests.conftest import named_page, sort
    from tests.test_scanner import text_pdf as page_pdf
    from tracker.filer import read_index
    from tracker.layout import inbox_of

    in_a_child.setattr(content_check, "_CHILD_READER", _a_reader_only_the_pass_has(in_a_child))
    row = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                      min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(tmp_path, [row])
    drop = page_pdf(inbox_of(engagement) / "w2.pdf", named_page("Form W-2 Wage and Tax Statement 2025"))
    content = drop.read_bytes()

    judged = content_check.judge_bounded(drop, _questions_about(row))
    assert judged.extraction.transient and judged.extraction.text is None
    assert judged.rows == {} and judged.names == {} and not judged.has_words

    report = sort(engagement, today=dt.date(2026, 7, 1))
    assert report.filed == [] and report.review == [] and read_index(engagement) == []
    _memos, verdicts = cache_rows(engagement)
    assert verdicts == {}

    loose = tmp_path / "loose.pdf"
    loose.write_bytes(content)
    cache = ContentCache()
    assert content_check.open_verdict(loose, cache, row) == judged.extraction.reason
    assert check_content(loose, row, cache).transient
    assert cache.get(loose, content_check.OPEN_TEST_FINGERPRINT) is None
    assert cache.get(loose, rules_fingerprint(row)) is None
    content_check.readers_that_could_not_start()


def test_a_file_in_flight_when_the_child_is_retired_or_killed_is_read_again_next_pass_and_never_marked_read(
        tmp_path, in_a_child):
    """UX motion 3's claim, through the judgment. A child that dies before
    "started" leaves the drop undecided and unrecorded, and the next pass
    judges it and files it; one that dies after "started" parks the file
    with its own sentence and never a verdict that says it was read; and a
    child retired after every document is retired between files, so each
    document is judged whole by one child and files."""
    import tracker.content_check as content_check
    from tests import child_readers
    from tests.conftest import named_page, sort
    from tests.test_scanner import text_pdf as page_pdf
    from tracker import ocr
    from tracker.filer import read_index
    from tracker.layout import inbox_of

    row = RequestItem(identifier="A01", document="W-2", period="TY2025", allowed_extensions=("pdf",),
                      min_size_kb=0, required_keywords=("W-2",))
    engagement = make_engagement(tmp_path, [row])
    inbox = inbox_of(engagement)

    # Killed before "started": nothing is decided, and the next pass files it.
    page_pdf(inbox / "first.pdf", named_page("Form W-2 Wage and Tax Statement 2025 first"))
    with pytest.MonkeyPatch.context() as broken:
        broken.setattr(content_check, "_CHILD_READER", _a_reader_only_the_pass_has(broken))
        waiting = sort(engagement, today=dt.date(2026, 7, 1))
    assert waiting.filed == [] and read_index(engagement) == []
    content_check.readers_that_could_not_start()
    again = sort(engagement, today=dt.date(2026, 7, 2))
    assert [e.original_name for e in again.filed] == ["first.pdf"]

    # Killed after "started": parked with its sentence, and no verdict says it was read.
    in_a_child.setattr(content_check, "_CHILD_READER", child_readers.a_reader_that_dies_on_a_crash)
    page_pdf(inbox / "a crash.pdf", named_page("Form W-2 Wage and Tax Statement 2025 crash"))
    crashed = sort(engagement, today=dt.date(2026, 7, 3))
    [parked] = crashed.review
    assert parked.original_name == "a crash.pdf" and parked.reason == CRASHED
    _memos, verdicts = cache_rows(engagement)
    assert all(v.get("reason") == CRASHED or v.get("ok")
               for by_rule in verdicts.values() for v in by_rule.values())
    assert not any(v.get("ok") for by_rule in verdicts.values() for v in by_rule.values()
                   if v.get("reason") == CRASHED)

    # Retired after every document: between files, never in one.
    in_a_child.setattr(content_check, "_CHILD_READER", content_check.open_and_read)
    in_a_child.setattr(ocr, "DOCUMENTS_PER_CHILD", 1)
    for n in range(3):
        page_pdf(inbox / f"w2 {n}.pdf", named_page(f"Form W-2 Wage and Tax Statement 2025 copy {n}"))
    with ocr.reading_session():
        retired = sort(engagement, today=dt.date(2026, 7, 4))
    assert sorted(e.original_name for e in retired.filed) == ["w2 0.pdf", "w2 1.pdf", "w2 2.pdf"]
    assert no_child_left()


def _the_pass_before_189(text: str, row: RequestItem):
    """What the pass computed in its own process before decision 189, call
    for call: the router's ordinary reading, its ``_required_matched`` and
    ``any_keyword_matched`` - each reading the form scan for itself."""
    from tracker.content_check import any_keyword_matched, own_forms, says, self_named_forms

    own = own_forms(self_named_forms(text)) if text else None
    required = bool(row.required_keywords) and all(says(text, k) for k in row.required_keywords)
    return evaluate_rules(text, row, own), required, any_keyword_matched(text, row)


def test_every_verdict_is_the_same_judged_in_the_child_as_in_the_pass(tmp_path, in_a_child):
    """The proof that CACHE_VERSION need not move. Every IRS form in the
    suite's corpus and the suite's own sample pages, judged against every
    row of every catalog: the judgment's answers - each row's verdict,
    reason and evidence, both leads, the page's own forms, its self-named
    forms, the 1099-B test and the name - are the answers the pass reached
    in its own process before, value for value. And a judgment made in the
    real child is the judgment made here.

    The forms are read to their first two pages - a title, a foot and a
    page of instructions each - so the claim costs the suite a minute and
    not three; the judgment of a longer text is the same code."""
    from pathlib import Path

    import tracker.content_check as content_check
    from tests import samples
    from tracker.content_check import (
        NameQuestion,
        carries_a_1099b_section,
        extract,
        judgment_of,
        own_forms,
        row_question,
        self_named_forms,
    )
    from tracker.manifest import validated
    from tracker.names import check_name
    from tracker.templates import FORM_TYPES, template_items

    rows = list(dict.fromkeys(
        one for form in FORM_TYPES for one in validated(template_items(form["id"], year=2025))
        if has_content_rules(one)
    ))
    distinct = list(dict.fromkeys(row_question(one) for one in rows))
    by_question = {row_question(one): one for one in rows}
    names = (NameQuestion(own=("Test Client",), others=(("the other return", ("Pat Sample",)),)),
             NameQuestion(own=("Pat Sample",), others=(("the first return", ("Test Client",)),)))
    questions = content_check.Questions(rows=tuple(distinct), names=names)

    irs = Path(__file__).parent / "irs"
    with pytest.MonkeyPatch.context() as shorter:
        shorter.setattr(content_check, "MAX_PAGES", 2)
        readings = {pdf.name: extract(pdf, ocr=False) for pdf in sorted(irs.glob("*.pdf"))}
    pages = {
        "w2": samples.w2_lines("Test Client", "Sample Widgets LLC", 2025),
        "1099-int": samples.lines_1099_int("Sample Bank", "Pat Sample", 2025),
        "prior return": samples.prior_return_lines("Test Client", 2024),
        "1098": samples.form_1098_lines("Sample Lending", "Test Client", 2025),
        "two forms": samples.scanned_w2_lines(2025) + samples.scanned_1098_lines(2025),
        "brokerage": samples.brokerage_cover_lines(2025),
        "scanned 1099-int": samples.scanned_1099_int_lines(2025),
    }
    for name, lines in pages.items():
        readings[name] = content_check.Extraction("\n".join(lines))
    assert len(readings) > 100

    for name, reading in readings.items():
        judged = judgment_of(reading, questions)
        words = "" if reading.needs_ocr else (reading.text or "")
        if not words:
            assert not judged.has_words, name
            continue
        own = own_forms(self_named_forms(words))
        assert judged.own == (frozenset(own) if own is not None else None), name
        assert judged.self_named == self_named_forms(words), name
        assert judged.carries_1099b == carries_a_1099b_section(words), name
        for question in distinct:
            before = _the_pass_before_189(words, by_question[question])
            now = judged.rows[question]
            assert (now.verdict, now.required_matched, now.any_matched) == before, (name, question)
        for question in names:
            assert judged.name(question) == check_name(words, question.own, dict(question.others)), name

    # And across the pipe: the child's judgment is this process's.
    for document in ("fw2.pdf", "f1099int.pdf", "f1098.pdf"):
        in_the_child = content_check.judge_bounded(irs / document, questions)
        here = judgment_of(content_check.open_and_read(irs / document), questions)
        assert in_the_child == replace_seconds(here, in_the_child), document
    assert no_child_left()


def replace_seconds(judged, like):
    """``judged`` with ``like``'s reading time, the one field two readings
    of the same bytes never share."""
    from dataclasses import replace

    return replace(judged, extraction=replace(judged.extraction, seconds=like.extraction.seconds))


def test_the_form_scan_runs_once_per_judgment(monkeypatch):
    """The first of the two costs. The form scan (``dominant_forms``, a whole
    walk of the text) ran once for every required keyword of every row the
    router asked, and again for every row's any-keywords and the 1099-B
    test: twenty rows at the reading budget was a quarter of an hour for
    one file. It runs once a judgment now, and ``self_named_forms`` once,
    however many rows are asked."""
    import tracker.content_check as content_check
    from tracker.content_check import Extraction, judgment_of

    counts = {"dominant_forms": 0, "self_named_forms": 0}
    for name in counts:
        real = getattr(content_check, name)

        def counted(text, _real=real, _name=name):
            counts[_name] += 1
            return _real(text)

        monkeypatch.setattr(content_check, name, counted)

    rows = [item(required_keywords=("W-2", f"box {n}"), any_keywords=("wages", "1099-B"),
                 date_pattern=r"2025") for n in range(20)]
    page = Extraction("Form W-2 Wage and Tax Statement 2025\nbox 1 wages\nbox 12 box 13 box 14\n")

    judged = judgment_of(page, _questions_about(*rows))

    assert counts == {"dominant_forms": 1, "self_named_forms": 1}
    assert len(judged.rows) == 20 and judged.row(rows[1]).required_matched


def test_an_ocr_reading_is_cut_to_the_budget_and_says_so(tmp_path, monkeypatch):
    """Decision 178 cut every reading to the budget but OCR's, which was
    handed back whole. OCR's words are cut there too now, and the reading
    says it was cut, as a text layer's does; one inside the budget is
    whole."""
    import tracker.content_check as content_check

    monkeypatch.setattr(content_check, "READING_CHAR_BUDGET", 60)
    photo = tmp_path / "photo.jpg"
    photo.write_bytes(b"not read: the reader is a stand-in")
    said = {"text": "Form W-2 Wage and Tax Statement 2025\n" + "box 1 wages 1000.00\n" * 10}
    monkeypatch.setattr(content_check, "_ocr_image", lambda _path: said["text"])

    long = content_check.extract_by_ocr(photo)
    assert long.from_ocr and long.cut and long.text == said["text"][:60]

    said["text"] = "Form W-2 Wage and Tax Statement 2025"
    short = content_check.extract_by_ocr(photo)
    assert short.from_ocr and not short.cut and short.text == said["text"]


# --------------------------------------- the cache's key (P216) ----


def test_a_files_cache_key_is_its_folders_held_answer_and_its_own_name(tmp_path, monkeypatch):
    """P216: inside a household's hold the folder is resolved once and every
    file in it is keyed by that answer and its own name - not a whole
    resolve a file."""
    from tracker import settings

    folder = tmp_path / "Inbox"
    folder.mkdir()
    files = [folder / name for name in ("One.pdf", "two.PDF", "three.jpg")]
    for file in files:
        file.write_bytes(b"fabricated")
    resolves = []
    real = type(folder).resolve

    def counted(self, *args, **kwargs):
        resolves.append(self)
        return real(self, *args, **kwargs)

    monkeypatch.setattr(type(folder), "resolve", counted)
    with settings.one_household():
        keys = [ContentCache._key(file) for file in files]
    assert keys == [str(real(folder) / file.name).lower() for file in files]
    assert resolves == [folder]


def test_a_file_that_is_a_link_is_keyed_by_where_it_leads(tmp_path):
    target = tmp_path / "elsewhere" / "Real Name.pdf"
    target.parent.mkdir()
    target.write_bytes(b"fabricated")
    link = tmp_path / "Inbox" / "Link.pdf"
    link.parent.mkdir()
    try:
        link.symlink_to(target)
    except OSError:
        pytest.skip("this machine makes no symbolic links")
    assert ContentCache._key(link) == str(target.resolve()).lower()


def test_the_key_is_the_one_a_whole_resolve_gives(tmp_path):
    """The same key string as before (P216): a file in a folder reached
    through a link, a file not there yet, a name in another case, and the
    names that are not plain names."""
    from tracker import settings

    real = tmp_path / "Real Folder"
    real.mkdir()
    (real / "Statement.PDF").write_bytes(b"fabricated")
    through = tmp_path / "Through"
    try:
        through.symlink_to(real, target_is_directory=True)
    except OSError:
        through = real
    cases = [real / "Statement.PDF", through / "Statement.PDF", real / "not yet.pdf",
             tmp_path / "Real Folder" / "..", tmp_path / "Real Folder" / "."]
    for case in cases:
        assert ContentCache._key(case) == str(case.resolve()).lower(), case
    with settings.one_household():
        for case in cases:
            assert ContentCache._key(case) == str(case.resolve()).lower(), case
