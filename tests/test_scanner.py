"""Tests for tracker/scanner.py — full-loop orchestration on local folders."""

import datetime as dt
import json
import os

import pytest
from openpyxl import load_workbook

from tests.samples import col
from tracker import reasons
from tracker.locking import LOCK_FILENAME, STALE_LOCK_SECONDS
from tracker.manifest import (
    COL_EXPECTED_COUNT,
    COL_MANUAL_OVERRIDE,
    ENGAGEMENT_SHEET_NAME,
    SHEET_NAME,
    SUMMARY_SEPARATOR,
    Override,
    RequestItem,
    Status,
    create_template,
    load_manifest,
)
from tracker.scaffold import (
    MANIFEST_FILENAME,
    PREPARED_DIR_NAME,
    REVIEW_DIR_NAME,
    scaffold_engagement,
)
from tracker.scanner import (
    CACHE_FILENAME,
    DUPLICATES_NOTE,
    OVERRIDE_NOTE,
    PARTIAL_NOTE,
    REGRESSION_COUNT_RAISED,
    REGRESSION_FILES_CHANGED,
    REGRESSION_NOTE,
    SYNCING_NOTE,
    ScanLockedError,
    scan_engagement,
)

DAY1 = dt.date(2026, 7, 1)
DAY2 = dt.date(2026, 7, 9)

ITEMS = [
    RequestItem(
        identifier="A01", document="Bank Statement", allowed_extensions=("pdf",),
        min_size_kb=0, required_keywords=("Chase",),
    ),
    RequestItem(
        identifier="A02", document="Monthly Statements", expected_count=3,
        allowed_extensions=("csv",), min_size_kb=0,
    ),
    RequestItem(
        identifier="B01", document="Payroll Register", allowed_extensions=("xlsx",),
        min_size_kb=10,
    ),
]


def text_pdf(path, text: str, pages: int = 1):
    """A valid PDF whose every page says ``text`` (nothing, for a scan)."""
    def _pdf_string(line: str) -> str:
        return "(" + line.replace(chr(92), chr(92) * 2).replace("(", chr(92) + "(").replace(")", chr(92) + ")") + ")"
    lines = text.split(chr(10)) if text else [""]
    body = " ".join(f"{_pdf_string(line)} Tj T*" for line in lines)
    content = f"BT /F1 12 Tf 14 TL 72 720 Td {body} ET".encode("latin-1", "replace")
    stream, font = 3 + pages, 4 + pages          # objects 3..2+pages are the pages
    kids = b" ".join(b"%d 0 R" % (3 + i) for i in range(pages))
    bodies = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, pages),
        stream: b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        font: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for i in range(pages):
        bodies[3 + i] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents %d 0 R /Resources << /Font << /F1 %d 0 R >> >> >>" % (stream, font)
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for n in sorted(bodies):
        offsets[n] = len(out)
        out += b"%d 0 obj\n" % n + bodies[n] + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(bodies) + 1)
    for n in sorted(bodies):
        out += b"%010d 00000 n \n" % offsets[n]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(bodies) + 1, xref_at,
    )
    path.write_bytes(bytes(out))
    return path


@pytest.fixture
def engagement(tmp_path):
    eng = tmp_path / "TY2025 1040"
    eng.mkdir()
    create_template(eng / MANIFEST_FILENAME, ITEMS)
    scaffold_engagement(eng)
    return eng


def folder(engagement, prefix):
    """The Prepared/ working folder for one request row."""
    prepared = engagement / PREPARED_DIR_NAME
    return next(p for p in prepared.iterdir() if p.is_dir() and p.name.startswith(prefix))


def statuses(engagement):
    return {i.identifier: i for i in load_manifest(engagement / MANIFEST_FILENAME)}


# ------------------------------------------------------------ resolution ----


def test_full_scan_statuses_and_writeback(engagement):
    text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec 2025")
    a02 = folder(engagement, "A02")
    (a02 / "jan.csv").write_text("jan data", encoding="utf-8")
    (a02 / "feb.csv").write_text("feb data", encoding="utf-8")
    # B01 left empty

    report = scan_engagement(engagement, today=DAY1)
    assert report.written and not report.deferred

    rows = statuses(engagement)
    assert rows["A01"].status == Status.RECEIVED
    assert rows["A01"].received_date == DAY1
    assert rows["A01"].file_count == 1
    assert rows["A02"].status == Status.PARTIAL
    assert PARTIAL_NOTE.format(count=2, expected=3) in rows["A02"].validation_notes
    assert rows["B01"].status == Status.MISSING
    assert rows["B01"].file_count == 0


def test_failed_validation_with_reasons(engagement):
    # Wrong content: tier 2 passes, tier 3 keyword check fails.
    text_pdf(folder(engagement, "A01") / "wrong.pdf", "Wells Fargo Statement Dec 2025")
    scan_engagement(engagement, today=DAY1)
    row = statuses(engagement)["A01"]
    assert row.status == Status.FAILED
    assert reasons.WRONG_DOCUMENT.matches(row.validation_notes) and "'Chase'" in row.validation_notes


def test_received_date_sticky_across_scans(engagement):
    text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec 2025")
    scan_engagement(engagement, today=DAY1)
    scan_engagement(engagement, today=DAY2)
    assert statuses(engagement)["A01"].received_date == DAY1  # first pass wins


def test_auto_revert_preserves_date_and_notes(engagement):
    pdf = text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec 2025")
    scan_engagement(engagement, today=DAY1)

    pdf.unlink()                                   # client deleted their file
    scan_engagement(engagement, today=DAY2)
    row = statuses(engagement)["A01"]
    assert row.status == Status.MISSING            # auto-revert
    assert row.received_date == DAY1               # original date preserved
    assert REGRESSION_NOTE.format(status=Status.RECEIVED, date="2026-07-01", why="").rstrip("; ") in row.validation_notes


def test_manual_override_status_untouched(tmp_path):
    items = [
        RequestItem(
            identifier="A01", document="Bank Statement",
            allowed_extensions=("pdf",), min_size_kb=0,
            required_keywords=("Chase",), manual_override=Override.ACCEPTED,
            status=Status.RECEIVED, received_date=DAY1,
        ),
    ]
    eng = tmp_path / "Eng"
    eng.mkdir()
    create_template(eng / MANIFEST_FILENAME, items)
    # seed scanner columns the accountant "kept" via override
    from tracker.manifest import StatusUpdate, write_statuses
    write_statuses(
        eng / MANIFEST_FILENAME,
        {"A01": StatusUpdate(status=Status.RECEIVED, file_count=1, received_date=DAY1)},
    )
    scaffold_engagement(eng)
    text_pdf(folder(eng, "A01") / "wrong.pdf", "Wells Fargo Statement December")

    scan_engagement(eng, today=DAY2)
    row = statuses(eng)["A01"]
    assert row.status == Status.RECEIVED           # override kept it
    assert row.received_date == DAY1
    assert row.validation_notes.startswith(OVERRIDE_NOTE.format(override=Override.ACCEPTED))
    assert reasons.WRONG_DOCUMENT.matches(row.validation_notes) and "'Chase'" in row.validation_notes  # facts still recorded


def test_duplicates_do_not_inflate_count(engagement):
    a02 = folder(engagement, "A02")
    (a02 / "jan.csv").write_text("identical bytes", encoding="utf-8")
    (a02 / "jan (1).csv").write_text("identical bytes", encoding="utf-8")
    (a02 / "feb.csv").write_text("different bytes", encoding="utf-8")

    scan_engagement(engagement, today=DAY1)
    row = statuses(engagement)["A02"]
    assert row.file_count == 2                     # 3 files, 2 distinct
    assert row.status == Status.PARTIAL
    assert DUPLICATES_NOTE.format(n=1) in row.validation_notes


def test_pending_sync_status(engagement, monkeypatch):
    ghost = folder(engagement, "A01") / "cloud.pdf"
    ghost.write_bytes(b"unsynced placeholder bytes")
    monkeypatch.setattr(
        "tracker.validators.is_cloud_placeholder", lambda p: p.name == "cloud.pdf"
    )
    scan_engagement(engagement, today=DAY1)
    row = statuses(engagement)["A01"]
    assert row.status == Status.PENDING_SYNC
    assert SYNCING_NOTE.format(n=1) in row.validation_notes


def test_deleted_folder_reported_missing(engagement):
    folder(engagement, "A01").rmdir()
    scan_engagement(engagement, today=DAY1)
    row = statuses(engagement)["A01"]
    assert row.status == Status.MISSING
    assert reasons.NO_REQUEST_FOLDER.matches(row.validation_notes)


# --------------------------------------------------------------- unfiled ----


def test_strays_in_prepared_are_warnings_and_parked_files_are_not(engagement):
    # Parked documents are the index's record (with why they were parked);
    # the scan only warns about what nothing else knows: loose files in
    # Prepared/ and folders matching no request.
    prepared = engagement / PREPARED_DIR_NAME
    review = prepared / REVIEW_DIR_NAME
    (review / "scan0012.pdf").write_bytes(b"x" * 100)
    (prepared / "loose_notes.txt").write_text("oops", encoding="utf-8")
    rogue = prepared / "misc uploads"
    rogue.mkdir()
    (rogue / "something.pdf").write_bytes(b"x")

    report = scan_engagement(engagement, today=DAY1)
    assert report.warnings == [
        f"loose_notes.txt is loose in {PREPARED_DIR_NAME}/; it belongs in a request folder",
        f"folder 'misc uploads' in {PREPARED_DIR_NAME}/ matches no request (1 file(s) inside)",
    ]
    wb = load_workbook(engagement / MANIFEST_FILENAME)
    assert wb.sheetnames == [SHEET_NAME, ENGAGEMENT_SHEET_NAME]     # no second record
    wb.close()

    (prepared / "loose_notes.txt").unlink()
    (rogue / "something.pdf").unlink()
    rogue.rmdir()
    assert scan_engagement(engagement, today=DAY2).warnings == []


def test_a_document_a_person_said_nothing_asks_for_is_not_a_warning_either(engagement):
    """A dismissal rewrites a row; the folder it leaves is the one it was.

    The index is the record of a parked document, with why it was parked and
    now who set it aside. A warning of its own would be the scan disagreeing
    with the index about a file the scan cannot see the decision on.
    """
    from tracker.filer import (
        INDEX_FILENAME,
        NOT_REQUESTED,
        dismiss_review_file,
        file_drops,
        read_index,
    )
    from tracker.scaffold import SHARED_DIR_NAME

    text_pdf(engagement / SHARED_DIR_NAME / "irs-notice.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    dismiss_review_file(engagement, parked.pbc_location, today=DAY1)

    report = scan_engagement(engagement, today=DAY2)
    assert report.warnings == []
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NOT_REQUESTED
    assert (engagement / row.prepared_location).is_file()


def test_the_scan_report_carries_the_one_summary(engagement):
    text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec 2025")
    report = scan_engagement(engagement, today=DAY1)
    assert report.summary.received == 1
    assert report.summary.outstanding == 2
    assert report.summary.line == SUMMARY_SEPARATOR.join([f"{Status.MISSING}: 2", f"{Status.RECEIVED}: 1"])


# ------------------------------------------------------- lock and dry-run ----


def test_dry_run_writes_nothing(engagement):
    text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec 2025")
    report = scan_engagement(engagement, today=DAY1, dry_run=True)
    assert report.dry_run and not report.written
    assert report.updates["A01"].status == Status.RECEIVED  # facts computed
    assert statuses(engagement)["A01"].status == ""          # nothing written
    assert not (engagement / CACHE_FILENAME).exists()
    assert not (engagement / LOCK_FILENAME).exists()


def test_fresh_lock_blocks_scan(engagement):
    (engagement / LOCK_FILENAME).write_text(f"pid={os.getpid()}", encoding="utf-8")
    with pytest.raises(ScanLockedError, match="another scan"):
        scan_engagement(engagement, today=DAY1)


def test_stale_lock_replaced_and_released(engagement):
    import os

    lock = engagement / LOCK_FILENAME
    lock.write_text(f"pid={os.getpid()}", encoding="utf-8")
    old = (dt.datetime.now() - dt.timedelta(seconds=STALE_LOCK_SECONDS + 1)).timestamp()
    os.utime(lock, (old, old))

    report = scan_engagement(engagement, today=DAY1)   # takes over stale lock
    assert report.written
    assert not lock.exists()                           # released afterwards


def test_the_scanner_reads_the_manifest_only_under_the_lock(engagement, monkeypatch):
    # A sort that finished between an early read and the lock would be
    # invisible to the scan, so the manifest is read after the lock is held.
    import tracker.scanner as scanner_module

    real = scanner_module.load_manifest

    def under_the_lock(path):
        assert (engagement / LOCK_FILENAME).exists(), "manifest read before the lock was taken"
        return real(path)

    monkeypatch.setattr(scanner_module, "load_manifest", under_the_lock)
    assert scan_engagement(engagement, today=DAY1).written
    assert not (engagement / LOCK_FILENAME).exists()


def test_content_cache_created_and_pruned(engagement):
    pdf = text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec")
    scan_engagement(engagement, today=DAY1)
    cache_file = engagement / CACHE_FILENAME
    assert cache_file.exists()
    data = json.loads(cache_file.read_text(encoding="utf-8"))
    assert any("chase.pdf" in k for k in data["files"])

    pdf.unlink()
    scan_engagement(engagement, today=DAY2)            # prune removes entry
    data = json.loads(cache_file.read_text(encoding="utf-8"))
    assert not any("chase.pdf" in k for k in data["files"])


def test_a_file_that_vanishes_mid_scan_does_not_crash_the_scan(engagement, monkeypatch):
    # The sync client replaces a file between the directory listing and the
    # stat. The scan must finish and write; the row waits for the next run.
    import tracker.scanner as scanner_module

    ghost = folder(engagement, "A01") / "ghost.pdf"
    ghost.write_bytes(b"%PDF-1.4 " + b"x" * 9000)
    real_listing = scanner_module.iter_candidate_files

    def listing_then_vanish(path):
        found = real_listing(path)
        if ghost in found:
            ghost.unlink()
        return found

    monkeypatch.setattr(scanner_module, "iter_candidate_files", listing_then_vanish)
    monkeypatch.setattr("tracker.validators.iter_candidate_files", listing_then_vanish)

    report = scan_engagement(engagement, today=DAY1)
    assert report.written
    assert report.updates["A01"].status == Status.PENDING_SYNC


def test_raising_expected_count_after_received_names_the_real_change(engagement):
    text_pdf(folder(engagement, "A01") / "chase.pdf", "Chase Bank Statement Dec 2025")
    scan_engagement(engagement, today=DAY1)
    wb = load_workbook(engagement / MANIFEST_FILENAME)
    wb[SHEET_NAME].cell(row=2, column=col(COL_EXPECTED_COUNT), value=2)   # A01 Expected Count 1 -> 2
    wb.save(engagement / MANIFEST_FILENAME)
    wb.close()
    report = scan_engagement(engagement, today=DAY2)
    update = report.updates["A01"]
    assert update.status == Status.PARTIAL
    assert update.received_date == DAY1
    assert REGRESSION_COUNT_RAISED.format(expected=2) in update.validation_notes
    assert REGRESSION_FILES_CHANGED not in update.validation_notes


def test_accepted_means_received_with_a_date(engagement):
    # Decision 2: Accepted = treat as Received despite the rules. The status
    # column, the date and every count that reads the column agree.
    from openpyxl import load_workbook as lw

    scan_engagement(engagement, today=DAY1)          # A01 Missing: no file at all
    wb = lw(engagement / MANIFEST_FILENAME)
    wb[SHEET_NAME].cell(row=2, column=col(COL_MANUAL_OVERRIDE), value=Override.ACCEPTED)
    wb.save(engagement / MANIFEST_FILENAME)
    report = scan_engagement(engagement, today=DAY2)
    update = report.updates["A01"]
    assert update.status == Status.RECEIVED
    assert update.received_date == DAY2
    assert update.validation_notes.startswith(OVERRIDE_NOTE.format(override=Override.ACCEPTED))
    assert statuses(engagement)["A01"].status == Status.RECEIVED
    # Stamped once: a later scan keeps the first date.
    assert scan_engagement(engagement, today=DAY2 + dt.timedelta(days=3)).updates["A01"].received_date == DAY2


def test_waived_rows_are_named_so_counts_can_leave_them_out(engagement):
    from openpyxl import load_workbook as lw

    wb = lw(engagement / MANIFEST_FILENAME)
    wb[SHEET_NAME].cell(row=4, column=col(COL_MANUAL_OVERRIDE), value=Override.WAIVED)   # B01
    wb.save(engagement / MANIFEST_FILENAME)
    report = scan_engagement(engagement, today=DAY1)
    assert report.summary.waived == 1
    assert report.summary.total == 2


def test_a_deferred_scan_is_what_the_next_scan_measures_from(engagement, monkeypatch):
    # Excel held the manifest through two scans. The Received Date the first
    # deferred scan stamped is "the first date all validations passed"; the
    # second must carry it, not re-stamp today. And when the file is then
    # gone and Excel closed, the row regresses from Received - with its date
    # and a note - instead of landing as plain Missing.
    from openpyxl.workbook.workbook import Workbook as WorkbookClass

    import tracker.manifest as manifest_module

    a01 = folder(engagement, "A01")
    text_pdf(a01 / "chase.pdf", "Chase Bank Statement Dec 2025")

    def locked_save(self, filename):
        raise PermissionError(f"[Errno 13] locked: {filename}")
    monkeypatch.setattr(WorkbookClass, "save", locked_save)
    monkeypatch.setattr(manifest_module, "LOCK_RETRY_DELAY", 0.001)

    first = scan_engagement(engagement, today=DAY1)
    assert first.deferred and first.updates["A01"].received_date == DAY1
    second = scan_engagement(engagement, today=DAY2)
    assert second.deferred and second.updates["A01"].received_date == DAY1

    (a01 / "chase.pdf").unlink()
    monkeypatch.undo()                                    # Excel closed
    third = scan_engagement(engagement, today=DAY2)
    assert third.written
    row = statuses(engagement)["A01"]
    assert row.status == Status.MISSING and row.received_date == DAY1
    assert REGRESSION_NOTE.format(status=Status.RECEIVED, date=DAY1.isoformat(), why="").rstrip("; ") \
        in row.validation_notes


def test_a_received_file_the_sync_client_dehydrated_is_not_a_regression(engagement, monkeypatch):
    # OneDrive "free up space" turns a Received file into a placeholder.
    # The row waits for the bytes; it did not lose the document.
    received = folder(engagement, "A01") / "chase.pdf"
    text_pdf(received, "Chase Bank Statement Dec 2025")
    scan_engagement(engagement, today=DAY1)
    assert statuses(engagement)["A01"].status == Status.RECEIVED

    monkeypatch.setattr("tracker.validators.is_cloud_placeholder", lambda p: p.name == "chase.pdf")
    scan_engagement(engagement, today=DAY2)
    row = statuses(engagement)["A01"]
    assert row.status == Status.PENDING_SYNC and row.received_date == DAY1
    assert REGRESSION_FILES_CHANGED not in row.validation_notes
    assert SYNCING_NOTE.format(n=1) in row.validation_notes


def test_an_empty_note_replaces_the_old_one(engagement):
    # openpyxl's cell(value=None) writes nothing. "request folder not found"
    # must not outlive the folder, or the reminder holds the row back for ever.
    folder(engagement, "A01").rmdir()
    scan_engagement(engagement, today=DAY1)
    assert reasons.NO_REQUEST_FOLDER.matches(statuses(engagement)["A01"].validation_notes)
    scaffold_engagement(engagement)                      # the next pass creates it
    scan_engagement(engagement, today=DAY2)
    assert statuses(engagement)["A01"].validation_notes == ""


def test_a_document_a_person_filed_is_not_second_guessed_by_the_rules(engagement):
    # A person filed it from Needs Review; the row's required keyword is not
    # in it. Their decision stands: Received, not "wrong document" and a
    # client asked for the right file.
    from tracker.filer import assign_review_file, file_drops
    from tracker.scaffold import SHARED_DIR_NAME
    from tracker.scanner import ACCEPTED_NOTE

    text_pdf(engagement / SHARED_DIR_NAME / "statement.pdf", "Annual account statement 2025 interest paid")
    parked = file_drops(engagement, today=DAY1).review[0]
    assign_review_file(engagement, parked.pbc_location, "A01", today=DAY1)
    scan_engagement(engagement, today=DAY1)
    row = statuses(engagement)["A01"]
    assert row.status == Status.RECEIVED and row.file_count == 1
    assert ACCEPTED_NOTE.format(n=1) in row.validation_notes


def test_a_persons_acceptance_covers_only_the_bytes_they_filed(engagement):
    # A person filed one document, and then a different one came to sit
    # under the canonical name theirs had. The acceptance was for their
    # bytes, not for that name: the newcomer faces the rules.
    #
    # Until decision 92 the newcomer arrived the way a client sends one -
    # a blank scan called "Chase scan.pdf", which the router filed on its
    # name. Nothing is filed on a name any more, and anything the rules
    # do route into that folder passes the same rules the scan applies, so
    # the only way a stranger reaches that name is a person putting it
    # there. That is what this writes, and the claim is unchanged.
    from tracker.filer import assign_review_file, file_drops
    from tracker.scaffold import SHARED_DIR_NAME

    text_pdf(engagement / SHARED_DIR_NAME / "statement.pdf", "Annual account statement 2025 interest paid")
    parked = file_drops(engagement, today=DAY1).review[0]
    filed = assign_review_file(engagement, parked.pbc_location, "A01", today=DAY1).entry
    (engagement / filed.prepared_location).unlink()
    text_pdf(engagement / filed.prepared_location, "Some other bank's statement for 2025")
    scan_engagement(engagement, today=DAY2)
    row = statuses(engagement)["A01"]
    assert row.status != Status.RECEIVED and "filed here by a person" not in row.validation_notes
