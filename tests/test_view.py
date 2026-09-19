"""Tests for tracker/view.py - the read-only workbook a pass regenerates.

The view's standing claim is not made here either: it is the autouse fixture
in tests/conftest.py, which compares every view the suite leaves behind with
what the readers say, cell for cell. What is made here is the view's own
behaviour - the four sheets, the stamp, the three words it answers "is this
still true?" with, the read-only attribute that does not stop the next pass,
and the one failure it is allowed to have: somebody had it open.
"""

from __future__ import annotations

import datetime as dt
import io
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.test_scanner import text_pdf
from tracker import view
from tracker.filer import INDEX_FILENAME, INDEX_SHEET, NEEDS_REVIEW, file_drops, read_index
from tracker.manifest import (
    COL_ANY_KEYWORDS,
    COL_STATUS,
    HEADERS,
    SHEET_NAME,
    RequestItem,
    Status,
    add_any_keyword,
    create_template,
    load_manifest,
)
from tracker.registry import RegistryError, discover_engagements, engagement_dirs
from tracker.runner import run_engagement
from tracker.scaffold import MANIFEST_FILENAME, SHARED_DIR_NAME, scaffold_engagement
from tracker.scanner import scan_engagement

REPO = Path(__file__).resolve().parents[1]
DAY1 = dt.date(2026, 7, 1)

ITEMS = [
    RequestItem(
        identifier="A01", document="W-2 Wage Statements", period="TY2025",
        allowed_extensions=("pdf",), min_size_kb=0, required_keywords=("W-2",),
        date_pattern=r"(?i)\b2025\b",
    ),
    RequestItem(
        identifier="C01", document="Mortgage Interest Statement", period="TY2025",
        allowed_extensions=("pdf",), min_size_kb=0, required_keywords=("1098",),
    ),
]


@pytest.fixture
def engagement(tmp_path):
    eng = tmp_path / "Smith Family 2025"
    eng.mkdir()
    create_template(eng / MANIFEST_FILENAME, ITEMS)
    scaffold_engagement(eng)
    return eng


def drop(engagement, name, text):
    return text_pdf(engagement / SHARED_DIR_NAME / name, text)


def a_pass(engagement, today=DAY1):
    """What a pass does to the folder, without a registry: sort, then scan."""
    file_drops(engagement, today=today)
    scan_engagement(engagement, today=today)


def sheets(engagement):
    return load_workbook(io.BytesIO(view.path_for(engagement).read_bytes()), data_only=True)


def rows_of(engagement, title, *, min_row=2):
    wb = sheets(engagement)
    try:
        return [["" if cell is None else str(cell) for cell in row]
                for row in wb[title].iter_rows(min_row=min_row, values_only=True)]
    finally:
        wb.close()


# ----------------------------------------------------------- the four sheets ----


def test_the_view_has_the_four_sheets_with_the_headers_their_owners_name(engagement):
    view.write_view(engagement)
    wb = sheets(engagement)
    try:
        assert wb.sheetnames == list(view.SHEETS)
        assert [c.value for c in wb[SHEET_NAME][1]] == list(HEADERS)
        assert [c.value for c in wb[INDEX_SHEET][1]] == list(view.INDEX_COLUMNS)
        assert [c.value for c in wb[NEEDS_REVIEW][1]] == [
            view.INDEX_LAYOUT[name][0] for name in view.NEEDS_REVIEW_FIELDS
        ]
    finally:
        wb.close()


def test_every_index_cell_is_the_readers_row_written_as_text(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "puzzle.pdf", "a page about nothing anybody asked for")
    a_pass(engagement)
    view.write_view(engagement)

    entries = read_index(engagement / INDEX_FILENAME)
    assert len(entries) == 2
    assert rows_of(engagement, INDEX_SHEET) == [view.index_row(e) for e in entries]
    # Every cell is a string: nothing in the view is a number or a date Excel
    # could re-type, and nothing is a formula.
    wb = sheets(engagement)
    try:
        for row in wb[INDEX_SHEET].iter_rows(min_row=2):
            for cell in row:
                assert cell.data_type != "f", (cell.coordinate, cell.value)
                assert cell.value is None or isinstance(cell.value, str), cell.coordinate
    finally:
        wb.close()


def test_the_requests_sheet_carries_a_learned_keyword_and_a_status_the_record_holds(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)
    add_any_keyword(engagement / MANIFEST_FILENAME, "C01", "lender")
    view.write_view(engagement)

    rows = rows_of(engagement, SHEET_NAME)
    by_identifier = {row[0]: dict(zip(HEADERS, row, strict=True)) for row in rows}
    assert by_identifier["A01"][COL_STATUS] == Status.RECEIVED
    assert "lender" in by_identifier["C01"][COL_ANY_KEYWORDS]
    # And that status is the one the record holds, not a second reading.
    assert by_identifier["A01"][COL_STATUS] == next(
        i.status for i in load_manifest(engagement / MANIFEST_FILENAME) if i.identifier == "A01"
    )


def test_needs_review_lists_exactly_the_parked_rows(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "puzzle.pdf", "a page about nothing anybody asked for")
    a_pass(engagement)
    view.write_view(engagement)

    parked = [e for e in read_index(engagement / INDEX_FILENAME) if e.decision == NEEDS_REVIEW]
    assert [e.original_name for e in parked] == ["puzzle.pdf"]
    assert rows_of(engagement, NEEDS_REVIEW) == [
        [getattr(e, name) for name in view.NEEDS_REVIEW_FIELDS] for e in parked
    ]


# ------------------------------------------------------------------ read-only ----


def test_the_file_is_read_only_afterwards_and_the_next_pass_replaces_it_anyway(engagement):
    first = view.write_view(engagement)
    assert not os.access(first.path, os.W_OK)
    before = first.path.read_bytes()

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)
    second = view.write_view(engagement)

    assert not second.stale
    assert second.path.read_bytes() != before
    assert not os.access(second.path, os.W_OK)
    assert second.rows == 1


# ---------------------------------------------------------------- the stamp ----


def test_the_summary_carries_the_rules_digest_and_the_record_head(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)
    view.write_view(engagement)

    from tracker import ledger

    stamp = view.read_stamp(engagement)
    assert stamp[view.LABEL_RULES_DIGEST] == view.rules_digest(engagement)
    assert stamp[view.LABEL_RECORD_DIGEST] == ledger.head(engagement)
    assert stamp[view.LABEL_ENGAGEMENT] == engagement.name
    assert view.VIEW_NOTE in [c.value for c in sheets(engagement)[view.SUMMARY_SHEET]["A"]]


def test_the_stamp_changes_when_the_rules_change_and_when_the_record_grows(engagement):
    view.write_view(engagement)
    first = view.read_stamp(engagement)

    add_any_keyword(engagement / MANIFEST_FILENAME, "C01", "lender")
    view.write_view(engagement)
    second = view.read_stamp(engagement)
    assert second[view.LABEL_RULES_DIGEST] != first[view.LABEL_RULES_DIGEST]
    assert second[view.LABEL_RECORD_DIGEST] == first[view.LABEL_RECORD_DIGEST]

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)
    view.write_view(engagement)
    third = view.read_stamp(engagement)
    assert third[view.LABEL_RECORD_DIGEST] != second[view.LABEL_RECORD_DIGEST]


# ---------------------------------------------------------------- the state ----


def test_there_is_no_view_until_one_is_written(engagement):
    assert view.view_state(engagement) == view.UNKNOWN
    view.write_view(engagement)
    assert view.view_state(engagement) == view.CURRENT


def test_a_stamp_that_cannot_be_read_is_unknown_rather_than_believed(engagement):
    view.write_view(engagement)
    path = view.path_for(engagement)
    os.chmod(path, 0o600)
    path.write_bytes(b"this is not a workbook")
    assert view.view_state(engagement) == view.UNKNOWN


def test_the_view_is_behind_after_a_person_edits_the_rules_and_after_a_new_event(engagement):
    view.write_view(engagement)
    add_any_keyword(engagement / MANIFEST_FILENAME, "C01", "lender")
    assert view.view_state(engagement) == view.BEHIND

    view.write_view(engagement)
    assert view.view_state(engagement) == view.CURRENT

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)
    assert view.view_state(engagement) == view.BEHIND

    view.write_view(engagement)
    assert view.view_state(engagement) == view.CURRENT


# --------------------------------------------------------- held open in Excel ----


def test_a_view_somebody_has_open_is_reported_stale_and_the_pass_still_succeeds(
    engagement, monkeypatch, caplog
):
    """Everywhere: the replace fails, the old view stands, the run is fine."""
    view.write_view(engagement)
    before = view.path_for(engagement).read_bytes()

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)

    def held(wb, path):
        raise PermissionError(13, "The process cannot access the file")

    monkeypatch.setattr(view, "save_workbook_atomically", held)
    with caplog.at_level("WARNING"):
        result = view.write_view(engagement)

    assert result.stale
    assert view.path_for(engagement).read_bytes() == before
    assert engagement.name in caplog.text
    # And the reader can still say, truthfully, that it is out of date.
    assert view.view_state(engagement) == view.BEHIND


def test_a_pass_whose_view_is_held_open_reports_it_and_still_succeeds(
    tmp_path, engagement, monkeypatch
):
    from tracker import runner
    from tracker.registry import engagement_from

    view.write_view(engagement)
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")

    def held(wb, path):
        raise PermissionError(13, "The process cannot access the file")

    monkeypatch.setattr(view, "save_workbook_atomically", held)
    run = run_engagement(engagement_from(engagement), today=DAY1)

    assert run.ok and not run.error
    assert run.view_stale
    assert run.filed == 1
    assert runner.VIEW_NOT_REGENERATED in run.summary()


@pytest.mark.skipif(sys.platform != "win32",
                    reason="Excel's share modes are a Windows file-system behaviour")
def test_a_view_held_the_way_excel_holds_one_reads_and_does_not_replace(engagement):
    """The first real evidence in the suite of Excel's share modes.

    Every other "locked" test in this repository monkeypatches
    ``Workbook.save`` or the atomic replace; none of them proves what
    Windows actually does when a workbook is open. This one opens the view
    with ``CreateFileW``, ``GENERIC_READ``, share mode ``FILE_SHARE_READ``
    and nothing else - what Excel does to a workbook it has open - and,
    with that handle held, shows the two halves of the design: reading the
    stamp still answers, and regenerating reports ``stale`` rather than
    raising or leaving a half-written file. Close the handle and the next
    pass replaces it.
    """
    import ctypes
    from ctypes import wintypes

    GENERIC_READ = 0x80000000
    FILE_SHARE_READ = 0x00000001
    OPEN_EXISTING = 3
    FILE_ATTRIBUTE_NORMAL = 0x80
    INVALID_HANDLE_VALUE = wintypes.HANDLE(-1).value

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                     wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD,
                                     wintypes.HANDLE)
    kernel32.CreateFileW.restype = wintypes.HANDLE
    kernel32.CloseHandle.argtypes = (wintypes.HANDLE,)

    view.write_view(engagement)
    path = view.path_for(engagement)
    before = path.read_bytes()

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)

    handle = kernel32.CreateFileW(str(path), GENERIC_READ, FILE_SHARE_READ, None,
                                  OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL, None)
    assert handle != INVALID_HANDLE_VALUE, ctypes.get_last_error()
    try:
        assert view.view_state(engagement) == view.BEHIND   # a read, and it answers
        held = view.write_view(engagement)
        assert held.stale
        assert path.read_bytes() == before                  # not a byte of it moved
    finally:
        kernel32.CloseHandle(handle)

    landed = view.write_view(engagement)
    assert not landed.stale
    assert path.read_bytes() != before
    assert view.view_state(engagement) == view.CURRENT


# -------------------------------------------------------- the pass and the app ----


def test_a_pass_regenerates_the_view_and_a_dry_run_writes_none(tmp_path, engagement):
    from tracker.registry import engagement_from

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    run_engagement(engagement_from(engagement), today=DAY1, dry_run=True)
    assert not view.path_for(engagement).exists()

    run = run_engagement(engagement_from(engagement), today=DAY1)
    assert not run.view_stale
    assert view.view_state(engagement) == view.CURRENT
    assert rows_of(engagement, INDEX_SHEET) == [
        view.index_row(e) for e in read_index(engagement / INDEX_FILENAME)
    ]


def test_the_state_the_app_reads_carries_the_view_and_its_path(engagement):
    from tracker.api import _state, _vocab

    state = _state(engagement)
    assert state["view"] == {"state": view.UNKNOWN, "path": str(view.path_for(engagement))}
    assert state["paths"]["view"] == str(view.path_for(engagement))

    view.write_view(engagement)
    assert _state(engagement)["view"]["state"] == view.CURRENT
    assert _vocab()["view"] == {"label": view.VIEW_LABEL, "states": list(view.VIEW_STATES)}


# ------------------------------------------------------------- not a manifest ----


def test_the_registry_never_mistakes_the_view_for_an_engagement(tmp_path, engagement):
    """A folder holding only a view is not an engagement, and the rollover's
    prior-year discovery - the same walk - never offers one."""
    view.write_view(engagement)
    lonely = tmp_path / "Not An Engagement"
    lonely.mkdir()
    (lonely / view.VIEW_FILENAME).write_bytes(view.path_for(engagement).read_bytes())

    assert engagement_dirs(tmp_path) == [engagement]
    assert [e.path for e in discover_engagements(tmp_path).engagements] == [engagement]

    alone = tmp_path / "alone"
    (alone / "Client 2025").mkdir(parents=True)
    (alone / "Client 2025" / view.VIEW_FILENAME).write_bytes(b"not a manifest")
    with pytest.raises(RegistryError):
        discover_engagements(alone)


# -------------------------------------------------------------------- CLI ----


def test_the_cli_writes_the_view_and_prints_its_state(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    a_pass(engagement)

    out = subprocess.run(
        [sys.executable, "-m", "tracker.view", str(engagement)],
        cwd=REPO, capture_output=True, text=True, check=True,
    ).stdout

    assert view.path_for(engagement).is_file()
    assert f"state:    {view.CURRENT}" in out
    assert view.LABEL_RECORD_DIGEST in out
    assert "rows:     1" in out
