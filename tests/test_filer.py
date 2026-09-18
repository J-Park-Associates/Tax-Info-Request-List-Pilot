"""Tests for tracker/filer.py — sorting the client's drop folder.

The promises being tested: the client's original is preserved byte-for-byte
under its own name in PBC, a renamed working copy appears in Prepared, every
decision is written to the index, and running twice changes nothing.
"""

import datetime as dt
import os
import sys

import pytest
from openpyxl import load_workbook

from tests.samples import col
from tests.test_scanner import text_pdf
from tracker.filer import (
    DUPLICATE,
    FILED,
    INDEX_COLUMNS,
    INDEX_FILENAME,
    INDEX_LAYOUT,
    INDEX_PENDING_FILENAME,
    INDEX_SHEET,
    NEEDS_REVIEW,
    IndexEntry,
    file_drops,
    prepared_name_for,
    read_index,
)
from tracker.manifest import COL_DOCUMENT, SHEET_NAME, RequestItem, create_template
from tracker.router import UNMATCHED
from tracker.scaffold import (
    MANIFEST_FILENAME,
    PBC_DIR_NAME,
    PREPARED_DIR_NAME,
    REVIEW_DIR_NAME,
    SHARED_DIR_NAME,
    scaffold_engagement,
)

DAY1 = dt.date(2026, 7, 1)
DAY2 = dt.date(2026, 7, 9)

ITEMS = [
    RequestItem(
        identifier="A01", document="W-2 Wage Statements", period="TY2025",
        expected_count=2, allowed_extensions=("pdf",), min_size_kb=0,
        required_keywords=("W-2",), date_pattern=r"(?i)\b2025\b",
    ),
    RequestItem(
        identifier="C01", document="Mortgage Interest Statement", period="TY2025",
        allowed_extensions=("pdf",), min_size_kb=0, required_keywords=("1098",),
    ),
]


@pytest.fixture
def engagement(tmp_path):
    eng = tmp_path / "Smith Family 2025"
    eng.mkdir()
    create_template(eng / MANIFEST_FILENAME, ITEMS)
    scaffold_engagement(eng)
    return eng


def drop(engagement, name, text):
    """Client drops one document into the single shared folder."""
    return text_pdf(engagement / SHARED_DIR_NAME / name, text)


def pbc(engagement):
    return engagement / SHARED_DIR_NAME / PBC_DIR_NAME


def prepared(engagement, folder_prefix):
    root = engagement / PREPARED_DIR_NAME
    return next(p for p in root.iterdir() if p.name.startswith(folder_prefix))


# ------------------------------------------------------------- the happy path ----


def test_file_is_sorted_renamed_and_indexed(engagement):
    original = drop(engagement, "scan0012.pdf", "Form W-2 Wage and Tax Statement 2025")
    before = original.read_bytes()

    report = file_drops(engagement, today=DAY1)

    # Original: moved out of the drop zone, name and bytes untouched.
    assert not original.exists()
    kept = pbc(engagement) / "scan0012.pdf"
    assert kept.read_bytes() == before

    # Working copy: renamed to the firm's convention.
    copies = list(prepared(engagement, "A01").iterdir())
    assert [p.name for p in copies] == ["A01 - W-2 Wage Statements - TY2025.pdf"]
    assert copies[0].read_bytes() == before

    # Index: one row that maps original -> filed.
    (entry,) = report.filed
    assert entry.original_name == "scan0012.pdf"
    assert entry.decision == FILED
    assert entry.identifier == "A01"
    assert entry.pbc_location == f"{SHARED_DIR_NAME}/{PBC_DIR_NAME}/scan0012.pdf"
    assert entry.prepared_location.endswith("A01 - W-2 Wage Statements - TY2025.pdf")

    rows = read_index(engagement / INDEX_FILENAME)
    assert [r.original_name for r in rows] == ["scan0012.pdf"]
    assert rows[0].digest == entry.digest


def test_second_file_for_the_same_request_is_numbered(engagement):
    drop(engagement, "john w2.pdf", "Form W-2 Wage and Tax Statement 2025 John")
    drop(engagement, "jane w2.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")
    file_drops(engagement, today=DAY1)

    names = sorted(p.name for p in prepared(engagement, "A01").iterdir())
    assert names == [
        "A01 - W-2 Wage Statements - TY2025 (2).pdf",
        "A01 - W-2 Wage Statements - TY2025.pdf",
    ]


def test_client_subfolders_are_flattened(engagement):
    nested = engagement / SHARED_DIR_NAME / "tax stuff"
    nested.mkdir()
    text_pdf(nested / "w2.pdf", "Form W-2 Wage and Tax Statement 2025")

    report = file_drops(engagement, today=DAY1)
    assert len(report.filed) == 1
    assert (pbc(engagement) / "w2.pdf").exists()


# --------------------------------------------------------------- needs review ----


def test_unroutable_file_is_preserved_and_parked(engagement):
    original = drop(engagement, "vacation.pdf", "Photos from Maui")
    before = original.read_bytes()

    report = file_drops(engagement, today=DAY1)

    assert (pbc(engagement) / "vacation.pdf").read_bytes() == before
    review = engagement / PREPARED_DIR_NAME / REVIEW_DIR_NAME
    assert (review / "vacation.pdf").exists()
    (entry,) = report.review
    assert entry.decision == NEEDS_REVIEW
    assert UNMATCHED in entry.reason
    assert not report.filed


def test_review_copies_keep_the_clients_own_name(engagement):
    drop(engagement, "IMG_4021.pdf", "unreadable receipt")
    file_drops(engagement, today=DAY1)
    review = engagement / PREPARED_DIR_NAME / REVIEW_DIR_NAME
    assert [p.name for p in review.iterdir()] == ["IMG_4021.pdf"]


# ------------------------------------------------------------------ re-running ----


def test_rerunning_changes_nothing(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    snapshot = {
        p: p.read_bytes()
        for p in (engagement / PREPARED_DIR_NAME).rglob("*") if p.is_file()
    }

    second = file_drops(engagement, today=DAY2)

    assert second.handled == 0
    assert {
        p: p.read_bytes()
        for p in (engagement / PREPARED_DIR_NAME).rglob("*") if p.is_file()
    } == snapshot


def test_same_document_dropped_twice_is_kept_but_filed_once(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    # Client re-sends the identical document under a different name.
    drop(engagement, "w2 again.pdf", "Form W-2 Wage and Tax Statement 2025")

    report = file_drops(engagement, today=DAY2)

    (dup,) = report.duplicates
    assert dup.decision == DUPLICATE
    assert "identical to w2.pdf" in dup.reason
    # Preserved, never deleted...
    assert (pbc(engagement) / "w2 again.pdf").exists()
    # ...but not filed a second time.
    assert len(list(prepared(engagement, "A01").iterdir())) == 1


def test_name_collision_in_pbc_never_overwrites(engagement):
    drop(engagement, "statement.pdf", "Form W-2 Wage and Tax Statement 2025 John")
    file_drops(engagement, today=DAY1)
    drop(engagement, "statement.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")

    file_drops(engagement, today=DAY2)

    kept = sorted(p.name for p in pbc(engagement).iterdir())
    assert kept == ["statement (2).pdf", "statement.pdf"]


# ------------------------------------------------------------------- guarded ----


def test_placeholders_are_left_where_they_are(engagement, monkeypatch):
    waiting = drop(engagement, "big.pdf", "Form W-2 Wage and Tax Statement 2025")
    monkeypatch.setattr("tracker.filer.is_cloud_placeholder", lambda p: True)

    report = file_drops(engagement, today=DAY1)

    assert report.waiting == [waiting]
    assert waiting.exists()                       # untouched, still syncing
    assert not any(pbc(engagement).iterdir())


def test_sync_junk_is_ignored(engagement):
    junk = engagement / SHARED_DIR_NAME / "w2.pdf.tmp.driveupload"
    junk.write_bytes(b"partial upload")

    report = file_drops(engagement, today=DAY1)

    assert report.handled == 0
    assert junk.exists()                          # not ours to move or delete


def test_pbc_contents_are_never_re_sorted(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    # The originals now sitting in PBC must not look like new drops.
    assert file_drops(engagement, today=DAY2).handled == 0


def test_dry_run_moves_nothing(engagement):
    original = drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")

    report = file_drops(engagement, today=DAY1, dry_run=True)

    assert len(report.filed) == 1                 # it still tells you the plan
    assert original.exists()
    assert not any(pbc(engagement).iterdir())
    assert not (engagement / INDEX_FILENAME).exists()


def test_a_dry_run_never_quarantines_a_corrupt_index_sidecar(engagement, caplog):
    # "Moves nothing" includes the sidecar: a preview must not rename a file
    # a real run would have moved aside as evidence.
    from tracker.manifest import CORRUPT_SUFFIX

    sidecar = engagement / INDEX_PENDING_FILENAME
    sidecar.write_text("{not json", encoding="utf-8")
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")

    report = file_drops(engagement, today=DAY1, dry_run=True)

    assert len(report.filed) == 1
    assert sidecar.read_text(encoding="utf-8") == "{not json"
    assert list(engagement.glob(f"*{CORRUPT_SUFFIX}")) == []
    assert "ignored for this read" in caplog.text


def test_filing_leaves_verdicts_the_scan_reuses(engagement, monkeypatch):
    # Route once, scan once, read the document once: the router's verdicts
    # are the scanner's, keyed by content so the working copy is a hit.
    from tests.test_content_check import counting_extractor
    from tracker.content_check import CACHE_FILENAME
    from tracker.scanner import scan_engagement

    calls = counting_extractor(monkeypatch)
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1, dry_run=True)
    assert not (engagement / CACHE_FILENAME).exists()      # a dry run writes nothing
    assert file_drops(engagement, today=DAY1).handled == 1
    assert (engagement / CACHE_FILENAME).exists()
    assert calls["n"] == 2                                  # the preview and the run

    report = scan_engagement(engagement, today=DAY1)
    assert report.updates["A01"].file_count == 1
    assert calls["n"] == 2                                  # the scan read nothing again


# --------------------------------------------------------------------- index ----


def test_index_workbook_is_readable_in_excel(engagement):
    from tracker.filer import INDEX_COLUMNS, INDEX_LAYOUT

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    wb = load_workbook(engagement / INDEX_FILENAME)
    ws = wb[INDEX_SHEET]
    headers = [c.value for c in ws[1]]
    assert headers == list(INDEX_COLUMNS)
    # Read the row by header name, the way the filer itself reads it back.
    row = dict(zip(headers, [c.value for c in ws[2]], strict=True))
    assert row[INDEX_LAYOUT["decision"][0]] == FILED
    assert row[INDEX_LAYOUT["identifier"][0]] == "A01"
    assert row[INDEX_LAYOUT["prepared_location"][0]].endswith(".pdf")
    assert ws.freeze_panes == "A2"
    wb.close()


def test_an_index_written_with_older_columns_still_reads(tmp_path):
    # Filed As and Document were stored copies and are gone; Candidates is
    # new. An index from before either change reads by header name.
    from openpyxl import Workbook

    from tracker.filer import read_index

    wb = Workbook()
    ws = wb.active
    ws.title = INDEX_SHEET
    ws.append(["Received", "Original Name", "Size KB", "SHA-256", "Identifier", "Document",
               "Filed As", "Prepared Location", "PBC Location", "Decision", "Reason"])
    ws.append(["2026-01-01", "w2.pdf", 9.4, "abc", "A01", "W-2 Wage Statements",
               "A01 - W-2 Wage Statements - TY2025.pdf",
               f"{PREPARED_DIR_NAME}/A01 - W-2 Wage Statements/A01 - W-2 Wage Statements - TY2025.pdf",
               f"{SHARED_DIR_NAME}/{PBC_DIR_NAME}/w2.pdf", FILED, "content matched"])
    path = tmp_path / INDEX_FILENAME
    wb.save(path)
    [entry] = read_index(path)
    assert entry.identifier == "A01" and entry.decision == FILED
    assert entry.filed_as == "A01 - W-2 Wage Statements - TY2025.pdf"   # derived, not stored
    assert entry.candidates == ""
    # Evidence is newer than Candidates; an index from before it reads too.
    assert entry.evidence == "" and entry.evidence_record == {}


def test_the_routers_candidates_travel_as_data_in_the_index(engagement):
    from tracker.filer import read_index

    drop(engagement, "old.pdf", "Form W-2 Wage and Tax Statement 2024")   # contested: wrong year
    report = file_drops(engagement, today=DAY1)
    [parked] = report.review
    assert parked.candidates == "A01"
    assert read_index(engagement / INDEX_FILENAME)[0].candidates == "A01"


def test_a_filed_row_and_a_parked_row_both_say_what_the_rules_saw(engagement):
    from tracker.content_check import RULE_REQUIRED, WHERE_TITLE
    from tracker.filer import read_index

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "old.pdf", "Form W-2 Wage and Tax Statement 2024")   # contested: wrong year
    file_drops(engagement, today=DAY1)

    by_name = {e.original_name: e for e in read_index(engagement / INDEX_FILENAME)}
    for name, decision in (("w2.pdf", FILED), ("old.pdf", NEEDS_REVIEW)):
        entry = by_name[name]
        assert entry.decision == decision
        # Read back through the workbook, parsed by the one parser.
        assert [(e.rule, e.term, e.where) for e in entry.evidence_record["A01"]][:1] == [
            (RULE_REQUIRED, "W-2", WHERE_TITLE),
        ], name
    # The cell carries the firm's own keyword and nothing the document said.
    assert "W-2@" in by_name["w2.pdf"].evidence


def test_the_evidence_column_round_trips_through_the_snapshot_sidecar(tmp_path):
    from tracker.content_check import RULE_ANY, WHERE_FOOTER, Evidence, format_evidence
    from tracker.filer import _read_pending_index, _save_pending_index

    entry = IndexEntry(
        received="2026-01-01", original_name="w2.pdf", size_kb=9.4, digest="abc",
        identifier="A01", prepared_location="", pbc_location="", decision=NEEDS_REVIEW,
        reason="matched no request", candidates="A01",
        evidence=format_evidence({"A01": (Evidence(RULE_ANY, "dividend", WHERE_FOOTER, 3),)}),
    )
    path = tmp_path / INDEX_FILENAME
    _save_pending_index(path, [entry])
    pending = _read_pending_index(path)
    assert pending is not None and pending.entries == [entry]
    assert pending.entries[0].evidence_record == {
        "A01": (Evidence(RULE_ANY, "dividend", WHERE_FOOTER, 3),),
    }


def test_prepared_name_for_collides_safely():
    item = ITEMS[0]
    taken: set[str] = set()
    first = prepared_name_for(item, "pdf", taken)
    second = prepared_name_for(item, "pdf", taken)
    assert first == "A01 - W-2 Wage Statements - TY2025.pdf"
    assert second == "A01 - W-2 Wage Statements - TY2025 (2).pdf"


def test_dry_run_previews_real_numbering(engagement):
    drop(engagement, "john w2.pdf", "Form W-2 Wage and Tax Statement 2025 John")
    drop(engagement, "jane w2.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")

    planned = [e.filed_as for e in file_drops(engagement, today=DAY1, dry_run=True).filed]

    assert sorted(planned) == [
        "A01 - W-2 Wage Statements - TY2025 (2).pdf",
        "A01 - W-2 Wage Statements - TY2025.pdf",
    ]


# ------------------------------------------- one bad file never costs the trail ----


def test_a_locked_index_never_orphans_files_already_moved(engagement, monkeypatch):
    # Originals are moved into PBC before the index is written. If Excel
    # holds _index.xlsx, the rows must still be recorded somewhere the next
    # run reads — a moved file that is nowhere in the index is lost to the
    # workflow, because PBC is never re-sorted.
    import tracker.filer as filer_module

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "mortgage.pdf", "Form 1098 Mortgage Interest Statement 2025")

    def locked(wb, path):
        raise PermissionError(f"[Errno 13] locked: {path}")

    monkeypatch.setattr(filer_module, "save_workbook_atomically", locked)
    monkeypatch.setattr(filer_module, "LOCK_RETRY_DELAY", 0.001)
    report = file_drops(engagement, today=DAY1)
    assert report.handled == 2
    assert report.index_deferred is True
    assert (engagement / filer_module.INDEX_PENDING_FILENAME).exists()
    assert not (engagement / INDEX_FILENAME).exists()
    # The deferred rows are already visible to whoever reads the index.
    assert {e.original_name for e in read_index(engagement / INDEX_FILENAME)} == {
        "w2.pdf", "mortgage.pdf",
    }

    # Excel closed; the next run has nothing to sort but folds the rows in.
    monkeypatch.undo()
    report = file_drops(engagement, today=DAY2)
    assert report.index_deferred is False
    assert (engagement / INDEX_FILENAME).exists()
    assert not (engagement / filer_module.INDEX_PENDING_FILENAME).exists()
    rows = read_index(engagement / INDEX_FILENAME)
    assert {e.original_name for e in rows} == {"w2.pdf", "mortgage.pdf"}
    assert all(e.decision == FILED for e in rows)


def _lock_the_index(monkeypatch):
    """Excel has _index.xlsx open: every save of it fails, quickly."""
    import tracker.filer as filer_module

    def locked(wb, path):
        raise PermissionError(f"[Errno 13] locked: {path}")

    monkeypatch.setattr(filer_module, "save_workbook_atomically", locked)
    monkeypatch.setattr(filer_module, "LOCK_RETRY_DELAY", 0.001)


def test_a_locked_index_keeps_a_persons_filing_decision(engagement, monkeypatch):
    # The row a person rewrote is an edit to a row the workbook already
    # holds. Deferring only the rows past the workbook's end lost it, while
    # the working copy had already been moved under the canonical name.
    from tracker.filer import ASSIGNED_BY_PERSON, assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    sidecar = engagement / INDEX_PENDING_FILENAME

    _lock_the_index(monkeypatch)
    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    assert result.index_deferred is True
    assert sidecar.exists()
    assert (engagement / result.entry.prepared_location).exists()
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and row.identifier == "C01"
    assert row.reason.startswith(f"{ASSIGNED_BY_PERSON} on {DAY2.isoformat()}; was: ")
    assert row.prepared_location == result.entry.prepared_location

    monkeypatch.undo()                            # Excel closed
    assert file_drops(engagement, today=DAY2).index_deferred is False
    assert not sidecar.exists()
    [row] = read_index(engagement / INDEX_FILENAME)        # the workbook alone now
    assert row.decision == FILED and row.identifier == "C01"


def test_the_pending_index_is_a_versioned_snapshot_of_every_row(engagement, monkeypatch):
    import json

    from tracker.filer import INDEX_SIDECAR_VERSION

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)                     # one row in the workbook
    drop(engagement, "mortgage.pdf", "Form 1098 Mortgage Interest Statement 2025")
    _lock_the_index(monkeypatch)
    file_drops(engagement, today=DAY2)

    payload = json.loads((engagement / INDEX_PENDING_FILENAME).read_text(encoding="utf-8"))
    assert payload["version"] == INDEX_SIDECAR_VERSION
    rows = read_index(engagement / INDEX_FILENAME)
    assert [e["original_name"] for e in payload["entries"]] == [e.original_name for e in rows]
    assert [e.original_name for e in rows] == ["w2.pdf", "mortgage.pdf"]


def test_mixed_append_and_edit_survive_a_locked_index(engagement, monkeypatch):
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]

    _lock_the_index(monkeypatch)
    assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)   # an edit
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY2)                                       # an append
    by_name = {e.original_name: e for e in read_index(engagement / INDEX_FILENAME)}
    assert by_name["scan0012.pdf"].decision == FILED and by_name["scan0012.pdf"].identifier == "C01"
    assert by_name["w2.pdf"].decision == FILED and by_name["w2.pdf"].identifier == "A01"

    monkeypatch.undo()
    file_drops(engagement, today=DAY2)
    assert not (engagement / INDEX_PENDING_FILENAME).exists()
    from openpyxl import load_workbook
    wb = load_workbook(engagement / INDEX_FILENAME, read_only=True)
    written = [r for r in wb[INDEX_SHEET].iter_rows(values_only=True)][1:]
    wb.close()
    assert sorted((r[1], r[4], r[7]) for r in written) == [
        ("scan0012.pdf", "C01", FILED), ("w2.pdf", "A01", FILED),
    ]


def test_a_bare_list_sidecar_from_an_older_version_is_appended_once(engagement):
    import json
    from dataclasses import asdict

    from tracker.filer import IndexEntry

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    old_row = IndexEntry(
        received=DAY1.isoformat(), original_name="older.pdf", size_kb=1.0, digest="abc",
        identifier="C01", prepared_location="Prepared/C01/x.pdf",
        pbc_location="Shared/PBC/older.pdf", decision=FILED, reason="",
    )
    (engagement / INDEX_PENDING_FILENAME).write_text(json.dumps([asdict(old_row)]), encoding="utf-8")

    assert [e.original_name for e in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf", "older.pdf"]
    assert file_drops(engagement, today=DAY2).index_deferred is False       # nothing to sort
    assert not (engagement / INDEX_PENDING_FILENAME).exists()
    assert [e.original_name for e in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf", "older.pdf"]


def test_a_sidecar_of_an_unknown_version_is_quarantined_not_guessed(engagement, caplog):
    import json

    from tracker.manifest import CORRUPT_SUFFIX

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    sidecar = engagement / INDEX_PENDING_FILENAME
    sidecar.write_text(json.dumps({"version": 99, "entries": []}), encoding="utf-8")

    assert [e.original_name for e in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]
    assert not sidecar.exists()
    assert sidecar.with_suffix(CORRUPT_SUFFIX).exists()
    assert "version 99" in caplog.text


def test_deferred_rows_are_folded_in_even_when_the_only_drop_is_still_syncing(engagement, monkeypatch):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    _lock_the_index(monkeypatch)
    assert file_drops(engagement, today=DAY1).index_deferred is True
    monkeypatch.undo()

    waiting = drop(engagement, "big.pdf", "Form 1098 Mortgage Interest Statement 2025")
    monkeypatch.setattr("tracker.filer.is_cloud_placeholder", lambda p: True)
    report = file_drops(engagement, today=DAY2)

    assert report.waiting == [waiting] and report.index_deferred is False
    assert not (engagement / INDEX_PENDING_FILENAME).exists()
    assert (engagement / INDEX_FILENAME).exists()
    assert [e.original_name for e in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]


def test_a_failure_after_the_move_is_recorded_and_the_rest_still_filed(engagement, monkeypatch):
    import shutil

    drop(engagement, "a-w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "b-mortgage.pdf", "Form 1098 Mortgage Interest Statement 2025")
    drop(engagement, "c-w2.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")
    real_copy = shutil.copy2

    def disk_full(src, dst, *args, **kwargs):
        if "b-mortgage" in str(src):
            raise OSError(28, "No space left on device")
        return real_copy(src, dst, *args, **kwargs)

    monkeypatch.setattr(shutil, "copy2", disk_full)
    report = file_drops(engagement, today=DAY1)

    assert [e.original_name for e in report.filed] == ["a-w2.pdf", "c-w2.pdf"]
    assert [e.name for e in report.errors] == ["b-mortgage.pdf"]
    assert report.errors[0].left_in_place is False
    # Every original is in PBC, and every one of them is in the index.
    assert {p.name for p in pbc(engagement).iterdir()} == {
        "a-w2.pdf", "b-mortgage.pdf", "c-w2.pdf",
    }
    rows = {e.original_name: e for e in read_index(engagement / INDEX_FILENAME)}
    assert set(rows) == {"a-w2.pdf", "b-mortgage.pdf", "c-w2.pdf"}
    assert rows["b-mortgage.pdf"].decision == NEEDS_REVIEW
    assert "No space left" in rows["b-mortgage.pdf"].reason
    assert f"{PBC_DIR_NAME}/b-mortgage.pdf" in rows["b-mortgage.pdf"].reason


def test_a_drop_still_held_open_is_left_for_the_next_run(engagement, monkeypatch):
    import os

    drop(engagement, "a-w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "b-mortgage.pdf", "Form 1098 Mortgage Interest Statement 2025")
    real_rename = os.rename

    def held_open(src, dst, *args, **kwargs):
        # What Windows does to a rename while another program has the file
        # open: refuses it. A copy of the same file would be allowed, and
        # is exactly what must not happen (the ninth reading found
        # shutil.move copying a half-written drop into PBC on this refusal).
        if "a-w2" in str(src):
            raise PermissionError("[WinError 32] used by another process")
        return real_rename(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "rename", held_open)
    report = file_drops(engagement, today=DAY1)

    assert [e.original_name for e in report.filed] == ["b-mortgage.pdf"]
    assert [e.name for e in report.errors] == ["a-w2.pdf"]
    assert report.errors[0].left_in_place is True
    assert (engagement / SHARED_DIR_NAME / "a-w2.pdf").exists()  # untouched
    assert not any("a-w2" in p.name for p in (engagement / SHARED_DIR_NAME / PBC_DIR_NAME).iterdir())
    assert {e.original_name for e in read_index(engagement / INDEX_FILENAME)} == {
        "b-mortgage.pdf",
    }

    # Released: the next run sorts it as if nothing had happened.
    monkeypatch.undo()
    report = file_drops(engagement, today=DAY2)
    assert [e.original_name for e in report.filed] == ["a-w2.pdf"]
    assert report.errors == []


# ------------------------------------------------------------ one run at a time ----


def test_the_filer_reads_nothing_before_it_holds_the_lock(engagement, monkeypatch):
    # What a run decides from - the manifest, the index, the drop folder -
    # must be read after the lock, or a run that finished in between is
    # invisible and its rows get rewritten from a stale picture.
    import tracker.filer as filer_module
    from tracker.locking import LOCK_FILENAME

    lock = engagement / LOCK_FILENAME
    seen = []
    for name in ("load_manifest", "read_index", "iter_drops"):
        real = getattr(filer_module, name)

        def under_the_lock(*args, _real=real, _name=name, **kwargs):
            seen.append((_name, lock.exists()))
            return _real(*args, **kwargs)

        monkeypatch.setattr(filer_module, name, under_the_lock)

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    assert file_drops(engagement, today=DAY1).handled == 1
    assert seen and all(held for _, held in seen), seen
    assert not lock.exists()

    seen.clear()                                  # a dry run reads without a lock
    file_drops(engagement, today=DAY1, dry_run=True)
    assert seen and not any(held for _, held in seen)


def test_a_persons_keyword_is_learned_under_the_lock(engagement, monkeypatch):
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file
    from tracker.locking import LOCK_FILENAME

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    real = filer_module.add_any_keyword

    def under_the_lock(*args, **kwargs):
        assert (engagement / LOCK_FILENAME).exists(), "keyword written outside the lock"
        return real(*args, **kwargs)

    monkeypatch.setattr(filer_module, "add_any_keyword", under_the_lock)
    result = assign_review_file(engagement, parked.pbc_location, "C01", keyword="lender")
    assert result.keyword == "lender"
    assert not (engagement / LOCK_FILENAME).exists()


def test_the_filer_holds_the_engagement_lock(engagement):
    # A second run mid-way used to move the remaining files and overwrite
    # the first run's index rows. Same lock as the scanner, same answer.
    from tracker.locking import LOCK_FILENAME, EngagementLockedError

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    (engagement / LOCK_FILENAME).write_text(f"pid={os.getpid()}", encoding="utf-8")
    with pytest.raises(EngagementLockedError):
        file_drops(engagement, today=DAY1)
    assert (engagement / SHARED_DIR_NAME / "w2.pdf").exists()   # nothing moved
    # A dry run writes nothing, so it needs no lock.
    assert file_drops(engagement, today=DAY1, dry_run=True).handled == 1
    (engagement / LOCK_FILENAME).unlink()
    assert file_drops(engagement, today=DAY1).handled == 1
    assert not (engagement / LOCK_FILENAME).exists()


def test_a_document_renamed_in_excel_keeps_filing_into_its_existing_folder(engagement):
    drop(engagement, "john.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    wb = load_workbook(engagement / MANIFEST_FILENAME)
    wb[SHEET_NAME].cell(row=2, column=col(COL_DOCUMENT), value="W-2s (all employers)")
    wb.save(engagement / MANIFEST_FILENAME)
    wb.close()
    drop(engagement, "jane.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")
    report = file_drops(engagement, today=DAY2)
    folders = [p.name for p in (engagement / PREPARED_DIR_NAME).iterdir() if p.name.startswith("A01")]
    assert folders == ["A01 - W-2 Wage Statements"]
    assert report.filed[0].prepared_location.startswith(f"{PREPARED_DIR_NAME}/A01 - W-2 Wage Statements/")


def test_a_file_dropped_straight_into_pbc_is_filed_and_indexed(engagement):
    # The client can see PBC/ and was told to drop things anywhere.
    original = text_pdf(pbc(engagement) / "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    before = original.read_bytes()
    report = file_drops(engagement, today=DAY1)
    assert [e.original_name for e in report.filed] == ["w2.pdf"]
    assert original.read_bytes() == before                    # not moved, not touched
    rows = read_index(engagement / INDEX_FILENAME)
    assert rows[0].pbc_location == f"{SHARED_DIR_NAME}/{PBC_DIR_NAME}/w2.pdf"
    assert (engagement / rows[0].prepared_location).exists()
    # Recorded now, so the next run leaves it alone.
    assert file_drops(engagement, today=DAY2).handled == 0


def test_a_resend_is_refiled_when_the_working_copy_was_deleted(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    first = file_drops(engagement, today=DAY1).filed[0]
    working = engagement / first.prepared_location
    working.unlink()                                          # a preparer's slip
    drop(engagement, "w2 again.pdf", "Form W-2 Wage and Tax Statement 2025")
    report = file_drops(engagement, today=DAY2)
    assert report.duplicates == []
    assert [e.original_name for e in report.filed] == ["w2 again.pdf"]
    assert "re-filed" in report.filed[0].reason
    assert working.exists()                                   # same canonical name again
    # A re-send whose working copy is still there is still just a duplicate.
    drop(engagement, "w2 third time.pdf", "Form W-2 Wage and Tax Statement 2025")
    assert [e.decision for e in file_drops(engagement, today=DAY2).duplicates] == [DUPLICATE]
    # ...and one more drop after the working copy goes again: the Duplicate
    # row must not shadow the Filed row, or this re-send is never re-filed.
    working.unlink()
    drop(engagement, "w2 fourth time.pdf", "Form W-2 Wage and Tax Statement 2025")
    fourth = file_drops(engagement, today=DAY2)
    assert fourth.duplicates == [] and [e.original_name for e in fourth.filed] == ["w2 fourth time.pdf"]
    assert working.exists()


def test_empty_client_folders_are_cleared_after_sorting(engagement):
    nested = engagement / SHARED_DIR_NAME / "from my phone" / "scans"
    nested.mkdir(parents=True)
    text_pdf(nested / "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    (engagement / SHARED_DIR_NAME / "keep" ).mkdir()
    (engagement / SHARED_DIR_NAME / "keep" / "later.txt.tmp").write_text("x")  # still uploading
    file_drops(engagement, today=DAY1)
    left = sorted(p.name for p in (engagement / SHARED_DIR_NAME).iterdir() if p.is_dir())
    assert left == [PBC_DIR_NAME, "keep"]


def test_a_hand_edited_index_cell_does_not_stop_the_next_run(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    wb = load_workbook(engagement / INDEX_FILENAME)
    size_col = INDEX_COLUMNS.index(INDEX_LAYOUT["size_kb"][0]) + 1
    wb.active.cell(row=2, column=size_col, value="about 9 KB")
    wb.save(engagement / INDEX_FILENAME)
    wb.close()
    drop(engagement, "jane.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")
    assert [e.original_name for e in file_drops(engagement, today=DAY2).filed] == ["jane.pdf"]
    assert read_index(engagement / INDEX_FILENAME)[0].size_kb == 0.0


# ------------------------------------------------- a person files a parked file ----


def test_assigning_a_parked_file_moves_it_under_the_canonical_name(engagement):
    from tracker.filer import ASSIGNED_BY_PERSON, assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    review_copy = engagement / parked.prepared_location
    assert review_copy.exists()

    result = assign_review_file(engagement, parked.pbc_location, "C01", keyword="Home Lending", today=DAY2)

    assert result.moved_review_copy is True
    assert not review_copy.exists()
    working = engagement / result.entry.prepared_location
    assert working.name == "C01 - Mortgage Interest Statement - TY2025.pdf"
    assert working.read_bytes() == (engagement / parked.pbc_location).read_bytes()
    assert (engagement / parked.pbc_location).exists()               # original untouched
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and row.identifier == "C01"
    assert row.reason.startswith(f"{ASSIGNED_BY_PERSON} on {DAY2.isoformat()}; was: ")
    assert result.keyword == "Home Lending" and result.keyword_note == ""
    from tracker.manifest import load_manifest
    assert next(i for i in load_manifest(engagement / MANIFEST_FILENAME) if i.identifier == "C01").any_keywords == ("Home Lending",)


def test_assigning_copies_from_pbc_when_the_review_copy_is_gone(engagement):
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    (engagement / parked.prepared_location).unlink()
    result = assign_review_file(engagement, "scan0012.pdf", "A01")   # by original name
    assert result.moved_review_copy is False
    assert (engagement / result.entry.prepared_location).exists()


def test_assigning_refuses_what_a_person_should_not_do(engagement):
    from tracker.filer import FilingError, assign_review_file
    from tracker.manifest import Override, RequestItem, create_template

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    with pytest.raises(FilingError, match="no request 'Z99'"):
        assign_review_file(engagement, "scan0012.pdf", "Z99")
    with pytest.raises(FilingError, match="is not waiting for review"):
        assign_review_file(engagement, "w2.pdf", "C01")               # already filed
    with pytest.raises(FilingError, match="nothing in the index is called"):
        assign_review_file(engagement, "ghost.pdf", "C01")

    waived = engagement.parent / "Waived"
    waived.mkdir()
    create_template(waived / MANIFEST_FILENAME, [
        RequestItem(identifier="A01", document="W-2", manual_override=Override.WAIVED)])
    scaffold_engagement(waived)
    drop(waived, "x.pdf", "nothing")
    file_drops(waived, today=DAY1)
    with pytest.raises(FilingError, match="waived"):
        assign_review_file(waived, "x.pdf", "A01")


def test_assigning_still_files_when_excel_holds_the_manifest(engagement, monkeypatch):
    import tracker.manifest as manifest_module
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    file_drops(engagement, today=DAY1)

    def locked(*args, **kwargs):
        raise PermissionError("[Errno 13] locked by Excel")

    monkeypatch.setattr(manifest_module, "add_any_keyword", locked)
    monkeypatch.setattr("tracker.filer.add_any_keyword", locked)
    result = assign_review_file(engagement, "scan0012.pdf", "C01", keyword="lender")
    assert result.entry.decision == FILED
    assert result.keyword == "" and "open in Excel" in result.keyword_note


# ---------------------------------------------- what a review found (decision 54) ----


def _age(path, seconds):
    """Push a file's mtime into the past so a later write is unmistakably newer."""
    old = path.stat().st_mtime - seconds
    os.utime(path, (old, old))


def test_a_snapshot_the_workbook_has_outlived_does_not_hide_the_workbooks_rows(engagement, monkeypatch):
    # The write landed but the sidecar could not be deleted (something held
    # it open). The workbook is the newer of the two; a later run must not
    # read the stale snapshot, re-route the original the workbook alone
    # knows, and copy it a second time.
    import tracker.filer as filer_module

    drop(engagement, "w2-a.pdf", "Form W-2 Wage and Tax Statement 2025 alpha")
    _lock_the_index(monkeypatch)
    file_drops(engagement, today=DAY1)                    # snapshot: [w2-a]
    monkeypatch.undo()
    sidecar = engagement / INDEX_PENDING_FILENAME
    assert sidecar.exists()

    def held(path):
        return False                                       # unlink refused
    monkeypatch.setattr(filer_module, "_discard_pending_index", held)
    drop(engagement, "w2-b.pdf", "Form W-2 Wage and Tax Statement 2025 beta")
    _age(sidecar, 60)
    report = file_drops(engagement, today=DAY2)           # workbook: [w2-a, w2-b]
    assert report.index_deferred is False and sidecar.exists()
    monkeypatch.undo()

    rows = read_index(engagement / INDEX_FILENAME)
    assert [r.original_name for r in rows] == ["w2-a.pdf", "w2-b.pdf"]
    assert file_drops(engagement, today=DAY2).handled == 0    # nothing re-sorted
    assert len(list(prepared(engagement, "A01").iterdir())) == 2


def test_a_persons_decision_survives_excel_saving_the_stale_workbook(engagement, monkeypatch):
    # Excel saved the (older) workbook after the filing was deferred: the
    # workbook is newer, but its parked row is not the truth.
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    _lock_the_index(monkeypatch)
    assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    monkeypatch.undo()
    _age(engagement / INDEX_PENDING_FILENAME, 60)
    os.utime(engagement / INDEX_FILENAME, None)           # Excel's save, no change to the row
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and row.identifier == "C01"


def test_a_sidecar_that_cannot_be_read_stops_the_run_rather_than_being_guessed_past(engagement, monkeypatch):
    import tracker.filer as filer_module

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    _lock_the_index(monkeypatch)
    file_drops(engagement, today=DAY1)
    monkeypatch.undo()
    sidecar = engagement / INDEX_PENDING_FILENAME
    real = filer_module.Path.read_text

    def unreadable(self, *args, **kwargs):
        if self == sidecar:
            raise OSError("[Errno 5] Input/output error")
        return real(self, *args, **kwargs)
    monkeypatch.setattr(filer_module.Path, "read_text", unreadable)
    with pytest.raises(OSError, match="could not read"):
        read_index(engagement / INDEX_FILENAME)
    monkeypatch.undo()
    assert sidecar.exists()                               # not quarantined


def test_an_original_replaced_under_its_own_name_is_said_out_loud_every_run(engagement):
    from tracker.filer import REPLACED_IN_PBC

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    first = file_drops(engagement, today=DAY1).filed[0]
    text_pdf(pbc(engagement) / "w2.pdf", "Form W-2 Wage and Tax Statement 2025 CORRECTED")
    for _ in range(2):
        report = file_drops(engagement, today=DAY2)
        assert report.handled == 0 and report.errors == []          # nothing was left unsorted
        [error] = report.attention
        assert error.name == "w2.pdf" and error.left_in_place
        assert error.error == REPLACED_IN_PBC.format(
            location=first.pbc_location, received=DAY1.isoformat(), prepared=first.prepared_location)
    assert len(read_index(engagement / INDEX_FILENAME)) == 1


def test_a_run_killed_after_copying_does_not_leave_a_second_copy_behind(engagement, monkeypatch):
    # The working copy was made, the process died before the index recorded
    # it. The next run finds the original unrecorded and reuses the copy.
    import tracker.filer as filer_module

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "scan0012.pdf", "nothing the rules recognise")

    def killed(*args, **kwargs):
        raise KeyboardInterrupt
    monkeypatch.setattr(filer_module, "write_index", killed)
    with pytest.raises(KeyboardInterrupt):
        file_drops(engagement, today=DAY1)
    monkeypatch.undo()
    assert not (engagement / INDEX_FILENAME).exists()

    report = file_drops(engagement, today=DAY2)
    assert len(report.filed) == 1 and len(report.review) == 1
    assert [p.name for p in prepared(engagement, "A01").iterdir()] == [report.filed[0].filed_as]
    assert [p.name for p in prepared(engagement, REVIEW_DIR_NAME).iterdir()] == ["scan0012.pdf"]


def test_a_failed_index_write_puts_the_filed_copy_back_where_the_index_says(engagement, monkeypatch):
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]

    def disk_full(*args, **kwargs):
        raise OSError(28, "No space left on device")
    # A full disk takes the workbook and the snapshot alike (the tenth
    # reading made every write failure retry and then snapshot, as Excel's
    # lock always did); with nowhere to write, the error is the caller's.
    monkeypatch.setattr(filer_module, "_save_index", disk_full)
    monkeypatch.setattr(filer_module, "_save_pending_index", disk_full)
    monkeypatch.setattr(filer_module.time, "sleep", lambda _: None)
    with pytest.raises(OSError):
        assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    monkeypatch.undo()
    assert (engagement / parked.prepared_location).exists()
    c01 = engagement / PREPARED_DIR_NAME / "C01 - Mortgage Interest Statement"
    assert not c01.exists() or not any(c01.iterdir())

    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    assert result.moved_review_copy is True
    assert [p.name for p in prepared(engagement, "C01").iterdir()] == [result.entry.filed_as]


def test_a_file_still_being_written_is_reported_as_waiting_not_passed_over(engagement):
    from tracker.manifest import TEMP_SUFFIX

    half = engagement / SHARED_DIR_NAME / f"upload{TEMP_SUFFIX}"
    half.write_bytes(b"%PDF-1.4 partial")
    report = file_drops(engagement, today=DAY1)
    assert report.waiting == [half] and report.handled == 0
    assert half.exists()


# ------------------------------------------- the second reading (decision 56) ----


def test_a_same_size_replacement_with_an_old_date_is_still_noticed(engagement):
    # One number changed in a CSV, copied in with its original timestamp:
    # same size, older mtime. Only the bytes can tell.
    from tracker.filer import REPLACED_IN_PBC

    (engagement / SHARED_DIR_NAME / "ledger.csv").write_text("a,1\nb,2\n", encoding="utf-8")
    first = file_drops(engagement, today=DAY1).review[0]
    original = pbc(engagement) / "ledger.csv"
    stamp = original.stat().st_mtime - 30 * 86400
    original.write_text("a,1\nb,3\n", encoding="utf-8")          # same length
    os.utime(original, (stamp, stamp))                            # older than its row
    [error] = file_drops(engagement, today=DAY2).attention
    assert error.error == REPLACED_IN_PBC.format(
        location=first.pbc_location, received=DAY1.isoformat(), prepared=first.prepared_location)


def test_an_original_the_sync_client_dehydrated_is_not_downloaded_to_be_checked(engagement, monkeypatch):
    import tracker.filer as filer_module

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    monkeypatch.setattr(filer_module, "is_cloud_placeholder", lambda p: p.name == "w2.pdf")

    def never(path):
        raise AssertionError(f"hashed {path.name}")
    monkeypatch.setattr(filer_module, "sha256_of", never)
    assert file_drops(engagement, today=DAY2).errors == []


def test_a_copy_that_fails_half_way_leaves_no_truncated_working_copy(engagement, monkeypatch):
    import tracker.filer as filer_module

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")

    def half(src, dst):
        filer_module.Path(dst).write_bytes(b"%PDF-1.4 half")
        raise OSError(28, "No space left on device")
    monkeypatch.setattr(filer_module.shutil, "copy2", half)
    report = file_drops(engagement, today=DAY1)
    monkeypatch.undo()
    assert len(report.review) == 1 and "No space left" in report.review[0].reason
    assert not any(prepared(engagement, "A01").iterdir())          # nothing half-written
    assert (pbc(engagement) / "w2.pdf").exists()                    # the record is safe


def test_an_interrupt_after_the_index_landed_does_not_undo_the_filing(engagement, monkeypatch):
    # write_index wrote the workbook; the interrupt hit while the sidecar
    # was being removed. The index says Filed at the request folder, so
    # the copy stays there - moving it back would leave the index lying.
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]

    def interrupted(path):
        raise KeyboardInterrupt
    monkeypatch.setattr(filer_module, "_discard_pending_index", interrupted)
    with pytest.raises(KeyboardInterrupt):
        assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    monkeypatch.undo()
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and (engagement / row.prepared_location).exists()
    assert not (engagement / parked.prepared_location).exists()


def test_every_unfinished_transfer_name_is_reported_as_waiting(engagement):
    from tracker.validators import UNFINISHED_SUFFIXES

    for suffix in UNFINISHED_SUFFIXES:
        (engagement / SHARED_DIR_NAME / f"upload{suffix}").write_bytes(b"partial")
    report = file_drops(engagement, today=DAY1)
    assert sorted(p.name for p in report.waiting) == sorted(f"upload{s}" for s in UNFINISHED_SUFFIXES)



def test_assigning_reuses_a_copy_an_earlier_attempt_left_and_leaves_no_half_copy(engagement, monkeypatch):
    # The same two guarantees the sort path has (decisions 54, 56), on the
    # path a person drives: a kill after the move made a copy the index
    # never learned of, and a copy that fails half-way leaves nothing.
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    c01 = engagement / PREPARED_DIR_NAME / "C01 - Mortgage Interest Statement"
    earlier = c01 / "C01 - Mortgage Interest Statement - TY2025.pdf"
    (engagement / parked.prepared_location).rename(earlier)      # the killed attempt's move
    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    assert [p.name for p in c01.iterdir()] == [earlier.name] and result.entry.filed_as == earlier.name

    drop(engagement, "scan0013.pdf", "nothing the rules recognise either")
    parked = file_drops(engagement, today=DAY2).review[0]
    (engagement / parked.prepared_location).unlink()             # a person removed the parked copy

    def half(src, dst):
        filer_module.Path(dst).write_bytes(b"%PDF-1.4 half")
        raise OSError(28, "No space left on device")
    monkeypatch.setattr(filer_module.shutil, "copy2", half)
    with pytest.raises(OSError):
        assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    monkeypatch.undo()
    assert [p.name for p in c01.iterdir()] == [earlier.name]     # no half copy beside it


def test_a_persons_filing_of_an_original_recorded_without_its_bytes_records_them(engagement, monkeypatch):
    # Decision 65 records an original the pass could not read back with no
    # digest. The ninth reading found that row carried through a person's
    # filing with no digest still, so the scanner never honoured the person
    # and the client was asked again. Assign records the bytes; and a later
    # pass that can read the original records them too.
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file
    from tracker.validators import sha256_of

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    real = filer_module.sha256_of

    def unreadable_once(path):
        if path.parent == pbc(engagement):
            raise PermissionError("held by the sync client")
        return real(path)
    monkeypatch.setattr(filer_module, "sha256_of", unreadable_once)
    parked = file_drops(engagement, today=DAY1).review[0]
    monkeypatch.undo()
    assert parked.digest == "" and parked.size_kb == 0.0

    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    original = engagement / parked.pbc_location
    assert result.entry.digest == sha256_of(original) and result.entry.size_kb > 0
    (row,) = read_index(engagement / INDEX_FILENAME)
    assert row.digest == result.entry.digest

    # And a pass, with nothing to sort, fills in the row the earlier pass could not.
    drop(engagement, "scan0013.pdf", "nothing the rules recognise either")
    monkeypatch.setattr(filer_module, "sha256_of", unreadable_once)
    file_drops(engagement, today=DAY2)
    monkeypatch.undo()
    assert [r.digest for r in read_index(engagement / INDEX_FILENAME) if r.original_name == "scan0013.pdf"] == [""]
    file_drops(engagement, today=DAY2)
    later = next(r for r in read_index(engagement / INDEX_FILENAME) if r.original_name == "scan0013.pdf")
    assert later.digest == sha256_of(engagement / later.pbc_location) and later.size_kb > 0


def test_bytes_recorded_after_the_fact_are_tied_to_the_row_or_not_recorded(engagement, monkeypatch):
    # The tenth reading: the digest filled in on a later pass came from the
    # original as it was THEN, so a client's replacement was adopted into
    # the old row. The eleventh: the working copy alone is no better - its
    # name is the client's and a freed name is taken by the next drop
    # called the same. A row is tied to its bytes only while its copy and
    # its original still agree; where they do not, nothing is adopted and
    # the row is said out loud, every pass.
    import tracker.filer as filer_module
    from tracker.filer import UNTIED_IN_PBC, FilingError, assign_review_file
    from tracker.validators import sha256_of

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    drop(engagement, "scan0013.pdf", "nothing the rules recognise either")
    real = filer_module.sha256_of

    def unreadable(path):
        if path.parent == pbc(engagement):
            raise PermissionError("held by the sync client")
        return real(path)
    monkeypatch.setattr(filer_module, "sha256_of", unreadable)
    report = file_drops(engagement, today=DAY1)
    monkeypatch.undo()
    filed = report.filed[0]
    parked = {r.original_name: r for r in report.review}
    assert filed.digest == "" and all(r.digest == "" for r in parked.values())
    first_bytes = (engagement / filed.prepared_location).read_bytes()

    # The client replaces the W-2 under its own name; a reviewer's app re-saves scan0012's parked copy.
    text_pdf(pbc(engagement) / "w2.pdf", "Form W-2 Wage and Tax Statement 2025 corrected")
    text_pdf(engagement / parked["scan0012.pdf"].prepared_location, "nothing the rules recognise, annotated")

    report = file_drops(engagement, today=DAY2)
    rows = {r.original_name: r for r in read_index(engagement / INDEX_FILENAME)}
    assert rows["w2.pdf"].digest == "" and rows["scan0012.pdf"].digest == ""              # nothing adopted
    assert rows["scan0013.pdf"].digest == sha256_of(pbc(engagement) / "scan0013.pdf")   # tied: copy and original agree
    assert (engagement / filed.prepared_location).read_bytes() == first_bytes            # untouched
    said = sorted(e.name for e in report.attention if e.error.startswith(UNTIED_IN_PBC.split("{")[0]))
    assert said == ["scan0012.pdf", "w2.pdf"] and report.errors == []   # for a person, not a failed pass

    # A person cannot file either untied row under it; the tied one files and carries its bytes.
    with pytest.raises(FilingError, match="look at both files first"):
        assign_review_file(engagement, parked["scan0012.pdf"].pbc_location, "C01", today=DAY2)
    result = assign_review_file(engagement, parked["scan0013.pdf"].pbc_location, "C01", today=DAY2)
    assert result.entry.digest == rows["scan0013.pdf"].digest


def test_a_persons_filing_moves_the_parked_copy_only_while_it_holds_the_rows_bytes(engagement):
    # The eleventh reading: a parked name is the client's, so once row A's
    # copy is gone the next drop called the same parks under that very
    # path, and filing row A carried document B into the request folder
    # under A's canonical name.
    from tracker.filer import assign_review_file
    from tracker.validators import sha256_of

    drop(engagement, "Scan.pdf", "document A, nothing the rules recognise")
    row_a = file_drops(engagement, today=DAY1).review[0]
    (engagement / row_a.prepared_location).unlink()                  # a person took it away to read
    drop(engagement, "Scan.pdf", "document B, unrelated")
    row_b = file_drops(engagement, today=DAY2).review[0]
    assert row_b.prepared_location == row_a.prepared_location         # the freed name, taken

    result = assign_review_file(engagement, row_a.pbc_location, "C01", today=DAY2)
    filed = engagement / result.entry.prepared_location
    assert result.moved_review_copy is False
    assert sha256_of(filed) == row_a.digest                            # document A, copied from its original
    assert (engagement / row_b.prepared_location).exists()            # document B still waits for a person
    assert sha256_of(engagement / row_b.prepared_location) == row_b.digest


def test_an_index_whose_header_row_moved_is_read_from_it_and_one_with_none_is_refused(engagement):
    # The twelfth reading: a header row that was not row 1 (a title typed
    # above it; Excel sorting with "my data has headers" unchecked) read as
    # no rows at all, and the next pass rebuilt the workbook without its
    # history - a person's filing included. The header is found; an index
    # with rows and no header is refused before anything is moved.
    from openpyxl import load_workbook

    from tracker.filer import FilingError

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    index = engagement / INDEX_FILENAME
    wb = load_workbook(index)
    ws = wb[INDEX_SHEET]
    ws.insert_rows(1, amount=2)
    ws["A1"] = "Smith Family - index of everything received"
    wb.save(index)
    wb.close()
    assert [r.original_name for r in read_index(index)] == ["w2.pdf"]
    drop(engagement, "1098.pdf", "Form 1098 Mortgage Interest Statement 2025")
    assert file_drops(engagement, today=DAY2).handled == 1
    rows = read_index(index)
    assert [r.original_name for r in rows] == ["w2.pdf", "1098.pdf"] and rows[0].received == DAY1.isoformat()

    wb = load_workbook(index)
    wb[INDEX_SHEET].delete_rows(1)                               # the header itself, gone
    wb.save(index)
    wb.close()
    with pytest.raises(FilingError, match="no header row"):
        read_index(index)
    drop(engagement, "later.pdf", "Form W-2 Wage and Tax Statement 2025 later")
    with pytest.raises(FilingError, match="no header row"):
        file_drops(engagement, today=DAY2)
    assert (engagement / SHARED_DIR_NAME / "later.pdf").exists()   # nothing was moved


def test_a_parked_path_a_later_row_claims_means_the_earlier_copy_is_gone(engagement, monkeypatch):
    # The twelfth reading: a digest-less row A whose parked copy a person
    # removed, then a same-named drop B parked under the freed name, was
    # "untied" every pass and refused for ever - though the index itself
    # says whose the file at that path is.
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file
    from tracker.validators import sha256_of

    drop(engagement, "Scan.pdf", "document A, nothing the rules recognise")
    real = filer_module.sha256_of

    def unreadable(path):
        if path.parent == pbc(engagement):
            raise PermissionError("held by the sync client")
        return real(path)
    monkeypatch.setattr(filer_module, "sha256_of", unreadable)
    row_a = file_drops(engagement, today=DAY1).review[0]
    monkeypatch.undo()
    assert row_a.digest == ""
    (engagement / row_a.prepared_location).unlink()
    drop(engagement, "Scan.pdf", "document B, unrelated")
    report = file_drops(engagement, today=DAY2)
    row_b = report.review[0]
    assert row_b.prepared_location == row_a.prepared_location
    assert report.attention == [] and report.errors == []          # not untied: its copy is gone

    result = assign_review_file(engagement, row_a.pbc_location, "C01", today=DAY2)
    assert result.entry.digest == sha256_of(pbc(engagement) / "Scan.pdf")   # document A, from its original
    assert sha256_of(engagement / result.entry.prepared_location) == result.entry.digest
    assert (engagement / row_b.prepared_location).exists()          # document B still waits


def test_an_annotated_parked_copy_is_left_behind_and_said_so(engagement):
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    text_pdf(engagement / parked.prepared_location, "nothing the rules recognise, with a reviewer's note")
    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    assert result.moved_review_copy is False and "left there" in result.left_in_review
    assert (engagement / parked.prepared_location).exists()
    assert (engagement / result.entry.prepared_location).read_bytes() == (pbc(engagement) / "scan0012.pdf").read_bytes()


def test_a_column_inserted_into_the_index_does_not_read_as_no_rows(engagement):
    # The eleventh reading: the blank-row test looked at the first physical
    # cell, so one column a person inserted at A read every row as blank,
    # and the next write dropped them all - a person's filing included.
    from openpyxl import load_workbook

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    wb = load_workbook(engagement / INDEX_FILENAME)
    wb[INDEX_SHEET].insert_cols(1)
    wb[INDEX_SHEET]["A1"] = "Checked"
    wb.save(engagement / INDEX_FILENAME)
    wb.close()

    assert [r.original_name for r in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]
    report = file_drops(engagement, today=DAY2)
    assert report.handled == 0
    assert [r.original_name for r in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]


def test_the_unlistable_walk_covers_the_same_ground_as_the_others(engagement, tmp_path):
    from tests.samples import listing_denied
    from tracker.filer import unlistable_folders

    shared = engagement / SHARED_DIR_NAME
    inside_pbc = shared / PBC_DIR_NAME / "inside"
    inside_pbc.mkdir(parents=True)
    staging = shared / ".tmp.driveupload"
    staging.mkdir()
    assert unlistable_folders(shared) == []
    with listing_denied(inside_pbc):                                # a client drops into PBC too (the twelfth reading)
        assert unlistable_folders(shared) == [inside_pbc]
    if sys.platform == "win32":
        import _winapi

        outside = tmp_path / "outside" / "deep"
        outside.mkdir(parents=True)
        _winapi.CreateJunction(str(tmp_path / "outside"), str(shared / "link"))
        assert unlistable_folders(shared) == []                       # never walks behind the junction


def test_a_folder_the_run_cannot_list_is_reported_every_pass(engagement):
    from tests.samples import listing_denied
    from tracker.filer import unlistable_folders

    denied = engagement / SHARED_DIR_NAME / "from my accountant"
    denied.mkdir()
    text_pdf(denied / "inside-w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "readable-w2.pdf", "Form W-2 Wage and Tax Statement 2025")

    with listing_denied(denied):
        assert unlistable_folders(engagement / SHARED_DIR_NAME) == [denied]
        report = file_drops(engagement, today=DAY1)
    assert [e.original_name for e in report.filed] == ["readable-w2.pdf"]
    assert [(e.name, e.left_in_place) for e in report.errors] == [(denied.name, True)]
    assert "cannot list" in report.errors[0].error


def test_a_cloud_placeholder_is_not_a_link(tmp_path):
    # A sync client's placeholder is a reparse point with the client's own
    # tag; only a mount point (junction) or a symlink is a link.
    import stat

    from tracker.filer import _LINK_TAGS, _is_link

    plain = tmp_path / "w2.pdf"
    plain.write_bytes(b"%PDF-1.4")
    assert not _is_link(plain)
    assert getattr(stat, "IO_REPARSE_TAG_MOUNT_POINT", 0xA0000003) in _LINK_TAGS or sys.platform != "win32"
    assert 0x9000001A not in _LINK_TAGS            # OneDrive's tag
    try:
        (tmp_path / "link.pdf").symlink_to(plain)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks need a privilege here")
    assert _is_link(tmp_path / "link.pdf")


@pytest.mark.skipif(sys.platform != "win32", reason="NTFS takes a lone surrogate in a name")
def test_a_name_the_index_cannot_hold_is_reported_not_allowed_to_take_the_index_down(engagement):
    # The tenth reading: one such name made every index write fail, after
    # the pass's originals had been moved, and nothing was recorded again.
    bad = engagement / SHARED_DIR_NAME / "bank statement \ud83d.pdf"
    bad.write_bytes(b"%PDF-1.4 a name with half an emoji")
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    report = file_drops(engagement, today=DAY1)
    assert [e.original_name for e in report.filed] == ["w2.pdf"]
    assert [e.left_in_place for e in report.errors] == [True] and "rename it" in report.errors[0].error
    assert bad.exists() and [r.original_name for r in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]
    # Dropped straight into PBC, it is reported there too and never sorted as a stray.
    bad.rename(pbc(engagement) / bad.name)
    report = file_drops(engagement, today=DAY2)
    assert report.filed == [] and len(report.errors) == 1
    assert [r.original_name for r in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]


@pytest.mark.skipif(sys.platform != "win32", reason="names Windows cannot open")
def test_a_name_windows_cannot_open_is_reported_not_passed_over(engagement):
    shared = engagement / SHARED_DIR_NAME
    # A Mac or a sync client can deliver these; only the \\?\ form creates them here.
    for name in ("dotted.pdf.", "spaced.pdf "):
        with open("\\\\?\\" + str(shared / name), "wb") as handle:
            handle.write(b"%PDF-1.4 not openable by its plain name")
    report = file_drops(engagement, today=DAY1)
    assert sorted(e.name for e in report.errors) == ["dotted.pdf.", "spaced.pdf "]
    assert all(e.left_in_place and "rename it" in e.error for e in report.errors)
    assert report.filed == [] and report.review == []


@pytest.mark.skipif(sys.platform != "win32", reason="junctions")
def test_what_lies_behind_a_junction_is_not_a_drop(engagement, tmp_path):
    import _winapi

    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "theirs.pdf").write_bytes(b"%PDF-1.4 somebody else's file")
    _winapi.CreateJunction(str(outside), str(engagement / SHARED_DIR_NAME / "link"))
    drop(engagement, "scan0012.pdf", "Form W-2 Wage and Tax Statement 2025")

    report = file_drops(engagement, today=DAY1)

    assert [e.original_name for e in report.filed] == ["scan0012.pdf"]
    assert (outside / "theirs.pdf").exists()                     # never moved
    assert (engagement / SHARED_DIR_NAME / "link").exists()      # never removed as an "empty folder"
    assert [p.name for p in pbc(engagement).iterdir()] == ["scan0012.pdf"]


def test_a_file_named_like_a_formula_is_recorded_as_its_name(engagement):
    drop(engagement, "=SUM scan.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.original_name == "=SUM scan.pdf"


def test_assigning_a_replaced_original_is_refused_not_recorded_under_the_old_bytes(engagement):
    from tracker.filer import FilingError, assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    (engagement / parked.prepared_location).unlink()             # the parked copy is gone
    text_pdf(pbc(engagement) / "scan0012.pdf", "the client replaced it with something else")
    with pytest.raises(FilingError, match="replaced after it arrived"):
        assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW and row.digest == parked.digest


# ------------------------------------------- the thirteenth reading (d72) ----


def excel_sorted(index, header):
    """Exactly what Excel does with "my data has headers" unchecked: every
    row of the sheet, the header row among them, ordered by one column's
    text. Returns the Original Name column in the order it now reads."""
    wb = load_workbook(index)
    ws = wb[INDEX_SHEET]
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    at = [str(value or "") for value in rows[0]].index(header)
    rows.sort(key=lambda row: str(row[at] or "").lower())   # digits before letters, as in Excel
    ws.delete_rows(1, ws.max_row)
    for row in rows:
        ws.append(row)
    wb.save(index)
    wb.close()
    return [str(row[1] or "") for row in rows]


def test_a_sheet_sorted_until_the_header_is_last_keeps_every_row_above_it(engagement):
    # Decision 69 found the header row but read only what was below it. A
    # sort on Received puts the header last - every date sorts before the
    # word - so the whole index read as no rows, and the next pass rebuilt
    # it, re-adopting each original as a stray and re-dating the trail.
    drop(engagement, "w2 employer a.pdf", "Form W-2 Wage and Tax Statement 2025 A")
    drop(engagement, "1098 from the bank.pdf", "Form 1098 Mortgage Interest Statement 2025")
    file_drops(engagement, today=DAY1)
    index = engagement / INDEX_FILENAME

    order = excel_sorted(index, INDEX_LAYOUT["received"][0])
    assert order[-1] == INDEX_LAYOUT["original_name"][0]         # the header sorted to the bottom

    rows = read_index(index)
    assert sorted(r.original_name for r in rows) == ["1098 from the bank.pdf", "w2 employer a.pdf"]
    drop(engagement, "w2 employer b.pdf", "Form W-2 Wage and Tax Statement 2025 B")
    report = file_drops(engagement, today=DAY2)
    assert report.handled == 1 and report.duplicates == []      # nothing was adopted a second time
    rows = read_index(index)
    assert sorted((r.original_name, r.received) for r in rows) == [
        ("1098 from the bank.pdf", DAY1.isoformat()),
        ("w2 employer a.pdf", DAY1.isoformat()),
        ("w2 employer b.pdf", DAY2.isoformat()),
    ]


def test_a_persons_filing_above_a_sorted_header_is_not_undone(engagement):
    # The same sort on Original Name leaves the header in the middle, so
    # only the rows above it were dropped - and with them the row a person
    # filed by hand, whose working copy then satisfied a request that no
    # index row named.
    from tracker.filer import ASSIGNED_BY_PERSON, assign_review_file

    drop(engagement, "a bank letter.pdf", "nothing the rules recognise")
    drop(engagement, "b w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "zz 1098.pdf", "Form 1098 Mortgage Interest Statement 2025")
    parked = file_drops(engagement, today=DAY1).review[0]
    assign_review_file(engagement, parked.pbc_location, "C01", today=DAY1)
    index = engagement / INDEX_FILENAME

    order = excel_sorted(index, INDEX_LAYOUT["original_name"][0])
    assert order.index(INDEX_LAYOUT["original_name"][0]) == 2    # the header sorted into the middle

    assert [r.original_name for r in read_index(index)] == [
        "a bank letter.pdf", "b w2.pdf", "zz 1098.pdf",
    ]
    file_drops(engagement, today=DAY2)
    rows = read_index(index)
    (person,) = [r for r in rows if r.reason.startswith(ASSIGNED_BY_PERSON)]
    assert person.original_name == "a bank letter.pdf" and person.decision == FILED
    named = {r.prepared_location for r in rows}
    assert all(
        p.relative_to(engagement).as_posix() in named
        for p in (engagement / PREPARED_DIR_NAME).rglob("*") if p.is_file()
    )                                                            # no working copy is an orphan


def test_the_index_is_read_from_the_sheet_carrying_its_header_not_the_selected_one(engagement):
    # `wb.active` is whatever tab a person left selected. One who renamed
    # the index's tab and added a tab of their own got an index that read
    # as no rows - decision 69's refusal never fired, because an empty
    # sheet holds nothing - and the next write rebuilt the workbook over
    # both sheets and the whole history.
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    index = engagement / INDEX_FILENAME
    wb = load_workbook(index)
    wb[INDEX_SHEET].title = "Index 2025"
    wb.create_sheet("Notes")["A1"] = "ask about the second W-2"
    wb.active = wb.sheetnames.index("Notes")
    wb.save(index)
    wb.close()

    assert [r.original_name for r in read_index(index)] == ["w2.pdf"]
    drop(engagement, "1098.pdf", "Form 1098 Mortgage Interest Statement 2025")
    assert file_drops(engagement, today=DAY2).handled == 1
    wb = load_workbook(index)
    try:
        assert wb.sheetnames == ["Index 2025", "Notes"]           # the person's sheets are still here
        assert wb["Notes"]["A1"].value == "ask about the second W-2"
    finally:
        wb.close()
    assert [(r.original_name, r.received) for r in read_index(index)] == [
        ("w2.pdf", DAY1.isoformat()), ("1098.pdf", DAY2.isoformat()),
    ]


def test_a_workbook_with_rows_and_no_header_on_any_sheet_is_refused_by_name(engagement):
    from tracker.filer import FilingError

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    index = engagement / INDEX_FILENAME
    wb = load_workbook(index)
    wb[INDEX_SHEET].delete_rows(1)                               # the header itself, gone
    wb.create_sheet("Notes")["A1"] = "a tab of my own"
    wb.save(index)
    wb.close()
    with pytest.raises(FilingError) as refused:
        read_index(index)
    assert "no header row" in str(refused.value) and "Notes" in str(refused.value)


def test_an_original_deleted_from_pbc_is_said_every_pass(engagement):
    # PBC is the provided-by-client record and the client can see it, so
    # Explorer will delete what is already there. Every other disagreement
    # between the index and the disk was said every pass; this one was said
    # by nothing at all, and the scan went on calling the request Received.
    from tracker.filer import MISSING_IN_PBC

    drop(engagement, "w2 employer a.pdf", "Form W-2 Wage and Tax Statement 2025 A")
    drop(engagement, "w2 employer b.pdf", "Form W-2 Wage and Tax Statement 2025 B")
    filed = {e.original_name: e for e in file_drops(engagement, today=DAY1).filed}
    (pbc(engagement) / "w2 employer a.pdf").unlink()

    for day in (DAY1, DAY2):
        report = file_drops(engagement, today=day)
        assert report.errors == []                               # a finding, not a failed pass
        assert [(e.name, e.left_in_place) for e in report.attention] == [("w2 employer a.pdf", True)]
        assert report.attention[0].error == MISSING_IN_PBC.format(
            location=filed["w2 employer a.pdf"].pbc_location, received=DAY1.isoformat())
    rows = read_index(engagement / INDEX_FILENAME)
    assert [r.decision for r in rows] == [FILED, FILED]
    # The working copy is the scan's business and is left exactly as it was.
    assert (engagement / filed["w2 employer a.pdf"].prepared_location).exists()


def test_an_original_the_client_renamed_in_pbc_is_followed_not_filed_again(engagement):
    from tracker.filer import MOVED_IN_PBC

    drop(engagement, "w2 employer a.pdf", "Form W-2 Wage and Tax Statement 2025 A")
    filed = file_drops(engagement, today=DAY1).filed[0]
    (pbc(engagement) / "w2 employer a.pdf").rename(pbc(engagement) / "employer A W2 2025.pdf")

    report = file_drops(engagement, today=DAY2)
    now = f"{SHARED_DIR_NAME}/{PBC_DIR_NAME}/employer A W2 2025.pdf"
    assert report.duplicates == [] and report.handled == 0 and report.errors == []
    assert [e.error for e in report.attention] == [
        MOVED_IN_PBC.format(location=filed.pbc_location, now=now)
    ]
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.pbc_location == now                                # the row followed the bytes
    assert row.original_name == filed.original_name               # what it arrived as is the record
    assert filed.pbc_location in row.reason
    assert row.prepared_location == filed.prepared_location and row.digest == filed.digest
    assert file_drops(engagement, today=DAY2).attention == []     # said once; then it is the record


def test_an_original_the_client_moved_into_a_subfolder_of_pbc_is_followed(engagement):
    drop(engagement, "w2 employer a.pdf", "Form W-2 Wage and Tax Statement 2025 A")
    filed = file_drops(engagement, today=DAY1).filed[0]
    (pbc(engagement) / "2025").mkdir()
    (pbc(engagement) / "w2 employer a.pdf").rename(pbc(engagement) / "2025" / "w2 employer a.pdf")

    report = file_drops(engagement, today=DAY2)
    assert report.duplicates == [] and report.handled == 0 and len(report.attention) == 1
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.pbc_location == f"{SHARED_DIR_NAME}/{PBC_DIR_NAME}/2025/w2 employer a.pdf"
    assert row.decision == FILED and row.prepared_location == filed.prepared_location


def test_a_duplicate_rows_own_original_is_watched_like_every_other(engagement):
    from tracker.filer import MISSING_IN_PBC

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    drop(engagement, "w2 again.pdf", "Form W-2 Wage and Tax Statement 2025")
    [twice] = file_drops(engagement, today=DAY1).duplicates
    (pbc(engagement) / "w2 again.pdf").unlink()

    report = file_drops(engagement, today=DAY2)
    assert [e.error for e in report.attention] == [
        MISSING_IN_PBC.format(location=twice.pbc_location, received=DAY1.isoformat())
    ]
    assert [r.decision for r in read_index(engagement / INDEX_FILENAME)] == [FILED, DUPLICATE]


def test_a_row_recorded_without_its_bytes_is_never_moved_onto_another_file(engagement, monkeypatch):
    # Decision 65: an empty digest is nobody's. Nothing proves a stray is
    # this row's original, so the row is said and left where it is.
    import tracker.filer as filer_module
    from tracker.filer import MISSING_IN_PBC

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    real = filer_module.sha256_of

    def unreadable(path):
        if path.parent == pbc(engagement):
            raise PermissionError("held by the sync client")
        return real(path)
    monkeypatch.setattr(filer_module, "sha256_of", unreadable)
    row = file_drops(engagement, today=DAY1).filed[0]
    monkeypatch.undo()
    assert row.digest == ""
    (pbc(engagement) / "w2.pdf").rename(pbc(engagement) / "renamed w2.pdf")

    report = file_drops(engagement, today=DAY2)
    gone = [e for e in report.attention
            if e.error == MISSING_IN_PBC.format(location=row.pbc_location, received=DAY1.isoformat())]
    assert len(gone) == 1
    rows = read_index(engagement / INDEX_FILENAME)
    assert rows[0].pbc_location == row.pbc_location               # never relocated
    assert [r.original_name for r in rows] == ["w2.pdf", "renamed w2.pdf"]


def test_an_original_the_sync_client_dehydrated_is_not_called_gone(engagement, monkeypatch):
    # A placeholder is a file the sync client has not brought down, which
    # is not the same as the client having deleted it.
    import tracker.filer as filer_module

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    monkeypatch.setattr(filer_module, "is_cloud_placeholder", lambda p: p.parent == pbc(engagement))
    report = file_drops(engagement, today=DAY2)
    assert report.attention == [] and report.errors == []


def test_a_name_the_workbook_cannot_hold_is_refused_like_one_utf8_cannot(engagement):
    # A control character in a name: NTFS refuses it, POSIX takes it, and
    # openpyxl refuses the cell with a ValueError the index write does not
    # retry - after the pass's originals had moved, with neither workbook
    # nor snapshot written. Half of CI is Linux.
    from tracker.filer import _storable

    shared = engagement / SHARED_DIR_NAME
    bad = shared / "scan\x072025.pdf"
    assert not _storable(bad)
    assert not _storable(shared / "scan\x0b2025.pdf")
    assert _storable(shared / "scan 2025.pdf")
    try:
        bad.write_bytes(b"%PDF-1.4 a name with a bell in it")
    except (OSError, ValueError):
        return                       # this filesystem refuses the name; the guard is the claim
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    report = file_drops(engagement, today=DAY1)
    assert [e.original_name for e in report.filed] == ["w2.pdf"]
    assert [e.left_in_place for e in report.errors] == [True] and "rename it" in report.errors[0].error
    assert bad.exists()
    assert [r.original_name for r in read_index(engagement / INDEX_FILENAME)] == ["w2.pdf"]


def test_a_failed_index_write_puts_back_a_parked_copy_a_reused_one_stood_in_for(engagement, monkeypatch):
    # Assign reuses a killed attempt's copy and removes the parked one it
    # stands in for (decision 69). The rollback covered the moved copy and
    # the fresh copy but not that removal, so a killed index write left the
    # row still naming a parked copy that was no longer there.
    import tracker.filer as filer_module
    from tracker.filer import assign_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    c01 = engagement / PREPARED_DIR_NAME / "C01 - Mortgage Interest Statement"
    c01.mkdir(parents=True, exist_ok=True)
    attempt = c01 / "C01 - Mortgage Interest Statement - TY2025.pdf"
    attempt.write_bytes((engagement / parked.prepared_location).read_bytes())

    def disk_full(*args, **kwargs):
        raise OSError(28, "No space left on device")
    monkeypatch.setattr(filer_module, "_save_index", disk_full)
    monkeypatch.setattr(filer_module, "_save_pending_index", disk_full)
    monkeypatch.setattr(filer_module.time, "sleep", lambda _: None)
    with pytest.raises(OSError):
        assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    monkeypatch.undo()

    back = engagement / parked.prepared_location
    assert back.is_file() and back.read_bytes() == attempt.read_bytes()
    assert [p.name for p in c01.iterdir()] == [attempt.name]      # the reused copy stays where it is
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW and (engagement / row.prepared_location).is_file()

    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)
    assert result.entry.filed_as == attempt.name and not back.exists()


def test_an_empty_folder_is_cleared_by_a_pass_with_nothing_to_sort(engagement):
    # The tidy-up used to sit past an early return, so a pass that found
    # nothing to sort left the folder the client dragged in behind for ever.
    left_behind = engagement / SHARED_DIR_NAME / "from my phone"
    left_behind.mkdir()
    report = file_drops(engagement, today=DAY1)
    assert report.handled == 0 and not left_behind.exists()


# ------------------------------- a person says no request asks for it (d76) ----


def test_dismissing_a_parked_file_rewrites_its_row_and_moves_nothing(engagement):
    """Not Requested is a decision about the request list, not about the file.

    An agency notice the client sent is still the client's: the working copy
    stays parked and the original stays in PBC. What the row says changes,
    and what it said before is kept after it.
    """
    from tracker.filer import DISMISSED_BY_PERSON, NOT_REQUESTED, dismiss_review_file

    drop(engagement, "irs-notice.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    copy = engagement / parked.prepared_location
    original = engagement / parked.pbc_location
    assert copy.is_file()

    result = dismiss_review_file(engagement, parked.pbc_location, "an IRS notice", today=DAY2)

    assert result.index_deferred is False
    assert result.entry.decision == NOT_REQUESTED
    assert result.entry.reason == (
        f"{DISMISSED_BY_PERSON} on {DAY2.isoformat()} (an IRS notice); was: {parked.reason}"
    )
    assert copy.is_file() and copy.read_bytes() == original.read_bytes()
    assert original.is_file()
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NOT_REQUESTED
    assert row.prepared_location == parked.prepared_location


def test_the_same_document_sent_again_after_a_dismissal_is_a_duplicate(engagement):
    """Parking it a second time would put back the warning a person just cleared."""
    from tracker.filer import NOT_REQUESTED, dismiss_review_file

    drop(engagement, "notice.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    dismiss_review_file(engagement, parked.pbc_location, today=DAY2)

    drop(engagement, "notice again.pdf", "nothing the rules recognise")   # same bytes
    report = file_drops(engagement, today=DAY2)

    assert [e.original_name for e in report.duplicates] == ["notice again.pdf"]
    assert report.review == []
    review_dir = engagement / PREPARED_DIR_NAME / REVIEW_DIR_NAME
    assert [p.name for p in review_dir.iterdir()] == ["notice.pdf"]
    assert (pbc(engagement) / "notice again.pdf").exists()        # preserved, as always
    decisions = [r.decision for r in read_index(engagement / INDEX_FILENAME)]
    assert decisions == [NOT_REQUESTED, DUPLICATE]


def test_filing_a_dismissed_document_is_how_the_decision_is_undone(engagement):
    """A person who was wrong files it; there is no second undo path to build."""
    from tracker.filer import ASSIGNED_BY_PERSON, assign_review_file, dismiss_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    dismiss_review_file(engagement, parked.pbc_location, today=DAY2)

    result = assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)

    assert result.moved_review_copy is True
    assert result.entry.filed_as == "C01 - Mortgage Interest Statement - TY2025.pdf"
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and row.identifier == "C01"
    assert row.reason.startswith(f"{ASSIGNED_BY_PERSON} on {DAY2.isoformat()}; was: ")
    assert (engagement / row.prepared_location).is_file()


def test_a_dismissal_excel_kept_out_of_the_workbook_is_not_lost(engagement, monkeypatch):
    """The snapshot is the index while Excel holds it, and it keeps the decision.

    A dismissal is an edit to a row the workbook already holds, so the same
    trap as a person's filing: deferring only what is past the workbook's
    end would lose it, and a stale workbook saved from Excel would undo it.
    """
    from tracker.filer import NOT_REQUESTED, dismiss_review_file

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    sidecar = engagement / INDEX_PENDING_FILENAME

    _lock_the_index(monkeypatch)
    result = dismiss_review_file(engagement, parked.pbc_location, today=DAY2)
    assert result.index_deferred is True and sidecar.exists()
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NOT_REQUESTED

    monkeypatch.undo()                            # Excel closed - having saved
    _age(sidecar, 60)
    os.utime(engagement / INDEX_FILENAME, None)   # the workbook is newer, and stale
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NOT_REQUESTED

    assert file_drops(engagement, today=DAY2).index_deferred is False
    assert not sidecar.exists()
    [row] = read_index(engagement / INDEX_FILENAME)       # the workbook alone now
    assert row.decision == NOT_REQUESTED


def test_dismissing_something_that_is_not_parked_says_what_it_is(engagement):
    """A filed document is not a person's to dismiss without unfiling it first."""
    from tracker.filer import FilingError, dismiss_review_file

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    with pytest.raises(FilingError, match=f"is not waiting for review \\(it is {FILED}"):
        dismiss_review_file(engagement, "w2.pdf", today=DAY2)
    with pytest.raises(FilingError, match="nothing in the index is called"):
        dismiss_review_file(engagement, "ghost.pdf", today=DAY2)


def test_dismissing_takes_the_engagement_lock(engagement):
    """A decision is written to the index, so it queues behind a run like every other."""
    from tracker.filer import NOT_REQUESTED, dismiss_review_file
    from tracker.locking import LOCK_FILENAME, EngagementLockedError

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]

    (engagement / LOCK_FILENAME).write_text(f"pid={os.getpid()}", encoding="utf-8")
    with pytest.raises(EngagementLockedError):
        dismiss_review_file(engagement, parked.pbc_location, today=DAY2)
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW               # nothing was written

    (engagement / LOCK_FILENAME).unlink()
    result = dismiss_review_file(engagement, parked.pbc_location, today=DAY2)
    assert result.entry.decision == NOT_REQUESTED
    assert not (engagement / LOCK_FILENAME).exists()


# ----------------------- a filed document goes back for review (d77) ----


def review_dir(engagement):
    return engagement / PREPARED_DIR_NAME / REVIEW_DIR_NAME


def test_unfiling_puts_the_working_copy_back_under_the_clients_own_name(engagement):
    """The parked name is the client's, whichever direction the document travels."""
    from tracker.filer import UNFILED_BY_PERSON, unfile_document

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025")
    filed = file_drops(engagement, today=DAY1).filed[0]
    working = engagement / filed.prepared_location
    assert working.is_file()

    result = unfile_document(engagement, filed.pbc_location, "wrong client", today=DAY2)

    assert result.moved_working_copy is True and result.left_filed == ""
    assert not working.exists()
    parked = engagement / result.entry.prepared_location
    assert parked.name == "w2 john.pdf" and parked.parent == review_dir(engagement)
    assert parked.read_bytes() == (engagement / filed.pbc_location).read_bytes()
    assert (engagement / filed.pbc_location).is_file()          # the original is untouched
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW and row.identifier == ""
    assert row.reason == (
        f"{UNFILED_BY_PERSON} on {DAY2.isoformat()} (wrong client); was: {filed.reason}"
    )


def test_filing_an_unfiled_document_puts_it_under_the_canonical_name_again(engagement):
    """Refiling is unfiling and then filing; there is no third path to keep honest."""
    from tracker.filer import (
        ASSIGNED_BY_PERSON,
        UNFILED_BY_PERSON,
        assign_review_file,
        unfile_document,
    )

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025")
    filed = file_drops(engagement, today=DAY1).filed[0]
    unfile_document(engagement, filed.pbc_location, today=DAY2)

    result = assign_review_file(engagement, filed.pbc_location, "C01", today=DAY2)

    assert result.moved_review_copy is True
    assert result.entry.filed_as == "C01 - Mortgage Interest Statement - TY2025.pdf"
    assert not list(review_dir(engagement).iterdir())
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and row.identifier == "C01"
    assert row.reason.startswith(
        f"{ASSIGNED_BY_PERSON} on {DAY2.isoformat()}; was: {UNFILED_BY_PERSON}"
    )


def test_unfiling_something_that_is_not_filed_says_what_it_is(engagement):
    """Each of the other decisions needs a different answer, so none of them is guessed."""
    from tracker.filer import (
        NOT_REQUESTED,
        FilingError,
        dismiss_review_file,
        unfile_document,
    )

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    with pytest.raises(FilingError, match=rf"is not filed \(it is {NEEDS_REVIEW}\)"):
        unfile_document(engagement, parked.pbc_location, today=DAY2)

    dismiss_review_file(engagement, parked.pbc_location, today=DAY2)
    with pytest.raises(FilingError, match=rf"is not filed \(it is {NOT_REQUESTED}\)"):
        unfile_document(engagement, parked.pbc_location, today=DAY2)

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    drop(engagement, "w2 again.pdf", "Form W-2 Wage and Tax Statement 2025")
    duplicate = file_drops(engagement, today=DAY2).duplicates[0]
    with pytest.raises(FilingError, match=rf"is not filed \(it is {DUPLICATE}\)"):
        unfile_document(engagement, duplicate.pbc_location, today=DAY2)

    with pytest.raises(FilingError, match="nothing in the index is called"):
        unfile_document(engagement, "ghost.pdf", today=DAY2)


def test_a_working_copy_somebody_re_saved_is_left_where_it_is_and_said_so(engagement):
    """A reviewer's notes are work. Which of the two files the firm wants is not ours."""
    from tracker.filer import LEFT_FILED, unfile_document

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025")
    filed = file_drops(engagement, today=DAY1).filed[0]
    working = engagement / filed.prepared_location
    working.write_bytes(working.read_bytes() + b"% a reviewer's note\n")

    result = unfile_document(engagement, filed.pbc_location, today=DAY2)

    assert result.moved_working_copy is False
    assert working.is_file()                                   # the notes are not thrown away
    parked = engagement / result.entry.prepared_location
    assert parked.is_file()
    assert parked.read_bytes() == (engagement / filed.pbc_location).read_bytes()
    assert result.left_filed == LEFT_FILED.format(
        location=filed.prepared_location, parked=parked.name)


def test_after_unfiling_the_request_is_not_received_and_the_note_says_why(engagement):
    """The status is put back in the same breath; nothing waits for the next pass."""
    from tracker.filer import unfile_document
    from tracker.manifest import Status, load_manifest
    from tracker.scanner import REGRESSION_FILES_CHANGED, REGRESSION_NOTE, scan_engagement

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025 John")
    drop(engagement, "w2 jane.pdf", "Form W-2 Wage and Tax Statement 2025 Jane")
    filed = file_drops(engagement, today=DAY1).filed
    assert {e.identifier for e in filed} == {"A01"}            # the row expects two
    scan_engagement(engagement, today=DAY1)
    a01 = next(i for i in load_manifest(engagement / MANIFEST_FILENAME) if i.identifier == "A01")
    assert a01.status == Status.RECEIVED

    result = unfile_document(engagement, filed[0].pbc_location, today=DAY2)

    assert result.scan_note == ""
    a01 = next(i for i in load_manifest(engagement / MANIFEST_FILENAME) if i.identifier == "A01")
    assert a01.status == Status.PARTIAL
    assert REGRESSION_NOTE.format(
        status=Status.RECEIVED, date=DAY1.isoformat(), why=REGRESSION_FILES_CHANGED,
    ) in a01.validation_notes


def test_the_next_pass_leaves_an_unfiled_original_where_the_row_says(engagement):
    """It is still in PBC with a row, so it is nobody's stray and no second copy is made."""
    from tracker.filer import unfile_document

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025")
    filed = file_drops(engagement, today=DAY1).filed[0]
    unfile_document(engagement, filed.pbc_location, today=DAY2)
    parked = sorted(p.name for p in review_dir(engagement).iterdir())

    report = file_drops(engagement, today=DAY2)

    assert report.handled == 0 and report.errors == [] and report.attention == []
    assert sorted(p.name for p in review_dir(engagement).iterdir()) == parked
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW

    # And the same document sent again is what a re-send of a parked one is.
    drop(engagement, "w2 again.pdf", "Form W-2 Wage and Tax Statement 2025")
    again = file_drops(engagement, today=DAY2)
    assert [e.original_name for e in again.duplicates] == ["w2 again.pdf"]
    assert sorted(p.name for p in review_dir(engagement).iterdir()) == parked


def test_an_unfiling_excel_kept_out_of_the_workbook_is_not_lost(engagement, monkeypatch):
    """The copy has already moved, so the row that says where it is cannot be."""
    from tracker.filer import UNFILED_BY_PERSON, unfile_document

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025")
    filed = file_drops(engagement, today=DAY1).filed[0]
    sidecar = engagement / INDEX_PENDING_FILENAME

    _lock_the_index(monkeypatch)
    result = unfile_document(engagement, filed.pbc_location, today=DAY2)
    assert result.index_deferred is True and sidecar.exists()
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW

    monkeypatch.undo()                            # Excel closed - having saved
    _age(sidecar, 60)
    os.utime(engagement / INDEX_FILENAME, None)   # the workbook is newer, and stale
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == NEEDS_REVIEW and row.reason.startswith(UNFILED_BY_PERSON)

    assert file_drops(engagement, today=DAY2).index_deferred is False
    assert not sidecar.exists()
    [row] = read_index(engagement / INDEX_FILENAME)       # the workbook alone now
    assert row.decision == NEEDS_REVIEW and row.reason.startswith(UNFILED_BY_PERSON)


def test_a_failed_index_write_puts_the_unfiled_copy_back_where_the_index_says(engagement, monkeypatch):
    """A retry does this once, not twice: the file goes where the index on disk says."""
    import tracker.filer as filer_module
    from tracker.filer import unfile_document

    drop(engagement, "w2 john.pdf", "Form W-2 Wage and Tax Statement 2025")
    filed = file_drops(engagement, today=DAY1).filed[0]
    working = engagement / filed.prepared_location

    def disk_full(*args, **kwargs):
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(filer_module, "_save_index", disk_full)
    monkeypatch.setattr(filer_module, "_save_pending_index", disk_full)
    monkeypatch.setattr(filer_module.time, "sleep", lambda _: None)
    with pytest.raises(OSError):
        unfile_document(engagement, filed.pbc_location, today=DAY2)
    monkeypatch.undo()

    assert working.is_file()
    assert not list(review_dir(engagement).iterdir())
    [row] = read_index(engagement / INDEX_FILENAME)
    assert row.decision == FILED and row.prepared_location == filed.prepared_location
