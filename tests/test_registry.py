"""Tests for tracker/registry.py — an engagement is a folder with a manifest in it.

The rule under test: nothing has to be registered. The scheduled run walks
the clients folder, finds every manifest, and reads each one's Engagement
sheet. A client nobody typed into a list is still found; a manifest nobody
can read is still reported.
"""

import datetime as dt
from pathlib import Path

import pytest

from tracker.manifest import (
    ENGAGEMENT_LABELS,
    ENGAGEMENT_SHEET_NAME,
    NO,
    YES,
    EngagementInfo,
    RequestItem,
    create_template,
)
from tracker.registry import (
    MAX_DEPTH,
    Engagement,
    RegistryError,
    discover_engagements,
    engagement_dirs,
)
from tracker.scaffold import MANIFEST_FILENAME, PREPARED_DIR_NAME, scaffold_engagement

ITEMS = [RequestItem(identifier="A01", document="W-2")]


def make(root, *parts, info=None, scaffold=False):
    folder = root.joinpath(*parts)
    folder.mkdir(parents=True)
    create_template(folder / MANIFEST_FILENAME, ITEMS, info)
    if scaffold:
        scaffold_engagement(folder)
    return folder


def test_a_rolled_forward_prior_stays_retired_when_the_clients_root_moves(tmp_path):
    # Rolled From is written as an absolute path. A root moved to another
    # drive (or renamed) would otherwise bring every prior back to life and
    # the draft day would chase last year's list (the ninth reading).
    from tracker.registry import engagement_from, mark_superseded

    root = tmp_path / "Clients"
    prior = make(root, "Smith", "Smith - 2025")
    make(root, "Smith", "Smith - 2026", info=EngagementInfo(rolled_from=str(prior.resolve())))
    moved = tmp_path / "Moved"
    root.rename(moved)

    found = [engagement_from(p) for p in engagement_dirs(moved)]
    retired = {e.path.name: e.superseded_by for e in mark_superseded(found)}
    assert retired["Smith - 2025"] and not retired["Smith - 2026"]


def test_a_year_level_above_the_client_never_retires_the_new_engagement_itself(tmp_path):
    # The tenth reading: with Clients/2025/Smith/1040 and Clients/2026/Smith/1040
    # the last two names are the same every year, and the two-name fallback
    # retired the new engagement into itself; the run then skipped both.
    from tracker.registry import engagement_from, mark_superseded

    root = tmp_path / "Clients"
    prior = make(root, "2025", "Smith", "1040")
    make(root, "2026", "Smith", "1040", info=EngagementInfo(rolled_from=str(prior.resolve())))
    make(root, "Chicago", "Smith", "1040")                 # another office's live client of the same name
    found = [engagement_from(p) for p in engagement_dirs(root)]
    retired = {str(e.path.relative_to(root)): e.superseded_by for e in mark_superseded(found)}
    assert retired == {str(Path("2025", "Smith", "1040")): "1040", str(Path("2026", "Smith", "1040")): "",
                       str(Path("Chicago", "Smith", "1040")): ""}

    root.rename(tmp_path / "Moved")                        # the root moves: the third name still decides
    moved = tmp_path / "Moved"
    found = [engagement_from(p) for p in engagement_dirs(moved)]
    retired = {str(e.path.relative_to(moved)): e.superseded_by for e in mark_superseded(found)}
    assert retired[str(Path("2025", "Smith", "1040"))] == "1040"
    assert retired[str(Path("2026", "Smith", "1040"))] == "" and retired[str(Path("Chicago", "Smith", "1040"))] == ""


def test_a_client_folder_the_walk_cannot_list_is_a_problem_row_not_a_silence(tmp_path):
    # The tenth reading: an ACL that denies the run's account made a whole
    # client vanish from the registry, and the run reported success.
    from tests.samples import listing_denied

    root = tmp_path / "Clients"
    make(root, "Jones", "Jones - 2025")
    smith = root / "Smith"
    make(root, "Smith", "Smith - 2025")

    with listing_denied(smith):
        registry = discover_engagements(root)
    by_name = {e.path.name: e for e in registry.engagements}
    assert by_name["Jones - 2025"].problem == ""
    assert "could not be listed" in by_name["Smith"].problem      # the run reports it as an error, by name
    assert [e.path.name for e in registry.engagements if not e.problem] == ["Jones - 2025"]


def test_every_folder_with_a_manifest_is_an_engagement(tmp_path):
    a = make(tmp_path, "Smith", "Smith 2025")
    b = make(tmp_path, "Acme Corp TY2025")
    c = make(tmp_path, "Office B", "Trusts", "Jones Trust 2025")
    assert engagement_dirs(tmp_path) == sorted([a, b, c])


def test_an_engagements_own_subfolders_are_never_engagements(tmp_path):
    outer = make(tmp_path, "Smith 2025", scaffold=True)
    # A stray manifest inside Prepared/ (or anywhere below) does not split
    # the engagement in two.
    create_template(outer / PREPARED_DIR_NAME / MANIFEST_FILENAME, ITEMS)
    assert engagement_dirs(tmp_path) == [outer]


def test_hidden_and_underscore_folders_are_skipped(tmp_path):
    make(tmp_path, ".tmp.driveupload", "Ghost 2025")
    make(tmp_path, "_archive", "Old 2019")
    real = make(tmp_path, "Real 2025")
    assert engagement_dirs(tmp_path) == [real]


def test_discovery_is_depth_limited(tmp_path):
    deep = make(tmp_path, *[f"level{i}" for i in range(MAX_DEPTH + 1)], "Too Deep 2025")
    shallow = make(tmp_path, "Shallow 2025")
    assert engagement_dirs(tmp_path) == [shallow]
    assert deep not in engagement_dirs(tmp_path)


def test_the_engagement_sheet_drives_the_run(tmp_path):
    make(tmp_path, "Smith 2025", info=EngagementInfo(
        client="John Smith", link="https://drive.example/abc",
        due=dt.date(2026, 3, 15), sender="Jason Park", firm="J Park",
    ))
    make(tmp_path, "Acme TY2025", info=EngagementInfo(client="Dana Lee", reminders=False))
    make(tmp_path, "Old 2024", info=EngagementInfo(active=False))

    registry = discover_engagements(tmp_path)
    by_label = {e.label: e for e in registry.engagements}
    smith = by_label["Smith 2025"]
    assert smith.client == "John Smith" and smith.due == dt.date(2026, 3, 15)
    assert smith.link == "https://drive.example/abc" and smith.sender == "Jason Park"
    assert by_label["Acme TY2025"].reminders is False
    assert by_label["Old 2024"].active is False
    assert [e.label for e in registry.active] == ["Acme TY2025", "Smith 2025"]
    assert registry.source == tmp_path


def test_a_manifest_without_the_sheet_is_still_an_engagement(tmp_path):
    from openpyxl import load_workbook

    folder = make(tmp_path, "Made Before The Sheet 2024")
    wb = load_workbook(folder / MANIFEST_FILENAME)
    del wb[ENGAGEMENT_SHEET_NAME]
    wb.save(folder / MANIFEST_FILENAME)
    [engagement] = discover_engagements(tmp_path).engagements
    assert engagement.active and engagement.reminders and engagement.client == ""
    assert engagement.problem == ""


def test_an_unreadable_manifest_is_listed_with_its_problem_not_dropped(tmp_path):
    folder = tmp_path / "Broken 2025"
    folder.mkdir()
    (folder / MANIFEST_FILENAME).write_bytes(b"not a workbook")
    make(tmp_path, "Fine 2025")
    registry = discover_engagements(tmp_path)
    broken = next(e for e in registry.engagements if e.path == folder)
    assert "Could not open" in broken.problem
    assert len(registry.engagements) == 2


def test_a_bad_yes_no_on_the_sheet_is_reported_not_guessed(tmp_path):
    from openpyxl import load_workbook

    folder = make(tmp_path, "Smith 2025")
    wb = load_workbook(folder / MANIFEST_FILENAME)
    ws = wb[ENGAGEMENT_SHEET_NAME]
    for row in ws.iter_rows(min_row=1, max_col=2):
        if row[0].value == ENGAGEMENT_LABELS["reminders"]:
            row[1].value = "maybe"
    wb.save(folder / MANIFEST_FILENAME)
    [engagement] = discover_engagements(tmp_path).engagements
    assert f"{ENGAGEMENT_LABELS['reminders']} must be {YES} or {NO}" in engagement.problem


def test_find_matches_label_or_path(tmp_path):
    make(tmp_path, "Smith Family 2025", info=EngagementInfo(name="Smiths"))
    make(tmp_path, "Acme 2025")
    registry = discover_engagements(tmp_path)
    assert [e.label for e in registry.find("smith")] == ["Smiths"]
    assert [e.label for e in registry.find("acme")] == ["Acme 2025"]
    assert registry.find("nobody") == []


def test_label_falls_back_to_the_folder_name(tmp_path):
    folder = tmp_path / "Smith 2025"
    assert Engagement(path=folder).label == "Smith 2025"
    assert Engagement(path=folder, info=EngagementInfo(name="The Smiths")).label == "The Smiths"


def test_a_root_that_is_not_a_folder_is_an_error(tmp_path):
    with pytest.raises(RegistryError, match="not a folder"):
        discover_engagements(tmp_path / "nowhere")


def test_a_root_with_no_engagement_is_an_error_not_a_quiet_no_op(tmp_path):
    (tmp_path / "Empty").mkdir()
    with pytest.raises(RegistryError, match="no engagement found"):
        discover_engagements(tmp_path)


def test_an_engagement_is_its_sheet(tmp_path):
    # Every field on the Engagement sheet is reachable on the Engagement
    # without being declared a second time.
    from dataclasses import fields

    make(tmp_path, "Smith 2025", info=EngagementInfo(client="John", sender="Jason"))
    [engagement] = discover_engagements(tmp_path).engagements
    for field in fields(EngagementInfo):
        assert getattr(engagement, field.name) == getattr(engagement.info, field.name)
    assert engagement.client == "John" and engagement.sender == "Jason"
    with pytest.raises(AttributeError):
        _ = engagement.no_such_field
