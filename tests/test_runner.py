"""Tests for tracker/runner.py — the weekly pass over every engagement.

The rules that matter here: reminders are drafted on Saturday and only on
Saturday, an edited draft survives the weekly run, one client's failure never
stops another client's run, and nothing is ever sent.
"""

import datetime as dt
import html
from pathlib import Path

import pytest

from tests.samples import DEMO_ITEMS, PRIOR_YEAR, YEAR, build_samples, col, row
from tracker.filer import INDEX_FILENAME, NEEDS_REVIEW, IndexEntry, read_index, write_index
from tracker.manifest import (
    COL_ALLOWED_EXTENSIONS,
    COL_ANY_KEYWORDS,
    COL_DATE_PATTERN,
    COL_DOCUMENT,
    COL_EXPECTED_COUNT,
    COL_IDENTIFIER,
    COL_MANUAL_OVERRIDE,
    COL_MIN_SIZE_KB,
    COL_PERIOD,
    SHEET_NAME,
    EngagementInfo,
    Override,
    Status,
    create_template,
    write_engagement_info,
)
from tracker.registry import SKIP_ROLLED_FORWARD, Engagement, Registry, discover_engagements
from tracker.reminder import DRAFT_BANNER, DRAFT_FILENAME, NEW_DRAFT_FILENAME
from tracker.runner import (
    DRAFT_WEEKDAY,
    LOG_FILENAME,
    MANIFEST_UNREADABLE,
    NOTHING_OUTSTANDING,
    REMINDERS_ALWAYS,
    REMINDERS_AUTO,
    REMINDERS_NEVER,
    STATUS_GENERATED,
    STATUS_PAGE_FILENAME,
    WEEKDAY_NAMES,
    EngagementRun,
    RunReport,
    append_log,
    format_report,
    is_draft_day,
    main,
    run_engagement,
    run_registry,
    should_draft,
    write_status_page,
)
from tracker.scaffold import (
    MANIFEST_FILENAME,
    PBC_DIR_NAME,
    PREPARED_DIR_NAME,
    SHARED_DIR_NAME,
    scaffold_engagement,
)
from tracker.settings import product_name

DRAFT_DAY = WEEKDAY_NAMES[DRAFT_WEEKDAY]

# The draft day is whatever the runner says it is; the dates follow it.
_WEEK_START = dt.date(2026, 3, 9)                      # a Monday
SATURDAY = _WEEK_START + dt.timedelta(days=DRAFT_WEEKDAY)  # the draft day
FRIDAY = SATURDAY - dt.timedelta(days=1)                # the day before
SUNDAY = SATURDAY + dt.timedelta(days=1)                # the day after


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    folder = tmp_path_factory.mktemp("samples")
    build_samples(folder)
    return folder


def test_the_runner_has_a_main_the_frozen_entry_can_call(tmp_path, samples, capsys):
    # api_entry.py runs the scheduled job through this function, so the
    # command line has to be one, not code under __main__.
    build_engagement(tmp_path, samples)
    assert main([str(tmp_path), "--dry-run", "--reminders", REMINDERS_NEVER]) == 0
    assert "Smith TY2025" in capsys.readouterr().out


def build_engagement(tmp_path, samples, drops=(f"W-2 John Smith {YEAR}.pdf",),
                     name="Smith TY2025", **kwargs):
    """A scaffolded engagement with files waiting in the client's drop folder."""
    folder = tmp_path / name
    folder.mkdir(parents=True)
    # The Engagement sheet is where the draft learns who the client is; the
    # Engagement dataclass only echoes it.
    create_template(folder / MANIFEST_FILENAME, DEMO_ITEMS,
                    EngagementInfo(client="John Smith", firm="J Park"))
    result = scaffold_engagement(folder)
    for drop in drops:
        (result.shared_dir / drop).write_bytes((samples / drop).read_bytes())
    return Engagement(path=folder, info=EngagementInfo(client="John Smith", firm="J Park", **kwargs))


# ------------------------------------------------------------ the Saturday ----


def test_saturday_is_the_draft_day():
    assert is_draft_day(SATURDAY) is True
    assert is_draft_day(FRIDAY) is False
    assert is_draft_day(SUNDAY) is False


@pytest.mark.parametrize("day, expected", [(FRIDAY, False), (SATURDAY, True),
                                           (SUNDAY, False)])
def test_auto_mode_drafts_only_on_saturday(day, expected):
    engagement = Engagement(path=Path("/x"))
    assert should_draft(engagement, day, REMINDERS_AUTO) is expected


def test_always_mode_drafts_on_any_day():
    assert should_draft(Engagement(path=Path("/x")), FRIDAY, REMINDERS_ALWAYS) is True


def test_never_mode_suppresses_even_on_saturday():
    assert should_draft(Engagement(path=Path("/x")), SATURDAY, REMINDERS_NEVER) is False


def test_reminders_off_beats_every_mode():
    """A standing decision not to chase this client by email is not a flag."""
    quiet = Engagement(path=Path("/x"), info=EngagementInfo(reminders=False))
    assert should_draft(quiet, SATURDAY, REMINDERS_AUTO) is False
    assert should_draft(quiet, SATURDAY, REMINDERS_ALWAYS) is False


def test_a_draft_day_the_machine_missed_is_caught_up_on_the_next_pass():
    # Weekly, not only on the day. The machine was off on Saturday: Sunday's
    # pass drafts, because the last draft is older than the Saturday that
    # went by. Monday after a Saturday that did draft: nothing, until next
    # week. Never drafted: wait for the first Saturday.
    from tracker.runner import last_draft_day

    engagement = Engagement(path=Path("/x"))
    assert last_draft_day(SUNDAY) == SATURDAY and last_draft_day(SATURDAY) == SATURDAY
    assert last_draft_day(FRIDAY) == SATURDAY - dt.timedelta(days=7)
    assert should_draft(engagement, SUNDAY, REMINDERS_AUTO, drafted=SATURDAY - dt.timedelta(days=7)) is True
    assert should_draft(engagement, SUNDAY, REMINDERS_AUTO, drafted=SATURDAY) is False
    assert should_draft(engagement, FRIDAY, REMINDERS_AUTO, drafted=SATURDAY - dt.timedelta(days=7)) is False
    assert should_draft(engagement, SUNDAY, REMINDERS_AUTO, drafted=None) is False


def test_a_pass_after_a_missed_saturday_writes_the_weeks_draft(tmp_path, samples):
    import os

    engagement = build_engagement(tmp_path, samples)
    drafted = run_engagement(engagement, today=SATURDAY - dt.timedelta(days=7)).drafted
    assert drafted is not None
    stamp = dt.datetime.combine(SATURDAY - dt.timedelta(days=7), dt.time(9)).timestamp()
    os.utime(drafted, (stamp, stamp))                  # written last Saturday
    assert run_engagement(engagement, today=FRIDAY).drafted is None      # this week not yet due
    run = run_engagement(engagement, today=SUNDAY)     # Saturday was missed
    assert run.drafted == engagement.path / DRAFT_FILENAME


def test_the_draft_day_can_be_moved():
    monday = dt.date(2026, 3, 16)
    assert should_draft(Engagement(path=Path("/x")), monday, REMINDERS_AUTO, weekday=0) is True
    assert should_draft(Engagement(path=Path("/x")), SATURDAY, REMINDERS_AUTO, weekday=0) is False


# ------------------------------------------------------------- the full pass ----


def test_a_weekday_run_files_and_scans_but_writes_no_draft(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    run = run_engagement(engagement, today=FRIDAY)

    assert run.ok and run.filed == 1
    assert run.statuses[Status.PARTIAL] == 1
    assert run.drafted is None
    assert f"not {DRAFT_DAY}" in run.draft_note.lower()
    assert not (engagement.path / DRAFT_FILENAME).exists()


def test_the_saturday_run_writes_a_draft(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    run = run_engagement(engagement, today=SATURDAY)

    assert run.drafted == engagement.path / DRAFT_FILENAME
    text = run.drafted.read_text(encoding="utf-8")
    assert text.startswith(DRAFT_BANNER)
    assert "Hi John Smith," in text


def test_a_manual_run_drafts_on_a_weekday(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    run = run_engagement(engagement, today=FRIDAY, reminders=REMINDERS_ALWAYS)
    assert run.drafted == engagement.path / DRAFT_FILENAME


def test_an_edited_draft_survives_the_next_weekly_run(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    first = run_engagement(engagement, today=SATURDAY).drafted
    # Read and write the bytes as they are: these drafts are CRLF for Windows,
    # and universal-newline translation would make this test lie.
    edited = first.read_bytes() + b"\r\nPS: ask about the rental.\r\n"
    first.write_bytes(edited)

    second = run_engagement(engagement, today=SATURDAY + dt.timedelta(days=7))

    assert first.read_bytes() == edited, "an edit must never be clobbered"
    assert second.drafted.name == NEW_DRAFT_FILENAME
    assert "has been edited" in second.draft_note


def test_an_edited_second_draft_is_never_clobbered_by_the_next_repeat(tmp_path, samples):
    # The draft day's repeat runs several times. The first made room for the
    # person's edits by writing NEW; the next must not take NEW's edits too.
    from tracker.reminder import BOTH_DRAFTS_EDITED

    engagement = build_engagement(tmp_path, samples)
    first = run_engagement(engagement, today=SATURDAY).drafted
    first.write_bytes(first.read_bytes() + b"\r\nPS: ask about the rental.\r\n")
    second = run_engagement(engagement, today=SATURDAY).drafted
    assert second.name == NEW_DRAFT_FILENAME
    edited_new = second.read_bytes() + b"\r\nPPS: and the boat.\r\n"
    second.write_bytes(edited_new)

    third = run_engagement(engagement, today=SATURDAY)
    assert third.drafted is None and third.ok
    assert third.draft_note == BOTH_DRAFTS_EDITED
    assert second.read_bytes() == edited_new


def test_a_draft_day_with_nothing_to_chase_refreshes_the_runs_own_stale_draft(tmp_path, samples):
    # Last week's draft asked for a document that has since arrived. The
    # run's own unedited draft is rewritten as today's (nothing to chase),
    # so it is neither stale text nor an old date; a draft a person edited
    # is theirs and stays exactly as it is.
    import os

    from tracker.runner import last_drafted

    only_the_return = [i for i in DEMO_ITEMS if i.identifier == "B01"]
    folder = tmp_path / "Settled TY2025"
    folder.mkdir()
    create_template(folder / MANIFEST_FILENAME, only_the_return)
    scaffolded = scaffold_engagement(folder)
    engagement = Engagement(path=folder, info=EngagementInfo(client="John Smith"))
    drafted = run_engagement(engagement, today=SATURDAY - dt.timedelta(days=7)).drafted
    stale = drafted.read_bytes()
    stamp = dt.datetime.combine(SATURDAY - dt.timedelta(days=7), dt.time(9)).timestamp()
    os.utime(drafted, (stamp, stamp))

    name = f"{PRIOR_YEAR} Form 1040 Tax Return.pdf"
    (scaffolded.shared_dir / name).write_bytes((samples / name).read_bytes())
    run = run_engagement(engagement, today=SATURDAY)
    assert run.draft_note == NOTHING_OUTSTANDING and run.drafted is None
    assert drafted.exists() and drafted.read_bytes() != stale          # today's words
    assert last_drafted(folder) == dt.date.today()
    assert run_engagement(engagement, today=SATURDAY + dt.timedelta(days=3)).drafted is None

    drafted.write_bytes(b"a person's own words")
    os.utime(drafted, (stamp, stamp))
    run_engagement(engagement, today=SATURDAY)
    assert drafted.read_bytes() == b"a person's own words"
    assert last_drafted(folder) == SATURDAY - dt.timedelta(days=7)


def test_an_untouched_draft_is_refreshed_in_place(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    first = run_engagement(engagement, today=SATURDAY).drafted
    second = run_engagement(engagement, today=SATURDAY + dt.timedelta(days=7))

    assert second.drafted == first
    assert not (engagement.path / NEW_DRAFT_FILENAME).exists()


def test_nothing_outstanding_means_no_draft_file(tmp_path, samples):
    """Everything in: there is nothing to chase, so no draft is written."""
    only_the_return = [i for i in DEMO_ITEMS if i.identifier == "B01"]
    folder = tmp_path / "Settled TY2025"
    folder.mkdir()
    create_template(folder / MANIFEST_FILENAME, only_the_return)
    scaffolded = scaffold_engagement(folder)
    name = f"{PRIOR_YEAR} Form 1040 Tax Return.pdf"
    (scaffolded.shared_dir / name).write_bytes((samples / name).read_bytes())

    run = run_engagement(Engagement(path=folder, info=EngagementInfo(client="John Smith")), today=SATURDAY)

    assert run.statuses == {Status.RECEIVED: 1}
    assert run.outstanding == 0
    assert run.drafted is None
    assert run.draft_note == NOTHING_OUTSTANDING
    assert not (folder / DRAFT_FILENAME).exists()


def test_a_dry_run_writes_nothing_at_all(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    run = run_engagement(engagement, today=SATURDAY, dry_run=True)

    assert run.drafted is None
    assert "would draft" in run.draft_note
    assert not (engagement.path / DRAFT_FILENAME).exists()
    assert (engagement.path / SHARED_DIR_NAME / f"W-2 John Smith {YEAR}.pdf").exists(), (
        "a dry run must not move the client's file"
    )


# --------------------------------------------------------- failure isolation ----


def test_a_missing_folder_is_recorded_not_raised(tmp_path):
    run = run_engagement(Engagement(path=tmp_path / "no-such-client"), today=SATURDAY)
    assert run.error.startswith("folder not found")


def test_an_unreadable_manifest_is_recorded_not_raised(tmp_path):
    folder = tmp_path / "Broken 2025"
    folder.mkdir()
    (folder / MANIFEST_FILENAME).write_text("this is not a workbook", encoding="utf-8")
    run = run_engagement(Engagement(path=folder), today=SATURDAY)
    assert run.error and run.ok is False


def test_one_broken_engagement_does_not_stop_the_others(tmp_path, samples):
    good = build_engagement(tmp_path, samples, name="Smith TY2025")
    broken = Engagement(path=tmp_path / "Ghost 2025")
    later = build_engagement(tmp_path, samples, name="Jones TY2025")

    report = run_registry(
        Registry(source=tmp_path,
                 engagements=[good, broken, later]),
        today=SATURDAY,
    )

    assert len(report.errors) == 1
    assert len(report.processed) == 2
    assert (good.path / DRAFT_FILENAME).exists()
    assert (later.path / DRAFT_FILENAME).exists(), (
        "an engagement after the broken one must still be processed"
    )


def test_inactive_engagements_are_skipped_not_failed(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples, active=False)
    run = run_engagement(engagement, today=SATURDAY)
    assert run.ok and run.skipped.startswith("inactive")
    assert run.filed == 0


def test_only_selects_a_subset(tmp_path, samples):
    smith = build_engagement(tmp_path, samples, name="Smith TY2025")
    jones = build_engagement(tmp_path, samples, name="Jones TY2025")
    registry = Registry(source=tmp_path,
                        engagements=[smith, jones])

    report = run_registry(registry, today=SATURDAY, only="jones")

    assert len(report.runs) == 1
    assert (jones.path / DRAFT_FILENAME).exists()
    assert not (smith.path / DRAFT_FILENAME).exists()


def test_an_unknown_reminder_mode_fails_loudly(tmp_path):
    registry = Registry(source=tmp_path, engagements=[])
    with pytest.raises(ValueError, match="reminders must be one of"):
        run_registry(registry, reminders="sometimes")


# ---------------------------------------------------------------- reporting ----


def test_the_report_says_what_day_it_is_and_never_claims_a_send(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    text = format_report(run_registry(
        Registry(source=tmp_path, engagements=[engagement]),
        today=SATURDAY,
    ))
    assert DRAFT_DAY in text
    assert "drafting reminders" in text
    assert "nothing has been sent to anyone" in text.lower()
    assert "1 processed, 0 failed, 1 draft(s) written" in text


def test_a_weekday_report_says_why_there_are_no_drafts(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    text = format_report(run_registry(
        Registry(source=tmp_path, engagements=[engagement]),
        today=FRIDAY,
    ))
    assert "no reminders today" in text
    assert "0 draft(s) written" in text


def test_the_log_takes_a_name_utf8_cannot_hold(tmp_path, samples):
    # The eleventh reading: a lone surrogate in a client's file name is
    # reported every pass by name, and the run log - the unattended run's
    # only trace - died on it, losing every engagement's line.
    engagement = build_engagement(tmp_path, samples)
    run = EngagementRun(engagement=engagement)
    run.error = "bank statement \ud83d.pdf: cannot be handled under this name"
    report = RunReport(today=FRIDAY, reminders=REMINDERS_AUTO, dry_run=False, runs=[run])
    log = tmp_path / LOG_FILENAME
    append_log(log, report)
    text = log.read_text(encoding="utf-8")
    assert "bank statement" in text and "cannot be handled" in text


def test_the_log_appends_rather_than_replaces(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples)
    registry = Registry(source=tmp_path, engagements=[engagement])
    log = tmp_path / LOG_FILENAME

    append_log(log, run_registry(registry, today=FRIDAY))
    append_log(log, run_registry(registry, today=SATURDAY))

    text = log.read_text(encoding="utf-8")
    assert text.count("Smith TY2025") == 2
    assert "2026-03-13" in text and "2026-03-14" in text


# ------------------------------------------------------- discovery end to end ----


def test_a_clients_folder_drives_a_real_run_with_nothing_registered(tmp_path, samples):
    engagement = build_engagement(tmp_path / "Clients" / "Smith", samples)
    write_engagement_info(engagement.path / MANIFEST_FILENAME, EngagementInfo(
        client="John Smith", due=dt.date(2026, 4, 15), firm="J Park & Associates, CPA",
    ))

    report = run_registry(discover_engagements(tmp_path / "Clients"), today=SATURDAY)

    assert len(report.processed) == 1
    draft = (engagement.path / DRAFT_FILENAME).read_text(encoding="utf-8")
    assert "Hi John Smith," in draft
    assert "April 15, 2026" in draft
    assert "J Park & Associates, CPA" in draft


def test_an_engagement_whose_manifest_cannot_be_read_fails_alone(tmp_path, samples):
    build_engagement(tmp_path / "Clients", samples, name="Good")
    bad = tmp_path / "Clients" / "Bad 2025"
    bad.mkdir()
    (bad / MANIFEST_FILENAME).write_bytes(b"not a workbook")

    report = run_registry(discover_engagements(tmp_path / "Clients"), today=FRIDAY)
    outcomes = {r.engagement.path.name: r for r in report.runs}
    assert outcomes["Good"].ok
    assert "manifest could not be read" in outcomes["Bad 2025"].error


def test_the_run_names_a_manifest_typo_with_its_row_before_touching_files(tmp_path, samples):
    from openpyxl import load_workbook

    engagement = build_engagement(tmp_path, samples)
    manifest = engagement.path / MANIFEST_FILENAME
    wb = load_workbook(manifest)
    wb[SHEET_NAME].cell(row=2, column=col(COL_DATE_PATTERN), value="(unclosed")
    wb.save(manifest)
    run = run_engagement(engagement, today=FRIDAY)
    assert run.error.startswith(f"Row 2: {COL_DATE_PATTERN} is not a valid regex")
    assert (engagement.path / SHARED_DIR_NAME / f"W-2 John Smith {YEAR}.pdf").exists()   # nothing moved


def test_rows_the_rules_cannot_act_on_are_reported_not_buried(tmp_path, samples):
    from tracker.manifest import RequestItem, create_template

    folder = tmp_path / "Loose 2025"
    folder.mkdir()
    create_template(folder / MANIFEST_FILENAME, [RequestItem(identifier="A01", document="Anything")])
    scaffold_engagement(folder)
    engagement = Engagement(path=folder)
    run = run_engagement(engagement, today=FRIDAY)
    assert run.ok
    assert any("never be filed automatically" in w for w in run.warnings)
    text = format_report(run_registry(Registry(source=tmp_path, engagements=[engagement]), today=FRIDAY))
    assert "! Row 2 (A01)" in text


def test_waived_and_accepted_rows_are_not_outstanding(tmp_path, samples):
    from openpyxl import load_workbook

    engagement = build_engagement(tmp_path, samples, drops=())
    manifest = engagement.path / MANIFEST_FILENAME
    wb = load_workbook(manifest)
    ws = wb[SHEET_NAME]
    ws.cell(row=2, column=col(COL_MANUAL_OVERRIDE), value=Override.ACCEPTED)   # A01
    ws.cell(row=3, column=col(COL_MANUAL_OVERRIDE), value=Override.WAIVED)     # A02
    wb.save(manifest)
    run = run_engagement(engagement, today=FRIDAY)
    assert run.ok
    assert run.statuses.get(Status.RECEIVED) == 1          # the accepted row
    assert Override.WAIVED not in run.statuses
    assert run.outstanding == len(DEMO_ITEMS) - 2


def test_a_row_added_in_excel_has_its_folder_by_the_next_run(tmp_path, samples):
    from openpyxl import load_workbook

    from tracker.scaffold import README_NAME

    engagement = build_engagement(tmp_path, samples, drops=())
    manifest = engagement.path / MANIFEST_FILENAME
    wb = load_workbook(manifest)
    wb[SHEET_NAME].append(row(**{
        COL_IDENTIFIER: "Z01", COL_DOCUMENT: "Rental Property Records", COL_PERIOD: "TY2025",
        COL_EXPECTED_COUNT: 1, COL_ALLOWED_EXTENSIONS: "pdf", COL_MIN_SIZE_KB: 5,
        COL_ANY_KEYWORDS: "schedule e",
    }))
    wb.save(manifest)
    run = run_engagement(engagement, today=FRIDAY)
    assert run.ok
    assert any(p.name.startswith("Z01") for p in (engagement.path / PREPARED_DIR_NAME).iterdir())
    assert "Z01 - Rental Property Records" in (engagement.path / SHARED_DIR_NAME / README_NAME).read_text(encoding="utf-8")
    assert run.statuses.get(Status.MISSING, 0) >= 1 and "folder not found" not in str(run.statuses)


def test_a_rolled_forward_engagement_is_retired_by_its_successor(tmp_path, samples):
    from tracker.manifest import EngagementInfo, write_engagement_info
    from tracker.registry import discover_engagements

    prior = build_engagement(tmp_path / "Clients", samples, name="Smith 2025")
    write_engagement_info(prior.path / MANIFEST_FILENAME, EngagementInfo(client="John"))
    new = build_engagement(tmp_path / "Clients", samples, name="Smith 2026", drops=())
    write_engagement_info(new.path / MANIFEST_FILENAME,
                          EngagementInfo(client="John", rolled_from=str(prior.path)))

    registry = discover_engagements(tmp_path / "Clients")
    by_name = {e.path.name: e for e in registry.engagements}
    assert by_name["Smith 2025"].active is False
    assert by_name["Smith 2025"].superseded_by == "Smith 2026"
    assert by_name["Smith 2026"].active is True

    report = run_registry(registry, today=SATURDAY)
    outcomes = {r.engagement.path.name: r for r in report.runs}
    assert outcomes["Smith 2025"].skipped == SKIP_ROLLED_FORWARD.format(successor="Smith 2026")
    assert not (prior.path / DRAFT_FILENAME).exists()          # last year is not chased
    assert (prior.path / SHARED_DIR_NAME / f"W-2 John Smith {YEAR}.pdf").exists()   # and not touched
    assert outcomes["Smith 2026"].ok


def test_strays_in_prepared_reach_the_run_report(tmp_path, samples):
    engagement = build_engagement(tmp_path, samples, drops=())
    (engagement.path / PREPARED_DIR_NAME / "loose.txt").write_text("x", encoding="utf-8")
    run = run_engagement(engagement, today=FRIDAY)
    assert run.ok
    assert any(f"loose.txt is loose in {PREPARED_DIR_NAME}/" in w for w in run.warnings)


# ----------------------------------------------------- the practice on a page ----


def test_a_real_pass_writes_the_status_page_into_the_root_and_a_dry_run_does_not(tmp_path, samples, capsys):
    """The page is the pass's standing answer, so a pass that changed nothing
    on disk must not leave one - a dry run writes nothing, this included."""
    build_engagement(tmp_path, samples)
    page = tmp_path / STATUS_PAGE_FILENAME

    assert main([str(tmp_path), "--date", FRIDAY.isoformat(), "--dry-run"]) == 0
    assert not page.exists()

    assert main([str(tmp_path), "--date", FRIDAY.isoformat()]) == 0
    text = page.read_text(encoding="utf-8")
    assert product_name() in text
    assert STATUS_GENERATED.split("{")[0].strip() in text
    assert "Smith TY2025" in text
    assert "<script" not in text and "http" not in text, "one file, nothing fetched to render it"
    capsys.readouterr()


def test_the_page_lists_every_engagement_the_registry_finds_including_an_unreadable_one(
    tmp_path, samples, capsys
):
    """An engagement missing from the page is an engagement nobody chases. One
    whose manifest cannot be read is on it by name, and named as a problem."""
    clients = tmp_path / "Clients"
    build_engagement(clients, samples, name="Good TY2025")
    bad = clients / "Bad TY2025"
    bad.mkdir(parents=True)
    (bad / MANIFEST_FILENAME).write_bytes(b"not a workbook")

    assert main([str(clients), "--date", FRIDAY.isoformat()]) == 1
    text = (clients / STATUS_PAGE_FILENAME).read_text(encoding="utf-8")

    assert "Good TY2025" in text
    assert text.count("Bad TY2025") >= 2            # its row, and the problems list
    assert MANIFEST_UNREADABLE.split("{")[0] in text
    capsys.readouterr()


def test_the_review_queue_lists_parked_files_newest_first_with_their_reasons(tmp_path, samples):
    """One queue across the practice, because the work is one queue: the file
    that arrived last night is at the top, whichever client sent it."""
    older = build_engagement(tmp_path, samples, drops=("vacation photo.jpg",), name="Older TY2025")
    newer = build_engagement(tmp_path, samples, drops=("Mortgage Notes.docx",), name="Newer TY2025")
    first = run_engagement(older, today=FRIDAY, reminders=REMINDERS_NEVER)
    second = run_engagement(newer, today=SATURDAY, reminders=REMINDERS_NEVER)
    assert first.review == 1 and second.review == 1

    page = write_status_page(tmp_path, RunReport(today=SATURDAY, runs=[first, second]))
    text = page.read_text(encoding="utf-8")

    assert text.index("Mortgage Notes.docx") < text.index("vacation photo.jpg")
    assert SATURDAY.isoformat() in text and FRIDAY.isoformat() in text
    for engagement in (older, newer):
        parked = [e for e in read_index(engagement.path / INDEX_FILENAME)
                  if e.decision == NEEDS_REVIEW]
        assert parked
        for entry in parked:
            assert html.escape(entry.reason) in text, entry.reason


def test_a_file_named_like_markup_is_shown_as_a_name_not_rendered(tmp_path):
    """Every value on the page goes through html.escape. The name cannot be
    made on Windows, so it arrives the way it would in life: an index row
    written elsewhere, from a client's folder that is not this machine's."""
    folder = tmp_path / "Evil TY2025"
    folder.mkdir()
    create_template(folder / MANIFEST_FILENAME, DEMO_ITEMS)
    name = "<b>evil</b>.pdf"
    reason = "<i>nothing matched</i>"
    write_index(folder / INDEX_FILENAME, [IndexEntry(
        received=FRIDAY.isoformat(), original_name=name, size_kb=1.0, digest="0" * 8,
        identifier="", prepared_location="",
        pbc_location=f"{SHARED_DIR_NAME}/{PBC_DIR_NAME}/{name}",
        decision=NEEDS_REVIEW, reason=reason,
    )])

    report = RunReport(today=FRIDAY, runs=[EngagementRun(engagement=Engagement(path=folder))])
    text = write_status_page(tmp_path, report).read_text(encoding="utf-8")

    assert name not in text and reason not in text
    assert "&lt;b&gt;evil&lt;/b&gt;.pdf" in text
    assert "&lt;i&gt;nothing matched&lt;/i&gt;" in text


def test_the_page_is_written_even_when_an_engagements_pass_failed(tmp_path, samples, capsys):
    """The pass that errored is exactly the one a person has to find out about,
    so the page it would have been written by is still written, with the error."""
    from openpyxl import load_workbook

    build_engagement(tmp_path, samples, name="Good TY2025")
    broken = build_engagement(tmp_path, samples, name="Broken TY2025")
    manifest = broken.path / MANIFEST_FILENAME
    wb = load_workbook(manifest)
    wb[SHEET_NAME].cell(row=2, column=col(COL_DATE_PATTERN), value="(unclosed")
    wb.save(manifest)

    assert main([str(tmp_path), "--date", FRIDAY.isoformat()]) == 1
    text = (tmp_path / STATUS_PAGE_FILENAME).read_text(encoding="utf-8")

    assert "Good TY2025" in text and text.count("Broken TY2025") >= 2
    assert f"Row 2: {COL_DATE_PATTERN}" in text
    capsys.readouterr()


def test_the_page_still_has_a_heading_where_nothing_can_say_what_the_product_is_called(
    tmp_path, monkeypatch
):
    """The packaged app's scheduled job runs without the Electron shell, which
    is the only thing that knows the product's name there. A pass that has
    already moved the client's files must not end in a traceback over a
    heading, so the firm's name - and then the folder - stand in."""
    from tracker.settings import SettingsError

    def unknown():
        raise SettingsError("no shell, no package.json")

    folder = tmp_path / "Clients"
    folder.mkdir()
    report = RunReport(today=FRIDAY, runs=[])

    monkeypatch.setattr("tracker.runner.product_name", unknown)
    monkeypatch.setattr("tracker.runner.firm", lambda: "J Park & Associates, CPA")
    assert "J Park &amp; Associates, CPA" in write_status_page(folder, report).read_text(encoding="utf-8")

    monkeypatch.setattr("tracker.runner.firm", lambda: "")
    assert folder.name in write_status_page(folder, report).read_text(encoding="utf-8")


def test_an_engagement_nobody_passed_is_read_rather_than_run(tmp_path, samples):
    """The app scans one engagement; the page is still about the practice. The
    others are read - manifest and sidecar, no lock, no scaffold, no scan -
    and say so by carrying no pass time."""
    from tracker.registry import discover_engagements
    from tracker.runner import STATUS_NOT_PASSED, status_report

    clients = tmp_path / "Clients"
    scanned = build_engagement(clients, samples, name="Scanned TY2025")
    untouched = build_engagement(clients, samples, name="Untouched TY2025")
    earlier = run_engagement(untouched, today=FRIDAY, reminders=REMINDERS_NEVER)
    run = run_engagement(scanned, today=FRIDAY, reminders=REMINDERS_NEVER)

    report = status_report(discover_engagements(clients), passed=[run])
    rows = {r.engagement.path.name: r for r in report.runs}
    assert set(rows) == {"Scanned TY2025", "Untouched TY2025"}
    assert rows["Scanned TY2025"].last_pass is not None
    # Read back from the manifest its own earlier pass wrote, without a pass.
    assert rows["Untouched TY2025"].outstanding == earlier.outstanding > 0
    assert rows["Untouched TY2025"].statuses == earlier.statuses
    assert rows["Untouched TY2025"].last_pass is None

    text = write_status_page(clients, report).read_text(encoding="utf-8")
    assert STATUS_NOT_PASSED in text
