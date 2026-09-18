"""Tests for tracker/api.py — the JSON command layer the desktop app calls.

The contract under test: every command answers with one JSON object and an
exit code, an error is a JSON object too (never a traceback on stdout), and
the create path refuses a bad request list with a sentence a person can act
on. The clients root is a temp folder recorded through settings, the way
the app records it, so no test touches a real one.
"""

import datetime as dt
import io
import json
import os
from pathlib import Path

import pytest

import tracker.api as api
from tests.samples import PRIOR_YEAR, col, row
from tracker.filer import FILED, NEEDS_REVIEW
from tracker.locking import STALE_LOCK_SECONDS, lock_line
from tracker.manifest import (
    COL_ALLOWED_EXTENSIONS,
    COL_ANY_KEYWORDS,
    COL_DATE_PATTERN,
    COL_DOCUMENT,
    COL_EXPECTED_COUNT,
    COL_IDENTIFIER,
    COL_MIN_SIZE_KB,
    COL_PERIOD,
    SHEET_NAME,
    ManifestError,
    Status,
    load_manifest,
)
from tracker.runner import DRAFT_WEEKDAY, LOG_FILENAME, STATUS_PAGE_FILENAME, WEEKDAY_NAMES
from tracker.scaffold import MANIFEST_FILENAME, PREPARED_DIR_NAME, SHARED_DIR_NAME
from tracker.scheduling import SCHEDULE_XML_ENCODING, TASK_NAME
from tracker.templates import BASE_YEAR, default_tax_year


@pytest.fixture
def demo_root(tmp_path, monkeypatch):
    """A clients root recorded the way the app records it: settings.json beside the app."""
    from tracker.settings import ENV_SETTINGS_DIR, set_clients_root, set_firm

    root = tmp_path / "Clients"
    root.mkdir()
    monkeypatch.setenv(ENV_SETTINGS_DIR, str(tmp_path / "app"))
    set_clients_root(root)
    set_firm("J Park & Associates, CPA")
    return root


def run(capsys, *argv, stdin=None):
    """Run one command the way main.js does: argv in, one JSON object out."""
    if stdin is not None:
        with pytest.MonkeyPatch.context() as mp:
            mp.setattr("sys.stdin", io.StringIO(json.dumps(stdin)))
            code = api.main(list(argv))
    else:
        code = api.main(list(argv))
    out = capsys.readouterr().out.strip().splitlines()
    assert len(out) == 1, f"expected exactly one JSON line, got {out!r}"
    return code, json.loads(out[0])


# ------------------------------------------------------------- the contract ----


def test_unknown_command_is_a_json_error(capsys):
    code, payload = run(capsys, "frobnicate")
    assert code == 1
    assert payload["error"].startswith("usage:")


def test_a_manifest_problem_is_a_json_error_not_a_traceback(capsys, demo_root):
    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(demo_root / "nowhere"))
    assert code == 1
    assert "Manifest not found" in payload["error"]


def test_the_engagement_flag_needs_a_folder_under_the_root(capsys, demo_root, tmp_path):
    hint = f"Pick an engagement first ({api.ENGAGEMENT_FLAG} <folder>)"
    assert run(capsys, "state")[1]["error"] == hint
    assert run(capsys, "state", api.ENGAGEMENT_FLAG)[1]["error"] == hint        # was an IndexError
    assert run(capsys, "state", api.ENGAGEMENT_FLAG, "  ")[1]["error"] == hint
    elsewhere = tmp_path / "Elsewhere" / "Smith"
    elsewhere.mkdir(parents=True)
    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(elsewhere))
    assert code == 1 and "not under the clients root" in payload["error"]


def test_an_unplugged_clients_root_is_not_a_licence_to_read_anywhere(capsys, demo_root, tmp_path):
    import shutil

    elsewhere = tmp_path / "Elsewhere" / "Smith"
    elsewhere.mkdir(parents=True)
    shutil.rmtree(demo_root)                          # the drive is gone
    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(elsewhere))
    assert code == 1 and "not under the clients root" in payload["error"]


def test_a_rollover_takes_its_prior_only_from_under_the_root(capsys, demo_root, tmp_path):
    from tracker.manifest import create_template
    from tracker.templates import template_items

    elsewhere = tmp_path / "Elsewhere" / "Smith TY2025"
    elsewhere.mkdir(parents=True)
    create_template(elsewhere / MANIFEST_FILENAME, template_items("1040", year=2025))
    code, payload = run(capsys, "rollover", stdin={"prior": str(elsewhere)})
    assert code == 1 and "not under the clients root" in payload["error"]


def test_templates_lists_every_form_with_its_checklist(capsys):
    code, payload = run(capsys, "templates")
    assert code == 0
    assert {f["id"] for f in payload["forms"]} == set(payload["templates"])
    for form, rows in payload["templates"].items():
        identifiers = [r["identifier"] for r in rows]
        assert len(identifiers) == len(set(identifiers)), form
        assert any(r["core"] for r in rows), form


# ------------------------------------------------------------------ create ----


def test_create_builds_manifest_and_folders(capsys, demo_root):
    spec = {
        "name": "Smith Family 2025 Form 1040",
        "form": "1040",
        "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]
        + [{"identifier": "X01", "document": "Rental Property Records",
            "extensions": "pdf, xlsx", "required_keywords": "Schedule E"}],
    }
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    engagement = demo_root / payload["created"]
    assert (engagement / MANIFEST_FILENAME).is_file()
    assert (engagement / SHARED_DIR_NAME).is_dir()
    assert any(p.name.startswith("X01") for p in (engagement / PREPARED_DIR_NAME).iterdir())
    assert [i["identifier"] for i in payload["state"]["items"]][-1] == "X01"


def test_create_with_no_name_builds_one_that_stays_under_the_root(capsys, demo_root):
    # The tenth reading: a blank name fell back to the raw client field, and
    # a client called "..\\..\\escaped" wrote the engagement above the root,
    # where discovery never finds it and nobody is chased.
    spec = {"name": "", "client": "..\\..\\escaped", "form": "1040", "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    created = demo_root / payload["created"]
    assert created.parent == demo_root and (created / MANIFEST_FILENAME).is_file()
    assert not (demo_root.parent.parent / "escaped TY2025 Form 1040").exists()
    for outside in ("../outside", "sub/child"):          # a separator either platform reads
        with pytest.raises(api.ManifestError, match="not a folder name"):
            api._new_engagement_dir(outside)


def test_the_tax_year_is_within_the_bounds_the_wizard_shows(capsys, demo_root):
    from tracker.manifest import YEAR_MAX, YEAR_MIN

    for year in (1000000000, YEAR_MIN - 1, YEAR_MAX + 1):
        code, payload = run(capsys, "create", stdin={"name": "Bad", "form": "1040", "year": year,
                                                     "items": [{"identifier": "A01", "document": "W-2"}]})
        assert code == 1 and f"between {YEAR_MIN} and {YEAR_MAX}" in payload["error"], year
        assert not (demo_root / "Bad").exists()
    code, payload = run(capsys, "create", stdin={"name": "Prior", "form": "1040",
                                                 "items": [{"identifier": "A01", "document": "W-2"}]})
    assert code == 0, payload
    code, payload = run(capsys, "rollover", stdin={"prior": payload["created"], "year": "abc"})
    assert code == 1 and "whole number" in payload["error"]


def test_create_refuses_an_identifier_that_cannot_name_a_folder(capsys, demo_root):
    spec = {"name": "Bad", "items": [{"identifier": "A:01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 1
    assert "A:01" in payload["error"] and "folder name" in payload["error"]
    assert not (demo_root / "Bad").exists()  # never leaves a half-built one


def test_create_refuses_a_non_numeric_count_with_a_sentence(capsys, demo_root):
    spec = {"name": "Bad", "items": [
        {"identifier": "A01", "document": "W-2", "expected_count": "two"},
    ]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 1
    assert payload["error"] == f"{COL_EXPECTED_COUNT} for A01 must be a whole number, got 'two'"


def test_create_refuses_a_duplicate_identifier(capsys, demo_root):
    spec = {"name": "Dup", "items": [
        {"identifier": "A01", "document": "One"},
        {"identifier": "a01", "document": "Two"},
    ]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 1
    assert "Duplicate identifier" in payload["error"]
    assert not (demo_root / "Dup").exists()


def test_create_will_not_overwrite_an_existing_engagement(capsys, demo_root):
    spec = {"name": "Twice", "items": [{"identifier": "A01", "document": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 1
    assert "already exists" in payload["error"]


# ----------------------------------------------------------- create + scan ----


def sample_engagement(capsys, demo_root, tmp_path, *names):
    """A 1040 engagement created through the API with sample documents dropped in."""
    from tests.samples import build_samples

    spec = {"name": "Smith Family 2025", "form": "1040", "client": "John Smith",
            "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    engagement = demo_root / "Smith Family 2025"
    samples = tmp_path / "samples"
    build_samples(samples)
    for sample in samples.iterdir():
        if not names or sample.name in names:
            (engagement / SHARED_DIR_NAME / sample.name).write_bytes(sample.read_bytes())
    return engagement


def test_create_then_scan_plays_a_whole_engagement_end_to_end(capsys, demo_root, tmp_path):
    # The client drags every sample into the one folder.
    engagement = sample_engagement(capsys, demo_root, tmp_path)

    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload
    run_result = payload["run"]
    assert run_result["ok"] and not run_result["skipped"]
    assert run_result["file_errors"] == [] and run_result["index_deferred"] is False
    assert run_result["manifest_deferred"] is False
    statuses = {i["identifier"]: i["status"] for i in payload["state"]["items"]}
    assert statuses["A01"] == Status.RECEIVED   # both current-year W-2s, duplicate ignored
    assert statuses["A02"] == Status.PARTIAL    # 2 of 3
    assert statuses["C01"] == Status.RECEIVED   # the 1098
    reviewed = {e["original_name"] for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW}
    assert f"W-2 Jane Smith {PRIOR_YEAR} - old.pdf" in reviewed   # wrong year, never guessed
    assert "vacation photo.jpg" in reviewed
    assert payload["state"]["summary"]["outstanding"] == run_result["outstanding"]
    assert f"{Status.RECEIVED}: 4" in payload["state"]["summary"]["line"]   # A01, B01, C01, D01
    # What the scanner wrote is what the state command reads back.
    rows = {i.identifier: i for i in load_manifest(engagement / MANIFEST_FILENAME)}
    assert rows["A01"].status == Status.RECEIVED


def test_the_apps_pass_appends_the_line_the_scheduled_run_appends(capsys, demo_root, tmp_path):
    """A pass made from the app used to leave no trace at all: the log is the
    command line's, and the button makes the same pass, so it writes the same
    line - appended, never replacing what earlier passes wrote."""
    engagement = sample_engagement(capsys, demo_root, tmp_path, "vacation photo.jpg")

    assert run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))[0] == 0
    first = (demo_root / LOG_FILENAME).read_text(encoding="utf-8")
    assert engagement.name in first

    assert run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))[0] == 0
    second = (demo_root / LOG_FILENAME).read_text(encoding="utf-8")
    assert second.startswith(first)
    assert second.count(engagement.name) == 2


def test_the_apps_pass_regenerates_the_practices_status_page(capsys, demo_root, tmp_path):
    """The page is about the practice, not about the engagement the button was
    pressed on: an engagement nobody scanned is on it too, read rather than run."""
    page = demo_root / STATUS_PAGE_FILENAME
    engagement = sample_engagement(capsys, demo_root, tmp_path, "vacation photo.jpg")
    assert run(capsys, "create", stdin={"name": "Jones Family 2025", "form": "1040",
                                        "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]})[0] == 0
    assert not page.exists()

    assert run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))[0] == 0

    text = page.read_text(encoding="utf-8")
    assert engagement.name in text and "Jones Family 2025" in text
    assert "vacation photo.jpg" in text          # the review queue, across the practice


def test_state_carries_the_path_of_the_practices_status_page(capsys, demo_root):
    """The shell opens only paths the API has reported, so the page's path is
    one of them; the app's Open Status button is that path and nothing else."""
    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(demo_root / "Smith"))
    assert code == 0
    assert payload["paths"]["status"] == str(demo_root / STATUS_PAGE_FILENAME)


def test_item_from_spec_normalizes_extensions():
    item = api.item_from_spec(
        {"identifier": "A01", "document": "W-2", "extensions": ".PDF, Csv"}
    )
    assert item.allowed_extensions == ("pdf", "csv")
    with pytest.raises(ManifestError):
        api.item_from_spec({"identifier": "", "document": "W-2"})


# ------------------------------------------------------- returning clients ----


def test_priors_then_rollover_is_the_desktop_returning_client_path(capsys, demo_root):
    # The wizard's default page: list what is on disk, roll one forward.
    spec = {"name": "Smith Family 2025", "form": "1040",
            "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]}
    assert run(capsys, "create", stdin=spec)[0] == 0

    code, payload = run(capsys, "priors")
    assert code == 0
    assert [p["name"] for p in payload["priors"]] == ["Smith Family 2025"]
    assert payload["priors"][0]["year"] == BASE_YEAR

    code, payload = run(capsys, "rollover", stdin={
        "prior": "Smith Family 2025", "form": "1040", "year": 2026,
    })
    assert code == 0, payload
    assert payload["created"] == "Smith Family 2025 - 2026"
    assert payload["rollover"]["prior_year"] == BASE_YEAR
    assert payload["rollover"]["target_year"] == 2026
    carried = {r["identifier"] for r in payload["rollover"]["carried"]}
    offered = {r["identifier"] for r in payload["rollover"]["offered"]}
    assert "A01" in carried and "E01" in offered and not carried & offered
    periods = {i["identifier"]: i["period"] for i in payload["state"]["items"]}
    assert periods["A01"] == "TY2026"          # shifted with the year
    engagement = demo_root / payload["created"]
    assert (engagement / SHARED_DIR_NAME).is_dir()  # scaffolded, ready to share


def test_rollover_refuses_a_missing_prior(capsys, demo_root):
    code, payload = run(capsys, "rollover", stdin={"prior": "Nobody 2020"})
    assert code == 1
    assert "No manifest found" in payload["error"]


def test_a_name_of_only_illegal_characters_falls_back(capsys, demo_root):
    spec = {"name": "///:::", "form": "1040", "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["created"] == api.default_engagement_name("", default_tax_year(), "1040")


def test_state_shows_statuses_a_locked_excel_deferred(capsys, demo_root):
    from tracker.manifest import StatusUpdate, _save_pending

    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    _save_pending(demo_root / "Smith" / MANIFEST_FILENAME,
                  {"A01": StatusUpdate(status=Status.RECEIVED, file_count=1)})
    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(demo_root / "Smith"))
    assert code == 0
    assert payload["pending_statuses"] == 1
    assert payload["items"][0]["status"] == Status.RECEIVED


# ------------------------------------------------------------ needs review ----


def test_assign_files_a_parked_document_and_rescans(capsys, demo_root, tmp_path):
    engagement = sample_engagement(capsys, demo_root, tmp_path,
                                   "Mortgage Notes.docx", "Form 1098 Mortgage Interest.pdf")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0
    parked = [e for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW]
    assert [e["original_name"] for e in parked] == ["Mortgage Notes.docx"]

    code, payload = run(capsys, "assign", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": parked[0]["pbc_location"], "identifier": "D01",
                               "keyword": "mortgage notes"})
    assert code == 0, payload
    assigned = payload["assigned"]
    assert assigned["identifier"] == "D01" and assigned["moved_review_copy"] is True
    assert assigned["keyword"] == "mortgage notes" and assigned["scan_note"] == ""
    assert not [e for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW]
    d01 = next(i for i in payload["state"]["items"] if i["identifier"] == "D01")
    assert "mortgage notes" in d01["any_keywords"]
    # The re-scan saw it straight away. D01 accepts pdf/xlsx and this is a
    # .docx, but a person looked at it and filed it: their decision stands,
    # the row is Received, and the note says the rules were not applied.
    from tracker.scanner import ACCEPTED_NOTE

    assert d01["status"] == Status.RECEIVED
    assert ACCEPTED_NOTE.format(n=1) in d01["validation_notes"]


def test_dismiss_records_that_nothing_asks_for_a_parked_document(capsys, demo_root, tmp_path):
    # The app's other half of Needs Review: a file no request asks for stops
    # being a thing to do without anything being deleted, and the word for it
    # comes from the vocabulary the renderer reads, not from the renderer.
    from tracker.filer import DISMISSED_BY_PERSON, NOT_REQUESTED

    engagement = sample_engagement(capsys, demo_root, tmp_path, "Mortgage Notes.docx")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload
    [parked] = [e for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW]

    code, payload = run(capsys, "dismiss", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": parked["pbc_location"], "note": "an IRS notice"})
    assert code == 0, payload
    dismissed = payload["dismissed"]
    assert dismissed["decision"] == NOT_REQUESTED
    assert dismissed["reason"].startswith(DISMISSED_BY_PERSON)
    assert "an IRS notice" in dismissed["reason"]
    assert dismissed["index_deferred"] is False
    # The row is rewritten in place and the working copy is still parked.
    assert not [e for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW]
    [row] = [e for e in payload["state"]["index"] if e["decision"] == NOT_REQUESTED]
    assert (engagement / row["prepared_location"]).is_file()
    assert api._vocab()["decisions"]["dismissed"] == NOT_REQUESTED


def test_unfile_sends_a_filed_document_back_for_review_and_the_status_with_it(capsys, demo_root, tmp_path):
    # The correction people used to make by dragging in Explorer, which the
    # index never learned: the row and the request both follow the file.
    from tracker.filer import UNFILED_BY_PERSON

    engagement = sample_engagement(capsys, demo_root, tmp_path, "Form 1098 Mortgage Interest.pdf")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload
    [filed] = [e for e in payload["state"]["index"] if e["decision"] == FILED]
    request = next(i for i in payload["state"]["items"] if i["identifier"] == filed["identifier"])
    assert request["status"] == Status.RECEIVED

    code, payload = run(capsys, "unfile", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": filed["pbc_location"], "note": "wrong request"})
    assert code == 0, payload
    unfiled = payload["unfiled"]
    assert unfiled["decision"] == NEEDS_REVIEW
    assert unfiled["reason"].startswith(UNFILED_BY_PERSON) and "wrong request" in unfiled["reason"]
    assert unfiled["moved_working_copy"] is True
    assert unfiled["left_filed"] == "" and unfiled["scan_note"] == ""
    # The working copy is parked under the client's own name...
    parked = engagement / unfiled["prepared_location"]
    assert parked.is_file() and parked.name == filed["original_name"]
    assert not (engagement / filed["prepared_location"]).exists()
    # ...and the request it was answering is not Received any more.
    request = next(i for i in payload["state"]["items"] if i["identifier"] == filed["identifier"])
    assert request["status"] != Status.RECEIVED
    assert payload["state"]["summary"]["received"] == 0


def test_unfile_refuses_what_is_not_filed(capsys, demo_root, tmp_path):
    engagement = sample_engagement(capsys, demo_root, tmp_path, "nothing")
    code, payload = run(capsys, "unfile", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": "ghost.pdf"})
    assert code == 1
    assert "nothing in the index is called 'ghost.pdf'" in payload["error"]
    code, payload = run(capsys, "unfile", api.ENGAGEMENT_FLAG, str(engagement), stdin={})
    assert code == 1 and payload["error"]


def test_dismiss_refuses_a_file_the_index_does_not_know(capsys, demo_root, tmp_path):
    engagement = sample_engagement(capsys, demo_root, tmp_path, "nothing")
    code, payload = run(capsys, "dismiss", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": "ghost.pdf"})
    assert code == 1
    assert "nothing in the index is called 'ghost.pdf'" in payload["error"]
    code, payload = run(capsys, "dismiss", api.ENGAGEMENT_FLAG, str(engagement), stdin={})
    assert code == 1 and payload["error"]


# ------------------------------------------------- the review queue, triaged ----


def triage_engagement(capsys, demo_root):
    """One engagement whose drop folder holds a document two requests want.

    Two rows, each with a plain required keyword, and one document that
    says A02's in its title and A01's far down the page: the router files
    neither (it matches more than one request) and the row is parked
    carrying both keywords, where they were said, through the filer's own
    writers. Nothing here types an evidence string.
    """
    from tests.test_scanner import text_pdf

    spec = {"name": "Reed Property 2025", "client": "Ada Reed", "items": [
        {"identifier": "A01", "document": "Mortgage Interest Statement", "period": "TY2025",
         "required_keywords": "mortgage interest", "min_size_kb": 0, "date_pattern": "*"},
        {"identifier": "A02", "document": "Rental Property Statements", "period": "TY2025",
         "required_keywords": "rental income", "min_size_kb": 0, "date_pattern": "*"},
    ]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    engagement = demo_root / "Reed Property 2025"
    text_pdf(engagement / SHARED_DIR_NAME / "packet.pdf", "\n".join([
        "Rental income summary for the year 2025",
        "Lakeside Property Management",
        *[f"Unit {n:03d} rent collected and expenses paid during the year" for n in range(60)],
        "The mortgage interest paid on the property is shown below.",
    ]))
    return engagement


def test_the_state_triages_each_parked_file_best_first_with_the_reason_behind_each(
    capsys, demo_root,
):
    from tracker.review import IDENTITY_UNKNOWN

    engagement = triage_engagement(capsys, demo_root)
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload
    state = payload["state"]

    [triaged] = state["review"]
    [parked] = [e for e in state["index"] if e["decision"] == NEEDS_REVIEW]
    assert triaged["pbc_location"] == parked["pbc_location"], "the queue joins to the index row"
    # Best first, and best is not the manifest's order: A02 said its
    # keyword in the title, A01 said its own far down the same page.
    assert [s["identifier"] for s in triaged["shortlist"]] == ["A02", "A01"]
    assert [s["rank"] for s in triaged["shortlist"]] == sorted(
        s["rank"] for s in triaged["shortlist"]
    )
    assert all(s["rank"][0] == IDENTITY_UNKNOWN for s in triaged["shortlist"])
    # The sentence behind each, whole, as review.py wrote it: the row's own
    # word and where the document said it.
    first, second = triaged["shortlist"]
    assert first["reason"].startswith("A02") and "'rental income' in the title" in first["reason"]
    assert second["reason"].startswith("A01") and "'mortgage interest'" in second["reason"]


def test_a_parked_file_the_evidence_says_nothing_about_is_offered_nothing(
    capsys, demo_root, tmp_path,
):
    # A photo: no request accepts a .jpg, so the record names no candidate
    # and the queue says so rather than nominating the nearest row.
    engagement = sample_engagement(capsys, demo_root, tmp_path, "vacation photo.jpg")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload

    [triaged] = payload["state"]["review"]
    assert triaged["original_name"] == "vacation photo.jpg"
    assert triaged["shortlist"] == [], "no evidence, no suggestion — the person reads it"


def test_dismissing_a_file_takes_it_out_of_the_review_queue(capsys, demo_root, tmp_path):
    from tracker.filer import NOT_REQUESTED

    engagement = sample_engagement(capsys, demo_root, tmp_path, "vacation photo.jpg")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    [parked] = payload["state"]["index"]
    assert [t["pbc_location"] for t in payload["state"]["review"]] == [parked["pbc_location"]]

    code, payload = run(capsys, "dismiss", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": parked["pbc_location"], "note": "a holiday snap"})
    assert code == 0, payload

    # The row is still in the index, said out loud; it is not work any more.
    [row] = [e for e in payload["state"]["index"] if e["decision"] == NOT_REQUESTED]
    assert row["pbc_location"] == parked["pbc_location"]
    assert payload["state"]["review"] == [], "a row nobody asks for is not triaged"


def test_assigning_a_shortlisted_request_files_it_and_the_queue_drops_it(capsys, demo_root):
    engagement = triage_engagement(capsys, demo_root)
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    [triaged] = payload["state"]["review"]
    best = triaged["shortlist"][0]["identifier"]

    code, payload = run(capsys, "assign", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": triaged["pbc_location"], "identifier": best})
    assert code == 0, payload

    assert payload["assigned"]["identifier"] == best
    assert payload["state"]["review"] == [], "filed is not parked"
    [filed] = [e for e in payload["state"]["index"] if e["decision"] == FILED]
    assert filed["identifier"] == best


def test_the_vocabulary_carries_every_word_the_review_card_shows(capsys, demo_root):
    from tracker.review import (
        IDENTIFIER_SEPARATOR,
        MAX_SUGGESTIONS,
        NOTHING_SUGGESTED,
        PLACE_WORDS,
    )

    vocab = run(capsys, "list")[1]["vocab"]

    # The card's own buttons and headings...
    labels = vocab["review_labels"]
    assert labels["file"] == api.FILE_LABEL and labels["dismiss"] == api.DISMISS_LABEL
    assert labels["file_anyway"] == api.FILE_ANYWAY_LABEL
    assert labels["suggested"] == api.SUGGESTED_HEADING
    assert labels["other_requests"] == api.OTHER_REQUESTS_HEADING
    # ...and every word tracker.review owns, from tracker.review.
    assert vocab["triage"] == {
        "nothing_suggested": NOTHING_SUGGESTED,
        "identifier_separator": IDENTIFIER_SEPARATOR,
        "places": dict(PLACE_WORDS),
        "max_suggestions": MAX_SUGGESTIONS,
    }
    # Nothing the card shows is typed in the renderer: every one of these
    # reaches the screen through vocab, never as a literal of its own.
    renderer = (Path(__file__).resolve().parent.parent / "app" / "renderer" / "app.js").read_text(
        encoding="utf-8"
    )
    for word in (*labels.values(), NOTHING_SUGGESTED, *PLACE_WORDS.values()):
        assert f'"{word}"' not in renderer and f"'{word}'" not in renderer, word


def test_assign_refuses_a_bad_request_with_a_sentence(capsys, demo_root, tmp_path):
    engagement = sample_engagement(capsys, demo_root, tmp_path, "nothing")
    code, payload = run(capsys, "assign", api.ENGAGEMENT_FLAG, str(engagement),
                        stdin={"original": "ghost.pdf", "identifier": "A01"})
    assert code == 1
    assert "nothing in the index is called 'ghost.pdf'" in payload["error"]


# --------------------------------------------------------- engagement sheet ----


def test_create_writes_the_engagement_sheet_the_scheduled_run_reads(capsys, demo_root):
    from tracker.registry import discover_engagements

    spec = {"name": "Smith Family 2025", "form": "1040", "client": "John Smith",
            "link": "https://drive.example/abc", "due": "2026-04-15",
            "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["state"]["engagement"]["client"] == "John Smith"
    assert payload["state"]["engagement"]["due"] == "2026-04-15"
    [found] = discover_engagements(demo_root).engagements
    assert found.client == "John Smith" and found.link == "https://drive.example/abc"
    from tracker.settings import firm
    assert found.firm == firm() and found.reminders is True
    assert found.name == ""          # the folder is the name; nothing to drift


def test_create_records_the_catalog_the_wizard_chose_and_state_carries_it(capsys, demo_root):
    """Decision 86: the engagement says which checklist it was cut from, on
    the sheet a person opens and in the state the app draws from."""
    from tracker.manifest import load_engagement_info

    spec = {"name": "Willow Inc 2025", "form": "1120S", "client": "Willow Inc",
            "items": [t for t in api.FORM_TEMPLATES["1120S"] if t["core"]]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["state"]["engagement"]["form"] == "1120S"
    assert load_engagement_info(demo_root / "Willow Inc 2025" / MANIFEST_FILENAME).form == "1120S"


def test_an_engagement_created_without_a_form_says_nothing_rather_than_guessing(capsys, demo_root):
    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["state"]["engagement"]["form"] == ""


def test_rollover_carries_the_catalog_the_prior_was_cut_from(capsys, demo_root):
    spec = {"name": "Smith 2025", "form": "1040", "client": "John Smith",
            "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    code, payload = run(capsys, "rollover", stdin={"prior": "Smith 2025", "year": 2026})
    assert code == 0, payload
    assert payload["state"]["engagement"]["form"] == "1040"


def test_rollover_carries_the_client_but_not_last_years_link_or_due(capsys, demo_root):
    spec = {"name": "Smith 2025", "form": "1040", "client": "John Smith",
            "link": "https://drive.example/old", "due": "2026-04-15", "sender": "Jason",
            "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    code, payload = run(capsys, "rollover", stdin={"prior": "Smith 2025", "year": 2026})
    assert code == 0, payload
    info = payload["state"]["engagement"]
    assert info["client"] == "John Smith" and info["sender"] == "Jason"
    assert info["link"] == "" and info["due"] == ""
    assert info["name"] == ""                      # the folder is the name
    assert payload["created"] == api.ROLLOVER_NAME_PATTERN.format(prior="Smith 2025", year=2026)


def test_a_bad_due_date_is_a_sentence(capsys, demo_root):
    spec = {"name": "X", "due": "next friday", "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 1
    assert payload["error"] == "Due date must be YYYY-MM-DD, got 'next friday'"
    assert not (demo_root / "X").exists()


# ---------------------------------------------------- check, lock and names ----


def test_check_reports_problems_and_warnings_with_rows(capsys, demo_root):
    from openpyxl import load_workbook

    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2", "required_keywords": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    manifest = demo_root / "Smith" / MANIFEST_FILENAME
    code, payload = run(capsys, "check", api.ENGAGEMENT_FLAG, str(demo_root / "Smith"))
    assert code == 0 and payload["ok"] and payload["warnings"] == []

    wb = load_workbook(manifest)
    wb[SHEET_NAME].cell(row=2, column=col(COL_EXPECTED_COUNT), value="two")
    wb.save(manifest)
    code, payload = run(capsys, "check", api.ENGAGEMENT_FLAG, str(demo_root / "Smith"))
    assert code == 0 and payload["ok"] is False
    assert payload["problems"] == [f"Row 2: {COL_EXPECTED_COUNT} must be a whole number, got 'two'"]


def test_state_shows_the_lock_and_unlock_clears_only_a_stale_one(capsys, demo_root):
    import os

    from tracker.locking import LOCK_FILENAME

    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    engagement = demo_root / "Smith"
    assert run(capsys, "state", api.ENGAGEMENT_FLAG, str(engagement))[1]["lock"] is None

    lock = engagement / LOCK_FILENAME
    lock.write_text(lock_line(os.getpid(), dt.datetime(2026, 3, 14, 7, 3)), encoding="utf-8")
    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(engagement))
    assert payload["lock"] == {"started": "2026-03-14T07:03:00", "age_minutes": 0, "stale": False,
                               "stale_after_minutes": STALE_LOCK_SECONDS // 60}
    code, payload = run(capsys, "unlock", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 1 and "may still be going" in payload["error"]

    old = (dt.datetime.now() - dt.timedelta(seconds=STALE_LOCK_SECONDS + 1)).timestamp()
    os.utime(lock, (old, old))
    code, payload = run(capsys, "unlock", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0 and payload["cleared"] and payload["state"]["lock"] is None


def test_state_reads_a_corrupt_sidecar_without_moving_it(capsys, demo_root):
    # Showing an engagement is a read. A sidecar the app cannot parse stays
    # where it is for the next real run to move aside as evidence.
    from tracker.filer import INDEX_PENDING_FILENAME
    from tracker.manifest import CORRUPT_SUFFIX, pending_path

    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    engagement = demo_root / "Smith"
    manifest_sidecar = pending_path(engagement / MANIFEST_FILENAME)
    index_sidecar = engagement / INDEX_PENDING_FILENAME
    manifest_sidecar.write_text("{not json", encoding="utf-8")
    index_sidecar.write_text("{not json", encoding="utf-8")

    code, payload = run(capsys, "state", api.ENGAGEMENT_FLAG, str(engagement))

    assert code == 0 and payload["pending_statuses"] == 0 and payload["index"] == []
    assert manifest_sidecar.exists() and index_sidecar.exists()
    assert list(engagement.glob(f"*{CORRUPT_SUFFIX}")) == []


def test_a_new_client_engagement_is_named_from_client_year_and_form(capsys, demo_root):
    spec = {"form": "1040", "client": "Smith Family",
            "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["created"] == api.default_engagement_name("Smith Family", default_tax_year(), "1040")
    code, payload = run(capsys, "templates")
    assert payload["default_year"] == default_tax_year()


# ------------------------------------------------- the calendar and the prior ----


def test_create_shifts_the_checklist_to_the_engagements_year(capsys, demo_root):
    spec = {"form": "1040", "client": "Smith", "year": 2027,
            "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["created"] == api.default_engagement_name("Smith", 2027, "1040")
    periods = {i["identifier"]: i["period"] for i in payload["state"]["items"]}
    assert periods["A01"] == "TY2027" and periods["B01"] == "TY2026"


def test_templates_carries_the_calendars_default_year(capsys):
    code, payload = run(capsys, "templates")
    assert payload["default_year"] == default_tax_year()
    assert "years" not in payload          # one number, not one per form


def test_rollover_retires_the_prior_in_the_priors_list(capsys, demo_root):
    spec = {"name": "Smith 2025", "form": "1040", "client": "John",
            "items": [t for t in api.FORM_TEMPLATES["1040"] if t["core"]]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    code, payload = run(capsys, "rollover", stdin={"prior": "Smith 2025", "year": 2026})
    assert code == 0, payload
    assert payload["state"]["engagement"]["rolled_from"].endswith("Smith 2025")
    code, payload = run(capsys, "priors")
    by_name = {p["name"]: p for p in payload["priors"]}
    assert by_name["Smith 2025"]["superseded_by"] == api.ROLLOVER_NAME_PATTERN.format(prior="Smith 2025", year=2026)
    assert by_name[api.ROLLOVER_NAME_PATTERN.format(prior="Smith 2025", year=2026)]["superseded_by"] == ""


def test_the_apps_pass_is_the_runners_pass(capsys, demo_root):
    # A row added in Excel gets its folder from the scan button, exactly as the
    # scheduled run would give it: one definition of a pass.
    from openpyxl import load_workbook

    spec = {"name": "Smith", "items": [{"identifier": "A01", "document": "W-2"}]}
    assert run(capsys, "create", stdin=spec)[0] == 0
    engagement = demo_root / "Smith"
    wb = load_workbook(engagement / MANIFEST_FILENAME)
    wb[SHEET_NAME].append(row(**{
        COL_IDENTIFIER: "Z01", COL_DOCUMENT: "Rental Records", COL_PERIOD: "TY2025",
        COL_EXPECTED_COUNT: 1, COL_ALLOWED_EXTENSIONS: "pdf", COL_MIN_SIZE_KB: 5,
        COL_ANY_KEYWORDS: "schedule e",
    }))
    wb.save(engagement / MANIFEST_FILENAME)
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload
    assert any(p.name.startswith("Z01") for p in (engagement / PREPARED_DIR_NAME).iterdir())
    assert payload["run"]["warnings"] == []

    # A lock held by another run is reported as skipped, not as an error.
    from tracker.locking import LOCK_FILENAME
    (engagement / LOCK_FILENAME).write_text(lock_line(os.getpid(), dt.datetime(2026, 3, 14, 7, 3)), encoding="utf-8")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0 and payload["run"]["skipped"].startswith("another run")
    (engagement / LOCK_FILENAME).unlink()

    # A manifest typo stops the pass with its row, and is a JSON error the app can show.
    wb = load_workbook(engagement / MANIFEST_FILENAME)
    wb[SHEET_NAME].cell(row=2, column=col(COL_DATE_PATTERN), value="(unclosed")
    wb.save(engagement / MANIFEST_FILENAME)
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 1 and payload["error"].startswith(f"Row 2: {COL_DATE_PATTERN}")


# ------------------------------------------------------------------ the root ----


def test_without_a_root_the_app_is_told_to_set_one(capsys, tmp_path, monkeypatch):
    from tracker.settings import ENV_SETTINGS_DIR

    monkeypatch.setenv(ENV_SETTINGS_DIR, str(tmp_path / "app"))
    code, payload = run(capsys, "list")
    assert code == 0 and payload["needs_root"] is True and payload["engagements"] == []
    code, payload = run(capsys, "create", stdin={"items": [{"identifier": "A01", "document": "W-2"}]})
    assert code == 1 and "where your clients live" in payload["error"]


def test_set_root_records_the_folder_the_job_will_walk(capsys, tmp_path, monkeypatch):
    from tracker.settings import ENV_SETTINGS_DIR, clients_root

    monkeypatch.setenv(ENV_SETTINGS_DIR, str(tmp_path / "app"))
    clients = tmp_path / "Clients"
    clients.mkdir()
    code, payload = run(capsys, "set-root", stdin={"root": str(clients)})
    assert code == 0 and payload["root"] == str(clients)
    assert clients_root() == clients
    code, payload = run(capsys, "set-root", stdin={"root": str(tmp_path / "nope")})
    assert code == 1 and "not a folder" in payload["error"]
    code, payload = run(capsys, "settings")
    assert payload["root"] == str(clients) and payload["exists"] is True


def test_the_firm_is_typed_once_in_settings_and_signs_every_engagement(capsys, demo_root):
    from tracker.settings import firm, set_firm

    set_firm("Park & Daughters CPA")
    assert firm() == "Park & Daughters CPA"
    spec = {"name": "First", "items": [{"identifier": "A01", "document": "W-2"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0 and payload["state"]["engagement"]["firm"] == "Park & Daughters CPA"
    # set-root can set it too, once, at first launch.
    code, payload = run(capsys, "set-root", stdin={"root": str(demo_root), "firm": "New Name LLP"})
    assert code == 0 and payload["firm"] == "New Name LLP" == firm()
    assert run(capsys, "list")[1]["vocab"]["firm"] == "New Name LLP"


def test_install_schedule_uses_the_same_root_as_the_app(capsys, demo_root, monkeypatch):
    import tracker.api as api_module

    calls = []
    monkeypatch.setattr(api_module, "install_task", lambda xml, name=TASK_NAME: (calls.append(xml), ["schtasks", "/create", "/xml", str(xml)])[1])
    code, payload = run(capsys, "install-schedule", stdin={"start": "06:30", "every": 60})
    assert code == 0, payload
    assert payload["root"] == str(demo_root)
    xml = Path(payload["xml"]).read_text(encoding=SCHEDULE_XML_ENCODING)
    assert str(demo_root) in xml and "T06:30:00" in xml and "PT60M" in xml
    assert calls == [Path(payload["xml"])]


def test_the_packaged_app_installs_a_schedule_against_its_own_executable(capsys, demo_root, monkeypatch, tmp_path):
    import sys

    import tracker.api as api_module
    from tracker.runner import LOG_FLAG, RUNNER_MODE_FLAG

    exe = tmp_path / "package" / "resources" / "api" / "api.exe"
    exe.parent.mkdir(parents=True)
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(exe))
    monkeypatch.setattr(api_module, "install_task", lambda xml, name=TASK_NAME: ["schtasks"])

    code, payload = run(capsys, "install-schedule", stdin={})
    assert code == 0, payload
    assert payload["frozen"] is True
    xml = Path(payload["xml"]).read_text(encoding=SCHEDULE_XML_ENCODING)
    assert f"<Command>{exe}</Command>" in xml
    assert f"<Arguments>{RUNNER_MODE_FLAG} " in xml and LOG_FLAG in xml and str(demo_root) in xml
    assert f"<WorkingDirectory>{exe.parent}</WorkingDirectory>" in xml
    assert "-m tracker.runner" not in xml


# ------------------------------------------------------------- vocabulary ----


def test_the_renderer_gets_its_vocabulary_from_the_api(capsys, demo_root):
    from tracker import STANDING_RULES
    from tracker.filer import DUPLICATE, FILED, NEEDS_REVIEW, NOT_REQUESTED
    from tracker.manifest import DEFAULT_EXTENSIONS, Override, Status
    from tracker.rollover import CARRIED_SHEET
    from tracker.runner import DRAFT_WEEKDAY, WEEKDAY_NAMES
    from tracker.scheduling import DEFAULT_REPEAT_MINUTES, DEFAULT_START

    code, payload = run(capsys, "list")
    vocab = payload["vocab"]
    assert [s["value"] for s in vocab["statuses"]] == list(Status.ALL)
    assert {s["key"] for s in vocab["statuses"]} == {api._slug(s) for s in Status.ALL}
    assert vocab["overrides"] == {"accepted": Override.ACCEPTED, "waived": Override.WAIVED}
    assert vocab["decisions"] == {"filed": FILED, "needs_review": NEEDS_REVIEW,
                                  "duplicate": DUPLICATE, "dismissed": NOT_REQUESTED}
    assert vocab["review_labels"] == {
        "dismiss": api.DISMISS_LABEL, "dismiss_note": api.DISMISS_NOTE_HINT,
        "dismissed_heading": api.DISMISSED_HEADING, "file": api.FILE_LABEL,
        "file_anyway": api.FILE_ANYWAY_LABEL,
        "unfile": api.UNFILE_LABEL, "unfile_note": api.UNFILE_NOTE_HINT,
        "filed_heading": api.FILED_HEADING,
        "suggested": api.SUGGESTED_HEADING, "other_requests": api.OTHER_REQUESTS_HEADING,
    }
    assert vocab["default_extensions"] == ", ".join(DEFAULT_EXTENSIONS)
    assert vocab["carried_sheet"] == CARRIED_SHEET
    assert vocab["unscanned_key"] == api._slug(vocab["unscanned_label"])
    assert vocab["commands"] == sorted(api.COMMANDS)
    assert [r["headline"] for r in vocab["rules"]] == [h for h, _ in STANDING_RULES]
    assert vocab["schedule"] == {"start": DEFAULT_START, "every": DEFAULT_REPEAT_MINUTES,
                                 "draft_day": WEEKDAY_NAMES[DRAFT_WEEKDAY],
                                 "task_name": TASK_NAME}


def test_every_chip_class_the_vocabulary_implies_exists_in_the_stylesheet(capsys, demo_root):
    css = (Path(__file__).resolve().parent.parent / "app" / "renderer" / "style.css").read_text(encoding="utf-8")
    vocab = run(capsys, "list")[1]["vocab"]
    for status in vocab["statuses"]:
        assert f".chip-{status['key']} " in css or f".chip-{status['key']}{{" in css.replace(" ", ""), status


def test_the_new_client_name_rule_lives_in_python_and_uses_the_form_label(capsys, demo_root):
    spec = {"form": "1120S", "client": "Acme", "year": 2026,
            "items": [{"identifier": "A01", "document": "Trial Balance", "any_keywords": "trial balance"}]}
    code, payload = run(capsys, "create", stdin=spec)
    assert code == 0, payload
    assert payload["created"] == api.default_engagement_name("Acme", 2026, "1120S")


def test_priors_carry_next_year_and_the_index_carries_candidates(capsys, demo_root, tmp_path):
    engagement = sample_engagement(capsys, demo_root, tmp_path, f"W-2 Jane Smith {PRIOR_YEAR} - old.pdf")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    [parked] = [e for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW]
    assert parked["candidates"] == ["A01"]
    assert parked["filed_as"] == f"W-2 Jane Smith {PRIOR_YEAR} - old.pdf"
    code, payload = run(capsys, "priors")
    [prior] = payload["priors"]
    assert prior["year"] == BASE_YEAR and prior["next_year"] == BASE_YEAR + 1


def test_the_state_carries_each_rows_evidence_as_data_not_as_a_string(capsys, demo_root, tmp_path):
    from tracker.content_check import EVIDENCE_PLACES, EVIDENCE_RULES, RULE_REQUIRED, WHERE_TITLE

    engagement = sample_engagement(capsys, demo_root, tmp_path, f"W-2 Jane Smith {PRIOR_YEAR} - old.pdf")
    code, payload = run(capsys, "scan", api.ENGAGEMENT_FLAG, str(engagement))
    assert code == 0, payload
    [parked] = [e for e in payload["state"]["index"] if e["decision"] == NEEDS_REVIEW]
    # Parsed by the tracker, never split by the renderer: a record keyed by
    # candidate, each entry naming its rule, its term and where it was said.
    assert [(e["rule"], e["term"], e["where"]) for e in parked["evidence"]["A01"]][:1] == [
        (RULE_REQUIRED, "W-2", WHERE_TITLE),
    ]
    assert all(e["rule"] in EVIDENCE_RULES for e in parked["evidence"]["A01"])
    code, listed = run(capsys, "list")
    vocab = listed["vocab"]["evidence"]
    assert vocab["rules"] == list(EVIDENCE_RULES) and vocab["places"] == list(EVIDENCE_PLACES)
    assert RULE_REQUIRED in vocab["rules"] and WHERE_TITLE in vocab["places"]


def test_install_schedule_defaults_come_from_scheduling(capsys, demo_root, monkeypatch):
    import tracker.api as api_module
    from tracker.scheduling import DEFAULT_REPEAT_MINUTES, DEFAULT_START

    monkeypatch.setattr(api_module, "install_task", lambda xml, name=TASK_NAME: ["schtasks"])
    code, payload = run(capsys, "install-schedule", stdin={})
    assert code == 0, payload
    assert payload["start"] == DEFAULT_START and payload["every"] == DEFAULT_REPEAT_MINUTES
    assert payload["draft_day"] == WEEKDAY_NAMES[DRAFT_WEEKDAY]
