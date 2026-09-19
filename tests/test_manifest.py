"""Tests for tracker/manifest.py — the manifest layer, no OneDrive needed."""

import datetime as dt
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookClass

from tests.samples import col
from tracker import reasons
from tracker.manifest import (
    COL_ALLOWED_EXTENSIONS,
    COL_DATE_PATTERN,
    COL_DOCUMENT,
    COL_EXPECTED_COUNT,
    COL_IDENTIFIER,
    COL_MANUAL_OVERRIDE,
    COL_STATUS,
    CORRUPT_SUFFIX,
    DEFAULT_MIN_SIZE_KB,
    ENGAGEMENT_LABELS,
    ENGAGEMENT_SHEET_NAME,
    HEADERS,
    NO,
    SHEET_NAME,
    SUMMARY_EMPTY,
    SUMMARY_SEPARATOR,
    TEMP_SUFFIX,
    UNSCANNED_LABEL,
    YES,
    ManifestError,
    Override,
    RequestItem,
    Status,
    StatusUpdate,
    check_narrowing_names,
    create_template,
    entity_keyword,
    load_manifest,
    narrowing_rows,
    pending_path,
    pending_updates,
    temp_path_for,
    write_statuses,
)
from tracker.scaffold import MANIFEST_FILENAME

SAMPLE_ITEMS = [
    RequestItem(
        identifier="A01",
        document="Dec 2025 Bank Statement",
        period="Dec 2025",
        allowed_extensions=("pdf",),
        required_keywords=("Chase",),
        date_pattern=r"(?i)december\s+2025|12/\d{1,2}/2025",
    ),
    RequestItem(
        identifier="A02",
        document="Monthly Bank Statements FY2025",
        period="FY2025",
        expected_count=12,
        allowed_extensions=("pdf", ".CSV"),  # dot + case normalize on load
        min_size_kb=10,
        any_keywords=("statement", "account summary"),
    ),
    RequestItem(
        identifier="B01",
        document="Payroll Register Q4",
        period="Q4 2025",
        allowed_extensions=("xlsx", "csv"),
        manual_override=Override.WAIVED,
    ),
]


@pytest.fixture
def manifest(tmp_path):
    return create_template(tmp_path / MANIFEST_FILENAME, SAMPLE_ITEMS)


# ---------------------------------------------------------------- loading ----


def test_template_load_roundtrip(manifest):
    items = load_manifest(manifest)
    assert [i.identifier for i in items] == ["A01", "A02", "B01"]

    a01 = items[0]
    assert a01.document == "Dec 2025 Bank Statement"
    assert a01.period == "Dec 2025"
    assert a01.expected_count == 1          # default
    assert a01.min_size_kb == DEFAULT_MIN_SIZE_KB
    assert a01.allowed_extensions == ("pdf",)
    assert a01.required_keywords == ("Chase",)
    assert a01.any_keywords == ()
    assert a01.date_pattern.startswith("(?i)december")
    assert a01.manual_override == ""
    assert a01.status == ""
    assert a01.received_date is None
    assert a01.row == 2

    a02 = items[1]
    assert a02.expected_count == 12
    assert a02.min_size_kb == 10
    assert a02.allowed_extensions == ("pdf", "csv")  # normalized
    assert a02.any_keywords == ("statement", "account summary")

    assert items[2].manual_override == Override.WAIVED


def test_template_refuses_overwrite(manifest):
    with pytest.raises(ManifestError, match="Refusing to overwrite"):
        create_template(manifest)


def test_missing_manifest(tmp_path):
    with pytest.raises(ManifestError, match="not found"):
        load_manifest(tmp_path / "nope.xlsx")


def test_missing_column_rejected(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for idx, header in enumerate(h for h in HEADERS if h != COL_STATUS):
        ws.cell(row=1, column=idx + 1, value=header)
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(ManifestError, match=f"missing column.*{COL_STATUS}"):
        load_manifest(path)


def _write_cell(path, row, header, value):
    wb = load_workbook(path)
    ws = wb[SHEET_NAME]
    ws.cell(row=row, column=col(header), value=value)
    wb.save(path)


def test_duplicate_identifier_rejected(manifest):
    _write_cell(manifest, 3, COL_IDENTIFIER, "A01")
    with pytest.raises(ManifestError, match="Duplicate identifier 'A01'.*rows 2 and 3"):
        load_manifest(manifest)


def test_bad_regex_rejected(manifest):
    _write_cell(manifest, 2, COL_DATE_PATTERN, "([unclosed")
    with pytest.raises(ManifestError, match="Row 2.*not a valid regex"):
        load_manifest(manifest)


def test_bad_expected_count_rejected(manifest):
    _write_cell(manifest, 2, COL_EXPECTED_COUNT, "twelve")
    with pytest.raises(ManifestError, match="Row 2.*whole number"):
        load_manifest(manifest)


def test_unknown_status_rejected(manifest):
    _write_cell(manifest, 2, COL_STATUS, "Done-ish")
    with pytest.raises(ManifestError, match="Row 2.*Status must be one of"):
        load_manifest(manifest)


def test_unknown_override_rejected(manifest):
    _write_cell(manifest, 2, COL_MANUAL_OVERRIDE, "Maybe")
    with pytest.raises(ManifestError, match="Row 2.*Manual Override"):
        load_manifest(manifest)


def test_blank_rows_skipped(manifest):
    _write_cell(manifest, 6, COL_IDENTIFIER, "C01")  # row 5 left entirely blank
    _write_cell(manifest, 6, COL_DOCUMENT, "Trial Balance")
    items = load_manifest(manifest)
    assert [i.identifier for i in items] == ["A01", "A02", "B01", "C01"]
    assert items[-1].row == 6


# ------------------------------------------------------------- write-back ----


def test_write_statuses_roundtrip(manifest):
    updates = {
        "A01": StatusUpdate(
            status=Status.RECEIVED,
            file_count=1,
            received_date=dt.date(2026, 7, 8),
        ),
        "A02": StatusUpdate(
            status=Status.PARTIAL,
            file_count=7,
            validation_notes="7 of 12 monthly statements received",
        ),
    }
    assert write_statuses(manifest, updates) is True

    items = {i.identifier: i for i in load_manifest(manifest)}
    assert items["A01"].status == Status.RECEIVED
    assert items["A01"].received_date == dt.date(2026, 7, 8)
    assert items["A01"].file_count == 1
    assert items["A02"].status == Status.PARTIAL
    assert items["A02"].file_count == 7
    assert "7 of 12" in items["A02"].validation_notes
    # untouched rows and accountant columns intact
    assert items["B01"].status == ""
    assert items["A01"].required_keywords == ("Chase",)


def test_write_unknown_identifier_dropped(manifest, caplog):
    assert write_statuses(manifest, {"Z99": StatusUpdate(status=Status.MISSING)}) is True
    assert "Z99" in caplog.text
    assert not pending_path(manifest).exists()


def test_locked_write_defers_to_sidecar(manifest, monkeypatch):
    def locked_save(self, filename):
        raise PermissionError(f"[Errno 13] locked: {filename}")

    monkeypatch.setattr(WorkbookClass, "save", locked_save)
    updates = {"A01": StatusUpdate(status=Status.MISSING, validation_notes="no files")}
    assert write_statuses(manifest, updates, retries=2, retry_delay=0.01) is False

    sidecar = pending_path(manifest)
    assert sidecar.exists()
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    assert payload["A01"]["status"] == Status.MISSING


def test_a_manifest_held_the_way_excel_holds_one_defers_and_lands_when_it_is_let_go(
    manifest, held_like_excel,
):
    """``test_locked_write_defers_to_sidecar``'s twin with a real handle.

    That test raises ``PermissionError`` from ``Workbook.save`` by hand and
    proves the retry and the sidecar; this one holds ``_manifest.xlsx`` the
    way Excel holds a workbook it has open and proves what Windows does. With
    the handle held, a read still answers, the write is refused, not a byte
    of the workbook moves, no temp file is left beside it, and the statuses
    wait in the sidecar where every reader overlays them. Let the handle go
    and the next write, with nothing new to say, lands what waited and
    clears the sidecar.
    """
    before = manifest.read_bytes()
    updates = {"A01": StatusUpdate(status=Status.MISSING, validation_notes="no files")}
    with held_like_excel(manifest):
        read = {i.identifier: i for i in load_manifest(manifest)}
        assert read["A01"].status == ""                             # a read, and it answers
        assert write_statuses(manifest, updates, retries=2, retry_delay=0.01) is False
        assert manifest.read_bytes() == before                      # not a byte of it moved
        assert list(manifest.parent.glob(f"*{TEMP_SUFFIX}")) == []
        assert pending_updates(manifest)["A01"].status == Status.MISSING

    assert write_statuses(manifest, {}, retries=2, retry_delay=0.01) is True
    assert not pending_path(manifest).exists()
    assert manifest.read_bytes() != before
    landed = {i.identifier: i for i in load_manifest(manifest)}["A01"]
    assert landed.status == Status.MISSING
    assert landed.validation_notes == "no files"


def test_pending_sidecar_merged_and_cleared(manifest):
    pending_path(manifest).write_text(
        json.dumps(
            {
                "A01": {  # stale: superseded by the new update below
                    "status": Status.MISSING,
                    "file_count": 0,
                    "received_date": None,
                    "validation_notes": "",
                },
                "B01": {  # only in sidecar: must still be applied
                    "status": Status.PENDING_SYNC,
                    "file_count": 1,
                    "received_date": None,
                    "validation_notes": reasons.PENDING_SYNC.format(),
                },
            }
        ),
        encoding="utf-8",
    )
    new = {
        "A01": StatusUpdate(
            status=Status.RECEIVED, file_count=1, received_date=dt.date(2026, 7, 8)
        )
    }
    assert write_statuses(manifest, new) is True
    assert not pending_path(manifest).exists()

    items = {i.identifier: i for i in load_manifest(manifest)}
    assert items["A01"].status == Status.RECEIVED  # new beat stale sidecar
    assert items["B01"].status == Status.PENDING_SYNC
    assert "placeholder" in items["B01"].validation_notes


def test_corrupt_sidecar_quarantined(manifest, caplog):
    pending_path(manifest).write_text("{not json", encoding="utf-8")
    assert write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)}) is True
    assert not pending_path(manifest).exists()
    assert pending_path(manifest).with_suffix(CORRUPT_SUFFIX).exists()
    assert "Unreadable pending sidecar" in caplog.text


# -------------------------------------------------------- identifier safety ----


def _manifest_with_identifiers(tmp_path, *identifiers):
    items = [
        RequestItem(identifier=ident, document=f"Doc {n}")
        for n, ident in enumerate(identifiers, start=1)
    ]
    return create_template(tmp_path / MANIFEST_FILENAME, items)


@pytest.mark.parametrize("identifier", ["A:01", "A/01", "A?1", 'B"1', "A01.", "A<1>"])
def test_identifier_that_cannot_name_a_folder_is_rejected(tmp_path, identifier):
    # The identifier is the folder-name prefix the scanner matches back on.
    # One the filesystem would alter is a permanent "folder not found".
    path = _manifest_with_identifiers(tmp_path, identifier)
    with pytest.raises(ManifestError, match="Identifier"):
        load_manifest(path)


def test_identifiers_differing_only_by_case_are_duplicates(tmp_path):
    # Windows folder names are case-insensitive: "A01" and "a01" would fight
    # over one folder, so they are the same identifier.
    path = _manifest_with_identifiers(tmp_path, "A01", "a01")
    with pytest.raises(ManifestError, match="Duplicate identifier"):
        load_manifest(path)


# ------------------------------------------------------------ atomic saves ----


def test_a_crash_mid_save_leaves_the_previous_manifest_intact(manifest, monkeypatch):
    # openpyxl streams straight into the target; a killed task mid-write
    # used to leave a manifest Excel could not open. The save now lands
    # beside the file and is swapped in whole, or not at all.
    def crash(self, filename):
        Path(str(filename)).write_bytes(b"PK\x03\x04 half a zip")
        raise KeyboardInterrupt

    monkeypatch.setattr(WorkbookClass, "save", crash)
    with pytest.raises(KeyboardInterrupt):
        write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)}, retries=1)

    assert not manifest.with_name(manifest.name + TEMP_SUFFIX).exists()
    assert [i.identifier for i in load_manifest(manifest)] == ["A01", "A02", "B01"]


def test_no_temp_file_is_left_beside_the_workbook_under_any_name(manifest, monkeypatch):
    # The temp name is unique per write, so the claim above has to hold for
    # every name ending in TEMP_SUFFIX, not just one fixed spelling.
    def crash(self, filename):
        Path(str(filename)).write_bytes(b"PK\x03\x04 half a zip")
        raise KeyboardInterrupt

    monkeypatch.setattr(WorkbookClass, "save", crash)
    with pytest.raises(KeyboardInterrupt):
        write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)}, retries=1)
    assert list(manifest.parent.glob(f"*{TEMP_SUFFIX}")) == []


def test_two_writers_never_share_a_temp_name_and_the_walk_ignores_it(manifest):
    from tracker.validators import is_ignored

    first, second = temp_path_for(manifest), temp_path_for(manifest)
    assert first != second
    assert first.parent == manifest.parent
    for temp in (first, second):
        assert temp.name.endswith(TEMP_SUFFIX)
        assert is_ignored(temp)                   # a stranded temp is never a document


def test_the_pending_sidecar_is_written_whole_or_not_at_all(manifest, monkeypatch):
    sidecar = pending_path(manifest)
    before = json.dumps({"B01": {"status": Status.MISSING, "file_count": 0,
                                 "received_date": None, "validation_notes": ""}})
    sidecar.write_text(before, encoding="utf-8")

    def locked_save(self, filename):
        raise PermissionError(f"[Errno 13] locked: {filename}")

    def refuse_replace(src, dst):
        raise OSError("disk full")

    monkeypatch.setattr(WorkbookClass, "save", locked_save)
    monkeypatch.setattr("tracker.manifest.os.replace", refuse_replace)
    with pytest.raises(OSError, match="disk full"):
        write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)},
                       retries=1, retry_delay=0.01)

    assert sidecar.read_text(encoding="utf-8") == before   # the old sidecar survived
    assert list(manifest.parent.glob(f"*{TEMP_SUFFIX}")) == []


def test_a_second_corrupt_sidecar_never_overwrites_the_first(manifest, caplog):
    first_evidence = "{not json, the first time"
    pending_path(manifest).write_text(first_evidence, encoding="utf-8")
    assert write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)}) is True
    pending_path(manifest).write_text("{not json, the second time", encoding="utf-8")
    assert write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)}) is True

    corrupt = sorted(p.name for p in manifest.parent.glob(f"*{CORRUPT_SUFFIX}"))
    assert len(corrupt) == 2
    kept = pending_path(manifest).with_suffix(CORRUPT_SUFFIX)
    assert kept.read_text(encoding="utf-8") == first_evidence


def test_reading_pending_updates_leaves_a_corrupt_sidecar_where_it_is(manifest, caplog):
    # A reader - the reminder, the app's state - must never move a file:
    # looking changes nothing. Only the write that will retry it moves it aside.
    sidecar = pending_path(manifest)
    sidecar.write_text("{not json", encoding="utf-8")
    assert pending_updates(manifest, quarantine=False) == {}
    assert sidecar.exists()
    assert list(manifest.parent.glob(f"*{CORRUPT_SUFFIX}")) == []
    assert "ignored for this read" in caplog.text


# --------------------------------------------------------- engagement sheet ----


def test_the_engagement_sheet_keeps_a_name_typed_like_a_formula_a_name(tmp_path):
    # Decision 58's rule - every cell the tracker fills from data is text -
    # had a hole on this sheet (the ninth reading).
    from tracker.manifest import EngagementInfo, load_engagement_info, write_engagement_info

    path = create_template(tmp_path / MANIFEST_FILENAME, SAMPLE_ITEMS, EngagementInfo(client="=1+1"))
    assert load_engagement_info(path).client == "=1+1"
    write_engagement_info(path, EngagementInfo(client="Jane", sender="=CMD|' /C calc'!A0"))
    assert load_engagement_info(path).sender == "=CMD|' /C calc'!A0"


def test_the_engagement_sheet_round_trips(tmp_path):
    from tracker.manifest import EngagementInfo, load_engagement_info, write_engagement_info

    info = EngagementInfo(client="John Smith", name="Smiths 2025", link="https://drive.example/abc",
                          due=dt.date(2026, 4, 15), sender="Jason Park", firm="J Park",
                          reminders=False, active=True)
    path = create_template(tmp_path / MANIFEST_FILENAME, SAMPLE_ITEMS, info)
    assert load_engagement_info(path) == info
    assert [i.identifier for i in load_manifest(path)] == ["A01", "A02", "B01"]  # Requests untouched

    write_engagement_info(path, EngagementInfo(client="Jane", active=False))
    loaded = load_engagement_info(path)
    assert loaded.client == "Jane" and loaded.active is False and loaded.reminders is True


def test_a_manifest_without_the_sheet_loads_as_defaults(manifest):
    from openpyxl import load_workbook as lw

    from tracker.manifest import EngagementInfo, load_engagement_info

    wb = lw(manifest)
    del wb[ENGAGEMENT_SHEET_NAME]
    wb.save(manifest)
    assert load_engagement_info(manifest) == EngagementInfo()


def test_yes_no_cells_are_forgiving_but_not_guessing(manifest):
    from openpyxl import load_workbook as lw

    from tracker.manifest import load_engagement_info

    def set_cell(label, value):
        wb = lw(manifest)
        for row in wb[ENGAGEMENT_SHEET_NAME].iter_rows(min_row=1, max_col=2):
            if row[0].value == label:
                row[1].value = value
        wb.save(manifest)

    set_cell(ENGAGEMENT_LABELS["reminders"], f"{NO.capitalize()} ")
    assert load_engagement_info(manifest).reminders is False
    set_cell(ENGAGEMENT_LABELS["active"], "TRUE")
    assert load_engagement_info(manifest).active is True
    set_cell(ENGAGEMENT_LABELS["active"], "later")
    with pytest.raises(ManifestError, match=f"{ENGAGEMENT_LABELS['active']} must be {YES} or {NO}"):
        load_engagement_info(manifest)


def test_create_records_the_catalog_the_request_list_was_cut_from(tmp_path):
    """Decision 86: the wizard's choice is written onto the sheet, so an
    engagement can say which checklist it came from without the folder name."""
    from tracker.manifest import load_engagement_info
    from tracker.templates import template_items

    path = create_template(tmp_path / MANIFEST_FILENAME,
                           template_items("1120S", year=2025), form="1120S")
    assert load_engagement_info(path).form == "1120S"


def test_the_form_a_caller_gives_never_clears_the_one_the_sheet_carries(tmp_path):
    from tracker.manifest import EngagementInfo, load_engagement_info

    info = EngagementInfo(client="John Smith", form="1065")
    path = create_template(tmp_path / MANIFEST_FILENAME, SAMPLE_ITEMS, info)
    assert load_engagement_info(path) == info          # no form= given; the info's stands


def test_a_manifest_that_never_recorded_a_form_loads_as_unknown(tmp_path):
    """An engagement made before the cell existed keeps every other cell
    where it was, and its form reads as blank rather than as a guess."""
    from openpyxl import load_workbook as lw

    from tracker.manifest import ENGAGEMENT_FIELDS, load_engagement_info

    path = create_template(tmp_path / MANIFEST_FILENAME, SAMPLE_ITEMS, form="1040")
    before = {label: None for label, _ in ENGAGEMENT_FIELDS}
    wb = lw(path)
    ws = wb[ENGAGEMENT_SHEET_NAME]
    for row in ws.iter_rows(min_row=1, max_col=2):
        if _cell_text(row[0]) in before:
            before[_cell_text(row[0])] = row[0].row
    ws.delete_rows(before[ENGAGEMENT_LABELS["form"]])      # the sheet as it was before
    wb.save(path)

    loaded = load_engagement_info(path)
    assert loaded.form == ""
    wb = lw(path)
    for label, _ in ENGAGEMENT_FIELDS:
        if label == ENGAGEMENT_LABELS["form"]:
            continue
        found = [c.row for c in wb[ENGAGEMENT_SHEET_NAME]["A"] if _cell_text(c) == label]
        assert found == [before[label]], label


def test_a_blank_form_cell_is_unknown_not_a_refusal(manifest):
    """Nothing reads the cell yet, and a blank must never fail a load."""
    from tracker.manifest import check_manifest, load_engagement_info

    assert load_engagement_info(manifest).form == ""      # created without a form
    assert check_manifest(manifest).problems == []


def _cell_text(cell) -> str:
    return "" if cell.value is None else str(cell.value).strip()


def test_add_any_keyword_appends_once_and_names_a_missing_row(manifest):
    from tracker.manifest import add_any_keyword

    assert add_any_keyword(manifest, "A01", "Schedule E") is True
    assert add_any_keyword(manifest, "A01", "schedule e") is False     # case-insensitive
    assert load_manifest(manifest)[0].any_keywords == ("Schedule E",)
    assert add_any_keyword(manifest, "A01", "   ") is False
    with pytest.raises(ManifestError, match="No request 'Z99'"):
        add_any_keyword(manifest, "Z99", "x")


# ------------------------------------------------------ allowed extensions ----


def test_a_blank_allowed_extensions_means_the_safe_default_not_anything(tmp_path):
    from openpyxl import load_workbook as lw

    from tracker.manifest import DEFAULT_EXTENSIONS

    path = create_template(tmp_path / MANIFEST_FILENAME, [RequestItem(identifier="A01", document="W-2")])
    wb = lw(path)
    wb[SHEET_NAME].cell(row=2, column=col(COL_ALLOWED_EXTENSIONS)).value = None   # the accountant cleared the cell
    wb.save(path)
    assert load_manifest(path)[0].allowed_extensions == DEFAULT_EXTENSIONS


def test_accepting_any_file_type_has_to_be_said_with_a_star(tmp_path):
    from openpyxl import load_workbook as lw

    path = create_template(tmp_path / MANIFEST_FILENAME, [RequestItem(identifier="A01", document="W-2")])
    # An item built with no extensions means anything; the template writes "*"
    # so the workbook says so and the loader reads it back the same way.
    wb = lw(path)
    assert wb[SHEET_NAME].cell(row=2, column=col(COL_ALLOWED_EXTENSIONS)).value == "*"
    wb.close()
    assert load_manifest(path)[0].allowed_extensions == ()


def test_a_spec_that_gives_its_extensions_as_a_list_gets_those_extensions(tmp_path):
    # The thirteenth reading: a spec typed as JSON holds ["pdf"], and
    # stringifying it made one extension called "['pdf']" - which no
    # document has, so every document parked against that row for ever.
    from tracker.manifest import DEFAULT_EXTENSIONS, parse_extensions
    from tracker.templates import item_from_spec

    assert parse_extensions(["pdf"]) == ("pdf",)
    assert parse_extensions(("PDF", ".Xlsx")) == ("pdf", "xlsx")
    assert parse_extensions(["*"]) == ()
    assert parse_extensions([]) == DEFAULT_EXTENSIONS
    item = item_from_spec({"identifier": "A01", "document": "W-2", "extensions": ["pdf", "xlsx"],
                           "required_keywords": ["W-2", "wage"]})
    assert item.allowed_extensions == ("pdf", "xlsx")
    assert item.required_keywords == ("W-2", "wage")


# ------------------------------------------------------------------ check ----


def test_check_manifest_reports_problems_with_their_row(manifest):
    from openpyxl import load_workbook as lw

    from tracker.manifest import check_manifest

    assert check_manifest(manifest).ok
    wb = lw(manifest)
    wb[SHEET_NAME].cell(row=3, column=col(COL_DATE_PATTERN), value="(unclosed")   # A02 Date Pattern
    wb.save(manifest)
    result = check_manifest(manifest)
    assert not result.ok
    assert result.problems[0].startswith(f"Row 3: {COL_DATE_PATTERN} is not a valid regex")


def test_check_manifest_warns_about_rows_the_rules_cannot_act_on(tmp_path):
    from tracker.manifest import _save_pending, check_manifest

    path = create_template(tmp_path / MANIFEST_FILENAME, [
        RequestItem(identifier="A01", document="Anything goes"),              # no rule, "*"
        RequestItem(identifier="A02", document="W-2", required_keywords=("W-2",),
                    allowed_extensions=("pdf",)),
        RequestItem(identifier="A03", document="Waived", manual_override=Override.WAIVED),
    ])
    _save_pending(path, {"A02": StatusUpdate(status=Status.RECEIVED)})
    result = check_manifest(path)
    assert result.ok
    assert [w[:14] for w in result.warnings] == ["Row 2 (A01): n", "Row 2 (A01): A", pending_path(path).name[:14]]
    assert "never be filed automatically" in result.warnings[0]
    assert "any file type counts" in result.warnings[1]
    assert "close Excel and re-scan" in result.warnings[2]


# --------------------------------------------------------------- summary ----


def test_summarize_is_the_one_count():
    from tracker.manifest import summarize

    items = [
        RequestItem(identifier="A01", document="a", status=Status.RECEIVED),
        RequestItem(identifier="A02", document="b", status=Status.MISSING),
        RequestItem(identifier="A03", document="c", status=Status.PARTIAL),
        RequestItem(identifier="A04", document="d", status=Status.FAILED,
                    manual_override=Override.ACCEPTED),          # signed off: counts as in
        RequestItem(identifier="A05", document="e", status=Status.MISSING,
                    manual_override=Override.WAIVED),            # nobody is waiting
        RequestItem(identifier="A06", document="f"),             # never scanned
        RequestItem(identifier="A07", document="g", status=Status.PENDING_SYNC),
    ]
    summary = summarize(items)
    assert summary.total == 6 and summary.waived == 1
    assert summary.received == 2
    assert summary.outstanding == 2
    assert summary.unscanned == 1
    assert summary.counts == {Status.RECEIVED: 2, Status.MISSING: 1, Status.PARTIAL: 1, Status.PENDING_SYNC: 1}
    assert summary.line == SUMMARY_SEPARATOR.join([
        f"{Status.MISSING}: 1", f"{Status.PARTIAL}: 1", f"{Status.PENDING_SYNC}: 1",
        f"{Status.RECEIVED}: 2", f"{UNSCANNED_LABEL}: 1", f"{Override.WAIVED}: 1",
    ])
    assert summarize([]).line == SUMMARY_EMPTY


# ------------------------------------------------------- the derived year ----


def test_a_period_with_a_year_implies_the_year_check(tmp_path):
    import re

    from tracker.manifest import derived_date_pattern

    assert derived_date_pattern("TY2025") == r"(?i)\b2025\b"
    assert derived_date_pattern("As of 12/31/2025") == r"(?i)\b2025\b"
    # A Period that names a month asks for that month: a November statement
    # is not the December one, and "12/31/2025" or "December 2025" both are.
    december = derived_date_pattern("Dec 2025")
    for text in ("Statement period 12/01/2025 - 12/31/2025", "December 2025 statement", "as of 12/31/2025"):
        assert re.search(december, text), text
    for text in ("Statement period 11/01/2025 - 11/30/2025", "Nov 2025", "Tax Year 2025"):
        assert not re.search(december, text), text
    assert derived_date_pattern("Current") == ""
    assert derived_date_pattern("Acct 120250") == ""      # not a year

    path = create_template(tmp_path / MANIFEST_FILENAME, [
        RequestItem(identifier="A01", document="W-2", period="TY2025", required_keywords=("W-2",)),
        RequestItem(identifier="A02", document="Trust deed", period="Current", required_keywords=("trust",)),
        RequestItem(identifier="A03", document="Typed", period="TY2025", required_keywords=("x",),
                    date_pattern=r"2025|2026"),
    ])
    rows = {i.identifier: i for i in load_manifest(path)}
    assert rows["A01"].date_pattern == r"(?i)\b2025\b" and rows["A01"].date_pattern_derived
    assert rows["A02"].date_pattern == "" and not rows["A02"].date_pattern_derived
    assert rows["A03"].date_pattern == r"2025|2026" and not rows["A03"].date_pattern_derived


def test_a_star_in_date_pattern_means_no_year_check(tmp_path):
    from openpyxl import load_workbook as lw

    path = create_template(tmp_path / MANIFEST_FILENAME, [
        RequestItem(identifier="A01", document="W-2", period="TY2025", required_keywords=("W-2",)),
    ])
    wb = lw(path)
    wb[SHEET_NAME].cell(row=2, column=col(COL_DATE_PATTERN), value="*")
    wb.save(path)
    [row] = load_manifest(path)
    assert row.date_pattern == "" and not row.date_pattern_derived


def test_a_derived_year_is_a_check_not_a_reason_to_route():
    from tracker.manifest import has_routing_rules

    typed = RequestItem(identifier="A01", document="x", date_pattern=r"\b2025\b")
    derived = RequestItem(identifier="A02", document="x", date_pattern=r"\b2025\b",
                          date_pattern_derived=True)
    keyed = RequestItem(identifier="A03", document="x", any_keywords=("w-2",))
    assert has_routing_rules(typed) and has_routing_rules(keyed)
    assert not has_routing_rules(derived)


def test_check_manifest_warns_when_a_keyword_names_a_family_of_forms(tmp_path):
    from tracker.manifest import BARE_FORM_NUMBER_WARNING, FORM_FAMILIES, check_manifest

    path = tmp_path / "m.xlsx"
    create_template(path, [
        RequestItem(identifier="B01", document="1099s", allowed_extensions=("pdf",), any_keywords=("1099",)),
        RequestItem(identifier="C01", document="Mortgage", allowed_extensions=("pdf",), required_keywords=("1098",)),
    ])
    check = check_manifest(path)
    assert check.ok
    assert check.warnings == [BARE_FORM_NUMBER_WARNING.format(
        row=2, identifier="B01", keyword="1099", example=FORM_FAMILIES["1099"])]


def test_a_failed_save_reports_its_own_error_not_a_locked_temp_file(tmp_path, monkeypatch):
    # openpyxl leaves the half-written zip open when save() raises; on
    # Windows the temp then cannot be deleted. That must not turn a full
    # disk into "open in Excel" retried five times.
    import tracker.manifest as manifest_module
    from tracker.manifest import atomic_replacement

    target = tmp_path / "x.xlsx"
    target.write_bytes(b"before")
    real_unlink = manifest_module.Path.unlink

    def held(self, *args, **kwargs):
        if self.name.endswith(manifest_module.TEMP_SUFFIX):
            raise PermissionError("[WinError 32] still open")
        return real_unlink(self, *args, **kwargs)
    monkeypatch.setattr(manifest_module.Path, "unlink", held)
    with pytest.raises(OSError, match="No space left"):
        with atomic_replacement(target) as temp:
            temp.write_bytes(b"half")
            raise OSError(28, "No space left on device")
    assert target.read_bytes() == b"before"


def test_a_formula_in_the_keyword_cell_is_refused_not_rewritten(manifest):
    from openpyxl import load_workbook

    from tracker.manifest import COL_ANY_KEYWORDS, add_any_keyword

    wb = load_workbook(manifest)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index(COL_ANY_KEYWORDS) + 1, value="=B2")
    wb.save(manifest)
    with pytest.raises(ManifestError, match="holds a formula"):
        add_any_keyword(manifest, "A01", "wages")
    assert load_workbook(manifest)[SHEET_NAME].cell(row=2, column=headers.index(COL_ANY_KEYWORDS) + 1).value == "=B2"


# ------------------------------------------- a row per issuer (decision 93) ----

K1_WORDS = ("partner's share of income", "member's share of income")


def k1_rows(*issuers):
    """The generic K-1 row and one row per (identifier, entity) after it."""
    rows = [RequestItem(identifier="F01", document="Schedule K-1s Received",
                        allowed_extensions=("pdf",), any_keywords=K1_WORDS)]
    rows += [RequestItem(identifier=identifier, document=f"Schedule K-1 - {entity}",
                         allowed_extensions=("pdf",), required_keywords=(entity,),
                         any_keywords=K1_WORDS)
             for identifier, entity in issuers]
    return rows


@pytest.mark.parametrize("typed, expected", [
    ("Ashford Holdings, L.P.", "Ashford Holdings LP"),
    ("Ashford Holdings LP", "Ashford Holdings LP"),
    ("  Ashford   Holdings  ", "Ashford Holdings"),
    ("Ashford Holdings | Birch Lane", "Ashford Holdings Birch Lane"),
    ("Ashford + Birch", "Ashford Birch"),
    ("O'Hara & Sons, Inc.", "O'Hara & Sons Inc"),
])
def test_an_entity_name_becomes_one_keyword_a_cell_can_hold(typed, expected):
    """A comma would split the cell into two required keywords; the grammar's
    own two characters would turn a name into alternatives."""
    assert entity_keyword(typed) == expected


def test_two_typings_of_one_entity_make_one_row():
    assert entity_keyword("Ashford Holdings, L.P.") == entity_keyword("Ashford Holdings LP")


def test_an_issuer_name_typed_with_commas_loads_as_one_keyword(tmp_path):
    """The point of normalising before the cell is written: read back, the
    row asks for one name and not for two."""
    from tracker.templates import issuer_row, item_from_spec

    path = create_template(tmp_path / MANIFEST_FILENAME,
                           [*k1_rows(), item_from_spec(issuer_row("F02", "Ashford Holdings, L.P."))])
    loaded = {i.identifier: i for i in load_manifest(path)}

    assert loaded["F02"].required_keywords == ("Ashford Holdings LP",)
    assert loaded["F02"].document == "Schedule K-1 - Ashford Holdings LP"


def test_an_issuer_row_narrows_the_generic_row():
    assert narrowing_rows(k1_rows(("F02", "Ashford Holdings"), ("F03", "Birch Lane"))) == {
        "F01": ("F02", "F03"),
    }


def test_a_row_sharing_no_word_with_the_generic_row_narrows_nothing():
    rows = [*k1_rows(), RequestItem(identifier="C01", document="Mortgage Interest",
                                    required_keywords=("1098",), any_keywords=("1098",))]
    assert narrowing_rows(rows) == {}


def test_two_issuer_rows_whose_names_nest_are_refused_when_the_manifest_is_read(tmp_path):
    """Both accept the same K-1, so every one of them would park with
    nothing said about why. The person renames one."""
    path = create_template(tmp_path / MANIFEST_FILENAME,
                           k1_rows(("F02", "Ashford"), ("F03", "Ashford Holdings")))

    with pytest.raises(ManifestError) as caught:
        load_manifest(path)

    assert "F02" in str(caught.value) and "F03" in str(caught.value)
    assert "F01" in str(caught.value), "the row they both narrow is named too"


def test_two_issuer_names_that_merely_share_a_word_are_allowed():
    check_narrowing_names(k1_rows(("F02", "Ashford Holdings"), ("F03", "Birch Holdings")))


def test_two_unrelated_rows_may_share_a_name(tmp_path):
    """Only rows narrowing the *same* row are compared: an ordinary request
    whose keywords sit inside another's is nobody's issuer."""
    rows = [
        RequestItem(identifier="A01", document="W-2", required_keywords=("Ashford",)),
        RequestItem(identifier="A02", document="1099", required_keywords=("Ashford Holdings",)),
    ]
    path = create_template(tmp_path / MANIFEST_FILENAME, rows)
    assert len(load_manifest(path)) == 2
