"""Tests for tracker/ledger.py - the engagement's own append-only record.

The record's main claim is not made here: it is the autouse fixtures in
tests/conftest.py, which after every test in the suite hold the record's
statuses to the manifest's and the store's rows to the record's. What is
made here is the record's own behaviour - one line per event, a torn tail
nobody trips over, the fold that is the index's own order, and the refusal
to write a word outside the engagement lock - and which of the record and
the manifest answered a read. Since decision 102 the index has only one
answer: the record. The manifest's scanner columns still have two, per
identifier, which is decision 88's rule and what decision 103 finishes.
"""

from __future__ import annotations

import datetime as dt
import subprocess
import sys
from dataclasses import asdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.test_scanner import text_pdf
from tracker import ledger, store
from tracker.filer import (
    ASSIGNED_BY_PERSON,
    FILED,
    INDEX_FILENAME,
    assign_review_file,
    file_drops,
    ledger_key,
    read_index,
)
from tracker.locking import engagement_lock
from tracker.manifest import (
    COL_DOCUMENT,
    COL_IDENTIFIER,
    COL_STATUS,
    SHEET_NAME,
    RequestItem,
    Status,
    create_template,
    load_manifest,
)
from tracker.scaffold import MANIFEST_FILENAME, SHARED_DIR_NAME, scaffold_engagement
from tracker.scanner import scan_engagement

REPO = Path(__file__).resolve().parents[1]
DAY1 = dt.date(2026, 7, 1)
DAY2 = dt.date(2026, 7, 9)

ITEMS = [
    RequestItem(
        identifier="A01", document="W-2 Wage Statements", period="TY2025",
        allowed_extensions=("pdf",), min_size_kb=0, required_keywords=("W-2",),
        date_pattern=r"(?i)\b2025\b",
    ),
    RequestItem(
        identifier="C01", document="Mortgage Interest Statement", period="TY2025",
        allowed_extensions=("pdf",), min_size_kb=0, required_keywords=("1098",),
    ),
]


@pytest.fixture
def engagement(tmp_path):
    eng = tmp_path / "Smith Family 2025"
    eng.mkdir()
    create_template(eng / MANIFEST_FILENAME, ITEMS)
    scaffold_engagement(eng)
    return eng


def drop(engagement, name, text):
    return text_pdf(engagement / SHARED_DIR_NAME / name, text)


def lines(engagement):
    return ledger.path_for(engagement).read_bytes().split(b"\n")


def a_keyword(n: int) -> dict:
    return ledger.new(ledger.KEYWORD_LEARNED, identifier=f"A{n:02d}", keyword="lender")


# ------------------------------------------------------------ one line each ----


def test_an_append_is_one_line_and_reads_back_as_it_was_written(engagement):
    with engagement_lock(engagement):
        ledger.append(engagement, a_keyword(1))

    assert lines(engagement)[-1] == b""             # newline-terminated, nothing after it
    assert len(lines(engagement)) == 2
    events = ledger.read_events(engagement)
    assert [e[ledger.EVENT_KEY] for e in events] == [ledger.KEYWORD_LEARNED]
    assert events[0]["identifier"] == "A01" and events[0]["keyword"] == "lender"
    assert events[0][ledger.AT_KEY].endswith("Z")


def test_two_appends_are_two_lines_in_the_order_they_happened(engagement):
    with engagement_lock(engagement):
        ledger.append(engagement, a_keyword(1))
        ledger.append(engagement, a_keyword(2))

    assert len([raw for raw in lines(engagement) if raw]) == 2
    assert [e["identifier"] for e in ledger.read_events(engagement)] == ["A01", "A02"]


def test_an_append_outside_the_engagement_lock_refuses_loudly(engagement):
    with pytest.raises(ledger.LedgerError, match="engagement lock"):
        ledger.append(engagement, a_keyword(1))
    assert not ledger.path_for(engagement).exists()


def test_an_event_this_version_does_not_know_is_refused(engagement):
    with pytest.raises(ledger.LedgerError, match="not an event"):
        ledger.new("invented")
    with engagement_lock(engagement), pytest.raises(ledger.LedgerError, match="not an event"):
        ledger.append(engagement, {ledger.EVENT_KEY: "invented", ledger.AT_KEY: "now"})


# ------------------------------------------------------------------- torn ----


def test_a_torn_last_line_is_ignored_by_the_reader_and_truncated_by_the_next_append(engagement):
    with engagement_lock(engagement):
        ledger.append(engagement, a_keyword(1))
    path = ledger.path_for(engagement)
    whole = path.read_bytes()
    # A run killed mid-append: bytes with no newline after them.
    path.write_bytes(whole + b'{"event": "keyword_lear')

    assert [e["identifier"] for e in ledger.read_events(engagement)] == ["A01"]

    with engagement_lock(engagement):
        ledger.append(engagement, a_keyword(2))
    assert path.read_bytes().startswith(whole)
    assert b"keyword_lear\n" not in path.read_bytes()
    assert [e["identifier"] for e in ledger.read_events(engagement)] == ["A01", "A02"]


def test_a_line_that_is_not_an_event_in_the_middle_is_refused(engagement):
    with engagement_lock(engagement):
        ledger.append(engagement, a_keyword(1))
        ledger.append(engagement, a_keyword(2))
    path = ledger.path_for(engagement)
    whole = path.read_bytes()
    path.write_bytes(whole.replace(b'"A01"', b'"A01', 1))

    with pytest.raises(ledger.LedgerError, match="line 1"):
        ledger.read_events(engagement)
    path.write_bytes(whole)    # put it back: the suite's own fixture reads this folder too


# ------------------------------------------------------------------- fold ----


def row(**fields) -> dict:
    base = {"received": "2026-07-01", "original_name": "w2.pdf", "size_kb": 1.0,
            "digest": "abc", "identifier": "", "prepared_location": "",
            "pbc_location": "Shared/PBC/w2.pdf", "decision": "Needs Review",
            "reason": "", "candidates": "", "evidence": ""}
    return {**base, **fields}


def test_the_fold_is_the_last_event_for_each_original():
    first = ledger.new(ledger.PARKED, key="p/one", row=row(decision="Needs Review"))
    later = ledger.new(ledger.ASSIGNED_BY_PERSON, key="p/one", row=row(decision="Filed"))
    other = ledger.new(ledger.FILED, key="p/two", row=row(original_name="1098.pdf"))

    folded = ledger.fold([first, other, later])
    assert set(folded) == {"p/one", "p/two"}
    assert folded["p/one"]["decision"] == "Filed"
    assert folded["p/two"]["original_name"] == "1098.pdf"


def test_a_row_that_followed_a_moved_original_leaves_no_ghost_behind():
    first = ledger.new(ledger.PARKED, key="p/one", row=row())
    moved = ledger.new(ledger.PRESERVED, key="p/moved", was="p/one",
                       row=row(pbc_location="p/moved"))

    assert set(ledger.fold([first, moved])) == {"p/moved"}


def test_a_row_that_changed_identity_keeps_the_place_the_index_keeps_it_in():
    # The index rewrites that row where it sits; the fold is what read_index()
    # returns and what the next write rebuilds the workbook from, so a row
    # that moved to the end here would reorder the audit trail.
    one = ledger.new(ledger.PARKED, key="p/one", row=row())
    two = ledger.new(ledger.FILED, key="p/two", row=row(original_name="1098.pdf"))
    three = ledger.new(ledger.FILED, key="p/three", row=row(original_name="1099.pdf"))
    moved = ledger.new(ledger.PRESERVED, key="p/elsewhere", was="p/one",
                       row=row(pbc_location="p/elsewhere"))

    assert list(ledger.fold([one, two, three, moved])) == ["p/elsewhere", "p/two", "p/three"]


def test_the_statuses_are_built_up_across_the_passes_that_changed_something():
    one = ledger.new(ledger.SCANNED, statuses={"A01": {"status": "Missing"}})
    two = ledger.new(ledger.SCANNED, statuses={"A01": {"status": "Received"}})

    assert ledger.statuses([one, two]) == {"A01": {"status": "Received"}}


# ------------------------------------------------------------------- head ----


def test_the_head_changes_with_every_append_and_not_otherwise(engagement):
    assert ledger.head(engagement) == ""
    with engagement_lock(engagement):
        ledger.append(engagement, a_keyword(1))
        first = ledger.head(engagement)
        assert first and ledger.head(engagement) == first    # looking changes nothing
        ledger.append(engagement, a_keyword(2))
        assert ledger.head(engagement) != first


# -------------------------------------------------------------- the writers ----


def test_a_pass_records_what_it_filed_and_what_it_parked(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    file_drops(engagement, today=DAY1)

    events = ledger.read_events(engagement)
    names = [e[ledger.EVENT_KEY] for e in events if e[ledger.EVENT_KEY] in ledger.ROW_EVENTS]
    assert sorted(names) == [ledger.FILED, ledger.PARKED]
    folded = ledger.fold(events)
    assert folded == {ledger_key(e): asdict(e) for e in read_index(engagement)}


def test_a_person_filing_a_parked_file_is_recorded_as_theirs(engagement):
    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]

    assign_review_file(engagement, parked.pbc_location, "C01", keyword="lender", today=DAY2)

    names = [e[ledger.EVENT_KEY] for e in ledger.read_events(engagement)]
    assert names[-2:] == [ledger.ASSIGNED_BY_PERSON, ledger.KEYWORD_LEARNED]
    assert ledger.read_events(engagement)[-1]["keyword"] == "lender"


def test_a_quiet_pass_appends_no_scan(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    scan_engagement(engagement, today=DAY1)
    scanned = [e for e in ledger.read_events(engagement) if e[ledger.EVENT_KEY] == ledger.SCANNED]
    assert scanned                                   # the first scan said something

    head = ledger.head(engagement)
    scan_engagement(engagement, today=DAY1)          # nothing has changed since
    assert ledger.head(engagement) == head


def test_a_scan_that_changes_a_status_records_it(engagement):
    scan_engagement(engagement, today=DAY1)
    head = ledger.head(engagement)
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY2)
    scan_engagement(engagement, today=DAY2)

    assert ledger.head(engagement) != head
    assert ledger.statuses(ledger.read_events(engagement))["A01"]["status"] == "Received"


# --------------------------------------------------------- the readers ----
#
# Decision 88: load_manifest() answers each identifier's status from the
# record wherever it has one and from the workbook wherever it has not.
# Decision 102 left the index with only the first of those two. The claims
# below are about which answered.


def test_the_index_is_the_record_folded_and_nothing_else(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    file_drops(engagement, today=DAY1)
    drop(engagement, "1098.pdf", "Form 1098 Mortgage Interest Statement 2025")
    file_drops(engagement, today=DAY2)

    rows = read_index(engagement)
    # The same rows, in the same order. The order is the audit trail's and
    # there is no second copy of it any more: the fold lays it and the
    # store's position column keeps it (decision 102).
    assert [asdict(e) for e in rows] == list(ledger.fold(ledger.read_events(engagement)).values())
    assert not (engagement / INDEX_FILENAME).exists()


def test_a_record_holding_only_a_scan_gives_an_index_of_no_rows(engagement):
    """A scan takes the lock and records what it applied; no row event has
    been written, so the engagement has no index rows - and that is an
    answer, not a gap, because there is nowhere else a row could be."""
    scan_engagement(engagement, today=DAY1)
    events = ledger.read_events(engagement)
    assert any(e[ledger.EVENT_KEY] == ledger.SCANNED for e in events)
    assert not any(e[ledger.EVENT_KEY] in ledger.ROW_EVENTS for e in events)

    assert read_index(engagement) == []


def test_a_recorded_identifier_reads_from_the_record_and_an_unrecorded_one_from_the_workbook(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    scan_engagement(engagement, today=DAY1)
    manifest = engagement / MANIFEST_FILENAME
    whole = manifest.read_bytes()
    assert ledger.statuses(ledger.read_events(engagement))["A01"]["status"] == Status.RECEIVED

    # A01's Status typed over in Excel, and a row beside it the record has
    # never seen a status for, with a status of its own.
    wb = load_workbook(manifest)
    ws = wb[SHEET_NAME]
    at = {str(cell.value): cell.column for cell in ws[1]}
    ws.cell(row=2, column=at[COL_STATUS], value=Status.MISSING)
    ws.cell(row=4, column=at[COL_IDENTIFIER], value="E01")
    ws.cell(row=4, column=at[COL_DOCUMENT], value="Charitable Contribution Receipts")
    ws.cell(row=4, column=at[COL_STATUS], value=Status.PARTIAL)
    wb.save(manifest)
    wb.close()

    by_id = {item.identifier: item for item in load_manifest(manifest)}
    assert by_id["A01"].status == Status.RECEIVED     # the record, not the cell typed over it
    assert by_id["E01"].status == Status.PARTIAL      # the workbook: the record has never seen E01
    assert by_id["A01"].required_keywords == ("W-2",)   # the request itself is always the person's

    manifest.write_bytes(whole)    # put it back: the suite's own fixture reads this folder too


def test_an_unreadable_record_refuses_the_index_and_falls_back_for_the_statuses(engagement, caplog):
    """Two different answers, because there are two different second copies.

    The index has none any more, so a record that does not read as one is
    refused by name and the line: guessing past it would answer a question
    about a client's documents with a shrug. The manifest's scanner columns
    still have the sheet behind them, so that reader says so loudly and
    reads it.
    """
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    scan_engagement(engagement, today=DAY1)
    path = ledger.path_for(engagement)
    whole = path.read_bytes()
    path.write_bytes(whole.replace(b'"event"', b'"even', 1))   # corruption in the middle
    store.close()                  # the rows this process already read are not the question

    caplog.clear()
    with pytest.raises(ledger.LedgerError, match="line 1"):
        read_index(engagement)
    assert {i.identifier for i in load_manifest(engagement / MANIFEST_FILENAME)} == {"A01", "C01"}
    assert engagement.name in caplog.text

    path.write_bytes(whole)        # put it back: the suite's own fixture reads this folder too


def test_the_rollover_reads_last_years_statuses_through_the_record_and_otherwise_the_workbook(engagement):
    from tracker.rollover import roll_forward

    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    scan_engagement(engagement, today=DAY1)
    manifest = engagement / MANIFEST_FILENAME
    whole = manifest.read_bytes()
    wb = load_workbook(manifest)
    ws = wb[SHEET_NAME]
    at = {str(cell.value): cell.column for cell in ws[1]}
    ws.cell(row=2, column=at[COL_STATUS], value=Status.MISSING)   # typed over in Excel
    wb.save(manifest)
    wb.close()

    def prior_status(report):
        return {r.item.identifier: r.prior_status for r in report.rolled}["A01"]

    assert prior_status(roll_forward(engagement)) == Status.RECEIVED      # the record
    ledger.path_for(engagement).unlink()                                 # last year, before the record
    assert prior_status(roll_forward(engagement)) == Status.MISSING      # the workbook

    manifest.write_bytes(whole)


def test_the_app_shows_a_persons_filing_the_moment_they_make_it(engagement):
    from tracker import api

    drop(engagement, "scan0012.pdf", "nothing the rules recognise")
    parked = file_drops(engagement, today=DAY1).review[0]
    assign_review_file(engagement, parked.pbc_location, "C01", today=DAY2)

    [row] = api._state(engagement)["index"]
    assert row["decision"] == FILED and row["identifier"] == "C01"
    # One copy, one answer: what the app shows is the row the person's
    # decision was recorded as, in the same transaction as the move.
    assert read_index(engagement)[0].reason.startswith(ASSIGNED_BY_PERSON)


# -------------------------------------------------------------------- CLI ----


def compare(engagement):
    return subprocess.run(
        [sys.executable, "-m", "tracker.ledger", str(engagement), "--compare"],
        cwd=REPO, capture_output=True, text=True,
    )


def test_the_compare_check_says_they_agree_and_names_a_disagreement(engagement):
    """The operator's check, now the manifest half only: the index has no
    workbook to disagree with since decision 102, and `python -m
    tracker.store <store> check <root>` is the check for the store."""
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)
    scan_engagement(engagement, today=DAY1)

    agreed = compare(engagement)
    assert agreed.returncode == 0 and "the record and the manifest agree" in agreed.stdout
    assert "rows:     1 recorded (the record is the only copy)" in agreed.stdout

    manifest = engagement / MANIFEST_FILENAME
    whole = manifest.read_bytes()
    wb = load_workbook(manifest)
    ws = wb[SHEET_NAME]
    header = [str(c.value or "") for c in ws[1]]
    ws.cell(row=2, column=header.index(COL_STATUS) + 1).value = Status.MISSING
    wb.save(manifest)
    wb.close()

    disagreed = compare(engagement)
    assert disagreed.returncode == 1
    assert "they do not agree" in disagreed.stdout
    assert "A01" in disagreed.stdout

    manifest.write_bytes(whole)    # put it back: the suite's own fixture reads this folder too


def test_the_cli_prints_what_the_folder_holds(engagement):
    drop(engagement, "w2.pdf", "Form W-2 Wage and Tax Statement 2025")
    file_drops(engagement, today=DAY1)

    out = subprocess.run(
        [sys.executable, "-m", "tracker.ledger", str(engagement)],
        cwd=REPO, capture_output=True, text=True, check=True,
    ).stdout
    assert f"events: {len(ledger.read_events(engagement))}" in out
    assert "rows:   1" in out
    assert ledger.head(engagement) in out
