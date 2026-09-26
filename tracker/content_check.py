"""Tier-3 content validation for the tracker (component 4).

Extracts text from client documents (read-only) and applies the deterministic
rules from the manifest row: Required Keywords (ALL must appear), Any
Keywords (at least ONE), and Date Pattern (regex). No rule logic is
hardcoded — everything comes from the manifest.

Extractors by extension:

- ``.pdf``            → pdfplumber, first ``MAX_PAGES`` pages only; if the
                        PDF has no text layer (a scan), fall back to OCR
                        (decision 169: RapidOCR through :mod:`tracker.ocr`,
                        with its models shipped in the app, so the reader
                        runs on every machine). When it cannot - a model
                        missing or broken - the file is reported as
                        unverifiable with a clear note, and nothing breaks.
- ``XLSX_EXTENSIONS``   → openpyxl (all sheets, cached formula values);
                        date cells are rendered in both ISO (2025-12-31)
                        and US (12/31/2025) forms so either pattern style
                        matches.
- ``TEXT_EXTENSIONS``   → plain text (utf-8, then cp1252 fallback).
- anything else       → no extractor; reported unverifiable, for a person (``reasons.UNCHECKABLE_TYPE``).

**One way of reading a document.** :func:`extract` is what both the router
(deciding where a drop goes) and the scanner (deciding a row's status) call,
so the two can never disagree about what a file says. The router asks for
the text layer first and OCR only when the file's name tells it nothing
(:mod:`tracker.router` explains why); the scanner always finishes the job.

Keyword matching is case-insensitive. Date Pattern is applied to the raw
text as-is, so authors control case sensitivity with inline flags (``(?i)``).

**A verdict keeps its evidence.** ``ok`` and ``reason`` say *what* was
decided; :class:`Evidence` says *why* — which rule found which of the row's
own words, where in the document, on which page. The verdict is unchanged
by it: nothing here reads the evidence back to decide anything, and the
reason sentence is worded exactly as before. It exists because a person
working the review queue needs "'1098' in the title, 'mortgage interest' on
page 1", not "matched no request", and because a later layer ranks parked
files from it. What is kept is the *firm's* words - the keyword the
manifest row asked for, the Period it asked for - and never a word of the
client's document, which is the same line :class:`ContentCache` draws.

Caching: :class:`ContentCache` stores only *verdicts* — pass/fail, reason
and evidence — keyed by the file's **content digest** and the row's rules
fingerprint, with a per-path memo of (size, mtime, digest) so an unchanged
file is not re-hashed. Keyed on content, not path, because the same bytes are read
twice in this system: once as a drop the router routes, once as the
working copy the scanner checks. A verdict the router reached is the
verdict the scanner finds. Extracted client text is deliberately never
persisted anywhere. Editing a row's rules changes the fingerprint and
triggers one re-extraction; unchanged files on unchanged rules are never
re-read, which keeps the scheduled cadence cheap.

**Where the cache lives** (decision 107): in the store on the designated
machine - two tables of :mod:`tracker.store`, the memos and the verdicts,
per engagement - and nowhere in the engagement folder. Until then it was a
JSON file beside the client's documents that every pass rewrote, in a
folder a cloud client syncs; the owner's rule is that nothing the machine
can derive lives in the synced folder, because a file rewritten every two
hours there is the synced-database failure one file over. It is **not
journalled**: a verdict is a derivation of the bytes in the folder and the
rules in the record, not a fact about the engagement, so a rebuild of the
store from the journals leaves the cache empty and the next pass reads
each document once - slower, never wrong. A save is one transaction of
the store's own, under the engagement lock every writer already holds,
and never ``store.record()``: there is no event. A cache built with no
engagement - the command line's, a test's - lives in memory and never
touches the store. The old file is never read; the next real pass
removes it (:data:`RETIRED_CACHE_FILENAME`).

:class:`Evidence` itself, the rule and place vocabularies and the cell format
it is written into live in :mod:`tracker.records` since decision 100 - they
are the shape of a verdict, not the reading of a document - and are
re-exported here for one release.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import time
import zipfile
from collections import OrderedDict
from dataclasses import MISSING, asdict, dataclass, fields, replace
from pathlib import Path

from tracker import ocr, reasons, store
from tracker.manifest import RequestItem, has_routing_rules, keyword_alternatives

# The Evidence record and the cell format it is written in live in
# tracker/records.py (decision 100): they are the shape of a verdict, not the
# reading of a document. The names this module does not use itself are
# re-exported from here so that every `from tracker.content_check import ...`
# still resolves to the same object; kept for one release; import from
# tracker.records.
from tracker.records import (
    EVIDENCE_PLACES,  # noqa: F401
    EVIDENCE_RULES,  # noqa: F401
    RULE_ANY,
    RULE_DATE,
    RULE_FILENAME,  # noqa: F401
    RULE_REFUSED,  # noqa: F401
    RULE_REQUIRED,
    WHERE_DEEP,
    WHERE_FIRST_PAGE,
    WHERE_FOOTER,
    WHERE_TITLE,
    Evidence,
    format_evidence,  # noqa: F401
    parse_evidence,  # noqa: F401
)
from tracker.validators import (
    IMAGE_EXTENSIONS,
    PDF_EXTENSION,
    TEXT_READ_CAP_MB,
    extension_of,
    open_test,
    picture_too_large_reason,
    sha256_of,
    too_large_reason,
)

log = logging.getLogger("tracker.content_check")

# pdfplumber/pdfminer warn loudly on ugly-but-parseable PDFs; failures are
# already surfaced through ContentResult.reason.
logging.getLogger("pdfminer").setLevel(logging.ERROR)

#: The file the verdict cache was, in the engagement folder, until decision
#: 107 moved it into the store. Never read again: the one use left is the
#: filer's tidy-up, which removes it from a folder a pass finds it in.
RETIRED_CACHE_FILENAME = "_content_cache.json"
#: The cache's layout, carried on every verdict row in the store; a row at
#: any other version is ignored on load and deleted on save (the cache is
#: disposable), so a matcher change invalidates without a store version.
#: 3: verdicts that were the machine's (no OCR, OCR failed) are no longer
#: stored; a cache written before that carried them for ever.
#: 4: says() changed what a keyword verdict means (a form number is title evidence).
#: 5: a form number in the title is weighed by its shape too; a keyword's words sit on one line.
#: 6: a keyword on a menu line, or one an ask-word asked for, is no longer said.
#: 7: a verdict carries the evidence behind it (an Evidence per matched rule).
#: 8: a form heading its line with its own printed title and its year names
#:    itself, so a one-copy W-2 says "W-2" where it said nothing before.
#: 9: a keyword may name alternatives ("|") and join the phrases one of
#:    them wants together ("+"), so a row can accept either of two
#:    documents; a verdict cached before that read the characters as words.
#: 10 (decision 127): the reading turns a page or a photo upright before
#:    OCR, so a verdict on a scan read at 9 may have been read sideways.
#: 11 (decision 141): a notice's header phrases (``FIRST_PAGE_PHRASES``)
#:    count on the first page only, so a verdict that found one deeper is void.
#: 12 (decision 169): RapidOCR reads instead of Tesseract, so every verdict
#:    on a scan or a photo was reached on another reader's words. The first
#:    pass after the install reads every scan and photo again, once.
CACHE_VERSION = 12

#: A "text" PDF with fewer stripped characters than this *per page read*
#: is a scan: what little it has is a scanner's stamp ("Scanned by
#: CamScanner", "Page 1 of 2"), not the document, and is not read as one.
#: Pages are counted from the page breaks ``_extract_pdf`` writes.
_MIN_TEXT_CHARS = 25
#: What one page of a reading ends with. Public since decision 128:
#: :mod:`tracker.names` counts the same breaks to say which page a name was
#: printed on, and a second module holding its own copy of the character
#: would be a second answer to "which page is this".
PAGE_BREAK = "\f"

#: Read at most this many pages of any PDF, with or without OCR. The words
#: that identify a document - its form number, the tax year, the payer -
#: are on its first pages; a 500-page general ledger dropped by a client
#: used to be read cover to cover on every route and scan, and that is what
#: stalled a run. A keyword deep in a long document is not evidence the
#: router should be acting on anyway.
MAX_PAGES = 10
#: The other file types whose text can be read (PDF_EXTENSION is the third).
XLSX_EXTENSIONS = ("xlsx", "xlsm")
TEXT_EXTENSIONS = ("csv", "tsv", "txt")
_MAX_OCR_PAGES = MAX_PAGES
#: The most pixels one page or one photo is read at (decision 137, B1).
#: Rendering has to stay bounded: a PDF page may be 14,400 points square,
#: which at the old fixed scale of 2 is some 830 million pixels, and
#: ``bitmap.to_pil()`` hands Pillow an image without its decompression-bomb
#: guard ever being asked. Forty million is a letter page at about 600 dpi -
#: far past what the reader needs - so an ordinary page is read exactly as
#: before, and a giant one is read smaller rather than refused.
PIXEL_BUDGET = 40_000_000
#: The scale an ordinary PDF page is rendered at for OCR (144 dpi).
RENDER_SCALE = 2.0
#: The reader's safety stop (decision 137, B1.2; the owner's Q-1 of
#: 2026-09-24): ten times the ceiling he approved for decision 127 (6 s a
#: photo, 60 s a ten-page scan). A page - a photo is one page - gets a
#: minute and a document ten, rendering included. A reading that reaches
#: either is **abandoned**: a kept verdict, ``reasons.READING_STOPPED``,
#: not read again until the file changes. Under the stop nothing changes:
#: a slow reading is finished and named (``runner.SLOW_READING_SECONDS``).
#: Before this, the only stop was Task Scheduler's two-hour kill, which
#: ended the whole pass and met the same file again next time. Since
#: decision 150 the same two numbers bound the **whole** reading as well -
#: text layer, render and OCR - because the pass reads in a child process
#: it ends at the stop (:func:`extract_bounded`): ten minutes a file, one
#: a photo.
READING_STOP_PAGE_SECONDS = 60.0
READING_STOP_DOCUMENT_SECONDS = 600.0
#: The most text one reading hands back, in characters, whatever it read
#: (decision 178): the number of megabytes a text file is read to
#: (``validators.TEXT_READ_CAP_MB``), counted in characters, so no reading
#: of any kind says more than a text file may. The workbook reader had no
#: bound at all: a 14.7 KB ``.xlsx`` whose one shared string of ten million
#: characters sat in twenty cells read to two hundred million, and every
#: rule the pass then asked of it ran in the pass's own process, outside
#: the stop. A reading past it is cut there and marked ``cut``, as a text
#: file past its cap always was.
READING_CHAR_BUDGET = TEXT_READ_CAP_MB * 1024 * 1024
#: The only ways a zip's part may be packed for a reading to unpack it
#: (decision 178): stored, and deflated - what Windows, macOS, Excel and
#: every mail client write. The standard library bounds a deflated read at
#: the size asked for; a bzip2 or LZMA part it inflates a whole compressed
#: chunk at a time, so one 64 KB read of 905 bytes of bzip2 grew past 2 GB
#: before any count could stop it. A workbook is a zip and is held to the
#: same list (:func:`_extract_xlsx`), and so is an attached zip
#: (:mod:`tracker.containers`).
PACKINGS = (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)
#: What a workbook packed any other way fails with.
UNKNOWN_PACKING = "a part of it is packed in a way only a person's zip program opens"
#: The clock the stop reads. A name of its own so the suite can move it
#: rather than wait ten minutes.
_clock = time.monotonic



# ---------------------------------------------------------------- evidence ----

#: How many non-blank lines at the end of a page count as its footer, where
#: a form repeats its own number on every copy. The Evidence record's own
#: ``WHERE_FOOTER`` is named from this reading.
_FOOTER_LINES = 2


def _evidence_from_json(raw: object) -> tuple[Evidence, ...]:
    """The evidence of a cached verdict, rebuilt from what ``asdict`` wrote.

    Total, like everything that reads the cache: the cache is disposable,
    so anything unrecognised costs a re-extraction, never an exception in
    a scheduled run. A tuple as well as a list: ``asdict`` leaves the
    evidence a tuple of dicts until the store's JSON round trip makes it a
    list, and a verdict read back in the same pass it was kept must be the
    verdict that was kept (decision 107's claim that the kept verdict and
    the miss verdict are one).
    """
    if not isinstance(raw, (list, tuple)):
        return ()
    return tuple(
        Evidence(str(item.get("rule", "")), str(item.get("term", "")),
                 str(item.get("where", "")), int(item.get("page", 0) or 0))
        for item in raw if isinstance(item, dict)
    )


@dataclass(frozen=True, slots=True)
class ContentResult:
    """Verdict of tier-3 rules for one file."""

    ok: bool
    reason: str = ""
    extractable: bool = True  # False: could not get text (no extractor / no OCR)
    #: True when the verdict is about the machine, not the document: the
    #: reader could not run, or it failed this once. Never cached - mending the reader
    #: or a second try must be able to change the answer.
    transient: bool = False
    #: Why the verdict went this way: one Evidence per rule that found its
    #: word, in the order the rules are applied. Collected on a failing
    #: verdict too - what *did* match is the lead a person works from - and
    #: read by nothing that decides anything.
    evidence: tuple[Evidence, ...] = ()


@dataclass(frozen=True, slots=True)
class Extraction:
    """What one document says, or why it could not be read.

    ``text`` is None when there is nothing usable and ``reason`` then says
    why in the words the scanner writes. ``needs_ocr`` marks a PDF whose
    text layer is empty when the caller asked not to OCR yet.
    """

    text: str | None
    from_ocr: bool = False       # the text came from OCR, not a text layer
    needs_ocr: bool = False      # a scan; OCR was not attempted (ocr=False)
    reason: str = ""             # why there is no usable text, if there is none
    extractable: bool = True     # False: no extractor, no OCR, or extraction failed
    error: str = ""              # the exception, when extraction raised
    transient: bool = False      # the machine's doing (no OCR, OCR failed), not the file's
    #: How long this reading took, in seconds (decision 127). Recorded on
    #: every reading, failed ones included, and never acted on: a slow
    #: reading is said in the run's summary and finished, never cut short.
    #: It is the number the reader benchmark and the owner's speed ceiling
    #: are measured in.
    seconds: float = 0.0
    #: Only the first ``validators.TEXT_READ_CAP_MB`` of a text file was
    #: read (decision 137). Said beside any verdict reached on it.
    cut: bool = False
    #: Tier 2's open test of the same file, made where the reading was made
    #: (decision 150, :func:`validators.open_test`): ``""`` when it opens
    #: cleanly, the refusal otherwise, and None when it was not made - a
    #: reading of :func:`extract` alone, or one that never finished.
    opened: str | None = None


# ------------------------------------------------------------------ rules ----


def has_content_rules(item: RequestItem) -> bool:
    """Whether tier 3 has anything to check on this row at all: a routing
    rule, or a year check derived from Period (which checks but never routes)."""
    return has_routing_rules(item) or bool(item.date_pattern)


def rules_fingerprint(item: RequestItem) -> str:
    """Stable hash of the row's tier-3 rules; part of the cache key."""
    blob = json.dumps(
        [list(item.required_keywords), list(item.any_keywords), item.date_pattern]
    )
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()[:16]


#: The variants the IRS joins to a form's number: ``1098-T`` is the tuition
#: form, not the mortgage form ``1098``; ``1099-R`` is not ``1099-INT``. A
#: keyword of the bare number means the bare form and must not match a
#: variant. Keyed by the number with its dashes removed (``w2`` for W-2), so
#: a keyword written either way is found. Any other hyphen-joined neighbour
#: - an employer, a bank, a person (``1098-Citi``, ``W2-Tom``) - is a
#: separator's hyphen and still matches, which is what a client's file
#: name usually carries. One list, in the one module that reads keywords.
FORM_VARIANTS: dict[str, tuple[str, ...]] = {
    "1098": ("t", "e", "c", "f", "q", "ma"),
    "1099": ("int", "div", "b", "r", "misc", "nec", "oid", "k", "g", "s", "sa",
             "q", "ltc", "patr", "cap", "c", "a", "h", "da", "ls", "sb", "qa"),
    "1095": ("a", "b", "c"),
    "1040": ("sr", "nr", "x", "es", "v", "ss", "c"),
    "1041": ("a", "es", "n", "qft", "t", "v"),
    "1065": ("x", "b"),
    "1120": ("s", "x", "f", "h", "w", "c", "l", "pc", "pol", "reit", "ric", "sf", "nd"),
    "940": ("pr", "ss"),
    "941": ("x", "ss", "pr"),
    "990": ("ez", "pf", "t", "n"),
    "5498": ("sa", "esa", "qa"),
    "w2": ("g", "gu", "as", "vi", "c"),
    "w3": ("c", "ss", "pr"),
}
#: What may sit between a number and its variant, or between a keyword's
#: words: nothing ("1098T"), spaces, or any dash a PDF or a keyboard yields
#: - not a line break, unless the phrase is a heading that wraps. A
#: keyword names one thing, and its words are printed together; a line
#: break is where one label ends and the next begins. A K-1's box "19
#: Distributions" above its footer "Schedule K-1" is not a distribution
#: schedule, and "Items affecting shareholder basis" above the same
#: footer is not a shareholder basis schedule. A heading set in its own
#: box wraps - "Balance Sheet" over "As of December 31, 2025" - and a
#: heading begins its line with two words or more; a phrase that begins
#: mid-line and runs on to the next is a sentence, and one word over
#: another is two labels. (The IRS sets the 1098's title one word per
#: line, "Mortgage / Interest / Statement"; that form is known by its
#: number, as every form in tests/irs/ is.)
_DASH_CHARS = "".join(("-", chr(0x2010), chr(0x2011), chr(0x2012), chr(0x2013), chr(0x2014), chr(0x2212), chr(0xAD)))   # hyphen, the Unicode dashes, minus, soft hyphen
#: An apostrophe as a keyboard types it, as a PDF prints it, and as OCR reads it.
_APOSTROPHES = "".join(("'", chr(0x2019), chr(0x2018), chr(0x02BC), chr(0x60)))
_SAME_LINE = r"[^\S\r\n\f\v]"   # whitespace that is not a line break (a no-break space included)
_JOINER = rf"(?:{_SAME_LINE}|[{re.escape(_DASH_CHARS)}])*"
_WRAPPING_JOINER = rf"[\s{re.escape(_DASH_CHARS)}]*"
#: A heading begins its line, after nothing but blank space, a bullet, a
#: pipe or a rule (OCR reads a box's edge as "|", a list as "-" or "•").
_LINE_START = rf"(?:^|(?<=[\r\n\f\v]))(?:{_SAME_LINE}|[{re.escape(_DASH_CHARS)}•*|_])*"
_DASHES = re.compile(rf"[\s{re.escape(_DASH_CHARS)}]")


def _joined_to_a_variant(text: str, keyword: str, match: re.Match[str]) -> bool:
    """True when the matched keyword ends in a form number that ``text``
    continues with one of that form's variants: ``1098`` in "1098-T" or
    "1098 T", ``form 1040`` in "Form-1040-SR"."""
    variants = FORM_VARIANTS.get(_DASHES.sub("", keyword.strip().lower().split()[-1]))
    if not variants:
        return False
    tail = re.compile(rf"{_JOINER}(?:{'|'.join(variants)})(?![a-z0-9])")
    return tail.match(text, match.end()) is not None


def keyword_pattern(keyword: str, *, wrapping: bool = False) -> str | None:
    """The regular expression one keyword is looked for with, or None if blank.

    Whole tokens only. Between a keyword's words, and on either side of a
    dash inside a word, anything a dash can become: ``interest income``
    matches "interest-income" in a file name, ``w-2`` matches "W2" and
    "W–2", ``1099-int`` matches "1099INT". A keyword's words are on one
    line (``_JOINER``); the ``wrapping`` pattern is the other reading, a
    heading that begins its line with two words or more of the keyword
    and wraps the rest on to the next, and is None for a keyword that
    cannot wrap. Two one-word lines are two labels, not one phrase.
    """
    # Only the letters and digits are the keyword; a "-" or "$" typed into
    # a keyword cell is nothing to look for, and must not match everything.
    keyword = keyword.strip().lower()
    if not any(ch.isalnum() for ch in keyword):
        return None
    words = [_JOINER.join(_seams(part) for part in _DASHES.split(w) if part) for w in keyword.split()]
    words = [w for w in words if w]
    if not wrapping:
        return rf"(?<![a-z0-9]){_JOINER.join(words)}(?![a-z0-9])"
    if len(words) < 3:
        return None
    first_line = _JOINER.join(words[:2])
    return rf"{_LINE_START}(?<![a-z0-9]){first_line}{_WRAPPING_JOINER}{_WRAPPING_JOINER.join(words[2:])}(?![a-z0-9])"


def _seams(part: str) -> str:
    """``part`` escaped, with a dash allowed where letters meet digits (a
    keyword typed ``w2`` finds "W-2" as ``w-2`` finds "W2"), a plural
    allowed on a word (``fixed asset`` finds "Fixed Assets"), and a number
    kept out of a larger amount (``704`` is not in "20,704" or "704.50")."""
    pieces = re.findall(r"[a-z]+|[0-9]+|[^a-z0-9]+", part)
    out = []
    for piece in pieces:
        if piece.isalpha() and len(piece) >= 3:
            out.append(re.escape(piece) + "(?:e?s)?")
        elif piece.isdigit():
            out.append(rf"(?<![0-9][,.]){re.escape(piece)}(?![,.][0-9])")
        elif all(ch in _APOSTROPHES for ch in piece):
            out.append(f"[{re.escape(_APOSTROPHES)}]?")   # typed straight, printed curly, or dropped
        else:
            out.append(re.escape(piece))
    return _JOINER.join(out)


def _occurrences(low: str, keyword: str):
    """Every place ``low`` (lower-cased text) says ``keyword`` as a whole token.

    Both readings of the keyword - its words on one line, and the heading
    that wraps - in the order they are written, so a caller that weighs
    *where* a keyword is said (:func:`says`) sees each occurrence.
    """
    for wrapping in (False, True):
        pattern = keyword_pattern(keyword, wrapping=wrapping)
        if pattern is None:
            continue
        for match in re.finditer(pattern, low):
            if not _joined_to_a_variant(low, keyword, match):
                yield match


def contains_keyword(text: str, keyword: str) -> bool:
    """True if ``keyword`` appears in ``text`` as a whole token.

    Matched on token boundaries rather than as a bare substring, so ``EIN``
    does not match "being", ``1098`` does not match "10983", and ``W-2``
    still matches "W-2 Wage and Tax Statement". A form's own variant is
    part of its name, not a boundary: ``1098`` does not match "1098-T"
    and ``1099`` does not match "1099-R" (``FORM_VARIANTS``), because a
    tuition statement filed as mortgage interest is exactly the misfiling
    the keywords exist to prevent. A row that means every 1099 lists them.
    Keywords drive both status and — via :mod:`tracker.router` — where a
    document gets filed, so a coincidental substring must never count as
    evidence.

    A keyword that names alternatives is there when one of them is, and
    an alternative that joins phrases is there when every phrase is
    (``keyword_alternatives``); a plain keyword is one phrase and reads
    exactly as it always did.

    This is the plain reading: the words are there. Whether they are the
    document's own words is :func:`says`, which the routing and status
    rules use; a file name, being all title, is read with this one.
    """
    low = text.lower()
    return any(all(next(_occurrences(low, phrase), None) is not None for phrase in alternative)
               for alternative in keyword_alternatives(keyword))


#: A form number a keyword names is content evidence only where a form
#: prints its own: in the title - the first ``TITLE_CHARS`` of the text -
#: or as the number the document prints most, on every copy and every
#: page's footer (a W-2 prints its title at the foot of the form, below
#: the boxes). Every 1040 says "Attach Form(s) W-2", every 1095-C's
#: instructions name "Form 1095-A", every 1099's say "Form 1040-ES", and
#: those are single mentions of other forms, not the form this is. In the
#: title as everywhere, a mention counts only when its shape is a form
#: naming itself, not a sentence about another form (``_weigh_mention``):
#: the IRS prints an "Attention" page ahead of every information return
#: that names Form 1099-NEC by way of example, an organizer's first lines
#: say "attach Form 1098", a bank's letter says "reported on Form
#: 1099-INT". A file name is read whole: it is all title.
TITLE_CHARS = 400
_FORM_NUMBER = re.compile(r"^(?:form\s+)?([a-z]?\d{3,4}|w\d)(?:[a-z]{1,4})?$")
_FORM_MENTION = re.compile(
    rf"(?<![a-z0-9])(?:form\s+)?(w[{re.escape(_DASH_CHARS)}]?[23]|\d{{3,4}})"
    rf"(?:[\s{re.escape(_DASH_CHARS)}]?([a-z]{{1,4}}))?(?![a-z0-9])"
)


def is_form_number(keyword: str) -> bool:
    """True for a keyword that is a form's number: ``1098``, ``form 1095-a``,
    ``W-2``, ``form 940``. A bare number counts only for a known family
    (``FORM_VARIANTS``); ``form`` in front makes any number one."""
    bare = _DASHES.sub("", keyword.strip().lower())
    if bare.startswith("form"):
        return _FORM_NUMBER.match("form " + bare[4:]) is not None
    match = _FORM_NUMBER.match(bare)
    return match is not None and match.group(1) in FORM_VARIANTS


def _form_key(number: str, variant: str | None) -> str:
    return _DASHES.sub("", number) + (variant or "")


def form_key(keyword: str) -> str | None:
    """The form a keyword names, keyed the way a mention is (``w2``,
    ``1099int``), or None for a keyword that is not a form number.

    The one normalisation, so a manifest row's own word and a number read
    off a document are compared as one thing: :func:`_one_says_where` asks
    it of the row's keyword, and :mod:`tracker.router` asks it of the
    keyword an :class:`Evidence` kept, to say which form a row was
    accepted because of (decision 94).
    """
    if not is_form_number(keyword):
        return None
    bare = _DASHES.sub("", keyword.strip().lower())
    return bare[4:] if bare.startswith("form") else bare


def _title(text: str) -> str:
    """The first ``TITLE_CHARS``, extended to the end of the word it cuts."""
    if len(text) <= TITLE_CHARS:
        return text
    space = re.compile(r"\s").search(text, TITLE_CHARS)
    return text[:space.start()] if space else text


#: How a form names itself, and how it names another form. Its own number
#: is followed by its year or revision - "Form 1040 (2025)", "Form 1099-DIV
#: (Rev. January 2024)", "941 for 2026:" - on the title and on every page's
#: footer. Another form's number is quoted: "(Form 1040)", "Form 1040 or
#: 1040-SR", "Form 1040 instructions", "Form 1040, line 8" - or told to
#: the reader: "attach Form 1098", "Forms W-2", "reported on Form
#: 1099-INT", "a 1098". The two shapes are told apart here, because a
#: W-2's instruction pages name Form 1040 thirty times and the W-2 itself
#: a dozen. A self-mention's year is on the same line as its number.
_SELF_WEIGHT, _PLAIN_WEIGHT, _REFERENCE_WEIGHT = 3, 1, 0
_SELF_AFTER = re.compile(
    rf"{_SAME_LINE}*(?:[{re.escape(_DASH_CHARS)}]{_SAME_LINE}*)?(?:\({_SAME_LINE}*)?(?:rev\b|(?:19|20)\d{{2}}\b)"
    rf"|{_SAME_LINE}+for{_SAME_LINE}+(?:19|20)\d{{2}}\b"
)
#: A form's number set off from what follows by a dash - "Form 1098 -
#: Mortgage Interest Statement", "Form 1099-INT - Interest Income" - is
#: a checklist's line where the page is a menu of forms. A payer's own
#: substitute form prints its title that way too, so the shape alone is
#: not the answer: see ``_lists_forms`` and ``_year_beside``.
_REFERENCE_AFTER = re.compile(
    rf"\s*[,.;)]|{_SAME_LINE}+[{re.escape(_DASH_CHARS)}]{_SAME_LINE}+|"
    r"\s+(?:or|and|line|lines|instructions?|to|if|is|are|was|were|schedule|box|boxes|page|"
    r"worksheet|for|with|at|by|filers?|must|may|should)\b"
)
#: The words a document prints before a form it is telling the reader
#: about; a form never names itself after them - on the same line or the
#: one above, because a wrapped sentence breaks where it will ("please
#: attach" / "Form 1098 (2025) from each lender", the twelfth reading).
#: A page footer whose prose above ends in one of these words is read as
#: a reference too; that costs nothing the corpus shows, since a footer
#: repeats on every page and the others count. A colon may stand between
#: the word and the form, because that is how a cover writes it:
#: "Enclosed: Form 1099-INT", "Re: Forms 1099-INT and 1099-DIV".
_REFERENCE_BEFORE = re.compile(
    r"(?<![a-z0-9])(?:forms|attach|attached|see|file|files|filed|use|of|on|from|with|to|and|or|a|an|the|"
    r"include|including|per|report|reported|enclosed|enclosure|enclosures|re)\s*:?\s+$"
)
#: A form number that continues a list - ", Form 5498", ", 5498" - is one
#: of the forms a cover names rather than the form the page is. The last
#: item of "Enclosed: Form 1099-INT, Form 1098, Form 5498" carries no
#: punctuation after it, and was read as named in its own right, which
#: blinded the 1099-INT printed behind it (the thirteenth reading).
_LIST_CONTINUES = re.compile(rf",{_SAME_LINE}*$")
_BEFORE_CHARS = 14   # room for the longest word above, its colon and the space after it

#: What decided a mention's weight. The number alone is enough to weigh a
#: document's forms; the shape is what tells a menu's line from a payer's
#: own title, and where that line's title begins (``_menu_tails``).
_BY_PAREN, _BY_TOLD, _BY_LIST, _BY_YEAR = "paren", "told", "list", "year"
_BY_DASH, _BY_COMMA, _BY_WORD, _BY_NOTHING = "dash", "comma", "word", ""


def _weigh_mention(text: str, match: re.Match[str], end: int) -> tuple[int, str]:
    """A mention's weight and the shape that decided it.

    ``end`` is where the form number (with its real variant) stops.
    """
    start = match.start()
    if start > 0 and text[start - 1] == "(":
        return _REFERENCE_WEIGHT, _BY_PAREN
    # What comes before is read first: "attach Form 1098 (2025)" is told to
    # the reader, year or no year (the tenth reading added the year).
    if _REFERENCE_BEFORE.search(text, max(0, start - _BEFORE_CHARS), start):
        return _REFERENCE_WEIGHT, _BY_TOLD
    if _LIST_CONTINUES.search(text, max(0, start - _BEFORE_CHARS), start):
        return _REFERENCE_WEIGHT, _BY_LIST
    if _SELF_AFTER.match(text, end):
        return _SELF_WEIGHT, _BY_YEAR
    separator = _REFERENCE_AFTER.match(text, end)
    if separator:
        return _REFERENCE_WEIGHT, _shape_after(separator.group(0))
    return _PLAIN_WEIGHT, _BY_NOTHING


def _shape_after(separator: str) -> str:
    """Which of ``_REFERENCE_AFTER``'s shapes matched: a dash, a comma, a word."""
    stripped = separator.strip()
    if len(stripped) == 1 and stripped in _DASH_CHARS:
        return _BY_DASH
    return _BY_COMMA if stripped == "," else _BY_WORD


def _mention_span(low: str, match: re.Match[str]) -> tuple[str, str | None, int] | None:
    """The number, its real variant and where it stops; None for a match
    that is no mention at all.

    A known family (``FORM_VARIANTS``) is a mention bare or after "Form":
    "W-2", "941 for 2026", "Form 1099-DIV". Any other number is one only
    after "Form" and takes a variant only joined by a dash - "Form 4562",
    "Form 1125-E" - because a bare 4562 is as likely a year, an OMB number
    or an amount, and "Form 4562 line" is not Form 4562-LINE.
    """
    number, variant = match.group(1), match.group(2)
    base = _DASHES.sub("", number)
    end = match.end()
    if base in FORM_VARIANTS:
        if variant and variant not in FORM_VARIANTS[base]:
            variant, end = None, match.end(1)   # "1040 line" is Form 1040, not a variant
    else:
        if not match.group(0).startswith("form"):
            return None
        if variant and low[match.start(2) - 1].isspace():
            variant, end = None, match.end(1)
    return number, variant, end


#: A menu lists. Two or more forms heading a line of their own title
#: within the title window is a checklist ("Form 1099-INT - Interest
#: Income", "Form 1099-DIV - Dividends and Distributions", ...); one such
#: line is a payer's own substitute form printing its title the way it
#: always has, and a year beside the number on its line is that form's
#: evidence, as a year after the number is (``_SELF_AFTER``). Decision
#: 69's dash rule alone parked "2025 Form 1099-INT - Interest Income".
_FORMS_IN_A_MENU = 2
_LINE_BREAKS = "\r\n\f\v"
_BREAK_RUN = re.compile(rf"[{re.escape(_LINE_BREAKS)}]+")
_A_YEAR = re.compile(r"(?<![0-9])(?:19|20)\d{2}(?![0-9])")
_SPACES = re.compile(rf"{_SAME_LINE}*")
#: What may stand between the start of a line and a form naming itself: a
#: bullet or a rule, its own year, and the word "Form" where the year ate
#: it (``_FORM_MENTION`` reads "2025 form" as a number and its would-be
#: variant, so the mention that follows begins at the number). A form
#: prints its number at the head of its title line; "2025 Consolidated
#: Form 1099 - Account 8812" is a broker's sentence about a family.
_HEADS_ITS_LINE = re.compile(
    rf"(?:{_SAME_LINE}|[{re.escape(_DASH_CHARS)}•*|_])*"
    rf"(?:(?:19|20)\d{{2}}{_SAME_LINE}+)?(?:forms?{_SAME_LINE}+)?"
)


_A_BREAK = re.compile(rf"[{re.escape(_LINE_BREAKS)}]")
#: Everything up to and including the last break in what it is matched
#: against: the look back of :func:`_line_of`, one search over a window.
_TO_THE_LAST_BREAK = re.compile(rf"(?s:.*)[{re.escape(_LINE_BREAKS)}]")
#: How far the look back of :func:`_line_of` first reaches; each window
#: that holds no break doubles it.
_LOOK_BACK = 256
#: The line :func:`_line_of` found last and the text it was found in, held
#: by reference: a mention inside it is on it, with nothing searched.
_last_line: tuple[str, int, int] | None = None


def _line_of(low: str, start: int, end: int) -> tuple[int, int]:
    """Where the line holding ``low[start:end]`` begins and ends.

    Each end is the nearest break of any kind, found without walking past
    it (decision 178). Looking for the nearest of *each* kind walked the
    whole text for every kind a page never uses - a form feed, a vertical
    tab - so every mention cost the length of the reading and a long one
    cost its square: a 1.9 MB text took a minute a return. The look back
    is one search over a window that doubles from ``start`` until it holds
    a break or reaches the start of the text; the look forward stops at
    the first break there is; and a mention on the line found last is on
    it, so many mentions on one long line pay for the line once. The
    answer is the old one on every text, whatever its line ends.
    """
    global _last_line
    last = _last_line
    if last is not None and last[0] is low and last[1] <= start and end <= last[2]:
        return last[1], last[2]
    left, reach = 0, _LOOK_BACK
    while True:
        lo = max(0, start - reach)
        found = _TO_THE_LAST_BREAK.match(low, lo, start)
        if found is not None:
            left = found.end()
            break
        if lo == 0:
            break
        reach *= 2
    ahead = _A_BREAK.search(low, end)
    right = ahead.start() if ahead else len(low)
    if _A_BREAK.search(low, start, end) is None:     # a mention across a break is on no one line
        _last_line = (low, left, right)
    return left, right


def _titles_its_own_line(low: str, start: int, end: int) -> bool:
    """A form number at the head of its line with a year beside it, before
    or after: "2025 Form 1099-INT - Interest Income", "Form 1098 -
    Mortgage Interest Statement 2025", "W-2 Wage and Tax Statement 2025
    Department of the Treasury". That is how a payer prints its own
    substitute form and how the IRS prints the foot of a one-copy
    information return, and the only shape a dash-set-off number is not a
    checklist's line in."""
    left, right = _line_of(low, start, end)
    if not _HEADS_ITS_LINE.fullmatch(low, left, start):
        return False
    return bool(_A_YEAR.search(low, left, start) or _A_YEAR.search(low, end, right))


#: A form's own printed title between its number and its year, with
#: nothing to set it off: "W-2 Wage and Tax Statement 2025", "1099-DIV
#: Dividends and Distributions 2025", "Form 1040 U.S. Individual Income
#: Tax Return 2025". The IRS prints that line at the foot of every copy of
#: an information return, and the 2024 and 2025 W-2 revisions print one
#: copy to the page, so Copy B says "W-2" exactly once and says it that
#: way. Words only, on the form's own line, and the year is a year rather
#: than the tail of a date - "Form 941 04/30/2025" on a CP 575 is a filing
#: deadline, and the notice is not a 941. A title is short.
_TITLE_WORDS = 8
_SELF_TITLED = re.compile(
    rf"(?:{_SAME_LINE}+[a-z][a-z&/.{re.escape(_APOSTROPHES)}{re.escape(_DASH_CHARS)}]*){{1,{_TITLE_WORDS}}}"
    rf"{_SAME_LINE}+\(?(?:19|20)\d{{2}}(?![0-9])"
)


def _names_itself(low: str, start: int, end: int, shape: str) -> bool:
    """Whether the mention at ``low[start:end]`` is a form printing its own
    name on a line of its own, in either shape a printed title wears: a
    dash between the number and the title, with a year beside them
    (``_titles_its_own_line``), or nothing between them and the year after
    the title (``_SELF_TITLED``). Both shapes head their line, because
    "2025 Consolidated Form 1099 - Account 8812-4455" is a broker's
    sentence about a family, not a form.
    """
    if shape == _BY_DASH:
        return _titles_its_own_line(low, start, end)
    if shape != _BY_NOTHING:
        return False
    left, _right = _line_of(low, start, end)
    return (_HEADS_ITS_LINE.fullmatch(low, left, start) is not None
            and _SELF_TITLED.match(low, end) is not None)


def _lists_forms(low: str) -> bool:
    """Whether the title window heads ``_FORMS_IN_A_MENU`` lines or more
    with a form naming itself (``_names_itself``): a menu of forms, not a
    form. A dash-set-off number counts wherever it stands, as decision 69
    counted it."""
    stop = len(_title(low))
    listed = 0
    for match in _FORM_MENTION.finditer(low, 0, stop):
        span = _mention_span(low, match)
        if span is None:
            continue
        shape = _weigh_mention(low, match, span[2])[1]
        if shape == _BY_DASH or _names_itself(low, match.start(), span[2], shape):
            listed += 1
            if listed >= _FORMS_IN_A_MENU:
                return True
    return False


def _scan(low: str):
    """Every form number ``low`` (lower-cased text) mentions, as
    ``(key, start, end, weight, shape)``.

    A form heading its own title line with its year beside it names
    itself (``_titles_its_own_line``), whether a dash sets the title off
    or nothing does: "W-2 Wage and Tax Statement 2025" is the dated
    self-mention "W-2 2025" is, printed the way the IRS prints a one-copy
    return's foot. A page that heads two such lines is a menu, and the
    mention that is not the first on the page is not the page's own.
    """
    menu = _lists_forms(low)
    first = True
    for match in _FORM_MENTION.finditer(low):
        span = _mention_span(low, match)
        if span is None:
            continue
        number, variant, end = span
        weight, shape = _weigh_mention(low, match, end)
        if (first and not menu
                and not _asked_for(low, match.start())
                and _names_itself(low, match.start(), end, shape)):
            weight, shape = _SELF_WEIGHT, _BY_YEAR
        first = False
        yield _form_key(number, variant), match.start(), end, weight, shape


def _mentions(low: str):
    """Every form number ``low`` mentions, as ``(key, start, weight)``."""
    for key, start, _end, weight, _shape in _scan(low):
        yield key, start, weight


def dominant_forms(text: str) -> set[str]:
    """The form number ``text`` is about, normalised (``w2``, ``1099int``),
    as a set of at most one.

    Mentions are weighed by their shape (``_weigh_mention``): a form's
    own number, dated, on the title and every footer outweighs the other
    forms its instructions quote. The heaviest is its own; among equals,
    the first mentioned. A number that never appears in its own right and
    only once in passing is nobody's - an IRS notice that lists three
    forms to file is about none of them.
    """
    scores: dict[str, int] = {}
    first: dict[str, int] = {}
    for key, start, weight in _mentions(text.lower()):
        scores[key] = scores.get(key, 0) + weight
        first.setdefault(key, start)
    if not scores:
        return set()
    top = max(scores.values())
    if top < 2:
        return set()
    return {min((key for key, n in scores.items() if n == top), key=first.get)}


#: A title that names this many form families in their own right is a
#: list of forms - an organizer's checklist ("Form W-2 - Wage and Tax
#: Statement", "Form 1098 - Mortgage Interest Statement", "Form 1099-INT
#: - Interest Income"), a transmittal's "Enclosed: Form W-2 2025, Form
#: 1098 2025, Form 1099-INT 2025" - and is none of them, as decision 63
#: says of a notice that lists three forms to file. Families, not
#: numbers: a broker's consolidated 1099 names "Form 1099", "1099-INT",
#: "1099-DIV" and "1099-B" and is one family's document, which the rows
#: that ask for its parts then contest (the eleventh reading found it
#: blinded, parked with no candidate at all).
_LIST_OF_FORMS = 3
_FAMILY = re.compile(r"[a-z]?\d{3,4}|w\d")


def _title_forms(low: str) -> set[str]:
    """The form numbers the title (``_title``) of ``low`` (lower-cased
    text) names in their own right: a mention there that is not a
    sentence about another form. None when the title names
    ``_LIST_OF_FORMS`` families or more."""
    stop = len(_title(low))
    forms = set()
    for key, start, weight in _mentions(low):
        if start >= stop:
            break
        if weight > _REFERENCE_WEIGHT:
            forms.add(key)
    families = {_FAMILY.match(key).group(0) for key in forms}
    return forms if len(families) < _LIST_OF_FORMS else set()


def form_family(key: str) -> str:
    """The family a form key belongs to: ``1099int`` and ``1099div`` are
    both ``1099``, ``w2`` is ``w2``. A broker's consolidated statement
    prints several of one family's numbers and is one document
    (``_title_forms``' note); two families on one page are two documents
    (decision 94)."""
    return _FAMILY.match(key).group(0)


#: How many distinct form families naming themselves make one page two
#: documents. A broker's consolidated statement prints "1099-INT" and
#: "1099-DIV" and is one family's document (``_title_forms`` says so of a
#: title); two families is two forms on one sheet (decision 94). Here, and
#: not in the router, since decision 107: the scanner reads through the
#: same rule on a cache miss.
MULTI_FORM_FAMILIES = 2


def own_forms(text: str) -> set[str] | None:
    """The forms ``text`` counts as its own when it names two or more
    families as itself, else ``None`` - the ordinary reading.

    **The one reading of which forms a page is** (decision 107). Decision
    94 files a page that prints two forms' own names under both requests,
    and the verdict the second copy needs is the page read with both forms
    counted as its own; decision 105 kept that verdict in the cache, and
    the cache alone. The step-0 re-run of 107 found what that meant: a
    rebuilt store - the runbook's own upgrade path, a new machine,
    ``rebuild`` - read the second copy the ordinary way, called it the
    other form's page, and the Saturday draft asked the client for a form
    already sent. The cache was carrying a decision, not a derivation. So
    the router's split (:mod:`tracker.router`) and the scanner's miss
    path (:func:`check_content`) both read through this, and the cache
    holds nothing a reader cannot recompute from the bytes.
    """
    named = self_named_forms(text)
    if len({form_family(key) for key in named}) >= MULTI_FORM_FAMILIES:
        return set(named)
    return None


#: The form whose section makes a statement a broker's (decision 146):
#: ``1099-B``, keyed the way :func:`form_key` keys a mention.
BROKER_FORM = "1099b"


def carries_a_1099b_section(text: str) -> bool:
    """Whether ``text`` carries a 1099-B section - which is what makes a
    broker's consolidated 1099 a brokerage statement (decision 146).

    **Not the word "consolidated".** A bank's "Consolidated Statement", the
    consolidated-return blanks and a broker's "your Consolidated Form 1099
    is available" email all say it, and none of them is a brokerage
    statement. What a consolidated statement with a 1099-B section has and
    a bank's combined 1099-INT and 1099-DIV page does not is the 1099-B,
    read as a form label exactly the way a ``1099-b`` keyword is read
    (:func:`_one_says_where`): named in its own right in the title
    (``_title_forms``, which keeps every 1099 variant a consolidated
    statement names because they are one family), or as the page's own
    dominant number. So this is true of a page exactly when a row asking
    for ``1099-b`` could be accepted *because of* the 1099-B - a checklist
    that lists it (a menu names nothing in its title) and a notice that
    only mentions it are not. A statement with only interest and dividend
    sections carries none and still files with the 1099-INT/DIV row: what
    decides is the 1099-B section, not who issued the statement.

    It says what the page carries and nothing more. Which request the page
    then files under is :mod:`tracker.router`'s: exactly one of the rows
    that accepted it must have been accepted because of the 1099-B, or it
    parks.
    """
    low = text.lower()
    return BROKER_FORM in _title_forms(low) or BROKER_FORM in dominant_forms(text)


def self_named_forms(text: str) -> tuple[str, ...]:
    """Every form ``text`` prints its *own* name on, first mention first.

    Decision 85's predicate asked of every mention rather than only of the
    page's first: a form number heading its own line with its year after
    it, with the form's printed title in between (``_names_itself``, both
    shapes a title wears) or with nothing in between (``_SELF_AFTER``'s
    dated self-mention, "Form 1040 (2025)", "941 for 2026"). Anything a
    page says *about* another form - quoted, told to the reader, asked for
    - names nothing, and a page the menu rules call a menu
    (``_lists_forms``, decision 73) names none at all: a checklist lists
    forms, it is not one, which is what stops a cover letter buying itself
    a filing per line.

    Where :func:`dominant_forms` asks which one form the page *is* and
    answers with at most one - weighing every mention against every other
    - this asks which forms print their own name on it. Two families of
    them (:func:`form_family`) is two documents on one sheet, which is the
    owner's decision 94: the page files a copy under each, but only when
    each of those forms is asked for by exactly one request. Nothing here
    decides that; :mod:`tracker.router` does, and this says only what the
    page names.
    """
    low = text.lower()
    if _lists_forms(low):
        return ()
    first: dict[str, int] = {}
    for match in _FORM_MENTION.finditer(low):
        span = _mention_span(low, match)
        if span is None:
            continue
        number, variant, end = span
        shape = _weigh_mention(low, match, end)[1]
        if _asked_for(low, match.start()):
            continue
        if shape == _BY_YEAR:
            # A dated self-mention still has to head its own line to be
            # this page's own: "the 2025 Form 1098 from each lender" in a
            # cover's prose is a sentence about a form.
            left, _right = _line_of(low, match.start(), end)
            if _HEADS_ITS_LINE.fullmatch(low, left, match.start()) is None:
                continue
        elif not _names_itself(low, match.start(), end, shape):
            continue
        first.setdefault(_form_key(number, variant), match.start())
    return tuple(sorted(first, key=first.__getitem__))


#: The words a document asks for a document with. What one of them
#: introduces is what the page wants, not what the page is: the firm's own
#: organizer prints "If yes, attach your brokerage statement and any
#: realized gain and loss report", and a row that keys on a broker's words
#: filed the organizer as the client's brokerage statement (the thirteenth
#: reading). An ask governs the rest of its own sentence, because one ask
#: names several documents, and the line below it when the ask is what
#: that line ends with - a request wraps where it will.
_ASK_WORDS = ("attached", "attach", "provide", "send", "include", "enclosed", "enclose",
              "upload", "submit", "bring", "list", "see", "copies of", "copy of")
#: What may stand between an ask and the document it asks for and still
#: leave the two joined ("attach your", "see the latest", "copies of all").
_ASK_FILLER = ("most recent", "applicable", "either", "latest", "each", "both", "your",
               "and", "any", "all", "the", "an", "or", "of", "a")
_ASKED_FOR = re.compile(rf"(?<![a-z0-9])(?:{'|'.join(_ASK_WORDS)})(?![a-z0-9])")
#: An ask that ends the line above, with nothing but filler after it: the
#: request wrapped, and what begins the next line is what it asked for.
_ASK_WRAPPED = re.compile(
    rf"(?<![a-z0-9])(?:{'|'.join(_ASK_WORDS)})(?![a-z0-9])(?:\s+(?:{'|'.join(_ASK_FILLER)})(?![a-z0-9]))*\s*$"
)
#: Where the sentence an ask began ends. A colon is not an end: it is how
#: a page introduces what it is asking for ("Please send the following:").
_SENTENCE_END = re.compile(r"[.?!;]")
#: How far back an ask is looked for: within its own sentence on its own
#: line (so a "see" earlier in a form's dense column is not one), and, for
#: the wrapped shape, the words the line above ends with.
_ASK_WITHIN = 200
_WRAP_WITHIN = 40


def _ask_window(low: str, start: int) -> int:
    """Where to start looking back from ``start`` for an ask-word: at most
    ``_ASK_WITHIN`` characters, no further back than this line, and never
    past the end of an earlier sentence."""
    edge = max(0, start - _ASK_WITHIN, _line_of(low, start, start)[0])
    ends = list(_SENTENCE_END.finditer(low, edge, start))
    return ends[-1].end() if ends else edge


def _asked_for(low: str, start: int) -> bool:
    """True when what begins at ``start`` is what an ask-word asked for.

    On its own line the ask governs the rest of its sentence, because one
    ask names several documents; across a line break it governs only the
    wrap - the ask and its filler ending the line above - because a line
    further up is another label, not the rest of this request. A tax
    return's "Attach Forms W-2G and 1099-R if tax was withheld" two lines
    above its jurat asks for nothing on the jurat's line.
    """
    if _ASKED_FOR.search(low, _ask_window(low, start), start):
        return True
    return _ASK_WRAPPED.search(low, max(0, start - _WRAP_WITHIN), start) is not None


def _menu_tails(low: str) -> set[int]:
    """Where each menu line's title begins: just past a form number set off
    from what follows by a dash or a comma.

    A checklist's "Form 1099-B - Proceeds From Broker and Barter Exchange
    Transactions" names the form and prints its title, and the whole line
    is the one reference: a keyword that starts there is the menu's word,
    not a word the document says about itself.
    """
    tails = set()
    for _key, _start, end, weight, shape in _scan(low):
        if weight != _REFERENCE_WEIGHT or shape not in (_BY_DASH, _BY_COMMA):
            continue
        separator = _REFERENCE_AFTER.match(low, end)
        if separator is None:                      # the shape says it matched
            continue
        tail = _SPACES.match(low, separator.end()).end()
        if not _BREAK_RUN.search(low, end, tail):  # the title is on the form's own line
            tails.add(tail)
    return tails


def _in_its_own_words_at(low: str, keyword: str) -> int | None:
    """Where ``low`` first says ``keyword`` the way a document says what it
    is, or None.

    An occurrence that is a menu line's title (``_menu_tails``) or what an
    ask-word asked for (``_asked_for``) is a page talking about a document
    rather than being one. One occurrence that is neither is enough: a
    document prints its own words plainly somewhere.
    """
    tails = None
    for match in _occurrences(low, keyword):
        if tails is None:                          # only worth scanning if it is said at all
            tails = _menu_tails(low)
        if match.start() in tails or _asked_for(low, match.start()):
            continue
        return match.start()
    return None


def _in_the_footer(low: str, at: int) -> bool:
    """Whether ``at`` falls in the last ``_FOOTER_LINES`` non-blank lines of
    its page - where a form repeats its own number on every copy."""
    start = low.rfind(PAGE_BREAK, 0, at) + 1
    end = low.find(PAGE_BREAK, at)
    end = len(low) if end < 0 else end
    spans: list[tuple[int, int]] = []
    offset = start
    for line in low[start:end].splitlines(keepends=True):
        if line.strip():
            spans.append((offset, offset + len(line)))
        offset += len(line)
    return any(left <= at < right for left, right in spans[-_FOOTER_LINES:])


def _where_said(text: str, at: int) -> tuple[str, int]:
    """The place and the 1-based page of the character at ``at``.

    The title wins over everything - a form printing its own name at the
    top is the strongest place a word can be said - then the footer, which
    is a real place on page 1 as much as on page 7, and only then the rest
    of the first page.
    """
    page = text.count(PAGE_BREAK, 0, at) + 1
    if at < len(_title(text)):
        return WHERE_TITLE, page
    if _in_the_footer(text, at):
        return WHERE_FOOTER, page
    return (WHERE_FIRST_PAGE if page == 1 else WHERE_DEEP), page


def _first_said_at(low: str, keyword: str, within: int | None = None) -> int | None:
    """Where ``low`` says ``keyword`` as a whole token, first or within the
    first ``within`` characters; None if it does not."""
    starts = (match.start() for match in _occurrences(low, keyword))
    if within is None:
        return next(starts, None)
    return next((start for start in starts if start < within), None)


def _says_where(text: str, keyword: str, dominant: set[str] | None = None) -> tuple[str, int] | None:
    """Where ``text`` says ``keyword`` as content evidence, reading the
    keyword's alternatives (``keyword_alternatives``): the first
    alternative whose every phrase is said, placed where its first phrase
    is. A plain keyword is one alternative of one phrase, so this is
    :func:`_one_says_where` for every keyword written before decision 90.

    The place named is the leading phrase's because that is the row's own
    order - the phrase the accountant wrote first is the one that says
    what the document is, and the others confirm it.
    """
    if dominant is None:
        dominant = dominant_forms(text)
    for alternative in keyword_alternatives(keyword):
        places = [_one_says_where(text, phrase, dominant) for phrase in alternative]
        if all(place is not None for place in places):
            return places[0]
    return None


#: The phrases a notice prints in its header block, which count as evidence
#: only on the first page (decision 141, the designer's ruling on the
#: review, N2). A notice says who sent it, the tax year and the notice date
#: at its head; the same words deep in another document - the back of a
#: county tax bill, a letter's closing paragraph - are that document talking
#: about a notice. Keyed by the phrase, as ``FORM_VARIANTS`` is by the
#: number, in the one module that reads keywords; the catalog's notices row
#: (``tracker.templates.SHARED["notices"]``) is the only row that uses them,
#: and ``tests/test_irs_forms.py`` holds it to that. A reading with no page
#: breaks (OCR reads a scan as one run of text) is all first page.
FIRST_PAGE_PHRASES = frozenset({"notice date", "tax year", "tax period", "irs", "internal revenue service"})


def _one_says_where(text: str, keyword: str, dominant: set[str]) -> tuple[str, int] | None:
    """Where ``text`` says one phrase as content evidence - the place and
    the page - or None where it does not say it at all.

    The verdict is :func:`says`'s, unchanged: a form number counts only
    where the title names it in its own right (``_title_forms``) or as the
    document's own (dominant) number, and every other keyword counts where
    the document says it in its own words (``_in_its_own_words_at``)
    rather than on a menu line or after an ask. All this adds is *where*,
    so an :class:`Evidence` can name it.
    """
    low = text.lower()
    if not is_form_number(keyword):
        at = _in_its_own_words_at(low, keyword)
        if at is None:
            return None
        place = _where_said(low, at)
        # The first place a phrase is said in its own words; a header phrase
        # first said past page 1 is not said on page 1 at all.
        if keyword.strip().lower() in FIRST_PAGE_PHRASES and place[1] > 1:
            return None
        return place
    at = _first_said_at(low, keyword)
    if at is None:
        return None
    key = form_key(keyword)
    if key in _title_forms(low):
        # It is the title that accepted it, so it is the title's occurrence
        # the evidence names, not whichever came first in a long document.
        in_title = _first_said_at(low, keyword, within=len(_title(low)))
        return _where_said(low, at if in_title is None else in_title)
    if key in dominant:
        return _where_said(low, at)
    return None


def says(text: str, keyword: str, dominant: set[str] | None = None) -> bool:
    """``contains_keyword`` as content evidence.

    A form number counts only where the title names it in its own right
    (``_title_forms``) or as the document's own (dominant) number. Every
    other keyword counts where the document says it in its own words
    rather than on a menu line or after an ask. A keyword naming
    alternatives is said when one of them is, every phrase of it
    (``keyword_alternatives``). The thin wrapper over :func:`_says_where`
    is deliberate: one reading of a document, whether the caller wants
    the verdict or the evidence behind it.
    """
    return _says_where(text, keyword, dominant) is not None


def any_keyword_matched(text: str, item: RequestItem) -> bool:
    """Whether the row's any-keywords accept ``text``, apart from every other rule.

    The year a Period implies is a *check* on a document a keyword already
    matched, never evidence on its own (decision 40). When only that check
    fails, the row is still the lead a person needs, and this is how
    :mod:`tracker.router` asks for it without reading a verdict's sentence.
    """
    if not item.any_keywords:
        return False
    dominant = dominant_forms(text)
    return any(says(text, k, dominant) for k in item.any_keywords)


def _found(rule: str, text: str, keyword: str, dominant: set[str]) -> Evidence | None:
    """The Evidence for one of the row's words under one rule, or None."""
    place = _says_where(text, keyword, dominant)
    return None if place is None else Evidence(rule, keyword, place[0], place[1])


#: The longest line a Date Pattern is run over (decision 137, L3). A
#: pattern is a person's typing, and a typo such as ``(\d+)+x`` backtracks
#: exponentially in the length of what it is run over. Run over the whole
#: of an OCR reading, one such pattern could hold a pass for hours; run
#: line by line, with long lines left out, the worst it can do is bounded
#: by one short line. A date on a real document sits on a short line.
DATE_LINE_MAX = 500


def date_pattern_at(pattern: str, text: str) -> int | None:
    """Where in ``text`` the row's Date Pattern first matches, or None.

    Compiled once, then run **line by line** (decision 137, L3), and a
    line longer than :data:`DATE_LINE_MAX` is skipped. The offset is into
    the whole text, so where the date was said is still found by
    :func:`_where_said`. A pattern that does not compile was refused when
    it was saved; one that reaches here anyway matches nothing.
    """
    try:
        compiled = re.compile(pattern)
    except re.error:
        return None
    offset = 0
    for line in text.splitlines(keepends=True):
        if len(line) <= DATE_LINE_MAX and (match := compiled.search(line)):
            return offset + match.start()
        offset += len(line)
    return None


def evaluate_rules(text: str, item: RequestItem, dominant: set[str] | None = None) -> ContentResult:
    """Apply the manifest row's content rules to extracted text.

    The verdict and the reason are what they have always been; what is new
    is that every rule that *did* find its word leaves an :class:`Evidence`
    behind, on a failing verdict as much as on a passing one, because what
    matched is the lead a person works a parked file from.

    ``dominant`` is which form numbers count as the document's own, and is
    :func:`dominant_forms`'s answer unless a caller says otherwise. A page
    that names two forms as itself is two documents
    (:func:`self_named_forms`, decision 94), and the ordinary reading calls
    at most one of them the page's own; the router reads such a page a
    second time with both, and files only under the strict rule decision
    94 sets. Since decision 107 the scanner's miss path reads it the same
    way, through :func:`own_forms`, so every verdict the cache keeps is
    one the miss path reproduces; for a page naming at most one family
    :func:`own_forms` is ``None`` and this is the reading it always was.
    """
    if dominant is None:
        dominant = dominant_forms(text)
    found: list[Evidence] = []
    missing: list[str] = []
    for keyword in item.required_keywords:
        evidence = _found(RULE_REQUIRED, text, keyword, dominant)
        if evidence is None:
            missing.append(keyword)
        else:
            found.append(evidence)
    if missing:
        listed = ", ".join(f"'{k}'" for k in missing)
        return ContentResult(ok=False, reason=reasons.WRONG_DOCUMENT.format(listed=listed),
                             evidence=tuple(found))

    if item.any_keywords:
        hits = [e for e in (_found(RULE_ANY, text, k, dominant) for k in item.any_keywords)
                if e is not None]
        if not hits:
            listed = ", ".join(item.any_keywords)
            return ContentResult(ok=False, reason=reasons.NO_EXPECTED_KEYWORD.format(listed=listed),
                                 evidence=tuple(found))
        found.extend(hits)

    if item.date_pattern:
        # A pattern a person typed runs line by line (decision 137, L3); the
        # one derived from the Period is the firm's own and runs over the
        # whole text as it always did (the review's F7), so "December" and
        # "2025" split by a line break still say the year.
        if item.date_pattern_derived:
            hit = re.search(item.date_pattern, text)
            at = hit.start() if hit else None
        else:
            at = date_pattern_at(item.date_pattern, text)
        if at is None:
            return ContentResult(ok=False,
                                 reason=reasons.WRONG_PERIOD.format(pattern=item.date_pattern),
                                 evidence=tuple(found))
        # The row's Period, not the regex it derives and not a word of the
        # document: what a person reading the index needs is "TY2025".
        where, page = _where_said(text, at)
        found.append(Evidence(RULE_DATE, item.period or item.date_pattern, where, page))

    return ContentResult(ok=True, evidence=tuple(found))


# ------------------------------------------------------------- extraction ----


def _extract_pdf(path: Path) -> str:
    import pdfplumber  # deferred: heavy import, only needed for PDFs

    parts: list[str] = []
    said = 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:MAX_PAGES]:
            parts.append(page.extract_text() or "")
            said += len(parts[-1]) + 1
            if said > READING_CHAR_BUDGET:
                break           # :func:`_extract` cuts it and says so
    return PAGE_BREAK.join(parts)


def _extract_xlsx(path: Path) -> str:
    import datetime as dt

    from openpyxl import load_workbook

    # A sheet's row is a line and its cells are set apart by a tab: a
    # keyword's words may run across a row's cells ("Fixed Asset" beside
    # "Schedule") and never down its rows.
    with zipfile.ZipFile(path) as archive:
        if any(info.compress_type not in PACKINGS for info in archive.infolist()):
            raise ValueError(UNKNOWN_PACKING)
    parts: list[str] = []
    said = 0
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            parts.append(str(ws.title))
            said += len(parts[-1]) + 1
            for row in ws.iter_rows(values_only=True):
                if said > READING_CHAR_BUDGET:
                    return "\n".join(parts)      # :func:`_extract` cuts it and says so
                cells: list[str] = []
                for value in row:
                    if value is None:
                        continue
                    if isinstance(value, (dt.datetime, dt.date)):
                        # Both forms, so ISO- and US-style patterns match.
                        d = value.date() if isinstance(value, dt.datetime) else value
                        cells.append(f"{d.isoformat()} {d.month}/{d.day}/{d.year}")
                    else:
                        cells.append(str(value))
                if cells:
                    parts.append("\t".join(cells))
                    said += len(parts[-1]) + 1
    finally:
        wb.close()
    return "\n".join(parts)


def _text_cap_bytes() -> int:
    return TEXT_READ_CAP_MB * 1024 * 1024


def _is_cut(path: Path) -> bool:
    """Whether :func:`_extract_textfile` reads less than the whole file."""
    try:
        return path.stat().st_size > _text_cap_bytes()
    except OSError:
        return False


def _extract_textfile(path: Path) -> str:
    """The text of a ``.csv``/``.tsv``/``.txt``, **up to its first**
    ``TEXT_READ_CAP_MB`` (decision 137): enough for any real statement, and
    a file past it is read that far and no further. A UTF-8 character the
    cut split in two is left off rather than turning the whole reading
    into another encoding."""
    cap = _text_cap_bytes()
    with path.open("rb") as handle:
        raw = handle.read(cap + 1)
    cut = len(raw) > cap
    raw = raw[:cap]
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as exc:
        if cut and exc.start >= len(raw) - 3:
            try:
                return raw[:exc.start].decode("utf-8")
            except UnicodeDecodeError:
                pass
    try:
        return raw.decode("cp1252")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="replace")


def extract_text(path: Path) -> str | None:
    """Extract the text layer from a supported file; None if no extractor exists.

    A photo comes back None **on purpose** (decision 127): an image has no
    text layer to extract, ever, and :func:`extract` reads that None as
    "this is a scan" rather than "no extractor for this type" - the one
    place the two answers have to be told apart.
    """
    extension = extension_of(path)
    if extension == PDF_EXTENSION:
        return _extract_pdf(path)
    if extension in XLSX_EXTENSIONS:
        return _extract_xlsx(path)
    if extension in TEXT_EXTENSIONS:
        return _extract_textfile(path)
    return None


class ReadingStopped(Exception):
    """A reading reached the safety stop (decision 137, B1.2). ``seconds``
    is how long it had been reading the file when it stopped."""

    def __init__(self, seconds: float):
        super().__init__(f"the reading stopped after {seconds:.0f} s")
        self.seconds = seconds


class _SafetyStop:
    """The time one document's reading may take: a document deadline and,
    within it, one page's (decision 137, B1.2)."""

    def __init__(self) -> None:
        self.started = _clock()
        self.deadline = self.started + READING_STOP_DOCUMENT_SECONDS
        self.page_deadline = self.deadline

    def page(self) -> None:
        """A page is about to start: none starts past the stop - the page
        before it overran its minute, or the document its ten (decision
        169, ruling 5; the reader takes no timeout of its own, so a page
        under way is bounded by the child's end at the document's stop).
        Then it has a minute, or what is left of the ten."""
        self.remaining()
        self.page_deadline = min(_clock() + READING_STOP_PAGE_SECONDS, self.deadline)

    def remaining(self) -> float:
        """Seconds left for the page now being read, or the stop."""
        left = self.page_deadline - _clock()
        if left <= 0:
            raise ReadingStopped(_clock() - self.started)
        return left


#: The stop of the reading under way, set by :func:`extract_by_ocr`. None
#: outside a reading (a test's call straight into a reader).
_STOP: _SafetyStop | None = None


def render_scale(width: float, height: float) -> float:
    """The scale a PDF page of ``width`` x ``height`` points is rendered at
    for OCR: :data:`RENDER_SCALE`, or less, so the page never passes
    :data:`PIXEL_BUDGET` (decision 137, B1). At scale 1 a point is a pixel,
    so the budget over the page's own area is the square of the most it
    may be scaled by."""
    area = max(float(width), 1.0) * max(float(height), 1.0)
    return min(RENDER_SCALE, (PIXEL_BUDGET / area) ** 0.5)


def _within_budget(image):
    """``image``, reduced to :data:`PIXEL_BUDGET` if it is larger (decision
    137, B1): a photo past it is read smaller, never refused."""
    pixels = image.width * image.height
    if pixels <= PIXEL_BUDGET:
        return image
    factor = (PIXEL_BUDGET / pixels) ** 0.5
    size = (max(1, int(image.width * factor)), max(1, int(image.height * factor)))
    return image.resize(size)


def _ocr_pdf(path: Path) -> str | None:
    """OCR the first pages of a PDF. None when the reader cannot run on
    this machine (:class:`tracker.ocr.ReaderUnavailable`: a model missing or
    broken) - a normal, reported condition, never an error.

    Each page is rendered at :func:`render_scale` of its own size (decision
    137), so no page is ever drawn past :data:`PIXEL_BUDGET`: an ordinary
    page at the scale it always had, a giant one smaller. It is read as it
    comes (decision 169, ruling 6): the reader keeps the words of a page
    scanned sideways, and nothing turns it. The pages are joined with a
    line break, not :data:`PAGE_BREAK`, as they always were (SPEC-127.2X
    D9.2): the footer and first-page rules read an OCR reading as one page.
    """
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return None

    try:
        parts: list[str] = []
        doc = pdfium.PdfDocument(path)
        try:
            for index in range(min(len(doc), _MAX_OCR_PAGES)):
                if _STOP is not None:
                    _STOP.page()              # no page starts past the stop
                page_of = doc[index]
                bitmap = page_of.render(scale=render_scale(*page_of.get_size()))
                if _STOP is not None:
                    _STOP.remaining()         # the render counts against the budget
                parts.append(ocr.read_page(bitmap.to_pil().convert("RGB"), name=path.name))
        finally:
            doc.close()
        return "\n".join(parts)
    except ocr.ReaderUnavailable as exc:
        log.warning("The reader cannot run on this machine (%s)", exc)
        return None
    except ReadingStopped:
        raise
    except MemoryError:
        raise           # the child's memory limit: a crash, not a retry (SPEC-169 section 9)
    except Exception as exc:
        log.warning("OCR failed on %s: %s", path.name, exc)
        raise OcrError(f"{exc.__class__.__name__}: {exc}") from exc


def _ocr_image(path: Path) -> str | None:
    """OCR one photo. None when the reader cannot run on this machine.

    One photo is one page. The phone's own rotation flag is undone first
    (``exif_transpose``) - a photo held upright and recorded sideways is
    the commonest case there is, and EXIF alone settles it - and the photo
    is read as it then stands (decision 169, ruling 6).

    Nothing is written: the photo on disk is the client's original.

    Its size is checked against :data:`PIXEL_BUDGET` **before** it is
    decoded (decision 137): a JPEG past it is decoded smaller by its own
    decoder (``draft``), any other photo is reduced to the budget the moment
    it is decoded, and Pillow's own decompression-bomb guard stays on for
    both. A photo too large even for that is refused by Pillow rather than
    decoded, and that refusal is a kept "Too large to read" verdict for a
    person (:class:`TooLargeToRead`), not a retry every pass.
    """
    try:
        from PIL import Image, ImageOps
    except ImportError:
        return None

    try:
        with Image.open(path) as opened:
            if opened.width * opened.height > PIXEL_BUDGET and opened.format == "JPEG":
                factor = (PIXEL_BUDGET / (opened.width * opened.height)) ** 0.5
                opened.draft("RGB", (int(opened.width * factor), int(opened.height * factor)))
            image = _within_budget(ImageOps.exif_transpose(opened).convert("RGB"))
        return ocr.read_page(image, name=path.name)
    except ocr.ReaderUnavailable as exc:
        log.warning("The reader cannot run on this machine (%s)", exc)
        return None
    except ReadingStopped:
        raise
    except Image.DecompressionBombError as exc:
        # Pillow's own guard, which stays on (decision 137, B1): a picture
        # too large even to decode smaller. A size rule, so it is a kept
        # verdict for a person, not a retry every pass.
        raise TooLargeToRead(picture_too_large_reason(exc)) from exc
    except MemoryError:
        raise           # the child's memory limit: a crash, not a retry (SPEC-169 section 9)
    except Exception as exc:
        log.warning("OCR failed on %s: %s", path.name, exc)
        raise OcrError(f"{exc.__class__.__name__}: {exc}") from exc


class OcrError(RuntimeError):
    """The reader ran but failed on this file this time; try again later."""


class TooLargeToRead(Exception):
    """The file is too large to read at all (decision 137): its reason is
    ``reasons.TOO_LARGE``, and it is the file's, not the machine's - a kept
    verdict, not retried until the file changes."""


def extract(path: Path, *, ocr: bool = True) -> Extraction:
    """What ``path`` says - the one reading both the router and the scanner use.

    With ``ocr=False`` a scan (a PDF with no text layer, or a photo, which
    never has one) comes back with ``needs_ocr`` set and whatever little
    text there was, so the caller can decide - the router first tries the
    file's name - and then call :func:`extract_by_ocr` if it still needs
    the words.

    Every reading is timed and the seconds ride back on the
    :class:`Extraction` (decision 127). Nothing acts on the number: no
    reading is cut short for being slow - a pass says its slowest reading
    and finishes it. The one exception is the safety stop, ten times the
    owner's ceiling (decision 137, B1.2, :data:`READING_STOP_PAGE_SECONDS`
    and :data:`READING_STOP_DOCUMENT_SECONDS`): a reading that reaches it
    is abandoned, and its verdict kept. **What the stop covers here:** no
    OCR page starts past it (decision 169, ruling 5), and an overrun is
    noticed between pages. It cannot interrupt a page being read, the PDF
    text layer (pdfplumber), a workbook's reading, or a single page render,
    which run inside the reading's own process and can only be seen to
    have overrun once they return. The bound on the whole reading is the process
    itself: the pass reads through :func:`extract_bounded`, which runs
    this function in a child it ends at the stop (decision 150).
    """
    started = time.perf_counter()
    # A file past the ceiling is never opened (decision 137, M5): the
    # reason is the whole reading, and it is the file's, not the machine's,
    # so the verdict is kept and the file is not tried again next pass.
    if too_large := too_large_reason(path):
        return Extraction(None, reason=too_large, extractable=False)
    reading = _extract(path, ocr=ocr)
    return replace(reading, seconds=time.perf_counter() - started)


def _extract(path: Path, *, ocr: bool) -> Extraction:
    """:func:`extract`'s reading, untimed."""
    extension = extension_of(path)
    if extension in IMAGE_EXTENSIONS:
        # A photo is a scan with no text layer to be disappointed by: it
        # goes the way a scanned PDF goes, and with ocr=False it is
        # unrouted for the same reason a scan is - there are no words yet.
        return Extraction(None, needs_ocr=True) if not ocr else extract_by_ocr(path)
    try:
        text = extract_text(path)
    except MemoryError:
        raise           # the child's memory limit: a crash, not a corrupt file (SPEC-169 section 9)
    except Exception as exc:  # a corrupt file is a reason, not a crash
        error = f"{exc.__class__.__name__}: {exc}"
        return Extraction(
            None, reason=reasons.EXTRACTION_FAILED.format(error=error),
            extractable=False, error=error,
        )
    if text is None:
        return Extraction(
            None, reason=reasons.UNCHECKABLE_TYPE.format(extension=extension), extractable=False,
        )
    pages = text.count(PAGE_BREAK) + 1
    if extension == PDF_EXTENSION and len(text.strip()) < _MIN_TEXT_CHARS * pages:
        if not ocr:
            return Extraction(text, needs_ocr=True)
        return extract_by_ocr(path)
    cut = extension in TEXT_EXTENSIONS and _is_cut(path)
    if len(text) > READING_CHAR_BUDGET:
        text, cut = text[:READING_CHAR_BUDGET], True
    return Extraction(text, cut=cut)


def extract_by_ocr(path: Path) -> Extraction:
    """Read a scan or a photo by OCR; says why when it cannot (the reader
    cannot run on this machine, no words, or the safety stop)."""
    global _STOP

    reader = _ocr_image if extension_of(path) in IMAGE_EXTENSIONS else _ocr_pdf
    # The safety stop (decision 137, B1.2): one per document, and a photo
    # is its own one page.
    _STOP = _SafetyStop()
    _STOP.page()
    try:
        ocr_text = reader(path)
    except ReadingStopped as exc:
        return abandoned(exc.seconds)
    except TooLargeToRead as exc:
        return Extraction(None, reason=str(exc), extractable=False)
    except OcrError as exc:
        # Ours to retry, not the client's to resend: the file may be fine.
        return Extraction(
            None, reason=reasons.OCR_FAILED.format(error=str(exc)),
            extractable=False, error=str(exc), transient=True,
        )
    finally:
        _STOP = None
    if ocr_text is None:
        # The reader cannot run on this machine: a fact about the machine,
        # remembered by nobody, so the day it is mended the scan reads the file.
        return Extraction(None, reason=reasons.NO_TEXT_LAYER.format(), extractable=False, transient=True)
    if not ocr_text.strip():
        return Extraction(None, reason=reasons.NO_TEXT_AFTER_OCR.format(), extractable=False)
    return Extraction(ocr_text, from_ocr=True)


def abandoned(seconds: float) -> Extraction:
    """The reading the safety stop abandoned after ``seconds`` (decision
    137, B1.2): no text, and a verdict that is the file's, so it is kept
    and the file is not read again until it changes. One sentence, whether
    the stop was noticed between pages inside the reading or was the end of
    the reading's own process (decision 150)."""
    minutes = max(1, round(seconds / 60))
    said = "1 minute" if minutes == 1 else f"{minutes} minutes"
    return Extraction(None, reason=reasons.READING_STOPPED.format(minutes=said),
                      extractable=False, seconds=seconds)


# ----------------------------------------------- a reading the pass can stop ----
#
# Decision 150. The stop above is noticed between pages, and nothing in
# the reading's own process can interrupt a page being read, the PDF text
# layer, a page's render or a workbook: a compressed PDF of a few megabytes
# can hold hours of drawing commands for pdfplumber or pdfium, far under
# the size ceiling. So the pass reads in a child process and waits for each
# document at most the document's stop. On time, the child hands back the
# Extraction it made, exactly what extract() returns in the pass's own
# process. Over time, the child is ended - with everything it started - and
# the reading is abandoned: the same kept verdict as above. A child that
# ends without answering - pdfium or the reader crashing, memory running
# out - is a reading that failed, kept the same way: the file waits for a
# person, and the pass goes on to the next one.
#
# Since decision 169 (R-4) the child is **one per pass**, serving documents
# one at a time, and it lives in tracker.ocr with the reader it runs: the
# reader's engine takes about a second to load on either device, and a
# child per document paid it every time. It is replaced after a stop, a
# crash, a fault on the graphics card and every hundred documents
# (ocr.ReadingChild, ocr.reading_session). This module words what became
# of each document.
#
# The child makes tier 2's open test too, before it reads (the designer's
# ruling on 150's review): opening a PDF is parsing it, and a page tree of a
# few kilobytes can keep pypdf counting pages for ever. So nothing of a
# client's file is parsed in the pass's own process: the open test's
# verdict comes back with the reading (Extraction.opened), and the router
# and the scanner take it from there (open_verdict, kept by fingerprint).
#
# A child that cannot start at all is the machine's fault and not the
# file's. It says "started" before it does anything with a file, and an
# end before that - or no word by the stop - is a reader that could not
# start (reasons.READER_UNAVAILABLE): transient, nothing kept, the drop
# neither decided nor recorded (tracker.filer leaves it for the next pass)
# and the pass warned once. Only an end after "started" is kept against
# the file.
#
# Opening an email or a zip is parsing it too (decision 154): olefile, the
# standard library's email and zipfile, on bytes a client sent. So it is
# one more job for the same child (in_a_child, the one mechanism), under
# the stop for a file: tracker.containers.open_bounded hands the child the
# container, and the child hands back the parts as plain data. The pass
# writes every attachment, and only once the child has answered.
#
# And the child never outlives its pass: a lifeline pipe and, on Windows, a
# job object (tracker.ocr says how).

#: Whether :func:`extract_bounded` reads in a child process (decision 150).
#: Always, in the tracker. The suite turns it off for every test but the
#: ones about the child (``tests/conftest.py``): its stand-in readers are
#: patched into the test's own process, which a child never shares.
READ_IN_A_CHILD = True
#: What the child runs: the open test and the one reading
#: (:func:`open_and_read`). A name of its own, handed to the child by
#: reference, so the suite can give the child a reader that never finishes
#: or dies - a patch made in the pass's process is not in the child's,
#: which imports this module afresh.
_CHILD_READER = None     # open_and_read, set below it
#: The fingerprint the open test's verdict is kept under in the verdict
#: cache (decision 150). It is about the file, whatever the row, so it
#: sits beside the rows' verdicts under a name no rules fingerprint - a
#: SHA-256 - can ever be.
OPEN_TEST_FINGERPRINT = "open-test"
#: The files whose reader could not start this pass, by name (decision
#: 150). The pass says it once (``runner``), taking the list as it goes.
_COULD_NOT_START: list[str] = []
def reading_stop_seconds(path: Path) -> float:
    """How long the whole reading of ``path`` may take (decision 150): a
    photo is one page - :data:`READING_STOP_PAGE_SECONDS` - and any other
    file is a document - :data:`READING_STOP_DOCUMENT_SECONDS`."""
    if extension_of(path) in IMAGE_EXTENSIONS:
        return READING_STOP_PAGE_SECONDS
    return READING_STOP_DOCUMENT_SECONDS


def extract_bounded(path: Path, *, ocr: bool = True) -> Extraction:
    """:func:`extract`, in a process the pass can stop (decision 150).

    What the pass reads through - the router's one reading of a drop
    (:func:`tracker.router.read_once`) and the scanner's reading on a
    cache miss - so the safety stop bounds the whole reading, text layer
    and render included, and a reader that crashes parks the file instead
    of ending the pass. A cached verdict never gets here, so it never
    starts a child. The child writes nothing anywhere: it hands back the
    reading, and the pass does every write, as before. The reader writes no
    temporary file either (SPEC-169 section 6), so the scratch folder
    decision 137 (L7) gave OCR is gone. The benchmark calls :func:`extract`
    itself: it measures the reader, not the stop.
    """
    if not READ_IN_A_CHILD:
        return open_and_read(Path(path), ocr=ocr)
    return _read_in_a_child(Path(path), ocr=ocr)


def open_and_read(path: Path, *, ocr: bool = True) -> Extraction:
    """Tier 2's open test, then the reading: what a reading's child does
    with the file (decision 150). The open test first, as tier 2 comes
    before tier 3; its verdict rides back on the reading as ``opened``."""
    opened = open_test(path)
    return replace(extract(path, ocr=ocr), opened=opened)


_CHILD_READER = open_and_read


def could_not_start(seconds: float, error: str, name: str) -> Extraction:
    """The reading whose child never started (decision 150): the machine's
    doing, so transient - nothing is kept - and the pass says it once."""
    _COULD_NOT_START.append(name)
    return Extraction(None, reason=reasons.READER_UNAVAILABLE.format(), extractable=False,
                      error=error, transient=True, seconds=seconds)


def readers_that_could_not_start() -> list[str]:
    """The files whose reader could not start since this was last asked,
    and forget them: the pass asks once, at its end, and says it once."""
    names = list(_COULD_NOT_START)
    _COULD_NOT_START.clear()
    return names


def open_verdict(path: Path, cache: ContentCache) -> str:
    """Tier 2's open test of ``path`` for the scanner: ``""`` or the
    refusal, never made in this process (decision 150).

    Kept in ``cache`` by the file's content digest, like a reading's
    verdicts: a hit starts no child. On a miss the file is read in a child
    (:func:`extract_bounded`) - the open test with the reading - and the
    reading is held in memory for the content check that follows, so the
    file is read once. A stopped or crashed child is the verdict, kept; a
    child that could not start, or a machine without the HEIC decoder, is
    the machine's, said and not kept.
    """
    digest = cache.digest_of(path)
    if digest:
        hit = cache.get_by_digest(digest, OPEN_TEST_FINGERPRINT)
        if hit is not None:
            return hit.reason
    reading = extract_bounded(path)
    verdict = reading.opened if reading.opened is not None else reading.reason
    if digest:
        cache.hold_reading(digest, reading)
        if not reading.transient and not reasons.HEIC_NOT_SUPPORTED.matches(verdict):
            cache.put_by_digest(digest, OPEN_TEST_FINGERPRINT,
                                ContentResult(ok=not verdict, reason=verdict))
    return verdict


def reading_failed(seconds: float, error: str) -> Extraction:
    """The reading whose process ended without an answer (decision 150):
    the file's verdict, kept like an abandoned one, for a person to read."""
    return Extraction(None, reason=reasons.READING_CRASHED.format(), extractable=False,
                      error=error, seconds=seconds)


def _read_in_a_child(path: Path, *, ocr: bool) -> Extraction:
    """Read ``path`` in a child process, waiting at most its stop."""
    answer, failed = in_a_child(path, _CHILD_READER, ocr=ocr)
    return failed if failed is not None else answer


def in_a_child(path: Path, job, *, stop: float | None = None, **kwargs) -> tuple[object, Extraction | None]:
    """Run ``job(path, **kwargs)`` in the pass's reading child
    (:func:`tracker.ocr.run_in_child`), waiting at most ``stop`` seconds (by
    default the file's own, :func:`reading_stop_seconds`). The one mechanism
    for every parse of a client's file (decision 150): the reading, and the
    opening of an email or a zip (decision 154,
    :func:`tracker.containers.open_bounded`).

    Gives ``(answer, None)`` on time, ``answer`` being what the job
    returned, and ``(None, why)`` otherwise, ``why`` the Extraction that
    says it: the stop (:func:`abandoned`), an end after "started"
    (:func:`reading_failed`), or a child that never started
    (:func:`could_not_start`, transient, nothing to keep). ``job`` is a
    module-level function, handed to the child by reference; what it
    returns crosses the pipe, so it is plain data.
    """
    path = Path(path)
    stop = reading_stop_seconds(path) if stop is None else stop
    outcome = ocr.run_in_child(job, (path,), kwargs, stop)
    if outcome.kind == "read":
        return outcome.answer, None
    if outcome.kind == "not_started":
        log.warning("The reader could not start for %s: %s", path.name, outcome.error)
        return None, could_not_start(outcome.seconds, outcome.error, path.name)
    if outcome.kind == "stopped":
        log.warning("Reading %s stopped at the safety stop (%.0f s)", path.name, stop)
        return None, abandoned(outcome.seconds)
    if outcome.kind == "failed":
        log.warning("The reader failed on %s: %s", path.name, outcome.trace)
        return None, reading_failed(outcome.seconds, outcome.error)
    log.warning("The reader stopped unexpectedly on %s: %s", path.name, outcome.error)
    return None, reading_failed(outcome.seconds, outcome.error)


# ---------------------------------------------------------------- checking ----


def check_content(
    path: Path, item: RequestItem, cache: ContentCache | None = None
) -> ContentResult:
    """Tier-3 verdict for one file, using the cache when possible.

    Files for items with no content rules pass immediately — no extraction,
    no file read. Callers must only pass files that already cleared the
    tier-2 placeholder check (reading a cloud-only file would hydrate it).
    """
    if not has_content_rules(item):
        return ContentResult(ok=True)

    fingerprint = rules_fingerprint(item)
    if cache is not None:
        hit = cache.get(path, fingerprint)
        if hit is not None:
            return hit

    result = _check_uncached(path, item, cache.held_reading(path) if cache is not None else None)

    if cache is not None and not result.transient:
        cache.put(path, fingerprint, result)
    return result


def _check_uncached(path: Path, item: RequestItem, reading: Extraction | None = None) -> ContentResult:
    """The scan's reading on a cache miss: the reading the router made.

    A page that names two forms as itself is read with both counted as its
    own (:func:`own_forms`), which is the verdict the router filed the
    second copy on; every other page reads exactly as it always did. So a
    verdict from an empty cache and a verdict the router kept are one
    verdict, and a rebuilt store cannot turn a filed copy into a failure.
    """
    reading = reading if reading is not None else extract_bounded(path)
    if reading.text is None:
        return ContentResult(
            ok=False, reason=reading.reason, extractable=False, transient=reading.transient,
        )
    verdict = evaluate_rules(reading.text, item, own_forms(reading.text))
    if reading.cut and not verdict.ok:
        # The rules were asked of the first part only, and a person reading
        # the reason must know that before believing "not found".
        verdict = replace(verdict, reason=f"{verdict.reason}; "
                                          f"{reasons.TEXT_CUT.format(limit=TEXT_READ_CAP_MB)}")
    return verdict


# ------------------------------------------------------------------- cache ----


#: How many readings a :class:`ContentCache` holds in memory for the
#: content check that follows an open test (decision 150).
_HELD_READINGS = 8


class ContentCache:
    """Verdict cache: ``(content digest, rules-fingerprint) -> ContentResult``.

    Stores only pass/fail verdicts and reasons — never extracted client
    text. A per-path memo of (size, mtime, digest) spares the hashing of a
    file that has not changed; the verdicts themselves are keyed by what
    the file *is*, so a working copy of a routed drop is a hit under its
    new name.

    **Backed by the store** (decision 107): built for an engagement, it
    loads that engagement's memos and verdicts at :data:`CACHE_VERSION`
    from the store's two tables once, works in memory exactly as it always
    did, and :meth:`save` writes back what changed - the memos and verdicts
    learned, the paths and digests pruned - in one transaction of the
    store's own, under the engagement lock. The engagement must already be
    in the store (both writers ``ensure()`` first); one it does not hold is
    refused by name here, at construction, rather than at the save. Built
    with no engagement it is memory only: the command line's cache, a dry
    run's, a test's - :meth:`save` then touches nothing. Disposable by
    design: a row written under another version simply means re-extraction.
    """

    def __init__(self, engagement_dir: Path | None = None):
        self._engagement = Path(engagement_dir) if engagement_dir is not None else None
        self._files: dict[str, dict] = {}                # path key -> size, mtime_ns, digest
        self._verdicts: dict[str, dict[str, dict]] = {}  # digest -> fingerprint -> verdict
        # What this cache has learned and forgotten since it was loaded, so
        # a save writes exactly those rows and no others.
        self._learned_files: set[str] = set()
        self._learned_verdicts: set[tuple[str, str]] = set()
        self._forgotten_files: set[str] = set()
        self._forgotten_digests: set[str] = set()
        #: Readings :func:`open_verdict` made in a child and the content
        #: check will want next, by digest (decision 150): in memory, a
        #: few at a time, for this pass only - never saved, like the
        #: reading the router holds for the drop it routes.
        self._readings: OrderedDict[str, Extraction] = OrderedDict()
        if self._engagement is not None:
            self._files, self._verdicts = store.cached_verdicts(
                store.connect(), self._engagement, version=CACHE_VERSION)

    @staticmethod
    def _key(file: Path) -> str:
        return str(file.resolve()).lower()  # Windows paths are case-insensitive

    def digest_of(self, file: Path) -> str | None:
        """The file's content digest, hashed once per (size, mtime); None if it vanished."""
        try:
            stat = file.stat()
        except OSError:
            return None
        key = self._key(file)
        memo = self._files.get(key)
        if memo and memo.get("size") == stat.st_size and memo.get("mtime_ns") == stat.st_mtime_ns:
            return str(memo["digest"])
        try:
            digest = sha256_of(file)
        except OSError:
            return None
        self._files[key] = {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "digest": digest}
        self._learned_files.add(key)
        self._forgotten_files.discard(key)
        return digest

    def remember_digest(self, file: Path, digest: str) -> None:
        """Remember ``digest`` for ``file`` as it now stands, without reading it.

        For a file whose bytes were proved under another name a moment ago
        and renamed to this one (decision 155's working copy: proved on its
        temp, then swapped in) - the rename keeps the size and the
        modification time, so the memo is what :meth:`digest_of` would have
        learned by reading it again. A file that has gone is not remembered.
        """
        try:
            stat = file.stat()
        except OSError:
            return
        key = self._key(file)
        self._files[key] = {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "digest": digest}
        self._learned_files.add(key)
        self._forgotten_files.discard(key)

    def get(self, file: Path, fingerprint: str) -> ContentResult | None:
        digest = self.digest_of(file)
        return self.get_by_digest(digest, fingerprint) if digest else None

    def get_by_digest(self, digest: str, fingerprint: str) -> ContentResult | None:
        entry = self._verdicts.get(digest, {}).get(fingerprint)
        if entry is None:
            return None
        defaults = {f.name: f.default for f in fields(ContentResult)}
        values = {name: entry.get(name, default) for name, default in defaults.items()
                  if name in entry or default is not MISSING}
        # asdict() flattened each Evidence on the way in; a verdict read back
        # must be the verdict that was stored, evidence and all.
        if "evidence" in values:
            values["evidence"] = _evidence_from_json(values["evidence"])
        return ContentResult(**values)

    def hold_reading(self, digest: str, reading: Extraction) -> None:
        """Keep a reading in memory for the content check about to ask
        (decision 150); never saved, and only the last few."""
        self._readings[digest] = reading
        while len(self._readings) > _HELD_READINGS:
            self._readings.popitem(last=False)

    def held_reading(self, file: Path) -> Extraction | None:
        """The reading :meth:`hold_reading` kept for ``file``'s bytes, if any."""
        digest = self.digest_of(file)
        return self._readings.get(digest) if digest else None

    def put(self, file: Path, fingerprint: str, result: ContentResult) -> None:
        digest = self.digest_of(file)
        if digest:
            self.put_by_digest(digest, fingerprint, result)

    def put_by_digest(self, digest: str, fingerprint: str, result: ContentResult) -> None:
        self._verdicts.setdefault(digest, {})[fingerprint] = asdict(result)
        self._learned_verdicts.add((digest, fingerprint))
        self._forgotten_digests.discard(digest)

    def prune(self, existing: set[Path]) -> None:
        """Forget files that are no longer there, and verdicts nothing there refers to."""
        keep = {self._key(p) for p in existing}
        for key in [k for k in self._files if k not in keep]:
            del self._files[key]
            self._learned_files.discard(key)
            self._forgotten_files.add(key)
        referenced = {str(memo.get("digest")) for memo in self._files.values()}
        for digest in [d for d in self._verdicts if d not in referenced]:
            del self._verdicts[digest]
            self._learned_verdicts = {(d, f) for d, f in self._learned_verdicts if d != digest}
            self._forgotten_digests.add(digest)

    def save(self) -> None:
        """Write what changed to the store, whole or nothing; a no-op for a
        cache nothing changed in, and for one built with no engagement."""
        if self._engagement is None:
            return
        if not (self._learned_files or self._learned_verdicts
                or self._forgotten_files or self._forgotten_digests):
            return
        learned: dict[str, dict[str, dict]] = {}
        for digest, fingerprint in self._learned_verdicts:
            learned.setdefault(digest, {})[fingerprint] = self._verdicts[digest][fingerprint]
        # One immediate transaction of the store's own - never store.record(),
        # which is for events - so a pass killed mid-save leaves the previous
        # rows rather than half of the new ones.
        store.remember_verdicts(
            store.connect(), self._engagement, version=CACHE_VERSION,
            memos={key: self._files[key] for key in self._learned_files},
            verdicts=learned,
            forget_paths=set(self._forgotten_files),
            forget_digests=set(self._forgotten_digests),
        )
        self._learned_files.clear()
        self._learned_verdicts.clear()
        self._forgotten_files.clear()
        self._forgotten_digests.clear()


# ------------------------------------------------------------------- CLI ----

if __name__ == "__main__":
    import argparse

    from tracker.manifest import load_manifest
    from tracker.page import tolerant_console

    tolerant_console()   # a client's name the console cannot encode is no traceback

    parser = argparse.ArgumentParser(
        description="What does this document say, and which rows accept it? "
                    "Read-only: no cache is written, nothing moves."
    )
    parser.add_argument("engagement_dir", help="the engagement folder")
    parser.add_argument("file", help="the document to read")
    ns = parser.parse_args()

    document = Path(ns.file)
    reading = extract(document)
    if reading.text is None:
        print(f"{document.name}: {reading.reason}")
        raise SystemExit(1)
    how = "OCR" if reading.from_ocr else "text layer"
    print(f"{document.name}: {len(reading.text)} characters from the {how}\n")
    for row in load_manifest(Path(ns.engagement_dir)):
        if not has_content_rules(row):
            continue
        verdict = evaluate_rules(reading.text, row)
        mark = "ok" if verdict.ok else f"no - {verdict.reason}"
        print(f"[{row.identifier}] {row.document}: {mark}")
