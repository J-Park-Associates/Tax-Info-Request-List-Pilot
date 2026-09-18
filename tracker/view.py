"""The workbook a person opens: derived from the record, holding no facts of its own.

Stage A step 3 of the record/view separation (decision 89). The record
(:mod:`tracker.ledger`) is what the machine decided; the two workbooks beside
it are still written exactly as before. What was missing is the third thing:
**one workbook a person opens that nobody has to be careful with.**

``VIEW_FILENAME`` is that workbook. :func:`write_view` regenerates it at the
end of every real pass from the readers - :func:`tracker.manifest.load_manifest`
for the request rows as the readers now answer them, and
:func:`tracker.filer.read_index` for the index - and sets the file's read-only
attribute afterwards. Four sheets:

- **Summary** - the rules workbook's digest, the record's head, when it was
  drawn, and the sentence that says it is regenerated every pass and edits
  here are lost.
- **Requests** - the person's rules, the keywords a filing taught, and each
  row's status.
- **Index** - every original, in the index's own column order.
- **Needs Review** - the rows parked for a person, and why.

**It holds no facts, and that is the whole point.** Nothing reads this file
back: not the pass, not the app, not the drafts. So the one failure mode
every other workbook has to survive - somebody has it open in Excel when the
pass wants to write it - costs nothing here. The replace fails, the run
carries ``view_stale`` and **succeeds**, and the view on disk stays one pass
behind until the next one lands. No sidecar, no retry, no deferred rows:
there is nothing in it that is not somewhere else.

**Current, behind, or unknown.** Because the view is a mirrored derivation, a
person has to be able to tell whether the thing in front of them is still
true. :func:`view_state` answers by comparing the Summary's two stamps - the
rules workbook's SHA-256 and the record's head as it was when the view was
drawn - with those same two things now: ``CURRENT`` when both match,
``BEHIND`` when the view is there and either has moved on, ``UNKNOWN`` when
there is no view or its stamp cannot be read. It takes no lock and opens the
file from its bytes, so it answers while Excel holds it open.

**Nothing here is a formula and nothing is a number.** Every cell is written
as text through :func:`tracker.manifest.as_text`, the helper the index
already uses, so a client file called like a formula is a name and a
reference number is not rounded into scientific notation by a spreadsheet
that thought it knew better.

**Imports downwards only.** This module reads the record, the manifest, the
index and the engagement's file names, and nothing in the package imports it
but :mod:`tracker.runner` and :mod:`tracker.api`, which sit above all of
them. Writing the view is the last thing a pass does.
"""

from __future__ import annotations

import datetime as dt
import io
import logging
import os
import stat
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from tracker import ledger
from tracker.filer import (
    INDEX_COLUMNS,
    INDEX_FILENAME,
    INDEX_LAYOUT,
    INDEX_SHEET,
    NEEDS_REVIEW,
    FilingError,
    IndexEntry,
    read_index,
)
from tracker.manifest import (
    ANY_EXTENSION,
    COL_ALLOWED_EXTENSIONS,
    COL_ANY_KEYWORDS,
    COL_DATE_PATTERN,
    COL_DOCUMENT,
    COL_EXPECTED_COUNT,
    COL_FILE_COUNT,
    COL_IDENTIFIER,
    COL_MANUAL_OVERRIDE,
    COL_MIN_SIZE_KB,
    COL_PERIOD,
    COL_RECEIVED_DATE,
    COL_REQUIRED_KEYWORDS,
    COL_STATUS,
    COL_VALIDATION_NOTES,
    COLUMN_WIDTHS,
    HEADERS,
    SHEET_NAME,
    ManifestError,
    RequestItem,
    as_text,
    load_manifest,
    save_workbook_atomically,
)
from tracker.scaffold import MANIFEST_FILENAME
from tracker.validators import sha256_of

log = logging.getLogger("tracker.view")

#: The workbook a person opens. Regenerated every pass; it holds no facts.
VIEW_FILENAME = "_status.xlsx"

#: The sheet carrying the stamp. The other three are named by the modules
#: that own what they show, so the view invents no tab name of its own.
SUMMARY_SHEET = "Summary"
#: Sheets, in the order they are written; the first is the one Excel opens on.
SHEETS = (SUMMARY_SHEET, SHEET_NAME, INDEX_SHEET, NEEDS_REVIEW)

#: What the Summary says about itself, in the first cell a person reads.
VIEW_NOTE = "This workbook is regenerated every pass; edits here are lost."

#: The Summary's labels. ``view_state`` reads the two digests back by these,
#: so they are the stamp's field names and are worded once.
LABEL_ENGAGEMENT = "Engagement folder"
LABEL_GENERATED = "Generated (UTC)"
LABEL_RULES_DIGEST = "Rules workbook SHA-256"
LABEL_RECORD_DIGEST = "Record SHA-256"
LABEL_REQUEST_ROWS = "Request rows"
LABEL_INDEX_ROWS = "Index rows"
LABEL_PARKED_ROWS = "Parked for a person"

#: What the app says about a view, and the three words it may say. The view
#: is a derivation of two things that both move, so "there is one" is not an
#: answer: a person has to know whether it still describes what is there.
CURRENT = "current"
BEHIND = "behind"
UNKNOWN = "unknown"
VIEW_STATES = (CURRENT, BEHIND, UNKNOWN)
#: What the chip beside the engagement calls it.
VIEW_LABEL = "Status workbook"

#: The index columns the review sheet repeats. The headers and widths still
#: come from the index's own table, so there is one owner for both.
NEEDS_REVIEW_FIELDS = ("received", "original_name", "pbc_location", "reason",
                       "candidates", "evidence")

_SUMMARY_LABEL_WIDTH = 24
_SUMMARY_VALUE_WIDTH = 72


class ViewError(RuntimeError):
    """The view could not be built from what the readers say."""


@dataclass(slots=True)
class ViewResult:
    """What one regeneration did. ``stale`` is the only interesting failure."""

    path: Path
    stale: bool = False              # somebody had it open; the old one stands
    requests: int = 0
    rows: int = 0
    parked: int = 0
    stamp: dict[str, str] = field(default_factory=dict)


# ------------------------------------------------------------------ cells ----


def _text(value: object) -> str:
    """One value as a view cell: a string, always.

    A date is written the way the manifest's own date hint spells one, and a
    blank is a blank rather than the word for nothing. Everything else is
    ``str()``: the view is read by a person and by nothing else, so a cell
    Excel might re-type as a number, a date or a formula buys nothing and
    can only lie.
    """
    if value is None:
        return ""
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def request_row(item: RequestItem) -> dict[str, str]:
    """One Requests row, column header to cell text.

    The one definition of what the view says about a request, so the suite's
    agreement fixture compares the sheet with the reader through the same
    function rather than a second copy of these rules.
    """
    return {
        COL_IDENTIFIER: _text(item.identifier),
        COL_DOCUMENT: _text(item.document),
        COL_PERIOD: _text(item.period),
        COL_EXPECTED_COUNT: _text(item.expected_count),
        # An item with no extensions accepts anything; the manifest says so
        # out loud rather than leaving a blank, and so does the view.
        COL_ALLOWED_EXTENSIONS: ", ".join(item.allowed_extensions) or ANY_EXTENSION,
        COL_MIN_SIZE_KB: _text(item.min_size_kb),
        COL_REQUIRED_KEYWORDS: ", ".join(item.required_keywords),
        COL_ANY_KEYWORDS: ", ".join(item.any_keywords),
        COL_DATE_PATTERN: _text(item.date_pattern),
        COL_MANUAL_OVERRIDE: _text(item.manual_override),
        COL_STATUS: _text(item.status),
        COL_RECEIVED_DATE: _text(item.received_date),
        COL_FILE_COUNT: _text(item.file_count),
        COL_VALIDATION_NOTES: _text(item.validation_notes),
    }


def index_row(entry: IndexEntry) -> list[str]:
    """One Index row as the view writes it: the reader's own row, as text."""
    return [_text(value) for value in entry.as_row()]


def _needs_review_row(entry: IndexEntry) -> list[str]:
    return [_text(getattr(entry, name)) for name in NEEDS_REVIEW_FIELDS]


# ------------------------------------------------------------------ write ----


def _sheet(ws, headers, widths, rows) -> int:
    """Head a sheet, fill it, size it. Returns how many rows were written."""
    ws.append(list(headers))
    written = 0
    for number, row in enumerate(rows, start=2):
        ws.append(list(row))
        for cell in ws[number]:
            as_text(cell)        # a file called like a formula is a name
        written += 1
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return written


def _summary_sheet(ws, stamp: dict[str, str]) -> None:
    ws.column_dimensions["A"].width = _SUMMARY_LABEL_WIDTH
    ws.column_dimensions["B"].width = _SUMMARY_VALUE_WIDTH
    as_text(ws.cell(row=1, column=1, value=VIEW_NOTE))
    for number, (label, value) in enumerate(stamp.items(), start=3):
        as_text(ws.cell(row=number, column=1, value=label))
        as_text(ws.cell(row=number, column=2, value=value))


def rules_digest(engagement_dir: Path | str) -> str:
    """The SHA-256 of the workbook the person's requests live in; "" if it
    cannot be read. Hashed with the digest every other reader of bytes in
    this package uses, so there is one answer to "are these the same bytes"."""
    try:
        return sha256_of(Path(engagement_dir) / MANIFEST_FILENAME)
    except OSError:
        return ""


def _make_writable(path: Path) -> None:
    """Clear the read-only attribute the last regeneration set.

    Windows refuses to replace a read-only file, and the view is written by
    replacing it, so the attribute this module sets would otherwise stop the
    next pass from ever landing one. Only the file this module wrote is ever
    touched, and only in the moment before it is replaced.
    """
    if path.exists():
        os.chmod(path, stat.S_IWRITE | stat.S_IREAD)


def _make_read_only(path: Path) -> None:
    """Say in the filesystem what the Summary says in words: do not type here."""
    try:
        os.chmod(path, stat.S_IREAD)
    except OSError as exc:
        # Cosmetic, and never a reason to fail a pass that has already
        # written every fact it had to write.
        log.warning("%s was written but could not be made read-only (%s)", path.name, exc)


def write_view(
    engagement_dir: Path | str,
    *,
    items: list[RequestItem] | None = None,
    entries: list[IndexEntry] | None = None,
    now: dt.datetime | None = None,
) -> ViewResult:
    """Regenerate the view for one engagement. Returns what it did.

    ``items`` and ``entries`` are the readers' answers, so a caller that has
    already read them (a pass has) does not pay for a second reading. Left
    out, they are read here exactly as the app reads them: nothing is moved,
    not even a sidecar that cannot be parsed - drawing the view changes
    nothing about what it describes.

    **Never raises because somebody had it open.** The replace goes through
    the same atomic swap the index uses, and a file Excel holds raises from
    it; that is reported as ``stale`` and the caller carries on. The view
    holds no fact that is not in the record and the workbooks, so a stale
    one costs a person one pass of freshness and nothing else.
    """
    engagement_dir = Path(engagement_dir)
    path = engagement_dir / VIEW_FILENAME
    try:
        if items is None:
            items = load_manifest(engagement_dir / MANIFEST_FILENAME)
        if entries is None:
            entries = read_index(engagement_dir / INDEX_FILENAME, quarantine=False)
    except (ManifestError, FilingError, OSError) as exc:
        # A manifest the loader refuses or an index the tracker will not
        # touch: there is nothing to draw, and the caller decides what that
        # means. A pass says it in a log line and stands.
        raise ViewError(f"{engagement_dir.name}: the view could not be built ({exc})") from exc

    parked = [entry for entry in entries if entry.decision == NEEDS_REVIEW]
    stamp = {
        LABEL_ENGAGEMENT: engagement_dir.name,
        LABEL_GENERATED: (now or dt.datetime.now(dt.UTC)).isoformat(timespec="seconds"),
        LABEL_RULES_DIGEST: rules_digest(engagement_dir),
        LABEL_RECORD_DIGEST: ledger.head(engagement_dir),
        LABEL_REQUEST_ROWS: str(len(items)),
        LABEL_INDEX_ROWS: str(len(entries)),
        LABEL_PARKED_ROWS: str(len(parked)),
    }
    result = ViewResult(path=path, requests=len(items), rows=len(entries),
                        parked=len(parked), stamp=stamp)

    wb = Workbook()
    try:
        _summary_sheet(wb.active, stamp)
        wb.active.title = SUMMARY_SHEET
        _sheet(wb.create_sheet(SHEET_NAME), HEADERS,
               [COLUMN_WIDTHS[header] for header in HEADERS],
               ([request_row(item)[header] for header in HEADERS] for item in items))
        _sheet(wb.create_sheet(INDEX_SHEET), INDEX_COLUMNS,
               [width for _, width in INDEX_LAYOUT.values()],
               (index_row(entry) for entry in entries))
        _sheet(wb.create_sheet(NEEDS_REVIEW),
               [INDEX_LAYOUT[name][0] for name in NEEDS_REVIEW_FIELDS],
               [INDEX_LAYOUT[name][1] for name in NEEDS_REVIEW_FIELDS],
               (_needs_review_row(entry) for entry in parked))
        wb.active = 0
        try:
            _make_writable(path)
            save_workbook_atomically(wb, path)
        except OSError as exc:
            # Somebody has it open. Nothing here is a fact, so this is the
            # one write in the system that may simply not happen.
            log.warning(
                "%s: %s was not regenerated (%s); it stays one pass behind and the run stands",
                engagement_dir.name, VIEW_FILENAME, exc,
            )
            result.stale = True
            return result
    finally:
        wb.close()
    _make_read_only(path)
    return result


# ------------------------------------------------------------------- read ----


def path_for(engagement_dir: Path | str) -> Path:
    """Where one engagement's view lives."""
    return Path(engagement_dir) / VIEW_FILENAME


def read_stamp(engagement_dir: Path | str) -> dict[str, str] | None:
    """The Summary's labelled values, or None if there is no readable stamp.

    Opened from the file's bytes rather than by name, so a view somebody has
    open in Excel - which shares reading and nothing else - is still read.
    Takes no lock: looking at a derivation cannot change anything.
    """
    path = path_for(engagement_dir)
    try:
        data = path.read_bytes()
    except OSError:
        return None
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:            # not a workbook any more; that is "unknown"
        return None
    try:
        if SUMMARY_SHEET not in wb.sheetnames:
            return None
        ws = wb[SUMMARY_SHEET]
        stamp = {}
        for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
            label = str(row[0] or "").strip()
            # A blank value is a value: an engagement with no record yet was
            # stamped with an empty head, and reading that as "no stamp"
            # would leave a fresh view permanently unknown.
            value = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
            if label and label != VIEW_NOTE:
                stamp[label] = value
        return stamp or None
    finally:
        wb.close()


def view_state(engagement_dir: Path | str) -> str:
    """Whether the view still describes the engagement: one of ``VIEW_STATES``.

    ``CURRENT`` when the stamp's two digests are the rules workbook's and
    the record's head *now*; ``BEHIND`` when there is a view and either has
    moved on since it was drawn; ``UNKNOWN`` when there is none, or its
    stamp cannot be read. The firm's rule for a mirrored derivation is that
    it says how old it is rather than looking authoritative, and these are
    the three honest answers.
    """
    engagement_dir = Path(engagement_dir)
    stamp = read_stamp(engagement_dir)
    if not stamp:
        return UNKNOWN
    if LABEL_RULES_DIGEST not in stamp or LABEL_RECORD_DIGEST not in stamp:
        return UNKNOWN
    matches = (
        stamp[LABEL_RULES_DIGEST] == rules_digest(engagement_dir)
        and stamp[LABEL_RECORD_DIGEST] == ledger.head(engagement_dir)
    )
    return CURRENT if matches else BEHIND


# -------------------------------------------------------------------- CLI ----

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(
        prog="python -m tracker.view",
        description="Regenerate the engagement's read-only status workbook and say "
                    "whether it is current. Writes nothing else and takes no lock.",
    )
    parser.add_argument("engagement_dir", help="the engagement folder")
    ns = parser.parse_args()

    folder = Path(ns.engagement_dir)
    written = write_view(folder)
    print(f"\n{written.path}")
    if written.stale:
        print("  not regenerated: it is open somewhere; the one on disk stands")
    print(f"  requests: {written.requests}")
    print(f"  rows:     {written.rows}")
    print(f"  parked:   {written.parked}")
    for label, value in written.stamp.items():
        print(f"  {label + ':':<26}{value}")
    print(f"  state:    {view_state(folder)}\n")
