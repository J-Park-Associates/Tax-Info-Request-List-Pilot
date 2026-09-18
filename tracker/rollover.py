"""Carry a returning client's request list into the next year (component 8).

A client who filed with us last year is not a blank form. What they actually
sent, what we waived, how many W-2s really turned up — that is better
information than any generic checklist, so for an existing client **the prior
year takes precedence and the form template does not get a vote.**

Precedence, precisely:

- Every field the prior year specifies is carried forward untouched. The
  template never overwrites a value that already exists.
- The template may only *fill blanks* — a keyword or extension the prior
  manifest simply never had. Filling an empty field overrides nothing.
- Periods, date rules and any year inside a document name are shifted by the
  same number of years, so ``TY2025`` becomes ``TY2026`` and the row asking
  for the ``TY2024`` prior-year return now asks for ``TY2025``.
- Counts learn from reality: a row that expected 2 W-2s and received 3 asks
  for 3 next year. Counts are never lowered — a client who under-delivered
  still owes what was asked.
- ``Override.WAIVED`` is a decision about the client, so it carries forward.
  ``Override.ACCEPTED`` is a judgment about specific files from one particular year,
  so it does not.

A template row the client has never had is **not** added. For a returning
client the list is last year's list; a generic checklist does not get to pad
it with nine requests they have never once needed. Those rows are reported as
*offers* instead — visible in the CLI output and on the Carried Forward sheet
— so a genuinely new requirement still surfaces for a person to accept.
``include_new=True`` adds them outright.

Nothing here writes to the prior year's engagement — it is read-only history.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from tracker.manifest import (  # shift_years/detect_year re-exported: they live in manifest
    COL_DOCUMENT,
    COL_IDENTIFIER,
    EngagementInfo,
    ManifestError,
    Override,
    RequestItem,
    Status,
    as_text,
    check_tax_year,
    create_template,
    detect_year,
    load_manifest,
    pending_updates,
    save_workbook_atomically,
    shift_item,
    shift_years,
    with_pending,
)
from tracker.scaffold import MANIFEST_FILENAME

#: Why a row is on the new list. Written to the Carried Forward sheet.
ORIGIN_PRIOR = "carried from last year"
ORIGIN_WAIVED = "waived last year"
ORIGIN_NEW = "new this year"

#: What stands in for the target year when the prior list gave none away.
UNKNOWN_YEAR_LABEL = "next year"
CARRIED_SHEET = "Carried Forward"
#: Heads the sheet's (and the CLI's) list of last year's unmatched documents.
UNFILED_HEADING = "Sent last year but never filed — check these are covered:"
#: The Carried Forward sheet, described once: header -> column width.
CARRIED_LAYOUT: dict[str, int] = {
    COL_IDENTIFIER: 12, COL_DOCUMENT: 40, "Origin": 22, "Why it is on the list": 52,
    "Last Year Status": 18, "Last Year Files": 14,
}
CARRIED_HEADERS = tuple(CARRIED_LAYOUT)

@dataclass(frozen=True, slots=True)
class RolledItem:
    """One row of next year's list, and where it came from."""

    item: RequestItem
    origin: str
    note: str = ""
    prior_status: str = ""
    prior_file_count: int | None = None


@dataclass(slots=True)
class RolloverReport:
    """Next year's request list, built from last year's engagement."""

    prior_dir: Path
    prior_year: int | None = None
    target_year: int | None = None
    rolled: list[RolledItem] = field(default_factory=list)
    offered: list[RolledItem] = field(default_factory=list)   # template rows not added
    unfiled_last_year: list[str] = field(default_factory=list)

    @property
    def items(self) -> list[RequestItem]:
        return [r.item for r in self.rolled]

    @property
    def carried(self) -> list[RolledItem]:
        return [r for r in self.rolled if r.origin != ORIGIN_NEW]

    @property
    def added(self) -> list[RolledItem]:
        return [r for r in self.rolled if r.origin == ORIGIN_NEW]

    @property
    def waived(self) -> list[RolledItem]:
        return [r for r in self.rolled if r.origin == ORIGIN_WAIVED]


# ------------------------------------------------------------- carry rules ----


def next_tax_year(prior_year: int) -> int:
    """The year a rolled engagement is for: the one after the prior's."""
    return prior_year + 1


def carry_engagement_info(prior: EngagementInfo, *, rolled_from: str) -> EngagementInfo:
    """Last year's Engagement sheet as this year's starting point.

    The client, the sender, the firm and the reminders decision are about
    the client and carry forward. The share link and the due date are this
    year's to set. The name is never written (the folder is the name), and
    Rolled From is what retires the prior (tracker.registry.mark_superseded).
    """
    return replace(prior, name="", link="", due=None, active=True, rolled_from=rolled_from)



def _carry(
    prior: RequestItem, template: RequestItem | None, delta: int, tmpl_delta: int
) -> tuple[RequestItem, str, str]:
    """Build next year's row from last year's, filling only true blanks."""
    # Counts learn from what actually arrived, and never shrink.
    expected = prior.expected_count
    if prior.status == Status.RECEIVED and (prior.file_count or 0) > expected:
        expected = prior.file_count

    # The template fills a blank; it never replaces a value.
    def fill(mine, theirs):
        return mine if mine else (theirs or ())

    # A year check derived from Period is not carried as text: the shifted
    # Period derives it again on load, so the new manifest stays as sparse
    # as the old one was.
    date_pattern = "" if prior.date_pattern_derived else shift_years(prior.date_pattern, delta)
    if not date_pattern and template and not template.date_pattern_derived:
        date_pattern = shift_years(template.date_pattern, tmpl_delta)

    item = RequestItem(
        identifier=prior.identifier,
        document=shift_years(prior.document, delta),
        period=shift_years(prior.period, delta) or (
            shift_years(template.period, tmpl_delta) if template else ""
        ),
        expected_count=expected,
        allowed_extensions=fill(prior.allowed_extensions,
                                template.allowed_extensions if template else ()),
        min_size_kb=prior.min_size_kb,
        required_keywords=fill(prior.required_keywords,
                               template.required_keywords if template else ()),
        any_keywords=fill(prior.any_keywords,
                          template.any_keywords if template else ()),
        date_pattern=date_pattern,
        # "Waived" is about the client and persists. "Accepted" was a call on
        # last year's particular files and must not pre-approve this year's.
        manual_override=(
            Override.WAIVED if prior.manual_override == Override.WAIVED else ""
        ),
    )

    if prior.manual_override == Override.WAIVED:
        return item, ORIGIN_WAIVED, f"{ORIGIN_WAIVED}; clear the override to request it again"
    if prior.status == Status.RECEIVED:
        note = f"received last year ({prior.file_count or 0} file(s))"
        if expected > prior.expected_count:
            note += f"; asking for {expected} this year to match"
        return item, ORIGIN_PRIOR, note
    if prior.status:
        return (
            item,
            ORIGIN_PRIOR,
            f"last year: {prior.status} — confirm it still applies",
        )
    return item, ORIGIN_PRIOR, "on last year's list"


# ---------------------------------------------------------------- rollover ----


def roll_forward(
    prior_engagement_dir: Path | str,
    *,
    target_year: int | None = None,
    template: Sequence[RequestItem] = (),
    include_new: bool = False,
) -> RolloverReport:
    """Build next year's request list from ``prior_engagement_dir``.

    ``template`` is the form's standard checklist. It is consulted only to
    fill blanks on carried rows; it never overrides prior-year data. Rows the
    client has never had are collected in ``report.offered`` rather than
    added, unless ``include_new`` is set.
    """
    prior_dir = Path(prior_engagement_dir)
    # A year is bounded wherever it is typed (decision 68 bounded the
    # wizard's box; the flag on this command line was the other door).
    # Unbounded, --year 20265 shifts every Period and every document name
    # by eighteen thousand years, writes the folders under those names and
    # retires the live engagement behind them.
    if target_year is not None:
        check_tax_year(target_year)
    # The prior year's last scan may have been deferred (Excel held the
    # manifest) and never landed - the usual end of an engagement. What
    # it saw is in the sidecar; read it, move nothing.
    prior_manifest = prior_dir / MANIFEST_FILENAME
    prior_items = with_pending(
        load_manifest(prior_manifest), pending_updates(prior_manifest, quarantine=False)
    )

    prior_year = detect_year(prior_items)
    if target_year is None and prior_year is not None:
        target_year = next_tax_year(prior_year)
    if prior_year and target_year and target_year <= prior_year:
        # A roll into the same year would retire the live engagement in
        # favour of a copy of itself; into an earlier one would shift every
        # period backwards. Neither is a rollover.
        raise ManifestError(
            f"Roll forward to a year after {prior_year}; {target_year} is not later"
        )
    delta = (target_year - prior_year) if (prior_year and target_year) else 0

    template_year = detect_year(template) if template else None
    tmpl_delta = (
        target_year - template_year if (template_year and target_year) else 0
    )

    report = RolloverReport(
        prior_dir=prior_dir, prior_year=prior_year, target_year=target_year
    )
    by_id = {t.identifier: t for t in template}

    for prior in prior_items:
        item, origin, note = _carry(prior, by_id.get(prior.identifier), delta, tmpl_delta)
        report.rolled.append(
            RolledItem(
                item=item, origin=origin, note=note,
                prior_status=prior.status, prior_file_count=prior.file_count,
            )
        )

    # By the identifier as load_manifest() compares it - without case - or
    # a template row differing from a prior's only in case would be written
    # beside it and the new manifest refused as a duplicate.
    seen = {p.identifier.upper() for p in prior_items}
    for spec in template:
        if spec.identifier.upper() in seen:
            continue
        offer = RolledItem(
            item=shift_item(spec, tmpl_delta),
            origin=ORIGIN_NEW,
            note="not on last year's list — confirm it applies",
        )
        (report.rolled if include_new else report.offered).append(offer)

    report.unfiled_last_year = _unfiled_last_year(prior_dir)
    return report


def _unfiled_last_year(prior_dir: Path) -> list[str]:
    """Documents the client sent last year that were never filed.

    Reported, never acted on: a request row needs a name and rules a person
    chooses. But a document that arrived and fitted nowhere is exactly the
    gap next year's list should close.
    """
    from tracker.filer import INDEX_FILENAME, NEEDS_REVIEW, read_index

    try:
        rows = read_index(prior_dir / INDEX_FILENAME, quarantine=False)  # a rollover only reads the prior year
    except Exception:  # an unreadable index must never block a rollover
        return []
    seen: dict[str, str] = {}
    for row in rows:
        if row.decision == NEEDS_REVIEW and row.original_name not in seen:
            seen[row.original_name] = row.reason
    return [f"{name} — {reason}" for name, reason in seen.items()]


# ------------------------------------------------------------------- write ----


def write_rollover_manifest(path: Path | str, report: RolloverReport) -> Path:
    """Write next year's manifest, plus a sheet saying where each row came from."""
    path = Path(path)
    create_template(path, report.items)

    wb = load_workbook(path)
    try:
        ws = wb.create_sheet(CARRIED_SHEET)
        ws.append(list(CARRIED_HEADERS))
        for rolled in report.rolled:
            ws.append([
                rolled.item.identifier,
                rolled.item.document,
                rolled.origin,
                rolled.note,
                rolled.prior_status or "",
                rolled.prior_file_count if rolled.prior_file_count is not None else "",
            ])
        if report.offered:
            ws.append([])
            ws.append(["Offered by the standard checklist, NOT added:"])
            for offer in report.offered:
                ws.append([offer.item.identifier, offer.item.document,
                           ORIGIN_NEW, offer.note])
        if report.unfiled_last_year:
            ws.append([])
            ws.append([UNFILED_HEADING])
            for line in report.unfiled_last_year:
                ws.append(["", line])
        for cells in ws.iter_rows():
            for cell in cells:
                as_text(cell)         # a client's file name in the unfiled list is a name
        for index, width in enumerate(CARRIED_LAYOUT.values(), start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"
        save_workbook_atomically(wb, path)
    finally:
        wb.close()
    return path


# --------------------------------------------------------------------- CLI ----

if __name__ == "__main__":
    import argparse

    from tracker.templates import require_form, template_items

    parser = argparse.ArgumentParser(
        description="Build next year's request list from a returning client's prior engagement"
    )
    parser.add_argument("prior_engagement_dir", help="last year's folder (read only)")
    parser.add_argument("new_engagement_dir", help="folder to create for the new year")
    parser.add_argument("--year", type=int, default=None, help="target tax year")
    parser.add_argument("--form", default="", help="form template used to fill blanks (e.g. 1040)")
    parser.add_argument(
        "--include-new", action="store_true",
        help="also add checklist rows this client has never had (default: just offer them)",
    )
    parser.add_argument("--scaffold", action="store_true", help="also build the folders")
    ns = parser.parse_args()

    template = []
    if ns.form:
        try:
            require_form(ns.form)
        except ManifestError as exc:
            parser.error(str(exc))
        template = template_items(ns.form)

    result = roll_forward(
        ns.prior_engagement_dir,
        target_year=ns.year,
        template=template,
        include_new=ns.include_new,
    )

    target = Path(ns.new_engagement_dir)
    target.mkdir(parents=True, exist_ok=True)
    manifest = write_rollover_manifest(target / MANIFEST_FILENAME, result)
    from tracker.manifest import load_engagement_info, write_engagement_info

    write_engagement_info(manifest, carry_engagement_info(
        load_engagement_info(result.prior_dir / MANIFEST_FILENAME),
        rolled_from=str(result.prior_dir.resolve()),   # the runner's cwd is not this one
    ))

    span = f"{result.prior_year} → {result.target_year}" if result.prior_year else UNKNOWN_YEAR_LABEL
    print(f"Rolled {result.prior_dir.name} forward ({span})\n")
    for rolled in result.carried:
        flag = "WAIVED " if rolled.origin == ORIGIN_WAIVED else "CARRIED"
        print(f"  {flag} {rolled.item.label}")
        print(f"          {rolled.note}")
    for rolled in result.added:
        print(f"  NEW     {rolled.item.label}")
        print(f"          {rolled.note}")
    if result.offered:
        print(f"\n  Not added — the standard {ns.form or 'checklist'} also has "
              f"{len(result.offered)} request(s) this client has never had.")
        print("  Add any that now apply with --include-new, or by hand:")
        for offer in result.offered:
            print(f"    + {offer.item.label}")
    if result.unfiled_last_year:
        print(f"\n  {UNFILED_HEADING}")
        for line in result.unfiled_last_year:
            print(f"    ? {line}")
    print(f"\n  Manifest: {manifest}")

    if ns.scaffold:
        from tracker.scaffold import scaffold_engagement

        scaffolded = scaffold_engagement(target)
        for line in scaffolded.describe():
            print(f"  {line}")
