"""JSON bridge for the demo UI (Electron) — python -m tracker.api <command>.

Commands print a single JSON object to stdout and exit 0, or {"error": ...}
and exit 1. All real logic lives in the tracker package; this module only
serializes it, so the UI can never disagree with the scanner.

Commands:
  state     current manifest rows + unfiled sheet + useful paths
  scaffold  build/refresh the Shared/ drop folder and Prepared/ tree
  sort      file the client's drops into PBC/ and Prepared/
  priors    list engagements a new year could be rolled forward from
  rollover  build next year's list from a returning client's prior year
  scan      run a full scan and write the manifest back
  assign    file one Needs Review document under a request (a person's call)
  check     validate the manifest now, with row numbers, instead of at the next scan
  unlock    clear a stale engagement lock (a fresh one is refused)
  reset     rebuild the entire marketing demo from scratch:
            engagement folder, manifest, folder tree, sample client docs
"""

from __future__ import annotations

import datetime as dt
import json
import os
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tracker.locking import clear_stale_lock, lock_status
from tracker.manifest import (
    EngagementInfo,
    ManifestError,
    UNFILED_SHEET_NAME,
    check_manifest,
    create_template,
    load_engagement_info,
    load_manifest,
    pending_updates,
    with_pending,
    write_engagement_info,
)
from tracker.registry import engagement_dirs
from tracker.filer import (
    INDEX_FILENAME,
    FilingError,
    assign_review_file,
    file_drops,
    read_index,
)
from tracker.rollover import detect_year, roll_forward, write_rollover_manifest
from tracker.scaffold import (
    MANIFEST_FILENAME,
    PBC_DIR_NAME,
    PREPARED_DIR_NAME,
    SHARED_DIR_NAME,
    sanitize_component,
    scaffold_engagement,
)
from tracker.scanner import ScanLockedError, scan_engagement
from tracker.templates import (  # the catalog; re-exported for the wizard and demo
    DEMO_FORM,
    FORM_TEMPLATES,
    FORM_TYPES,
    base_year,
    default_tax_year,
    item_from_spec,
    shift_item,
    template_items,
)

# Frozen (PyInstaller) builds live inside the portable package; the demo data
# root is then supplied by the Electron shell via TRACKER_DEMO_ROOT so it sits
# next to the packaged exe rather than inside the bundle.
if getattr(sys, "frozen", False):
    REPO_ROOT = Path(sys.executable).resolve().parent
else:
    REPO_ROOT = Path(__file__).resolve().parent.parent
DEMO_ROOT = Path(os.environ.get("TRACKER_DEMO_ROOT", REPO_ROOT / "demo-marketing"))
ENGAGEMENT_DIRNAME = "Engagement - Smith Family 2025 Form 1040"
SAMPLES_DIRNAME = "Sample Client Documents"
CONTACT = "J Park & Associates, CPA"

DEMO_ITEMS = template_items(DEMO_FORM, core_only=True)


# --------------------------------------------------------- sample documents ----


def _text_pdf(path: Path, lines: list[str]) -> Path:
    """Minimal but valid PDF with a real text layer (no extra deps)."""
    body = "\n".join(
        f"BT /F1 11 Tf 60 {740 - 14 * i} Td ({line}) Tj ET"
        for i, line in enumerate(lines[:48])
    )
    content = body.encode("ascii", "replace")
    padding = b" " * 8192  # unreferenced object: realistic file size, renders clean
    bodies = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        4: b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        5: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        6: b"<< /Length %d >>\nstream\n%s\nendstream" % (len(padding), padding),
    }
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for n in sorted(bodies):
        offsets[n] = len(out)
        out += b"%d 0 obj\n" % n + bodies[n] + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(bodies) + 1)
    for n in sorted(bodies):
        out += b"%010d 00000 n \n" % offsets[n]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(bodies) + 1,
        xref_at,
    )
    path.write_bytes(bytes(out))
    return path


def _w2_lines(employee: str, employer: str, year: int) -> list[str]:
    return [
        f"Form W-2 Wage and Tax Statement - Tax Year {year}",
        f"Employer: {employer}",
        f"Employee: {employee}",
        "",
        "Box 1  Wages, tips, other compensation       $84,500.00",
        "Box 2  Federal income tax withheld           $11,240.00",
        "Box 3  Social security wages                 $84,500.00",
        "Box 4  Social security tax withheld           $5,239.00",
        "Box 5  Medicare wages and tips               $84,500.00",
        "Box 6  Medicare tax withheld                  $1,225.25",
        "",
        f"Copy B - To Be Filed With Employee's Federal Tax Return, {year}",
    ]


def _1099_int_lines(payer: str, recipient: str, year: int) -> list[str]:
    return [
        f"Form 1099-INT Interest Income - {year}",
        f"Payer: {payer}",
        f"Recipient: {recipient}",
        "",
        "Box 1  Interest income                        $1,842.17",
        "Box 4  Federal income tax withheld                $0.00",
        "",
        "This is important tax information and is being furnished to the IRS.",
    ]


def _prior_return_lines(taxpayer: str, year: int) -> list[str]:
    return [
        f"Form 1040 - U.S. Individual Income Tax Return - Tax Year {year}",
        f"Taxpayer: {taxpayer}",
        "Filing status: Married filing jointly",
        "",
        "Line 1   Wages, salaries, tips                $161,300.00",
        "Line 11  Adjusted gross income                $168,455.00",
        "Line 24  Total tax                             $24,918.00",
        "Line 33  Total payments                        $26,102.00",
        "Line 34  Overpayment refunded                   $1,184.00",
    ]


def _form_1098_lines(lender: str, borrower: str, year: int) -> list[str]:
    return [
        f"Form 1098 Mortgage Interest Statement - {year}",
        f"Recipient/Lender: {lender}",
        f"Payer/Borrower: {borrower}",
        "",
        "Box 1  Mortgage interest received            $12,411.08",
        "Box 2  Outstanding mortgage principal       $342,900.00",
        "Box 5  Mortgage insurance premiums                $0.00",
        "Box 10 Real property taxes paid               $6,240.00",
    ]


def _donations_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Donations"
    ws.append(["Smith Family - Charitable Contributions 2025"])
    ws.append(["Date", "Organization", "Amount", "Receipt on file"])
    for i in range(1, 301):  # enough rows to clear the size minimum
        ws.append([f"0{(i % 9) + 1}/12/2025", f"Community Charity {i:03d}", 25 + i, "Yes"])
    wb.save(path)


def _build_samples(samples: Path) -> None:
    samples.mkdir(parents=True, exist_ok=True)

    good = _text_pdf(
        samples / "W-2 John Smith 2025.pdf",
        _w2_lines("John A. Smith", "Acme Manufacturing Inc.", 2025),
    )
    # Byte-identical duplicate — demonstrates content-hash de-duplication.
    shutil.copyfile(good, samples / "W-2 John Smith 2025 - Copy.pdf")

    _text_pdf(
        samples / "W-2 Jane Smith 2025.pdf",
        _w2_lines("Jane R. Smith", "Lakeside Medical Group", 2025),
    )
    # Wrong tax year — the content date check will flag it.
    _text_pdf(
        samples / "W-2 Jane Smith 2024 - old.pdf",
        _w2_lines("Jane R. Smith", "Lakeside Medical Group", 2024),
    )

    _text_pdf(
        samples / "1099-INT First National.pdf",
        _1099_int_lines("First National Bank", "John A. Smith", 2025),
    )
    (samples / "1099-DIV Vanguard 2025.csv").write_text(
        "Form 1099-DIV dividend summary - Vanguard Brokerage 2025\n"
        + "date,fund,ordinary dividends,qualified dividends\n" * 300,
        encoding="utf-8",
    )

    _text_pdf(
        samples / "2024 Form 1040 Tax Return.pdf",
        _prior_return_lines("John A. & Jane R. Smith", 2024),
    )

    _text_pdf(
        samples / "Form 1098 Mortgage Interest.pdf",
        _form_1098_lines("Home Lending Corp.", "John A. & Jane R. Smith", 2025),
    )
    (samples / "Mortgage Notes.docx").write_bytes(b"not a real docx " * 800)

    _donations_xlsx(samples / "Donation Receipts 2025.xlsx")

    # Google Drive realities. A client who keeps records in Google Sheets
    # shares a .gsheet shortcut, which is a link — not the spreadsheet; the
    # scanner rejects it with export instructions rather than a size error.
    (samples / "Donation Receipts 2025.gsheet").write_text(
        '{"url": "https://docs.google.com/spreadsheets/d/1aB2cD3eF4gH5iJ6kL7mN8oP/edit",'
        ' "doc_id": "1aB2cD3eF4gH5iJ6kL7mN8oP", "email": "client@example.com"}',
        encoding="utf-8",
    )
    # A Google Drive upload caught mid-flight: ignored, never counted as a
    # delivered document, and it disappears on its own once sync finishes.
    (samples / "W-2 Jane Smith 2025.pdf.tmp.driveupload").write_bytes(b"\x00" * 4096)

    (samples / "vacation photo.jpg").write_bytes(b"\xff\xd8\xff\xe0" + b"J" * 9000)


# ----------------------------------------------------------------- commands ----


def _engagement_dir(argv: list[str]) -> Path:
    if "--engagement" in argv:
        return Path(argv[argv.index("--engagement") + 1])
    return DEMO_ROOT / ENGAGEMENT_DIRNAME


def _read_unfiled(manifest_path: Path) -> list[dict]:
    if not manifest_path.exists():
        return []
    wb = load_workbook(manifest_path, data_only=True)
    try:
        if UNFILED_SHEET_NAME not in wb.sheetnames:
            return []
        return [
            {"name": r[0], "kind": r[1], "size_kb": r[2], "seen": str(r[3] or "")}
            for r in wb[UNFILED_SHEET_NAME].iter_rows(min_row=2, values_only=True)
            if r and r[0]
        ]
    finally:
        wb.close()


def _lock_payload(engagement: Path) -> dict | None:
    status = lock_status(engagement)
    if status is None:
        return None
    return {
        "started": status.started,
        "age_minutes": int(status.age_seconds // 60),
        "stale": status.stale,
    }


def _info_payload(info: EngagementInfo) -> dict:
    return {
        "client": info.client, "name": info.name, "link": info.link,
        "due": info.due.isoformat() if info.due else "", "sender": info.sender,
        "firm": info.firm, "reminders": info.reminders, "active": info.active,
        "rolled_from": info.rolled_from,
    }


def _info_from_spec(
    spec: dict, *, name: str, carry: EngagementInfo | None = None, rolled_from: str = ""
) -> EngagementInfo:
    """The Engagement sheet for a new engagement: what the wizard sent, over
    what last year's sheet said (rollover), over the firm default."""
    base = carry or EngagementInfo()

    def text(key: str, fallback: str) -> str:
        value = spec.get(key)
        return str(value).strip() if value not in (None, "") else fallback

    due_raw = str(spec.get("due", "") or "").strip()
    if due_raw:
        try:
            due = dt.date.fromisoformat(due_raw)
        except ValueError:
            raise ManifestError(f"Due date must be YYYY-MM-DD, got {due_raw!r}") from None
    else:
        due = None   # a new year never inherits last year's deadline
    return EngagementInfo(
        client=text("client", base.client),
        name=name,
        link=text("link", ""),          # a new folder has a new share link
        due=due,
        sender=text("sender", base.sender),
        firm=text("firm", base.firm or CONTACT),
        reminders=bool(spec.get("reminders", base.reminders)),
        active=True,
        rolled_from=rolled_from,
    )


def _engagement_name(requested: str, fallback: str) -> str:
    """A folder name from what the user typed, or the fallback if nothing usable is left."""
    name = sanitize_component(requested.strip())
    return name if any(ch.isalnum() for ch in name) else fallback


def _state(engagement: Path) -> dict:
    manifest_path = engagement / MANIFEST_FILENAME
    deferred = pending_updates(manifest_path)
    items = with_pending(load_manifest(manifest_path), deferred)
    info = load_engagement_info(manifest_path)
    return {
        "pending_statuses": len(deferred),
        "engagement": _info_payload(info),
        "lock": _lock_payload(engagement),
        "items": [
            {
                "identifier": i.identifier,
                "document": i.document,
                "period": i.period,
                "expected_count": i.expected_count,
                "allowed_extensions": list(i.allowed_extensions),
                "required_keywords": list(i.required_keywords),
                "any_keywords": list(i.any_keywords),
                "manual_override": i.manual_override,
                "status": i.status,
                "received_date": i.received_date.isoformat() if i.received_date else None,
                "file_count": i.file_count,
                "validation_notes": i.validation_notes,
            }
            for i in items
        ],
        "unfiled": _read_unfiled(manifest_path),
        "index": [
            {
                "received": e.received,
                "original_name": e.original_name,
                "identifier": e.identifier,
                "filed_as": e.filed_as,
                "prepared_location": e.prepared_location,
                "pbc_location": e.pbc_location,
                "decision": e.decision,
                "reason": e.reason,
            }
            for e in read_index(engagement / INDEX_FILENAME)
        ],
        "paths": {
            "engagement": str(engagement),
            "shared": str(engagement / SHARED_DIR_NAME),
            "pbc": str(engagement / SHARED_DIR_NAME / PBC_DIR_NAME),
            "prepared": str(engagement / PREPARED_DIR_NAME),
            "index": str(engagement / INDEX_FILENAME),
            "manifest": str(manifest_path),
            "samples": str(DEMO_ROOT / SAMPLES_DIRNAME),
        },
    }


def _cmd_state(argv: list[str]) -> dict:
    return _state(_engagement_dir(argv))


def _cmd_scaffold(argv: list[str]) -> dict:
    engagement = _engagement_dir(argv)
    result = scaffold_engagement(engagement, contact=CONTACT)
    return {
        "created": [p.name for p in result.created],
        "existing": result.existing,
        "waived": result.waived,
        "state": _state(engagement),
    }


def _cmd_sort(argv: list[str]) -> dict:
    """Sort the drop folder without validating — the filer on its own."""
    engagement = _engagement_dir(argv)
    filed = file_drops(engagement)
    return {"sorted": _sorted_payload(filed), "state": _state(engagement)}


def _sorted_payload(filed) -> dict:
    return {
        "filed": [
            {
                "original_name": e.original_name,
                "identifier": e.identifier,
                "filed_as": e.filed_as,
                "prepared_location": e.prepared_location,
                "reason": e.reason,
            }
            for e in filed.filed
        ],
        "review": [
            {"original_name": e.original_name, "reason": e.reason}
            for e in filed.review
        ],
        "duplicates": [
            {"original_name": e.original_name, "reason": e.reason}
            for e in filed.duplicates
        ],
        "waiting": [p.name for p in filed.waiting],
        "errors": [
            {"name": e.name, "error": e.error, "left_in_place": e.left_in_place}
            for e in filed.errors
        ],
        "index_deferred": filed.index_deferred,
    }


def _cmd_scan(argv: list[str]) -> dict:
    """What the scheduled job does: sort the drop folder, then validate."""
    engagement = _engagement_dir(argv)
    filed = file_drops(engagement)
    report = scan_engagement(engagement)
    return {
        "sorted": _sorted_payload(filed),
        "written": report.written,
        "deferred": report.deferred,
        "updates": {
            ident: {
                "status": u.status,
                "file_count": u.file_count,
                "received_date": u.received_date.isoformat() if u.received_date else None,
                "validation_notes": u.validation_notes,
            }
            for ident, u in report.updates.items()
        },
        "unfiled": [
            {"name": e.name, "kind": e.kind, "size_kb": e.size_kb, "seen": e.seen}
            for e in report.unfiled
        ],
        "state": _state(engagement),
    }


def _cmd_reset(argv: list[str]) -> dict:
    if DEMO_ROOT.exists():
        shutil.rmtree(DEMO_ROOT)
    engagement = DEMO_ROOT / ENGAGEMENT_DIRNAME
    engagement.mkdir(parents=True)
    create_template(engagement / MANIFEST_FILENAME, DEMO_ITEMS)
    scaffold_engagement(engagement, contact=CONTACT)
    _build_samples(DEMO_ROOT / SAMPLES_DIRNAME)
    return {"reset": True, "state": _state(engagement)}


def _cmd_templates(argv: list[str]) -> dict:
    """The catalog, plus the tax year a new engagement is for by default."""
    year = default_tax_year()
    return {
        "forms": FORM_TYPES,
        "templates": FORM_TEMPLATES,
        "years": {form: year for form in FORM_TEMPLATES},
        "default_year": year,
    }


def _cmd_check(argv: list[str]) -> dict:
    """Say now what the next scan would fail on, and what it would let slide."""
    engagement = _engagement_dir(argv)
    result = check_manifest(engagement / MANIFEST_FILENAME)
    return {"ok": result.ok, "problems": result.problems, "warnings": result.warnings}


def _cmd_unlock(argv: list[str]) -> dict:
    """Clear a stale engagement lock. A fresh one is refused with how long to wait."""
    engagement = _engagement_dir(argv)
    status = clear_stale_lock(engagement)
    return {"cleared": True, "age_minutes": int(status.age_seconds // 60),
            "state": _state(engagement)}


def _cmd_list(argv: list[str]) -> dict:
    """Every engagement under the root - the same discovery the scheduled run uses."""
    engagements = []
    if DEMO_ROOT.is_dir():
        for child in engagement_dirs(DEMO_ROOT):
            engagements.append({"name": child.name, "path": str(child)})
    return {"engagements": engagements}


def _cmd_create(argv: list[str]) -> dict:
    """Create a new engagement from a JSON spec on stdin:
    {"name": "...", "form": "1040", "year": 2026, "client": "...", "link": "...",
     "due": "YYYY-MM-DD", "items": [{identifier, document, extensions, ...}, ...]}
    year defaults to the most recently ended year; catalog rows are shifted to it.
    client/link/due land on the manifest's Engagement sheet, which is all the
    scheduled run needs - there is no registry to add the engagement to.
    """
    spec = json.loads(sys.stdin.read() or "{}")
    form = str(spec.get("form", "")).strip()
    if form and form not in FORM_TEMPLATES:
        raise ManifestError(f"Unknown tax form type '{form}'")
    client = str(spec.get("client", "") or "").strip()
    # The engagement's tax year: the calendar's default unless chosen.
    try:
        year = int(spec.get("year") or default_tax_year())
    except (TypeError, ValueError):
        raise ManifestError(f"Tax year must be a whole number, got {spec.get('year')!r}") from None
    fallback = " ".join(
        part for part in (
            client or "New",
            f"TY{year}",
            f"Form {form}" if form else "Engagement",
        ) if part
    )
    name = _engagement_name(str(spec.get("name", "")), fallback)
    engagement = DEMO_ROOT / name
    if engagement.exists():
        raise ManifestError(f"An engagement named '{name}' already exists")

    items = [item_from_spec(s) for s in spec.get("items", [])]
    if not items:
        raise ManifestError("Select at least one request item")
    # The wizard sends catalog rows as written (the base year); shift them
    # to the engagement's year so TY2025 does not get asked for in 2027.
    base = base_year(form) if form else None
    if base:
        items = [shift_item(item, year - base) for item in items]

    info = _info_from_spec(spec, name=name)
    engagement.mkdir(parents=True)
    try:
        create_template(engagement / MANIFEST_FILENAME, items, info)
        scaffold_engagement(engagement, contact=CONTACT)  # validates the manifest too
    except Exception:
        shutil.rmtree(engagement, ignore_errors=True)  # never leave a half-built one
        raise
    return {"created": name, "state": _state(engagement)}


def _cmd_priors(argv: list[str]) -> dict:
    """Engagements already on disk that a new year could be rolled from."""
    priors = []
    if DEMO_ROOT.is_dir():
        for child in engagement_dirs(DEMO_ROOT):
            manifest = child / MANIFEST_FILENAME
            try:
                items = load_manifest(manifest)
                info = load_engagement_info(manifest)
            except ManifestError:
                continue
            priors.append({
                "name": child.name,
                "path": str(child),
                "client": info.client,
                "rolled_from": info.rolled_from,
                "year": detect_year(items),
                "requests": len(items),
                "received": sum(1 for i in items if i.status == "Received"),
            })
    # Say which priors have already been rolled forward, so the wizard can
    # steer the person to the newest year instead of rolling the same
    # prior twice.
    names_by_path = {Path(p["path"]).resolve(): p["name"] for p in priors}
    for prior in priors:
        successor = ""
        if prior["rolled_from"]:
            pass
        for other in priors:
            if other["rolled_from"] and Path(other["rolled_from"]).resolve() == Path(prior["path"]).resolve():
                successor = other["name"]
        prior["superseded_by"] = successor
    return {"priors": priors}


def _cmd_rollover(argv: list[str]) -> dict:
    """Build next year's engagement from a returning client's prior one.

    JSON spec on stdin: {"prior": "<path or name>", "name": "...",
                         "form": "1040", "year": 2026, "include_new": false}
    Prior-year data wins on every field it specifies; the form template only
    fills blanks. Rows the client has never had are offered, not added.
    """
    spec = json.loads(sys.stdin.read() or "{}")
    prior_raw = str(spec.get("prior", "")).strip()
    if not prior_raw:
        raise ManifestError("Pick the engagement to roll forward")
    prior = Path(prior_raw)
    if not prior.is_absolute():
        prior = DEMO_ROOT / prior_raw
    if not (prior / MANIFEST_FILENAME).is_file():
        raise ManifestError(f"No manifest found in '{prior_raw}'")

    form = str(spec.get("form", "")).strip()
    template = template_items(form) if form else []

    report = roll_forward(
        prior,
        target_year=spec.get("year") or None,
        template=template,
        include_new=bool(spec.get("include_new")),
    )

    default_name = f"{prior.name} - {report.target_year}" if report.target_year else f"{prior.name} - next year"
    name = _engagement_name(str(spec.get("name", "")), default_name)
    engagement = DEMO_ROOT / name
    if engagement.exists():
        raise ManifestError(f"An engagement named '{name}' already exists")

    carried = load_engagement_info(prior / MANIFEST_FILENAME)
    # Naming the prior on the new sheet is what retires it (see
    # tracker.registry.mark_superseded); nothing is written to last year.
    info = _info_from_spec(spec, name=name, carry=carried, rolled_from=str(prior))
    engagement.mkdir(parents=True)
    try:
        write_rollover_manifest(engagement / MANIFEST_FILENAME, report)
        write_engagement_info(engagement / MANIFEST_FILENAME, info)
        scaffold_engagement(engagement, contact=CONTACT)
    except Exception:
        shutil.rmtree(engagement, ignore_errors=True)
        raise

    return {
        "created": name,
        "rollover": {
            "prior": prior.name,
            "prior_year": report.prior_year,
            "target_year": report.target_year,
            "carried": [
                {"identifier": r.item.identifier, "document": r.item.document,
                 "origin": r.origin, "note": r.note}
                for r in report.rolled
            ],
            "offered": [
                {"identifier": r.item.identifier, "document": r.item.document,
                 "note": r.note}
                for r in report.offered
            ],
            "unfiled_last_year": report.unfiled_last_year,
        },
        "state": _state(engagement),
    }


def _cmd_assign(argv: list[str]) -> dict:
    """File one parked document under a request, by a person's decision.

    JSON spec on stdin: {"original": "<PBC location or original name>",
                         "identifier": "A01", "keyword": "optional"}
    The working copy is filed under the canonical name, the index row is
    rewritten as Filed (attributed to a person), the keyword - if given - is
    added to the request so the next such file routes itself, and the
    engagement is re-scanned so the status reflects it straight away.
    """
    engagement = _engagement_dir(argv)
    spec = json.loads(sys.stdin.read() or "{}")
    original = str(spec.get("original", "")).strip()
    identifier = str(spec.get("identifier", "")).strip()
    if not original or not identifier:
        raise ManifestError("Pick the file and the request it belongs to")
    result = assign_review_file(
        engagement, original, identifier, keyword=str(spec.get("keyword", "") or "")
    )
    scan_note = ""
    try:
        scan = scan_engagement(engagement)
        if not scan.written:
            scan_note = "manifest open in Excel; status update saved to the sidecar"
    except ScanLockedError as exc:
        scan_note = f"not re-scanned: {exc}"
    return {
        "assigned": {
            "original_name": result.entry.original_name,
            "identifier": result.entry.identifier,
            "filed_as": result.entry.filed_as,
            "prepared_location": result.entry.prepared_location,
            "moved_review_copy": result.moved_review_copy,
            "keyword": result.keyword,
            "keyword_note": result.keyword_note,
            "index_deferred": result.index_deferred,
            "scan_note": scan_note,
        },
        "state": _state(engagement),
    }


COMMANDS = {
    "state": _cmd_state,
    "scaffold": _cmd_scaffold,
    "sort": _cmd_sort,
    "priors": _cmd_priors,
    "rollover": _cmd_rollover,
    "scan": _cmd_scan,
    "reset": _cmd_reset,
    "templates": _cmd_templates,
    "list": _cmd_list,
    "create": _cmd_create,
    "assign": _cmd_assign,
    "check": _cmd_check,
    "unlock": _cmd_unlock,
}


def main(argv: list[str]) -> int:
    if not argv or argv[0] not in COMMANDS:
        print(json.dumps({"error": f"usage: tracker.api {'|'.join(COMMANDS)}"}))
        return 1
    try:
        payload = COMMANDS[argv[0]](argv[1:])
    except (ManifestError, ScanLockedError, FilingError) as exc:
        print(json.dumps({"error": str(exc)}))
        return 1
    except Exception as exc:  # surface anything else as JSON, not a traceback
        print(json.dumps({"error": f"{exc.__class__.__name__}: {exc}"}))
        return 1
    print(json.dumps(payload))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
