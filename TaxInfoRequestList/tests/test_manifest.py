"""Tests for tracker/manifest.py — the manifest layer, no OneDrive needed."""

import datetime as dt
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookClass

from tracker.manifest import (
    HEADERS,
    SHEET_NAME,
    ManifestError,
    Override,
    RequestItem,
    Status,
    StatusUpdate,
    create_template,
    load_manifest,
    pending_path,
    write_statuses,
)

SAMPLE_ITEMS = [
    RequestItem(
        identifier="A01",
        document="Dec 2025 Bank Statement",
        period="Dec 2025",
        allowed_extensions=("pdf",),
        required_keywords=("Chase",),
        date_pattern=r"(?i)december\s+2025|12/\d{1,2}/2025",
    ),
    RequestItem(
        identifier="A02",
        document="Monthly Bank Statements FY2025",
        period="FY2025",
        expected_count=12,
        allowed_extensions=("pdf", ".CSV"),  # dot + case normalize on load
        min_size_kb=10,
        any_keywords=("statement", "account summary"),
    ),
    RequestItem(
        identifier="B01",
        document="Payroll Register Q4",
        period="Q4 2025",
        allowed_extensions=("xlsx", "csv"),
        manual_override=Override.WAIVED,
    ),
]


@pytest.fixture
def manifest(tmp_path):
    return create_template(tmp_path / "_manifest.xlsx", SAMPLE_ITEMS)


# ---------------------------------------------------------------- loading ----


def test_template_load_roundtrip(manifest):
    items = load_manifest(manifest)
    assert [i.identifier for i in items] == ["A01", "A02", "B01"]

    a01 = items[0]
    assert a01.document == "Dec 2025 Bank Statement"
    assert a01.period == "Dec 2025"
    assert a01.expected_count == 1          # default
    assert a01.min_size_kb == 5             # default
    assert a01.allowed_extensions == ("pdf",)
    assert a01.required_keywords == ("Chase",)
    assert a01.any_keywords == ()
    assert a01.date_pattern.startswith("(?i)december")
    assert a01.manual_override == ""
    assert a01.status == ""
    assert a01.received_date is None
    assert a01.row == 2

    a02 = items[1]
    assert a02.expected_count == 12
    assert a02.min_size_kb == 10
    assert a02.allowed_extensions == ("pdf", "csv")  # normalized
    assert a02.any_keywords == ("statement", "account summary")

    assert items[2].manual_override == Override.WAIVED


def test_template_refuses_overwrite(manifest):
    with pytest.raises(ManifestError, match="Refusing to overwrite"):
        create_template(manifest)


def test_missing_manifest(tmp_path):
    with pytest.raises(ManifestError, match="not found"):
        load_manifest(tmp_path / "nope.xlsx")


def test_missing_column_rejected(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for idx, header in enumerate(h for h in HEADERS if h != "Status"):
        ws.cell(row=1, column=idx + 1, value=header)
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(ManifestError, match="missing column.*Status"):
        load_manifest(path)


def _write_cell(path, row, header, value):
    wb = load_workbook(path)
    ws = wb[SHEET_NAME]
    col = HEADERS.index(header) + 1
    ws.cell(row=row, column=col, value=value)
    wb.save(path)


def test_duplicate_identifier_rejected(manifest):
    _write_cell(manifest, 3, "Identifier", "A01")
    with pytest.raises(ManifestError, match="Duplicate identifier 'A01'.*rows 2 and 3"):
        load_manifest(manifest)


def test_bad_regex_rejected(manifest):
    _write_cell(manifest, 2, "Date Pattern", "([unclosed")
    with pytest.raises(ManifestError, match="Row 2.*not a valid regex"):
        load_manifest(manifest)


def test_bad_expected_count_rejected(manifest):
    _write_cell(manifest, 2, "Expected Count", "twelve")
    with pytest.raises(ManifestError, match="Row 2.*whole number"):
        load_manifest(manifest)


def test_unknown_status_rejected(manifest):
    _write_cell(manifest, 2, "Status", "Done-ish")
    with pytest.raises(ManifestError, match="Row 2.*Status must be one of"):
        load_manifest(manifest)


def test_unknown_override_rejected(manifest):
    _write_cell(manifest, 2, "Manual Override", "Maybe")
    with pytest.raises(ManifestError, match="Row 2.*Manual Override"):
        load_manifest(manifest)


def test_blank_rows_skipped(manifest):
    _write_cell(manifest, 6, "Identifier", "C01")  # row 5 left entirely blank
    _write_cell(manifest, 6, "Document", "Trial Balance")
    items = load_manifest(manifest)
    assert [i.identifier for i in items] == ["A01", "A02", "B01", "C01"]
    assert items[-1].row == 6


# ------------------------------------------------------------- write-back ----


def test_write_statuses_roundtrip(manifest):
    updates = {
        "A01": StatusUpdate(
            status=Status.RECEIVED,
            file_count=1,
            received_date=dt.date(2026, 7, 8),
        ),
        "A02": StatusUpdate(
            status=Status.PARTIAL,
            file_count=7,
            validation_notes="7 of 12 monthly statements received",
        ),
    }
    assert write_statuses(manifest, updates) is True

    items = {i.identifier: i for i in load_manifest(manifest)}
    assert items["A01"].status == Status.RECEIVED
    assert items["A01"].received_date == dt.date(2026, 7, 8)
    assert items["A01"].file_count == 1
    assert items["A02"].status == Status.PARTIAL
    assert items["A02"].file_count == 7
    assert "7 of 12" in items["A02"].validation_notes
    # untouched rows and accountant columns intact
    assert items["B01"].status == ""
    assert items["A01"].required_keywords == ("Chase",)


def test_write_unknown_identifier_dropped(manifest, caplog):
    assert write_statuses(manifest, {"Z99": StatusUpdate(status=Status.MISSING)}) is True
    assert "Z99" in caplog.text
    assert not pending_path(manifest).exists()


def test_locked_write_defers_to_sidecar(manifest, monkeypatch):
    def locked_save(self, filename):
        raise PermissionError(f"[Errno 13] locked: {filename}")

    monkeypatch.setattr(WorkbookClass, "save", locked_save)
    updates = {"A01": StatusUpdate(status=Status.MISSING, validation_notes="no files")}
    assert write_statuses(manifest, updates, retries=2, retry_delay=0.01) is False

    sidecar = pending_path(manifest)
    assert sidecar.exists()
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    assert payload["A01"]["status"] == Status.MISSING


def test_pending_sidecar_merged_and_cleared(manifest):
    pending_path(manifest).write_text(
        json.dumps(
            {
                "A01": {  # stale: superseded by the new update below
                    "status": Status.MISSING,
                    "file_count": 0,
                    "received_date": None,
                    "validation_notes": "",
                },
                "B01": {  # only in sidecar: must still be applied
                    "status": Status.PENDING_SYNC,
                    "file_count": 1,
                    "received_date": None,
                    "validation_notes": "cloud-only placeholder",
                },
            }
        ),
        encoding="utf-8",
    )
    new = {
        "A01": StatusUpdate(
            status=Status.RECEIVED, file_count=1, received_date=dt.date(2026, 7, 8)
        )
    }
    assert write_statuses(manifest, new) is True
    assert not pending_path(manifest).exists()

    items = {i.identifier: i for i in load_manifest(manifest)}
    assert items["A01"].status == Status.RECEIVED  # new beat stale sidecar
    assert items["B01"].status == Status.PENDING_SYNC
    assert "placeholder" in items["B01"].validation_notes


def test_corrupt_sidecar_quarantined(manifest, caplog):
    pending_path(manifest).write_text("{not json", encoding="utf-8")
    assert write_statuses(manifest, {"A01": StatusUpdate(status=Status.MISSING)}) is True
    assert not pending_path(manifest).exists()
    assert pending_path(manifest).with_suffix(".corrupt.json").exists()
    assert "Unreadable pending sidecar" in caplog.text
