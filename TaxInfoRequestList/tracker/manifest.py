"""Manifest layer for the Client Document Tracker (component 1, docs/ROADMAP.md).

Owns everything about ``_manifest.xlsx``: the schema, loading and validating
rows into :class:`RequestItem` dataclasses, and writing scanner status back
with Excel-lock resilience (retry with backoff, then defer updates to a
``_manifest.pending.json`` sidecar that is merged on the next write).

This module never touches client files — only the manifest workbook and its
sidecar. All values are validated on load and fail loudly with row context so
a malformed manifest can never silently produce wrong statuses.
"""

from __future__ import annotations

import datetime as dt
import json
import logging
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

log = logging.getLogger("tracker.manifest")

SHEET_NAME = "Requests"
UNFILED_SHEET_NAME = "Unfiled"

# ---------------------------------------------------------------- schema ----

COL_IDENTIFIER = "Identifier"
COL_DOCUMENT = "Document"
COL_PERIOD = "Period"
COL_EXPECTED_COUNT = "Expected Count"
COL_ALLOWED_EXTENSIONS = "Allowed Extensions"
COL_MIN_SIZE_KB = "Min Size KB"
COL_REQUIRED_KEYWORDS = "Required Keywords"
COL_ANY_KEYWORDS = "Any Keywords"
COL_DATE_PATTERN = "Date Pattern"
COL_MANUAL_OVERRIDE = "Manual Override"
COL_STATUS = "Status"
COL_RECEIVED_DATE = "Received Date"
COL_FILE_COUNT = "File Count"
COL_VALIDATION_NOTES = "Validation Notes"

#: Columns the accountant edits; the scanner only ever reads these.
ACCOUNTANT_COLUMNS = (
    COL_IDENTIFIER,
    COL_DOCUMENT,
    COL_PERIOD,
    COL_EXPECTED_COUNT,
    COL_ALLOWED_EXTENSIONS,
    COL_MIN_SIZE_KB,
    COL_REQUIRED_KEYWORDS,
    COL_ANY_KEYWORDS,
    COL_DATE_PATTERN,
    COL_MANUAL_OVERRIDE,
)

#: Columns the scanner writes; the accountant should not hand-edit these.
SCANNER_COLUMNS = (
    COL_STATUS,
    COL_RECEIVED_DATE,
    COL_FILE_COUNT,
    COL_VALIDATION_NOTES,
)

HEADERS = ACCOUNTANT_COLUMNS + SCANNER_COLUMNS

DEFAULT_EXPECTED_COUNT = 1
DEFAULT_MIN_SIZE_KB = 5
DATE_FORMAT = "yyyy-mm-dd"


class Status:
    """Scanner-resolved status values (plain constants, stored as strings)."""

    MISSING = "Missing"
    PARTIAL = "Partial"
    FAILED = "Failed Validation"
    RECEIVED = "Received"
    PENDING_SYNC = "Pending Sync"

    ALL = (MISSING, PARTIAL, FAILED, RECEIVED, PENDING_SYNC)


class Override:
    """Accountant judgment values that beat the automated rules."""

    ACCEPTED = "Accepted"  # treat as Received despite failed checks
    WAIVED = "Waived"      # item no longer needed

    ALL = (ACCEPTED, WAIVED)


class ManifestError(Exception):
    """A manifest could not be read, parsed, or validated."""


@dataclass(frozen=True, slots=True)
class RequestItem:
    """One row of the Requests sheet, fully parsed and validated."""

    identifier: str
    document: str
    period: str = ""
    expected_count: int = DEFAULT_EXPECTED_COUNT
    allowed_extensions: tuple[str, ...] = ()   # lowercase, no leading dot
    min_size_kb: int = DEFAULT_MIN_SIZE_KB
    required_keywords: tuple[str, ...] = ()
    any_keywords: tuple[str, ...] = ()
    date_pattern: str = ""                     # validated to compile on load
    manual_override: str = ""                  # "", Override.ACCEPTED, Override.WAIVED
    status: str = ""                           # "", or a Status.ALL value
    received_date: dt.date | None = None
    file_count: int | None = None
    validation_notes: str = ""
    row: int = 0                               # Excel row this item came from


@dataclass(frozen=True, slots=True)
class StatusUpdate:
    """Scanner output for one identifier, destined for the scanner columns."""

    status: str
    file_count: int = 0
    received_date: dt.date | None = None
    validation_notes: str = ""


@dataclass(frozen=True, slots=True)
class UnfiledEntry:
    """Something in Shared/ that belongs to no request row — never ignored,
    never moved; reported on the Unfiled sheet for the accountant to triage."""

    name: str
    kind: str                 # e.g. "loose file in Shared root"
    size_kb: float | None = None
    seen: str = ""            # ISO date of the scan that noted it


# --------------------------------------------------------------- parsing ----


def _cell_str(value: object) -> str:
    return "" if value is None else str(value).strip()


def _csv_tuple(value: object) -> tuple[str, ...]:
    return tuple(p.strip() for p in _cell_str(value).split(",") if p.strip())


def _parse_int(value: object, default: int, column: str, row: int) -> int:
    text = _cell_str(value)
    if not text:
        return default
    try:
        number = float(text)
        if number != int(number):
            raise ValueError
        return int(number)
    except ValueError:
        raise ManifestError(
            f"Row {row}: {column} must be a whole number, got {value!r}"
        ) from None


def _parse_date(value: object, column: str, row: int) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _cell_str(value)
    if not text:
        return None
    try:
        return dt.date.fromisoformat(text)
    except ValueError:
        raise ManifestError(
            f"Row {row}: {column} must be a date, got {value!r}"
        ) from None


def _parse_enum(value: object, allowed: tuple[str, ...], column: str, row: int) -> str:
    text = _cell_str(value)
    if not text:
        return ""
    for candidate in allowed:
        if text.lower() == candidate.lower():
            return candidate
    raise ManifestError(
        f"Row {row}: {column} must be one of {', '.join(allowed)} (or blank), got {value!r}"
    )


def _header_map(ws) -> dict[str, int]:
    """Map canonical column names to 1-based column indexes; fail on missing."""
    found: dict[str, int] = {}
    for idx in range(1, (ws.max_column or 0) + 1):
        text = _cell_str(ws.cell(row=1, column=idx).value)
        if text:
            found[text.lower()] = idx
    missing = [h for h in HEADERS if h.lower() not in found]
    if missing:
        raise ManifestError(
            f"Sheet {SHEET_NAME!r} is missing column(s): {', '.join(missing)}"
        )
    return {h: found[h.lower()] for h in HEADERS}


# --------------------------------------------------------------- loading ----


def load_manifest(path: Path | str) -> list[RequestItem]:
    """Load and validate every request row from ``path``.

    Reads with ``data_only=True`` so formula cells yield their cached values.
    Raises :class:`ManifestError` with row context on any invalid data.
    """
    path = Path(path)
    if not path.exists():
        raise ManifestError(f"Manifest not found: {path}")
    try:
        wb = load_workbook(path, data_only=True)
    except ManifestError:
        raise
    except Exception as exc:  # zip/corruption errors from openpyxl
        raise ManifestError(f"Could not open {path}: {exc}") from exc
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ManifestError(f"{path.name} has no {SHEET_NAME!r} sheet")
        ws = wb[SHEET_NAME]
        columns = _header_map(ws)

        items: list[RequestItem] = []
        seen: dict[str, int] = {}
        for row in range(2, (ws.max_row or 1) + 1):
            values = {
                name: ws.cell(row=row, column=idx).value
                for name, idx in columns.items()
            }
            if all(_cell_str(v) == "" for v in values.values()):
                continue  # blank spacer row

            identifier = _cell_str(values[COL_IDENTIFIER])
            if not identifier:
                raise ManifestError(f"Row {row}: {COL_IDENTIFIER} is required")
            if identifier in seen:
                raise ManifestError(
                    f"Duplicate identifier {identifier!r} (rows {seen[identifier]} and {row})"
                )
            seen[identifier] = row

            document = _cell_str(values[COL_DOCUMENT])
            if not document:
                raise ManifestError(f"Row {row}: {COL_DOCUMENT} is required")

            expected_count = _parse_int(
                values[COL_EXPECTED_COUNT], DEFAULT_EXPECTED_COUNT, COL_EXPECTED_COUNT, row
            )
            if expected_count < 1:
                raise ManifestError(f"Row {row}: {COL_EXPECTED_COUNT} must be >= 1")

            min_size_kb = _parse_int(
                values[COL_MIN_SIZE_KB], DEFAULT_MIN_SIZE_KB, COL_MIN_SIZE_KB, row
            )
            if min_size_kb < 0:
                raise ManifestError(f"Row {row}: {COL_MIN_SIZE_KB} must be >= 0")

            date_pattern = _cell_str(values[COL_DATE_PATTERN])
            if date_pattern:
                try:
                    re.compile(date_pattern)
                except re.error as exc:
                    raise ManifestError(
                        f"Row {row}: {COL_DATE_PATTERN} is not a valid regex: {exc}"
                    ) from None

            items.append(
                RequestItem(
                    identifier=identifier,
                    document=document,
                    period=_cell_str(values[COL_PERIOD]),
                    expected_count=expected_count,
                    allowed_extensions=tuple(
                        e.lower().lstrip(".")
                        for e in _csv_tuple(values[COL_ALLOWED_EXTENSIONS])
                    ),
                    min_size_kb=min_size_kb,
                    required_keywords=_csv_tuple(values[COL_REQUIRED_KEYWORDS]),
                    any_keywords=_csv_tuple(values[COL_ANY_KEYWORDS]),
                    date_pattern=date_pattern,
                    manual_override=_parse_enum(
                        values[COL_MANUAL_OVERRIDE], Override.ALL, COL_MANUAL_OVERRIDE, row
                    ),
                    status=_parse_enum(values[COL_STATUS], Status.ALL, COL_STATUS, row),
                    received_date=_parse_date(
                        values[COL_RECEIVED_DATE], COL_RECEIVED_DATE, row
                    ),
                    file_count=(
                        _parse_int(values[COL_FILE_COUNT], -1, COL_FILE_COUNT, row)
                        if _cell_str(values[COL_FILE_COUNT])
                        else None
                    ),
                    validation_notes=_cell_str(values[COL_VALIDATION_NOTES]),
                    row=row,
                )
            )
        return items
    finally:
        wb.close()


# ------------------------------------------------------------- write-back ----


def pending_path(manifest_path: Path) -> Path:
    """Sidecar file that holds updates deferred by an Excel lock."""
    return manifest_path.with_name(manifest_path.stem + ".pending.json")


def _load_pending(manifest_path: Path) -> dict[str, StatusUpdate]:
    sidecar = pending_path(manifest_path)
    if not sidecar.exists():
        return {}
    try:
        raw = json.loads(sidecar.read_text(encoding="utf-8"))
        return {
            ident: StatusUpdate(
                status=u["status"],
                file_count=int(u.get("file_count", 0)),
                received_date=(
                    dt.date.fromisoformat(u["received_date"])
                    if u.get("received_date")
                    else None
                ),
                validation_notes=u.get("validation_notes", ""),
            )
            for ident, u in raw.items()
        }
    except (json.JSONDecodeError, KeyError, TypeError, ValueError) as exc:
        # Preserve the evidence rather than retry-looping on garbage forever.
        corrupt = sidecar.with_suffix(".corrupt.json")
        sidecar.replace(corrupt)
        log.error("Unreadable pending sidecar moved to %s: %s", corrupt.name, exc)
        return {}


def _save_pending(manifest_path: Path, updates: Mapping[str, StatusUpdate]) -> None:
    payload = {
        ident: {
            "status": u.status,
            "file_count": u.file_count,
            "received_date": u.received_date.isoformat() if u.received_date else None,
            "validation_notes": u.validation_notes,
        }
        for ident, u in updates.items()
    }
    pending_path(manifest_path).write_text(
        json.dumps(payload, indent=2), encoding="utf-8"
    )


def _write_unfiled_sheet(wb: Workbook, entries: list[UnfiledEntry]) -> None:
    """Replace the Unfiled sheet with a fresh snapshot (it is scanner-owned)."""
    if UNFILED_SHEET_NAME in wb.sheetnames:
        del wb[UNFILED_SHEET_NAME]
    ws = wb.create_sheet(UNFILED_SHEET_NAME)
    for idx, (header, width) in enumerate(
        [("Item", 50), ("Type", 34), ("Size KB", 10), ("Noted On", 12)], start=1
    ):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    for row, entry in enumerate(entries, start=2):
        ws.cell(row=row, column=1, value=entry.name)
        ws.cell(row=row, column=2, value=entry.kind)
        ws.cell(row=row, column=3, value=entry.size_kb)
        ws.cell(row=row, column=4, value=entry.seen or None)


def _apply_updates(
    path: Path,
    updates: Mapping[str, StatusUpdate],
    unfiled: list[UnfiledEntry] | None = None,
) -> None:
    """Open the workbook for writing and set scanner columns. May raise
    PermissionError if Excel holds the file locked (caller retries)."""
    wb = load_workbook(path)  # NOT data_only: preserves any formulas on save
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ManifestError(f"{path.name} has no {SHEET_NAME!r} sheet")
        ws = wb[SHEET_NAME]
        columns = _header_map(ws)

        rows_by_identifier = {
            _cell_str(ws.cell(row=r, column=columns[COL_IDENTIFIER]).value): r
            for r in range(2, (ws.max_row or 1) + 1)
        }

        for identifier, update in updates.items():
            row = rows_by_identifier.get(identifier)
            if row is None:
                # Row was deleted/renamed since the scan; nothing to update.
                log.warning(
                    "Identifier %r not found in %s; update dropped", identifier, path.name
                )
                continue
            ws.cell(row=row, column=columns[COL_STATUS], value=update.status)
            date_cell = ws.cell(
                row=row,
                column=columns[COL_RECEIVED_DATE],
                value=update.received_date,
            )
            if update.received_date is not None:
                date_cell.number_format = DATE_FORMAT
            ws.cell(row=row, column=columns[COL_FILE_COUNT], value=update.file_count)
            ws.cell(
                row=row,
                column=columns[COL_VALIDATION_NOTES],
                value=update.validation_notes or None,
            )
        if unfiled is not None:
            _write_unfiled_sheet(wb, unfiled)
        wb.save(path)
    finally:
        wb.close()


def write_statuses(
    path: Path | str,
    updates: Mapping[str, StatusUpdate],
    *,
    unfiled: list[UnfiledEntry] | None = None,
    retries: int = 5,
    retry_delay: float = 0.5,
) -> bool:
    """Write scanner columns for the given identifiers, lock-resiliently.

    Any updates deferred by a previous locked write (the pending sidecar) are
    merged first; new updates win for the same identifier. If the workbook is
    still locked after ``retries`` attempts with exponential backoff, the
    merged updates are saved to the sidecar and applied on the next call.

    ``unfiled`` (when not None) replaces the Unfiled sheet in the same save.
    It is a per-scan snapshot, so it is not sidecar-deferred — a locked write
    simply drops it and the next successful scan rewrites it.

    Returns True if the workbook was written, False if deferred to sidecar.
    """
    path = Path(path)
    merged: dict[str, StatusUpdate] = _load_pending(path)
    merged.update(updates)
    if not merged and unfiled is None:
        return True

    delay = retry_delay
    for attempt in range(1, retries + 1):
        try:
            _apply_updates(path, merged, unfiled)
            pending_path(path).unlink(missing_ok=True)
            return True
        except PermissionError as exc:
            log.warning(
                "Manifest locked (attempt %d/%d): %s", attempt, retries, exc
            )
            if attempt < retries:
                time.sleep(delay)
                delay *= 2

    _save_pending(path, merged)
    log.error(
        "Manifest still locked after %d attempts; %d update(s) deferred to %s",
        retries,
        len(merged),
        pending_path(path).name,
    )
    return False


# --------------------------------------------------------------- template ----

_COLUMN_WIDTHS = {
    COL_IDENTIFIER: 11,
    COL_DOCUMENT: 38,
    COL_PERIOD: 12,
    COL_EXPECTED_COUNT: 15,
    COL_ALLOWED_EXTENSIONS: 19,
    COL_MIN_SIZE_KB: 12,
    COL_REQUIRED_KEYWORDS: 24,
    COL_ANY_KEYWORDS: 24,
    COL_DATE_PATTERN: 28,
    COL_MANUAL_OVERRIDE: 16,
    COL_STATUS: 17,
    COL_RECEIVED_DATE: 14,
    COL_FILE_COUNT: 11,
    COL_VALIDATION_NOTES: 50,
}


def create_template(path: Path | str, items: Iterable[RequestItem] = ()) -> Path:
    """Create a fresh manifest workbook at ``path``, optionally seeded with rows.

    Refuses to overwrite an existing file — a live manifest carries scanner
    state and must never be clobbered by a re-run.
    """
    path = Path(path)
    if path.exists():
        raise ManifestError(f"Refusing to overwrite existing manifest: {path}")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(idx)].width = _COLUMN_WIDTHS[header]
    ws.freeze_panes = "A2"

    for row, item in enumerate(items, start=2):
        ws.cell(row=row, column=1, value=item.identifier)
        ws.cell(row=row, column=2, value=item.document)
        ws.cell(row=row, column=3, value=item.period or None)
        ws.cell(row=row, column=4, value=item.expected_count)
        ws.cell(row=row, column=5, value=", ".join(item.allowed_extensions) or None)
        ws.cell(row=row, column=6, value=item.min_size_kb)
        ws.cell(row=row, column=7, value=", ".join(item.required_keywords) or None)
        ws.cell(row=row, column=8, value=", ".join(item.any_keywords) or None)
        ws.cell(row=row, column=9, value=item.date_pattern or None)
        ws.cell(row=row, column=10, value=item.manual_override or None)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path
